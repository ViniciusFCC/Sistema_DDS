"""
Teste de integração da planilha (pneu, 10_TIRE) na versão vigente, V3.0. Recalcula pelo Excel (COM) ou pelo LibreOffice (EDR-044; recalculo_planilha.py).
As abas de suspensão da V3.0 têm teste próprio: tests/test_workbook_suspensao.py.
Uso: uv run python -X utf8 tests/test_workbook.py [<V3_0.xlsx> <V2.1.xlsx> <layout_V2_4.json> <tire.tir>]  (padrão: excel/ e data/)
Cada cenário altera uma CÓPIA, recalcula e lê; o arquivo entregue nunca é salvo pelo motor de recálculo. Namespace TC-WB24-xxx.
"""
import sys, os, shutil, subprocess, tempfile, json, math, random
_TMP = tempfile.TemporaryDirectory(prefix="dds_teste_", ignore_cleanup_errors=True); tempfile.tempdir = _TMP.name  # tudo que o teste cria em temp fica aqui e é apagado no fim
import openpyxl
import tire_model as t

import recalculo_planilha as rc
MOTOR, DETALHE = rc.escolher()
if MOTOR is None:
    print(f"NÃO EXECUTÁVEL: nenhum motor de recálculo ({DETALHE})"); sys.exit(3)
print(f"MOTOR: {MOTOR} ({DETALHE})")
_RAIZ = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
V22, V21, LAY, TIR = sys.argv[1:5] if len(sys.argv) >= 5 else [os.path.join(_RAIZ, *p) for p in (
    ("excel", "FSAE_EV_Suspension_Fase1_V3_0.xlsx"), ("excel", "FSAE_EV_Suspension_Fase1_V2.1.xlsx"),
    ("data", "layout_V2_4.json"), ("data", "Hoosier_R25B_13x7_TTC_provisional.tir"))]
L = json.load(open(LAY)); R, PR = L["rows"], L["prow"]
ds = t.load_dataset(TIR, source_axis_system="ISO_TYDEX")       # declarado na V3.0 (EDR-039 reaberto: .tir PAC2002 = TYDEX W, ISO)
ds_sae = t.load_dataset(TIR, source_axis_system="SAE_CALSPAN")  # modo alternativo, ainda suportado e testado
results = []
COLS = "BCDEF"; CORNER = {"B": "FL", "C": "FR", "D": "RL", "E": "RR", "F": "FR"}


def check(name, cond, detail=""):
    results.append((name, "PASS" if cond else "FAIL")); print(f"[{'PASS' if cond else 'FAIL'}] {name}  {detail}")


def scenario(mods):
    d = tempfile.mkdtemp(); f = os.path.join(d, "t.xlsx"); shutil.copy(V22, f)
    wb = openpyxl.load_workbook(f); ws = wb["10_TIRE"]
    for addr, val in mods.items(): ws[addr].value = val
    wb.save(f)
    info = rc.recalcular(f, MOTOR, 120)
    v = openpyxl.load_workbook(f, data_only=True)
    return info, v["10_TIRE"], v


def setcols(vals):
    """vals: col -> (fz, alpha_deg, camber_deg, kappa)"""
    m = {}
    for c, (fz, a, g, k) in vals.items():
        m[f"{c}{R['fz']}"] = fz; m[f"{c}{R['alpha']}"] = a; m[f"{c}{R['camber']}"] = g; m[f"{c}{R['kappa']}"] = k
    return m


def py(c, fz, a, g, k, dsx=None):
    fx, fy, _ = t.wheel_forces(dsx or ds, CORNER[c], fz, math.radians(a), math.radians(g), k)
    return fx, fy, t.mu_effective(fx, fy, fz)


def compare(ws, vals, tag, dsx=None):
    worst = 0.0; bad = []
    for c, (fz, a, g, k) in vals.items():
        e = (ws[f"{c}{R['o_fx']}"].value, ws[f"{c}{R['o_fy']}"].value, ws[f"{c}{R['o_mu']}"].value)
        try:
            p = py(c, fz, a, g, k, dsx)
        except t.TireRangeError:
            bad.append((c, "Python bloqueou mas Excel deu", e)); continue
        for ev, pv in zip(e, p):
            if not isinstance(ev, (int, float)):
                bad.append((c, "não numérico", ev)); break
            worst = max(worst, abs(ev - pv) / max(1.0, abs(pv)))
    return worst, bad


# ---------- preservação (sem recalcular) ----------
# V3.0: só mudam as células listadas em atualizar_planilha.alteracoes (EDR-045 a EDR-049); todo o resto da V2.1 fica igual.
import atualizar_planilha as ap
PERMITIDAS = {(x["aba"], x["cel"]) for x in ap.alteracoes(R["Ky"])}
o, n = openpyxl.load_workbook(V21), openpyxl.load_workbook(V22)
alt = [(s.title, c.coordinate) for s in o for row in s.iter_rows() for c in row
       if isinstance(c.value, str) and c.value.startswith("=") and n[s.title][c.coordinate].value != c.value
       and (s.title, c.coordinate) not in PERMITIDAS]
cnt21 = sum(1 for s in o for row in s.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
check("TC-WB24-001: as fórmulas da V2.1 permanecem idênticas na V3.0, fora das células da lista de alterações da V3.0", not alt and cnt21 == 88, f"alteradas fora da lista={alt}; V2.1={cnt21}")
check("TC-WB24-002: abas de suspensão (05, 06, 07, 08, 09, 11): toda célula fora da lista de alterações da V3.0 é idêntica à V2.1",
      all(o[s][c.coordinate].value == c.value for s in ("05_VEHICLE", "06_SUSPENSION", "07_KINEMATICS", "08_VEHICLE_STATE", "09_LOADS", "11_DESIGN")
          for row in n[s].iter_rows() for c in row if (s, c.coordinate) not in PERMITIDAS))

# ---------- cenário base ----------
info, ws, wbv = scenario({})
check("TC-WB24-003a: recálculo sem erros (#REF!, #VALUE!, …)", info["total_errors"] == 0, f"fórmulas={info['total_formulas']}")
base_vals = {c: (ws[f"{c}{R['fz']}"].value, ws[f"{c}{R['alpha']}"].value, ws[f"{c}{R['camber']}"].value, ws[f"{c}{R['kappa']}"].value) for c in COLS}
w, bad = compare(ws, base_vals, "base")
check("TC-WB24-003b: base — Excel == Python nos 4 cantos + VERIF (Fx, Fy, μ), tol. 1e-6 relativa", w < 1e-6 and not bad, f"pior dif relativa={w:.2e}; {bad}")
vv = wbv["13_VALIDATION"]
check("TC-WB24-003c: 13_VALIDATION: Fx/Fy/μ do caso VERIF = PASS; dataset OK; 4 cantos válidos", [vv["B40"].value, vv["B41"].value, vv["B42"].value, vv["B36"].value, vv["B43"].value] == ["PASS", "PASS", "PASS", "OK", 4])
check("TC-WB24-003d: resultados existentes de 13_VALIDATION preservados (TC-004/005, REQ-001..004 como na V2.1)",
      [vv[k].value for k in ("B5", "B7", "B19", "B21", "B22", "B30", "B32")] == ["PASS"] * 7)

# ---------- aleatórios (esquerda e direita, faixas completas) ----------
rnd = random.Random(20261012)
worst_all, bads = 0.0, []
for s in range(4):
    vals = {c: (rnd.uniform(100, 1521), rnd.uniform(-15, 15) * 0.999, rnd.uniform(-4, 4) * 0.999, rnd.uniform(-0.25, 0.25)) for c in COLS}
    info, ws, _ = scenario(setcols(vals))
    w, bad = compare(ws, vals, s)
    worst_all = max(worst_all, w); bads += bad
    if info["total_errors"]: bads.append(("erros", info["total_errors"]))
check("TC-WB24-004: 4 cenários aleatórios × 5 colunas (20 pontos; Fz 100–1521, α ±15°, γ ±4°, κ ±0,25; FL/RL = rodas esquerdas): Excel == Python (1e-6)",
      worst_all < 1e-6 and not bads, f"pior dif relativa={worst_all:.2e}; {bads}")

# ---------- coeficientes e fatores L* perturbados (detecta referência de célula trocada e L* ignorado) ----------
pr = random.Random(99)
p2 = dict(ds.p); mods = {}
for k in list(PR):
    if k in ("EPS", "IP", "IP_NOM") or k.startswith(("FZ", "ALP", "CAM", "KPU", "PP", "FNOMIN")):
        continue
    nv = ds.p[k] * pr.uniform(0.9, 1.1) if ds.p[k] != 0 else 0.0
    if k.startswith("L"):
        nv = pr.uniform(0.9, 1.1)
    p2[k] = nv; mods[f"B{PR[k]}"] = nv
ds2 = t.TireDataset(p=p2, strings=ds.strings, axis_system="ISO_TYDEX")
vals = {c: (rnd2 := random.Random(5 + i)).uniform(150, 1400) for i, c in enumerate(COLS)}
pts = {c: (vals[c], random.Random(11 + i).uniform(-12, 12), random.Random(21 + i).uniform(-3.5, 3.5), random.Random(31 + i).uniform(-0.2, 0.2)) for i, c in enumerate(COLS)}
info, ws, _ = scenario({**mods, **setcols(pts)})
worst = 0.0; badp = []
for c, (fz_, a_, g_, k_) in pts.items():
    fx_, fy_, _ = t.wheel_forces(ds2, CORNER[c], fz_, math.radians(a_), math.radians(g_), k_)
    ev = (ws[f"{c}{R['o_fx']}"].value, ws[f"{c}{R['o_fy']}"].value)
    if not all(isinstance(x, (int, float)) for x in ev): badp.append((c, ev)); continue
    worst = max(worst, abs(ev[0] - fx_) / max(1, abs(fx_)), abs(ev[1] - fy_) / max(1, abs(fy_)))
check("TC-WB24-004b: com TODOS os coeficientes P*/R* e fatores L* perturbados (±10%), Excel == Python nos 5 pontos (detecta referência de célula trocada e L* ignorado)",
      worst < 1e-6 and not badp and info["total_errors"] == 0, f"pior dif relativa={worst:.2e}; {badp}")

# ---------- sistema de eixos do arquivo (V2.4) ----------
AX = f"B{R['m_axis']}"
check("TC-WB24-004c: planilha declara ISO_TYDEX e s_α = s_Fy = s_γ = +1 no cenário base (EDR-039 reaberto)", ws0[AX].value == "ISO_TYDEX" and [ws0[f"B{R[k]}"].value for k in ("m_sa", "m_sf", "m_sg")] == [1, 1, 1]) if (ws0 := scenario({})[1]) else None
rnd3 = random.Random(77)
vals = {c: (rnd3.uniform(100, 1521), rnd3.uniform(-14, 14), rnd3.uniform(-3.9, 3.9), rnd3.uniform(-0.24, 0.24)) for c in COLS}
info, wsx, _ = scenario({**setcols(vals), AX: "SAE_CALSPAN"})
w, bad = compare(wsx, vals, "sae", ds_sae)
check("TC-WB24-004d: com eixo = SAE_CALSPAN (modo alternativo) o Excel == Python SAE nos 5 pontos aleatórios", w < 1e-6 and not bad and info["total_errors"] == 0, f"pior dif relativa={w:.2e}")
fy_alt = {c: wsx[f"{c}{R['o_fy']}"].value for c in COLS}
info, wsy, _ = scenario({**setcols(vals)})
diff_modes = max(abs(wsy[f"{c}{R['o_fy']}"].value - fy_alt[c]) for c in COLS)
check("TC-WB24-004e: SAE_CALSPAN e ISO_TYDEX dão Fy diferentes na planilha (a transformação está ativa)", diff_modes > 1.0, f"dif máx Fy={diff_modes:.1f} N")
info, wsz, wbz = scenario({AX: "UNKNOWN"})
check("TC-WB24-004f: eixo 'UNKNOWN' (ou qualquer outro texto) => dataset INVÁLIDO, todos os cantos N/A, 13_VALIDATION sem PASS, sem erros de fórmula",
      wsz[f"B{R['m_valid']}"].value == "INVÁLIDO" and all(str(wsz[f"{c}{R['status']}"].value).startswith("N/A") for c in COLS) and wbz["13_VALIDATION"]["B40"].value.startswith("N/A") and info["total_errors"] == 0)

# ---------- fronteiras e bloqueios ----------
ok = (100.0, 0.0, 0.0, 0.0), (1521.0, 0.0, 0.0, 0.0), (700.0, -14.9, 3.9, 0.249)
vals = dict(zip("BCD", ok)); vals["E"] = (700.0, 0.0, 0.0, 0.0); vals["F"] = (700.0, 0.0, 0.0, 0.0)
info, ws, _ = scenario(setcols(vals))
w, bad = compare(ws, {c: vals[c] for c in "BCD"}, "borda")
check("TC-WB24-005a: Fz=FZMIN, Fz=FZMAX e ponto perto dos limites de α/γ/κ são aceitos e batem com Python", all(ws[f"{c}{R['status']}"].value == "OK" for c in "BCD") and w < 1e-6 and not bad)
cases = {"B": ((0.0, 0, 0, 0), "liftoff"), "C": ((99.0, 0, 0, 0), "Fz fora"), "D": ((1522.0, 0, 0, 0), "Fz fora"),
         "E": ((700.0, 16.0, 0, 0), "α fora"), "F": ((700.0, 0, 4.2, 0), "camber fora")}
info, ws, _ = scenario(setcols({c: v[0] for c, v in cases.items()}))
sts = {c: ws[f"{c}{R['status']}"].value for c in cases}
check("TC-WB24-005b: Fz=0, Fz<FZMIN, Fz>FZMAX, α=16°, γ=4,2° => status N/A com motivo; saídas 'N/A' (sem extrapolar, sem erro)",
      all(sts[c].startswith("N/A") and cases[c][1] in sts[c] for c in cases) and all(ws[f"{c}{R['o_fx']}"].value == "N/A" and ws[f"{c}{R['o_mu']}"].value == "N/A" for c in cases) and info["total_errors"] == 0, str(sts))
info, ws, _ = scenario({**setcols({"B": (700.0, 0, 0, 0.30), "C": (700.0, None, 0, 0), "D": (700.0, 0, 0, 0), "E": (700.0, 0, 0, 0), "F": (700.0, 0, 0, 0)}), f"D{R['fz']}": "N/A"})
check("TC-WB24-005c: κ=0,30 => N/A; α ausente => 'entrada ausente'; Fz 'N/A' vindo de upstream => N/A (sem #VALUE!)",
      ws[f"B{R['status']}"].value.startswith("N/A - κ") and "ausente" in ws[f"C{R['status']}"].value and "ausente" in ws[f"D{R['status']}"].value and info["total_errors"] == 0)

# ---------- dataset inválido ----------
info, ws, wbv = scenario({f"B{PR['IP']}": 90000.0})
check("TC-WB24-006a: IP != IP_NOM => dataset INVÁLIDO e todos os cantos N/A; 13_VALIDATION não indica PASS",
      ws[f"B{R['m_valid']}"].value == "INVÁLIDO" and all("inválido" in ws[f"{c}{R['status']}"].value for c in COLS) and wbv["13_VALIDATION"]["B40"].value.startswith("N/A") and info["total_errors"] == 0)
info, ws, _ = scenario({f"B{PR['PPY1']}": 0.1})
check("TC-WB24-006b: PP* != 0 => dataset INVÁLIDO (efeito de pressão não implementado)", ws[f"B{R['m_valid']}"].value == "INVÁLIDO")
info, ws, _ = scenario({f"B{PR['FZMAX']}": 50.0})
check("TC-WB24-006c: FZMAX <= FZMIN => dataset INVÁLIDO", ws[f"B{R['m_valid']}"].value == "INVÁLIDO")

# ---------- coerência de espelhamento na planilha ----------
vals = {"B": (700.0, 3.0, 1.5, 0.05), "C": (700.0, -3.0, 1.5, 0.05), "D": (700.0, 0, 0, 0), "E": (700.0, 0, 0, 0), "F": (700.0, 0, 0, 0)}
info, ws, _ = scenario(setcols(vals))
fyl, fyr = ws[f"B{R['o_fy']}"].value, ws[f"C{R['o_fy']}"].value
check("TC-WB24-007: no Excel, Fy_FL(α) = −Fy_FR(−α) e Fx igual (mesmo γ do projeto)", abs(fyl + fyr) < 1e-6 * max(1, abs(fyl)) and abs(ws[f"B{R['o_fx']}"].value - ws[f"C{R['o_fx']}"].value) < 1e-6, f"Fy_FL={fyl:.3f} Fy_FR(-α)={fyr:.3f}")

print("\n--- Resumo ---"); p = sum(1 for _, s in results if s == "PASS"); f = len(results) - p
print(f"{p} PASS, {f} FAIL, de {len(results)} testes")
sys.exit(1 if f else 0)
