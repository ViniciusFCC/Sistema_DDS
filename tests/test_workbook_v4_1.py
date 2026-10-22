"""
Teste da V4.1 no motor de recálculo (incremento 4, equilíbrio na planilha; EDR-044): passo 2 (análise de par) × Python.
Recalcula CÓPIAS (aba_equilibrio.copia_v41: V4.0 da tag v4.0 + abas 16-18 + nomes da V4.1 + valores de teste) pelo Excel
(COM) ou pelo LibreOffice (recalculo_planilha.py; DDS_RECALC=excel|libreoffice força um). A V4.0 de referência em excel/
e a V3.0 viva nunca são abertas para escrita. Sem motor: NÃO EXECUTÁVEL (código 3), que nunca conta como aprovado.
Uso: uv run python -X utf8 tests/test_workbook_v4_1.py
Namespace TC-WB41-xxx (IDs autorizados pelo usuário: 001 a 006; este script tem, por enquanto, o 004 e o 005).

TC-WB41-004: em cada caso (valores de TESTE, só nas cópias; não são o carro), o caso que a 16_EQUILIBRIO calcula, lido
por nomes (caso_planilha.caso_da_planilha_v41), vai a equilibrio.analise_par, e a planilha tem de dar o mesmo:
- por linha da grade pedida (17) e da refinada (18, contra equilibrio._varredura com passo/2 e 2·n − 1): Fz e γ das
  quatro rodas, capacidade e motivo de cada eixo e, na 17, a folga (erro relativo <= 1e-9);
- no resumo: S_f, S_r (17 e 18), resolução, veredito, motivo do N/A (texto exato), a_y limite, C_N, limitações (texto
  exato, uma por roda e limite), tempo de volta no skidpad (com R) e pior origem (pior_origem_do_caso_v41).
Casos: C1 (V4.1 como gerada, demonstrativa); C2 simétrico com LMUY 0,45 (esperado NEUTRO); C3 = C2 com K_ARB dianteiro × 3
(PUSH) e C3b com o traseiro × 3 (LOOSE); C4 = C3 com as hipóteses ligadas (R 8 m, h2, h3, CG lateral 10 mm com h7,
sentido "Direita", h8 com ρ 1,2 e CzA 2,0 e downforce fixa 0, h9 com CdA 1,0 e Crr 0,015, toe na frente); C4b = C4
com δi = 22° na 07_KINEMATICS (% Ackermann ≠ 0, exercita a bisseção do Ackermann); C5 = C2 com h_CG 0,7 m (a roda interna
descola antes de saturar: N/A por descolamento).
O C4 usa toe 0,25°/0,10° e camber estático −1,5°/−1,0° em FL/FR (assimétricos), para o espelho ser testado.
Variantes do Review Focus: RF1 = C4 sem CzA (N/A "CzA a confirmar" na planilha e o mesmo ValueError no Python, nunca o
resultado da h8 desligada); RF3 = C3 com LKY × 0,9 e LMUY escalado até o eixo traseiro saturar na grade pedida e não
na refinada, perto de 2 g (resolução = passo da grade pedida, regra de equilibrio._resolucao); RF4 está no C1 (limitação que só a grade refinada
teria não entra no aviso: a planilha agrega só a 17, como equilibrio._limitacoes).
TC-WB41-005: nas MESMAS cópias recalculadas (sem recalcular de novo), o passo 1 da 19_EQ_LINEAR contra
equilibrio_linear.tabela_linear(caso, pior_origem=...) (a_y = k·0,05 g, k = 0…8): parcelas k_pneus, k_transferência,
k_deslocamento, k_camber, k_roll_steer e k_total por linha (|Δ| <= 1e-6 °/g; N/A onde o Python dá None), motivo do N/A
(texto exato), k_inclinacao e o motivo dela; os ângulos necessários do Newton (decisão 5 do plano) contra a bissecção de
angulos_necessarios nos 11 pontos (0…0,4 g, 0,005 g e 0,01 g; |Δα| <= 1e-8·max(|α|, 1e-3): outra raiz é falha, Review
Focus 2); a igualdade da 1ª linha k_pneus + k_deslocamento + k_camber = k_inclinacao, com o resíduo da planilha igual ao
do Python (|Δ| <= 1e-6 °/g; a identidade com rel. <= 1e-6 vai ao log como ACHADO: o próprio Python não chega lá, ver
comparar_k); na 16, o
K da 1ª linha (res_16_k_total e as parcelas), V_char e V_crit (velocidade_caracteristica/critica) e a pior origem do K
(pior_origem_do_caso_v41(..., linear=True), que inclui o roll steer da 07_KINEMATICS). Caso com entradas inválidas: todo K
é N/A com o motivo da situação das entradas. Roll steer (07_KINEMATICS B82/B83): RS1 = C2 com ε_f 0,05 e ε_r −0,03 °/°
(TESTE; k_roll_steer ≠ 0 bate com o Python); RS2 = C2 com ε_f vazio (caso_da_planilha_v41 levanta o motivo; a planilha dá
k_roll_steer e K N/A com "… a confirmar" e o resto das parcelas igual ao Python). O passo 2 usa o caso com passo1=False.
O log traz o tempo de recálculo de cada cópia e o número de células com fórmula de cada aba (insumo do portão da Task 4).
"Passou" = a planilha reproduz o oráculo; não é validação física.
"""
import dataclasses, math, sys, tempfile, time
_TMP = tempfile.TemporaryDirectory(prefix="dds_teste_", ignore_cleanup_errors=True); tempfile.tempdir = _TMP.name
from pathlib import Path
import openpyxl
import aba_equilibrio as ae
import caso_planilha as cp
import equilibrio as eq
import equilibrio_linear as el
import recalculo_planilha as rc
from suspension_model import G

MOTOR, DETALHE = rc.escolher()
if MOTOR is None:
    print(f"NÃO EXECUTÁVEL: nenhum motor de recálculo ({DETALHE})"); sys.exit(3)
print(f"MOTOR: {MOTOR} ({DETALHE})")
_RAIZ = Path(__file__).resolve().parent.parent
TIR = _RAIZ / "data" / "Hoosier_R25B_13x7_TTC_provisional.tir"
TOL_REL = 1e-9
results = []


def check(name, cond, detail=""):
    results.append((name, "PASS" if cond else "FAIL")); print(f"[{'PASS' if cond else 'FAIL'}] {name}  {detail}")


# --- Casos (valores de TESTE, só nas cópias) ---
ENT = ae.NOME_ENT
C2 = {"ent_05_rear_track_t_r": 1.277, "ent_05_fracao_de_peso_no_eixo_dianteiro_f": 0.5,
      "ent_05_massa_nao_suspensa_traseira": 15, "ent_06_spring_rate_traseiro": 49909.24,
      "ent_06_k_arb_traseiro": 5012.32, "ent_06_roll_center_traseiro_rh_estatico": 0.0675,
      "ent_07_camber_gain_traseiro_dcamber_dz": -0.6, "ent_10_camber_gama_projeto_topo_para_fora_d": -1.5,
      "ent_10_camber_gama_projeto_topo_para_fora_e": -1.5, "ent_08_aero_balance_fracao_dianteira": 0.5,
      "ent_10_lmuy": 0.45}   # t_R = t_F, a = b, massas, molas, MR, barras, RC, camber e aero iguais nos eixos
C3 = {**C2, "ent_06_k_arb_dianteiro": 5012.32 * 3}
C3b = {**C2, "ent_06_k_arb_traseiro": 5012.32 * 3}
C4 = {**C3, ENT["R"]: 8.0, ENT["h2"]: 1, ENT["h3"]: 1, ENT["cg"]: 0.010, ENT["h7"]: 1, ENT["sentido"]: "Direita",
      ENT["h8"]: 1, ENT["rho"]: 1.2, ENT["cza"]: 2.0, "ent_08_downforce_total": 0, ENT["h9"]: 1, ENT["cda"]: 1.0,
      ENT["crr"]: 0.015, ENT["toeFL"]: 0.25, ENT["toeFR"]: 0.10,
      # toe e camber estático com FL ≠ FR (TESTE): o "Direita" tem de trocar L↔R, não só inverter o CG lateral
      "ent_10_camber_gama_projeto_topo_para_fora_b": -1.5, "ent_10_camber_gama_projeto_topo_para_fora_c": -1.0}
C4b = {**C4, "ent_07_angulo_da_roda_dianteira_interna_deltai": 22}
C5 = {**C2, "ent_05_altura_do_cg_total": 0.7}
RF1 = {**C4, ENT["cza"]: None}
# Roll steer (SUS-415, 07_KINEMATICS B82/B83, °/°; valores de TESTE, só nas cópias): RS1 com ε_f ≠ ε_r (k_roll_steer ≠ 0,
# pega sinal trocado em ε_f − ε_r; sobre o C2, cujo K sem roll steer é 0); RS2 com ε_f vazio (K N/A "… a confirmar").
EPS_F, EPS_R = "ent_07_roll_steer_dianteiro_epsilon_f", "ent_07_roll_steer_traseiro_epsilon_r"
RS1 = {**C2, EPS_F: 0.05, EPS_R: -0.03}
RS2 = {**C2, EPS_F: None}
# nome, alterações, veredito esperado (None = só a comparação com o Python), motivo esperado (prefixo) ou None
CASOS = [("C1", {}, None, None), ("C2", C2, eq.NEUTRO, None), ("C3", C3, eq.PUSH, None), ("C3b", C3b, eq.LOOSE, None),
         ("C4", C4, None, None), ("C4b", C4b, None, None), ("C5", C5, eq.NA, "Fz_"), ("RF1", RF1, eq.NA, "CzA a confirmar"),
         ("RS1", RS1, None, None), ("RS2", RS2, None, None)]


def num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


DIF_MAX = [0.0, ""]   # maior |a − b|/escala visto (para o log: diferença entre motores é relatada, não ajustada)


def perto(a, b, escala=None, onde="") -> bool:
    """|a − b| <= TOL_REL·escala (escala padrão: max(|a|, |b|)); a da planilha, b do Python."""
    if not (num(a) and num(b)):
        return False
    e = max(abs(a), abs(b)) if escala is None else escala
    if e > 0 and abs(a - b) / e > DIF_MAX[0]:
        DIF_MAX[:] = [abs(a - b) / e, onde]
    return a == b or abs(a - b) <= TOL_REL * e


def valor(wb, nome):
    aba, cel = ae._destino(wb, nome)
    return wb[aba][cel].value


def recalcular(caminho: Path) -> tuple[dict, float]:
    t0 = time.perf_counter()
    if MOTOR == "excel":
        rc._excel_recalc(str(caminho.resolve()), 900)
    else:
        rc._libreoffice(str(caminho.resolve()), 900)
    t_motor = time.perf_counter() - t0
    return rc._resumo(str(caminho.resolve())), t_motor


def comparar_linhas(ws, linhas, n_alfa, pedida, caso, falhas, rotulo):
    cl = ae._colunas(n_alfa, pedida)
    for k, ln in enumerate(linhas):
        r = ae.PRIMEIRA_LINHA + k
        gam = eq.camber(caso, ln.ay_ms2)
        for c in eq.CANTOS:
            fz, gm = ws[f"{cl['Fz_' + c]}{r}"].value, ws[f"{cl['gam_' + c]}{r}"].value
            if not perto(fz, ln.fz[c], onde=f"{rotulo} Fz"):
                falhas.append(f"{rotulo} k={k} Fz_{c}: planilha {fz!r} Python {ln.fz[c]!r}")
            if not perto(gm, gam[c], max(abs(gam[c]), 1e-3), f"{rotulo} γ"):
                falhas.append(f"{rotulo} k={k} γ_{c}: planilha {gm!r} Python {gam[c]!r}")
        for e, cap, req in (("F", ln.cap_f, ln.req_f), ("R", ln.cap_r, ln.req_r)):
            v, mot = ws[f"{cl['cap_' + e]}{r}"].value, ws[f"{cl['mot_' + e]}{r}"].value
            if (cap.motivo or "-") != mot:
                falhas.append(f"{rotulo} k={k} motivo {e}: planilha {mot!r} Python {cap.motivo!r}")
            if cap.fy_max_N is None:
                if v != "N/A":
                    falhas.append(f"{rotulo} k={k} cap {e}: planilha {v!r}, Python None")
                continue
            if not perto(v, cap.fy_max_N, onde=f"{rotulo} cap"):
                falhas.append(f"{rotulo} k={k} cap {e}: planilha {v!r} Python {cap.fy_max_N!r}")
            if pedida:
                f_sh = ws[f"{cl['folga_' + e]}{r}"].value
                if cap.motivo is None:
                    if not perto(f_sh, cap.fy_max_N - req, max(abs(cap.fy_max_N), abs(req)), f"{rotulo} folga"):
                        falhas.append(f"{rotulo} k={k} folga {e}: planilha {f_sh!r} Python {cap.fy_max_N - req!r}")
                elif f_sh != "N/A":
                    falhas.append(f"{rotulo} k={k} folga {e}: planilha {f_sh!r}, Python sem folga ({cap.motivo})")


def sat_igual(sh, py) -> bool:
    return sh == "N/A" if py is None else perto(sh, py, onde="resumo")


def comparar_caso(nome, copia, esperado, prefixo, falhas):
    """Compara a cópia recalculada com o Python; devolve o resumo (texto) do caso."""
    wb = openpyxl.load_workbook(copia, data_only=True)
    r16 = {k: valor(wb, n) for k, n in ae.NOME16.items()}
    try:
        caso = cp.caso_da_planilha_v41(copia, TIR, passo1=False)
    except ValueError as e:   # entradas inválidas: a planilha tem de mostrar o mesmo motivo, como N/A
        wb.close()
        ok = r16["status"] == str(e) and r16["ver"] == eq.NA and r16["mot"] == str(e)
        if prefixo is not None:
            ok = ok and str(e).startswith(prefixo)
        if not ok:
            falhas.append(f"{nome}: Python ValueError({e!s}); planilha situação={r16['status']!r} "
                          f"veredito={r16['ver']!r} motivo={r16['mot']!r}")
        return f"{nome}: N/A pelas entradas ({e})"
    res = eq.analise_par(caso)
    fina, sf2, sr2 = eq._varredura(caso, ae.AY_MAX_G, ae.PASSO_G / 2.0, 2 * ae.N_ALFA - 1)
    comparar_linhas(wb[ae.ABA17], res.linhas, ae.N_ALFA, True, caso, falhas, f"{nome} 17")
    comparar_linhas(wb[ae.ABA18], fina, 2 * ae.N_ALFA - 1, False, caso, falhas, f"{nome} 18")
    pares = [("S_F", res.ay_sat_f), ("S_R", res.ay_sat_r), ("S2_F", sf2.ay_sat), ("S2_R", sr2.ay_sat),
             ("resol", res.resolucao_ms2), ("aylim", res.ay_lim), ("cn", res.c_n)]
    for k, py in pares:
        if not sat_igual(r16[k], py):
            falhas.append(f"{nome} {k}: planilha {r16[k]!r} Python {py!r}")
    if r16["ver"] != res.veredito:
        falhas.append(f"{nome} veredito: planilha {r16['ver']!r} Python {res.veredito!r}")
    if r16["mot"] != (res.motivo or "-"):
        falhas.append(f"{nome} motivo: planilha {r16['mot']!r} Python {res.motivo!r}")
    lims = [r16[f"lim_{c}_{lim}"] for c in ae.CANTOS for lim in ae.LIMITES]
    lims = tuple(x for x in lims if x != "-")
    if lims != res.limitacoes or r16["lims"] != ("; ".join(res.limitacoes) or "nenhuma"):
        falhas.append(f"{nome} limitações: planilha {lims!r} / {r16['lims']!r} Python {res.limitacoes!r}")
    R = valor(wb, ENT["R"])   # o tempo de volta usa o R digitado, com ou sem chave ligada
    if num(R) and R > 0 and res.ay_lim is not None:
        t_py = 2 * math.pi * math.sqrt(R / res.ay_lim)
        if not perto(r16["tvolta"], t_py):
            falhas.append(f"{nome} tempo de volta: planilha {r16['tvolta']!r} Python {t_py!r}")
    elif not (isinstance(r16["tvolta"], str) and r16["tvolta"].startswith("N/A")):
        falhas.append(f"{nome} tempo de volta: planilha {r16['tvolta']!r}, esperado N/A")
    wb.close()
    pior = cp.pior_origem_do_caso_v41(copia)
    if r16["pior"] != pior:
        falhas.append(f"{nome} pior origem: planilha {r16['pior']!r} Python {pior!r}")
    if esperado is not None and res.veredito != esperado:
        falhas.append(f"{nome}: veredito do Python {res.veredito} ≠ esperado no plano {esperado} (achado, não ajuste)")
    if prefixo is not None and not (res.motivo or "").startswith(prefixo):
        falhas.append(f"{nome}: motivo do Python {res.motivo!r} não começa por {prefixo!r}")
    g = lambda x: f"{x / G:.4f} g" if x is not None else "None"
    return (f"{nome}: {res.veredito} (motivo {res.motivo!r}); S_f {g(res.ay_sat_f)}, S_r {g(res.ay_sat_r)}, "
            f"resolução {g(res.resolucao_ms2)}, C_N {res.c_n}, {len(res.limitacoes)} limitação(ões)")


LKY_RF3 = 0.9   # escala só de teste: põe o pico traseiro em 2 g entre dois α da grade pedida (a refinada fica acima)


def lmuy_rf3(caso_c3) -> float:
    """LMUY (valor da 10_TIRE) em que, com LKY × LKY_RF3, o eixo traseiro do C3 satura na grade pedida e não na
    refinada. Com mais α, a refinada tem capacidade >= em cada linha; o fator de LMUY abaixo do qual cada grade satura
    até 2 g sai por bisseção, e o caso usa o meio da janela entre os dois limiares (folga longe do arredondamento dos
    motores). Valor de TESTE."""
    def satura(f, passo, n):
        cs = dataclasses.replace(caso_c3, ds=eq.com_fatores(caso_c3.ds, {"LMUY": f, "LKY": LKY_RF3}))
        return eq._varredura(cs, ae.AY_MAX_G, passo, n)[2].ay_sat is not None

    def limiar(passo, n):
        lo, hi = 1.3, 2.0
        if not satura(lo, passo, n) or satura(hi, passo, n):
            raise RuntimeError("RF3: o intervalo do fator de LMUY não cerca a saturação traseira em 2 g")
        for _ in range(40):
            mid = 0.5 * (lo + hi)
            lo, hi = (mid, hi) if satura(mid, passo, n) else (lo, mid)
        return lo
    f_ped, f_fina = limiar(ae.PASSO_G, ae.N_ALFA), limiar(ae.PASSO_G / 2.0, 2 * ae.N_ALFA - 1)
    if not f_fina < f_ped:
        raise RuntimeError(f"RF3: sem janela entre as grades (fator pedida {f_ped}, refinada {f_fina})")
    return caso_c3.ds.p["LMUY"] * 0.5 * (f_ped + f_fina)


# --- TC-WB41-005: passo 1 (19_EQ_LINEAR e K na 16) × equilibrio_linear.tabela_linear ---
TOL_K = 1e-6        # °/g (parcelas, k_total, k_inclinacao)
TOL_IGUAL = 1e-6    # relativo: k_pneus + k_deslocamento + k_camber = k_inclinacao na 1ª linha
TOL_ANG = 1e-8      # relativo a max(|α|, 1e-3): mesma raiz do Newton e da bissecção
DIF5 = {"k": [0.0, ""], "ang": [0.0, ""], "igual_planilha": [0.0, ""], "igual_python": [0.0, ""]}
ACHADOS5 = []       # identidade da 1ª linha acima de TOL_IGUAL (vai ao log; não decide o check: ver comparar_k)


def _maior(chave, x, onde):
    if x > DIF5[chave][0]:
        DIF5[chave][:] = [x, onde]


def k_igual(sh, py, onde) -> bool:
    """sh (planilha, rad/(m/s²) ou "N/A") e py (Python, rad/(m/s²) ou None) iguais a TOL_K em °/g."""
    if py is None:
        return sh == "N/A"
    if not num(sh):
        return False
    d = abs(el.para_graus_por_g(sh) - el.para_graus_por_g(py))
    _maior("k", d, onde)
    return d <= TOL_K


def comparar_k(nome, copia, falhas):
    """Compara o passo 1 da cópia recalculada com equilibrio_linear; devolve o resumo (texto) do caso."""
    wb = openpyxl.load_workbook(copia, data_only=True)
    lin = [{q: valor(wb, ae.NOME19[(q, k)]) for q in ae.QTD19} for k in range(ae.N_LIN19)]
    r16 = {k: valor(wb, ae.NOME16[k]) for k in ae.CHAVES16_K}
    pior16 = wb[ae.ABA16][f"E{ae.LINHAS16['kTot']}"].value
    pior19 = [wb[ae.ABA19][ae.celula_pior19(k)].value for k in range(ae.N_LIN19)]
    ws19 = wb[ae.ABA19]
    angs = [(ws19[ae.celula_aux19("ang_F", p)].value, ws19[ae.celula_aux19("ang_R", p)].value)
            for p in range(len(ae.PONTOS19))]
    wb.close()
    m_rs = None   # motivo do roll steer sem dado (só o passo 1 o usa): k_roll_steer e k_total N/A, o resto compara
    e = None      # motivo das entradas inválidas (texto)
    try:
        caso = cp.caso_da_planilha_v41(copia, TIR)
    except ValueError as erro:
        try:
            caso = cp.caso_da_planilha_v41(copia, TIR, passo1=False)
            m_rs = str(erro)
        except ValueError as erro2:
            caso, e = None, str(erro2)
    if caso is None:   # entradas inválidas: todo K N/A com o motivo da situação das entradas
        for k, ln in enumerate(lin):
            if not (ln["kTot"] == "N/A" and ln["mot"] == str(e) and ln["kI"] == "N/A" and ln["motI"] == str(e)):
                falhas.append(f"{nome} linha {k}: entradas inválidas ({e}); planilha K={ln['kTot']!r} motivo={ln['mot']!r} "
                              f"inclinação={ln['kI']!r} motivo={ln['motI']!r}")
        if not (r16["kTot"] == "N/A" and r16["kmot"] == str(e)):
            falhas.append(f"{nome} 16: K={r16['kTot']!r} motivo={r16['kmot']!r}, esperado N/A com {e}")
        return f"{nome}: K N/A pelas entradas ({e})"
    pior = cp.pior_origem_do_caso_v41(copia, linear=True)
    tab = el.tabela_linear(caso, pior_origem=pior)
    if len(tab) != ae.N_LIN19:
        falhas.append(f"{nome}: tabela_linear com {len(tab)} linhas, planilha com {ae.N_LIN19}")
    for k, (ln, row) in enumerate(zip(lin, tab)):
        p = row["parcelas"]
        if m_rs is not None:   # o que o Python daria sem o dado: as duas parcelas None e o motivo no fim da lista
            p = dataclasses.replace(p, k_roll_steer=None, k_total=None,
                                    motivo="; ".join(m for m in (p.motivo, m_rs) if m))
        for q, py in (("kP", p.k_pneus), ("kT", p.k_transferencia), ("kD", p.k_deslocamento), ("kC", p.k_camber),
                      ("kRS", p.k_roll_steer), ("kTot", p.k_total), ("kI", row["k_inclinacao"])):
            if not k_igual(ln[q], py, f"{nome} linha {k} {q}"):
                falhas.append(f"{nome} linha {k} {q}: planilha {ln[q]!r} Python {py!r}")
        if p.k_total is None:
            if ln["kTotg"] != "N/A":
                falhas.append(f"{nome} linha {k} K em °/g: planilha {ln['kTotg']!r}, Python None")
        elif not (num(ln["kTotg"]) and abs(ln["kTotg"] - el.para_graus_por_g(p.k_total)) <= TOL_K):
            falhas.append(f"{nome} linha {k} K em °/g: planilha {ln['kTotg']!r} Python {el.para_graus_por_g(p.k_total)!r}")
        if ln["mot"] != (p.motivo or "-"):
            falhas.append(f"{nome} linha {k} motivo: planilha {ln['mot']!r} Python {p.motivo!r}")
        if ln["motI"] != (row["motivo_inclinacao"] or "-"):
            falhas.append(f"{nome} linha {k} motivo da inclinação: planilha {ln['motI']!r} Python {row['motivo_inclinacao']!r}")
        if not perto(ln["ay_g"], row["ay_g"], onde=f"{nome} a_y"):
            falhas.append(f"{nome} linha {k} a_y: planilha {ln['ay_g']!r} Python {row['ay_g']!r}")
        if pior19[k] != pior:
            falhas.append(f"{nome} linha {k} pior origem do K: planilha {pior19[k]!r} Python {pior!r}")
    # Newton (planilha) × bissecção (Python) nos 11 pontos: a mesma raiz, ou N/A onde o Python dá None
    for p_i, ay in enumerate(ae.PONTOS19_AY):
        py = el.angulos_necessarios(caso, ay)
        sh = angs[p_i]
        if py is None:
            if all(num(x) for x in sh):
                falhas.append(f"{nome} ponto {p_i}: Python sem ângulos (None), planilha {sh!r}")
            continue
        for e, a_sh, a_py in (("F", sh[0], py[0]), ("R", sh[1], py[1])):
            if not num(a_sh):
                falhas.append(f"{nome} ponto {p_i} α_{e}: planilha {a_sh!r}, Python {a_py!r}")
                continue
            d = abs(a_sh - a_py) / max(abs(a_py), 1e-3)
            _maior("ang", d, f"{nome} ponto {p_i} α_{e}")
            if d > TOL_ANG:
                falhas.append(f"{nome} ponto {p_i} α_{e}: planilha {a_sh!r} Python {a_py!r} (outra raiz?)")
    # Igualdade da 1ª linha (spec §6): na planilha e, para registro, no Python
    l0, p0 = lin[0], tab[0]["parcelas"]
    if m_rs is not None:
        p0 = dataclasses.replace(p0, k_roll_steer=None, k_total=None)
        if not (l0["kRS"] == "N/A" and l0["mot"] == m_rs and m_rs.endswith("a confirmar")):
            falhas.append(f"{nome}: roll steer sem dado; planilha k_roll_steer={l0['kRS']!r} motivo={l0['mot']!r}, "
                          f"esperado N/A com {m_rs!r}")
    if nome == "RS1" and not (num(l0["kRS"]) and l0["kRS"] != 0.0):
        falhas.append(f"RS1: k_roll_steer da planilha {l0['kRS']!r} deveria ser ≠ 0 (ε_f ≠ ε_r)")
    rel_sh = rel_py = None
    # O check é que a planilha reproduz o resíduo do Python: |resíduo_planilha − resíduo_Python| <= TOL_K em °/g, com
    # resíduo = k_inclinacao − (k_pneus + k_deslocamento + k_camber). A identidade em si, rel. <= TOL_IGUAL, só vai ao log
    # (achado da Task 3): o próprio Python não chega a 1e-6 (resíduo de truncamento do Richardson; com a h8, a downforce
    # 0,5·ρ·CzA·|a_y|·R tem a derivada lateral em a_y = 0 que a diferença central do deslocamento cancela e a inclinação,
    # de um lado só, vê). Com K(0) abaixo de TOL_K (carro simétrico: C2, C5) o relativo não tem sentido e o log mostra o
    # absoluto.
    def residuo(kp, kd, kc, ki):
        soma = kp + kd + kc
        if abs(el.para_graus_por_g(soma)) < TOL_K:
            return ki - soma, abs(el.para_graus_por_g(ki - soma)), "abs °/g"
        return ki - soma, abs(ki - soma) / abs(soma), "rel"
    crit_sh = crit_py = ""
    res_sh = res_py = None
    if all(num(l0[q]) for q in ("kP", "kD", "kC", "kI")):
        res_sh, rel_sh, crit_sh = residuo(l0["kP"], l0["kD"], l0["kC"], l0["kI"])
        if crit_sh == "rel":
            _maior("igual_planilha", rel_sh, nome)
    if None not in (p0.k_pneus, p0.k_deslocamento, p0.k_camber, tab[0]["k_inclinacao"]):
        res_py, rel_py, crit_py = residuo(p0.k_pneus, p0.k_deslocamento, p0.k_camber, tab[0]["k_inclinacao"])
        if crit_py == "rel":
            _maior("igual_python", rel_py, nome)
    if (res_sh is None) != (res_py is None):
        falhas.append(f"{nome}: igualdade da 1ª linha calculável só de um lado (planilha {res_sh!r}, Python {res_py!r})")
    elif res_sh is not None:
        d = abs(el.para_graus_por_g(res_sh) - el.para_graus_por_g(res_py))
        _maior("k", d, f"{nome} resíduo da igualdade")
        if d > TOL_K:
            falhas.append(f"{nome}: resíduo da igualdade da 1ª linha na planilha difere do Python em {d:.2e} °/g")
        if crit_sh == "rel" and rel_sh > TOL_IGUAL:
            ACHADOS5.append(f"{nome}: identidade da 1ª linha com rel. {rel_sh:.2e} > {TOL_IGUAL:g} (Python {rel_py:.2e})")
    # 16: o K da 1ª linha, V_char e V_crit
    for q in ("kP", "kT", "kD", "kC", "kRS", "kTot", "kI", "kTotg"):
        if r16[q] != l0[q]:
            falhas.append(f"{nome} 16 {q}: {r16[q]!r} ≠ 19 linha 0 {l0[q]!r}")
    if r16["kmot"] != l0["mot"]:
        falhas.append(f"{nome} 16 motivo do K: {r16['kmot']!r} ≠ 19 linha 0 {l0['mot']!r}")
    L = caso.cfg.wheelbase
    for q, py in (("vchar", None if p0.k_total is None else el.velocidade_caracteristica(L, p0.k_total)),
                  ("vcrit", None if p0.k_total is None else el.velocidade_critica(L, p0.k_total))):
        if not sat_igual(r16[q], py):
            falhas.append(f"{nome} 16 {q}: planilha {r16[q]!r} Python {py!r}")
    if pior16 != pior:
        falhas.append(f"{nome} 16 pior origem do K: planilha {pior16!r} Python {pior!r}")
    kt = p0.k_total
    krs = '' if p0.k_roll_steer in (None, 0.0) else f", k_roll_steer {el.para_graus_por_g(p0.k_roll_steer):+.4f} °/g"
    return (f"{nome}: K(0) = {'N/A' if kt is None else f'{el.para_graus_por_g(kt):+.4f} °/g'}{krs} "
            f"(motivo {l0['mot']!r}); "
            f"igualdade da 1ª linha: planilha {rel_sh if rel_sh is None else f'{rel_sh:.2e} ({crit_sh})'}, "
            f"Python {rel_py if rel_py is None else f'{rel_py:.2e} ({crit_py})'}")


# --- TC-WB41-004 e TC-WB41-005 (as mesmas cópias recalculadas) ---
falhas, resumos, tempos = [], [], []
falhas5, resumos5 = [], []
base = ae.extrair_v40(Path(tempfile.mkdtemp()))
info_b, t_b = recalcular(base)
print(f"  V4.0 da tag (referência de tempo): {t_b:.1f} s no motor, {info_b['total_formulas']} fórmulas, "
      f"{info_b['total_errors']} erros")
caso_c3 = None
for nome, alt, esperado, prefixo in CASOS + [("RF3", None, None, None)]:
    if nome == "RF3":
        alt = {**C3, "ent_10_lmuy": lmuy_rf3(caso_c3), "ent_10_lky": caso_c3.ds.p["LKY"] * LKY_RF3}
    copia = ae.copia_v41(Path(tempfile.mkdtemp()), alt, f"{nome}.xlsx")
    info, t = recalcular(copia)
    tempos.append(t)
    print(f"  {nome}: recalculado em {t:.1f} s no motor; {info['total_formulas']} fórmulas, {info['total_errors']} erros")
    if info["total_errors"]:
        falhas.append(f"{nome}: {info['total_errors']} erro(s) de fórmula {info['error_summary']}")
    if nome == "C1":
        wbf = openpyxl.load_workbook(copia)
        for aba in (ae.ABA16, ae.ABA17, ae.ABA18, ae.ABA19):
            n = sum(1 for lin in wbf[aba].iter_rows() for c in lin if isinstance(c.value, str) and c.value.startswith("="))
            print(f"  CÉLULAS com fórmula em {aba}: {n}")
        wbf.close()
    if nome == "C3":
        caso_c3 = cp.caso_da_planilha_v41(copia, TIR, passo1=False)
    resumos.append(comparar_caso(nome, copia, esperado, prefixo, falhas))
    resumos5.append(comparar_k(nome, copia, falhas5))
    if nome == "RF3":
        caso = cp.caso_da_planilha_v41(copia, TIR, passo1=False)
        r = eq.analise_par(caso)
        _, _, sr2 = eq._varredura(caso, ae.AY_MAX_G, ae.PASSO_G / 2.0, 2 * ae.N_ALFA - 1)
        if not (r.ay_sat_r is not None and sr2.ay_sat is None and r.resolucao_ms2 == ae.PASSO_G * G):
            falhas.append(f"RF3: o caso não ficou com a traseira saturando só na grade pedida (S_r {r.ay_sat_r}, "
                          f"S_r refinada {sr2.ay_sat}, resolução {r.resolucao_ms2})")
for s in resumos:
    print("  " + s)
print(f"  MAIOR diferença relativa planilha × Python ({MOTOR}): {DIF_MAX[0]:.2e} ({DIF_MAX[1]})")
print(f"  TEMPO de recálculo ({MOTOR}): por cópia da V4.1 mín {min(tempos):.1f} s, máx {max(tempos):.1f} s, "
      f"média {sum(tempos) / len(tempos):.1f} s; V4.0 da tag {t_b:.1f} s")
check("TC-WB41-004: análise de par em fórmulas (16-18) == equilibrio.analise_par nos casos C1-C5, C3b, C4b e nas "
      "variantes RF1/RF3 (Fz, γ, capacidade, motivo e folga por linha, rel. <= 1e-9; S_f, S_r, resolução, veredito, "
      "motivo, a_y limite, C_N, limitações, tempo de volta, pior origem); 0 erros de fórmula",
      not falhas, f"{len(falhas)} divergência(s): {falhas[:8]}")
for s5 in resumos5:
    print("  " + s5)
print(f"  MAIOR diferença do K planilha × Python ({MOTOR}): {DIF5['k'][0]:.2e} °/g ({DIF5['k'][1]}); ângulos (rel.): "
      f"{DIF5['ang'][0]:.2e} ({DIF5['ang'][1]}); igualdade da 1ª linha: planilha {DIF5['igual_planilha'][0]:.2e} "
      f"({DIF5['igual_planilha'][1]}), Python {DIF5['igual_python'][0]:.2e} ({DIF5['igual_python'][1]})")
for a5 in ACHADOS5:
    print(f"  ACHADO (não decide o check; o Python dá o mesmo): {a5}")
check("TC-WB41-005: passo 1 em fórmulas (19_EQ_LINEAR e o K na 16) == equilibrio_linear.tabela_linear nos casos C1-C5, "
      "C3b, C4b e nas variantes RF1/RF3/RS1/RS2 (parcelas, k_total e k_inclinacao por linha, |Δ| <= 1e-6 °/g; motivos; ângulos "
      "do Newton = bissecção, rel. <= 1e-8; resíduo da igualdade da 1ª linha = o do Python, |Δ| <= 1e-6 °/g; K, V_char, "
      "V_crit e pior origem na 16)",
      not falhas5, f"{len(falhas5)} divergência(s): {falhas5[:8]}")

print("\n--- Resumo ---"); p = sum(1 for _, s in results if s == "PASS"); f = len(results) - p
print(f"{p} PASS, {f} FAIL, de {len(results)} testes")
sys.exit(1 if f else 0)
