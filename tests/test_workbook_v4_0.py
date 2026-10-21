"""
Teste da V4.0 no motor de recálculo (incremento 3, fundação; EDR-044): equivalência por valor e nomes válidos.
Recalcula CÓPIAS pelo Excel (COM) ou pelo LibreOffice (recalculo_planilha.py); a V4.0 de referência em excel/ e a
V3.0 nunca são salvas nem recalculadas. A V3.0 é o arquivo da tag v3.0, extraído para uma pasta temporária.
Uso: uv run python -X utf8 tests/test_workbook_v4_0.py [<V4_0.xlsx>]   (padrão: excel/FSAE_EV_Suspension_Fase1_V4_0.xlsx)
Namespace TC-WB40-xxx (IDs autorizados pelo usuário: 001 a 006; este script tem 005 e 006).
TC-WB40-005: as cópias recalculadas da V4.0 e da V3.0 da tag têm 0 erros de fórmula e 387 fórmulas cada; para cada
nome `res` da tabela, o valor da V4.0 lido pelo nome definido (nome -> aba/célula da própria V4.0) é igual ao valor
da V3.0 no mesmo endereço (erro relativo <= 1e-12 para números; igualdade para texto, bool e vazio).
TC-WB40-006: numa cópia da V4.0, a aba temporária ZZ_NOMES recebe `=<nome>` para cada nome da tabela (um por linha);
depois do recálculo, nenhuma célula dela tem erro (#NAME? ou outro) e cada valor é igual ao do destino do nome na
mesma cópia recalculada. Destino vazio: o motor devolve 0 para `=nome`, e isso conta como igual (regra do Excel
para referência a célula vazia). A aba existe só na cópia.
"Passou" = a planilha recalcula igual à V3.0 e os nomes resolvem no motor; não é validação física.
"""
import math, os, shutil, sys, tempfile, time
_TMP = tempfile.TemporaryDirectory(prefix="dds_teste_", ignore_cleanup_errors=True); tempfile.tempdir = _TMP.name
from pathlib import Path
import openpyxl
import gerar_v4_0 as gv
import nomes_v4_0 as nv
import recalculo_planilha as rc

MOTOR, DETALHE = rc.escolher()
if MOTOR is None:
    print(f"NÃO EXECUTÁVEL: nenhum motor de recálculo ({DETALHE})"); sys.exit(3)
print(f"MOTOR: {MOTOR} ({DETALHE})")
_RAIZ = Path(__file__).resolve().parent.parent
V40 = Path(sys.argv[1]) if len(sys.argv) > 1 else _RAIZ / "excel" / "FSAE_EV_Suspension_Fase1_V4_0.xlsx"
TABELA = nv.ler_tabela(_RAIZ / "data" / "nomes_V4_0.csv")
N_FORMULAS = 387
TOL_REL = 1e-12
ABA_TEMP = "ZZ_NOMES"
results = []


def check(name, cond, detail=""):
    results.append((name, "PASS" if cond else "FAIL")); print(f"[{'PASS' if cond else 'FAIL'}] {name}  {detail}")


def copia(origem: Path, nome: str) -> Path:
    """Cópia de `origem` numa pasta temporária nova; só a cópia é recalculada."""
    destino = Path(tempfile.mkdtemp()) / nome
    shutil.copy(origem, destino)
    return destino


def recalcular(caminho: Path) -> dict:
    t0 = time.perf_counter()
    info = rc.recalcular(str(caminho), MOTOR, 300)
    print(f"  recalculado {caminho.name} em {time.perf_counter() - t0:.1f} s: "
          f"{info['total_formulas']} fórmulas, {info['total_errors']} erros")
    return info


def destinos(wb) -> dict[str, tuple[str, str]]:
    """nome -> (aba, célula) dos nomes da tabela definidos no escopo do workbook (um destino de uma célula por nome).
    Nomes fora da tabela (o `Origens` da V3.0, um intervalo) ficam de fora."""
    da_tabela = {n.nome for n in TABELA}
    out = {}
    for nome, dn in wb.defined_names.items():
        if nome not in da_tabela:
            continue
        dest = list(dn.destinations)
        if len(dest) != 1 or ":" in dest[0][1]:
            raise ValueError(f"nome {nome}: destino inesperado ({dn.attr_text}); esperado uma célula")
        out[nome] = (dest[0][0], dest[0][1].replace("$", ""))
    return out


def iguais(a, b) -> bool:
    """Números (fora bool): erro relativo <= TOL_REL; o resto: mesmo tipo e mesmo valor."""
    num = (int, float)
    if isinstance(a, num) and isinstance(b, num) and not isinstance(a, bool) and not isinstance(b, bool):
        return a == b or abs(a - b) <= TOL_REL * max(abs(a), abs(b))
    return type(a) is type(b) and a == b


def tem_erro(v) -> bool:
    return isinstance(v, str) and any(e in v for e in rc.ERROS)


# --- TC-WB40-005: 0 erros, 387 fórmulas e equivalência por valor dos `res` ---
v30 = gv.extrair_v30(Path(tempfile.mkdtemp()))
c40, c30 = copia(V40, "V4_0.xlsx"), copia(v30, "V3_0_tag.xlsx")
info40, info30 = recalcular(c40), recalcular(c30)
wb40, wb30 = openpyxl.load_workbook(c40, data_only=True), openpyxl.load_workbook(c30, data_only=True)
dest40 = destinos(wb40)
res = [n for n in TABELA if n.papel == "res"]
sem_nome = [n.nome for n in res if n.nome not in dest40]
fora_do_lugar = [n.nome for n in res if n.nome in dest40 and dest40[n.nome] != (n.aba, n.celula)]
difs = []
for n in res:
    if n.nome not in dest40:
        continue
    aba, cel = dest40[n.nome]
    a, b = wb40[aba][cel].value, wb30[n.aba][n.celula].value
    if not iguais(a, b):
        difs.append(f"{n.nome} ({n.aba}!{n.celula}): V4.0={a!r} V3.0={b!r}")
wb40.close(); wb30.close()
check("TC-WB40-005: cópias recalculadas da V4.0 e da V3.0 (tag) com 0 erros e 387 fórmulas; cada `res` da V4.0 "
      "lido pelo nome == V3.0 no endereço (rel. <= 1e-12; igualdade para texto/bool/vazio)",
      info40["total_errors"] == 0 and info30["total_errors"] == 0
      and info40["total_formulas"] == N_FORMULAS and info30["total_formulas"] == N_FORMULAS
      and len(res) > 0 and not sem_nome and not fora_do_lugar and not difs,
      f"{len(res)} res; erros V4.0={info40['total_errors']} V3.0={info30['total_errors']}; fórmulas "
      f"V4.0={info40['total_formulas']} V3.0={info30['total_formulas']}; sem nome={sem_nome[:5]}; "
      f"destino diferente da tabela={fora_do_lugar[:5]}; {len(difs)} diferença(s) {difs[:5]}")

# --- TC-WB40-006: `=<nome>` para cada nome da tabela numa aba temporária da cópia ---
c40n = copia(V40, "V4_0_nomes.xlsx")
wb = openpyxl.load_workbook(c40n)
if ABA_TEMP in wb.sheetnames:
    raise RuntimeError(f"a V4.0 já tem a aba {ABA_TEMP}")
ws = wb.create_sheet(ABA_TEMP)
for i, n in enumerate(TABELA, start=1):
    ws.cell(i, 1, n.nome)
    ws.cell(i, 2, f"={n.nome}")
wb.save(c40n); wb.close()
info_n = recalcular(c40n)
wbv = openpyxl.load_workbook(c40n, data_only=True)
dest_n = destinos(wbv)
erros_zz, difs_zz, vazios = [], [], 0
for i, n in enumerate(TABELA, start=1):
    v = wbv[ABA_TEMP].cell(i, 2).value
    if tem_erro(v):
        erros_zz.append(f"{n.nome}: {v}")
        continue
    if n.nome not in dest_n:
        difs_zz.append(f"{n.nome}: sem nome definido")
        continue
    aba, cel = dest_n[n.nome]
    alvo = wbv[aba][cel].value
    if alvo is None:
        vazios += 1
        if not (iguais(v, 0) or v is None):
            difs_zz.append(f"{n.nome} ({aba}!{cel} vazio): {v!r}")
    elif not iguais(v, alvo):
        difs_zz.append(f"{n.nome} ({aba}!{cel}): ZZ={v!r} destino={alvo!r}")
wbv.close()
wb_ref = openpyxl.load_workbook(V40, read_only=True)
viva_sem_aba = ABA_TEMP not in wb_ref.sheetnames
wb_ref.close()
check("TC-WB40-006: `=<nome>` para todos os nomes da tabela numa aba temporária da cópia: nenhum #NAME? (nem outro "
      "erro) e cada valor == destino na mesma cópia; a V4.0 de referência segue sem a aba",
      len(TABELA) > 0 and not erros_zz and not difs_zz and info_n["total_errors"] == 0
      and info_n["total_formulas"] == N_FORMULAS + len(TABELA) and viva_sem_aba,
      f"{len(TABELA)} nomes ({vazios} com destino vazio); erros={len(erros_zz)} {erros_zz[:5]}; "
      f"diferenças={len(difs_zz)} {difs_zz[:5]}; fórmulas={info_n['total_formulas']}")

print("\n--- Resumo ---"); p = sum(1 for _, s in results if s == "PASS"); f = len(results) - p
print(f"{p} PASS, {f} FAIL, de {len(results)} testes")
sys.exit(1 if f else 0)
