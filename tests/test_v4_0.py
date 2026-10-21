"""
Testes da fundação da V4.0 (incremento 3): nomes definidos e dados por chave.
TC-WB40-001 (ID autorizado pelo usuário): tabela de nomes gerada da V3.0 da tag v3.0 é válida e completa
(nomes únicos e aceitos pelo Excel, células existentes, entradas e resultados cobertos, CSV commitado em dia) e
`verificar_cobertura` recusa entrada sem nome, listando as células.
Roda como script: uv run python -X utf8 tests/test_v4_0.py
A V3.0 de partida é o arquivo da tag v3.0, extraído para uma pasta temporária (nunca a planilha viva).
"""

import os
import re
import subprocess
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.formula import ArrayFormula

import atualizar_planilha as ap
import nomes_v4_0 as nv

HERE = os.path.dirname(os.path.abspath(__file__))
RAIZ = Path(HERE).parent
CSV_NOMES = RAIZ / "data" / "nomes_V4_0.csv"
TAG_V30 = "v3.0"
CAMINHO_V30 = "excel/FSAE_EV_Suspension_Fase1_V3_0.xlsx"
results = []


def raises(fn, exc=ValueError, needle=None):
    try:
        fn()
        return False
    except exc as e:
        return needle is None or needle in str(e)


def check(name, condition, detail=""):
    status = "PASS" if condition else "FAIL"
    results.append((name, status, detail))
    print(f"[{status}] {name}  {detail}")


def extrair_v30(pasta: Path) -> Path:
    """Arquivo da tag v3.0 numa pasta temporária (binário, sem passar por texto)."""
    destino = pasta / "V3_0_tag.xlsx"
    with open(destino, "wb") as f:
        subprocess.run(["git", "show", f"{TAG_V30}:{CAMINHO_V30}"], cwd=RAIZ, stdout=f, check=True)
    return destino


def eh_formula(v) -> bool:
    return isinstance(v, ArrayFormula) or (isinstance(v, str) and v.startswith("="))


with tempfile.TemporaryDirectory() as tmp:
    v30 = extrair_v30(Path(tmp))
    tabela = nv.tabela_nomes(v30)
    wb = openpyxl.load_workbook(v30)
    try:
        formulas = {(ws.title, c.coordinate) for ws in wb.worksheets if ws.title[:2] in nv.ABAS_RES
                    for lin in ws.iter_rows() for c in lin if eh_formula(c.value)}
        abas = set(wb.sheetnames)
        dims = {ws.title: (ws.max_row, ws.max_column) for ws in wb.worksheets}
    finally:
        wb.close()

    # ---------------- TC-WB40-001 ----------------
    nomes = [n.nome for n in tabela]
    unicos = len(set(nomes)) == len(nomes)
    invalidos = [n for n in nomes if not nv.NOME_VALIDO.fullmatch(n)]
    referencias = [n for n in nomes if nv.parece_referencia(n)]

    def existe(aba, cel):
        if aba not in abas:
            return False
        m = re.fullmatch(r"([A-Z]+)(\d+)", cel)
        return bool(m) and int(m.group(2)) <= dims[aba][0] and column_index_from_string(m.group(1)) <= dims[aba][1]

    inexistentes = [(n.aba, n.celula) for n in tabela if not existe(n.aba, n.celula)]
    por_celula: dict = {}
    for n in tabela:
        por_celula.setdefault((n.aba, n.celula), []).append(n)
    duplicadas = [k for k, v in por_celula.items() if len(v) > 1]
    entradas = ap.celulas_de_entrada()
    ent_mal = sorted(k for k in entradas
                     if len(por_celula.get(k, [])) != 1 or por_celula[k][0].papel not in ("ent", "org", "fte"))
    ent_sobra = sorted(k for k, v in por_celula.items() if v[0].papel in ("ent", "org", "fte") and k not in entradas)
    res_mal = sorted(k for k in formulas if len(por_celula.get(k, [])) != 1 or por_celula[k][0].papel != "res")
    res_sobra = sorted(k for k, v in por_celula.items() if v[0].papel == "res" and k not in formulas)
    papeis_ok = all(n.papel in ("ent", "org", "fte", "res") for n in tabela)
    convencao = [n.nome for n in tabela if not n.nome.startswith(f"{n.papel}_{n.aba[:2]}_")]
    sus_mal = [n.nome for n in tabela if n.sus_id and not all(re.fullmatch(r"SUS-\d{3}(-[FR])?", s)
                                                               for s in n.sus_id.split(";"))]
    ordem_ok = tabela == sorted(tabela, key=lambda n: (n.aba, int(re.sub(r"[A-Z]", "", n.celula)),
                                                         column_index_from_string(re.sub(r"\d", "", n.celula))))
    deterministica = nv.tabela_nomes(v30) == tabela
    csv_igual = CSV_NOMES.exists() and nv.ler_tabela(CSV_NOMES) == tabela
    csv_lf = CSV_NOMES.exists() and b"\r" not in CSV_NOMES.read_bytes()
    csv_cab = CSV_NOMES.exists() and CSV_NOMES.read_bytes().split(b"\n", 1)[0] == b"nome,aba,celula,papel,sus_id,rotulo"
    # Ida e volta do CSV numa pasta temporária (UTF-8, LF, rótulos com vírgula, aspas e acento).
    with tempfile.TemporaryDirectory() as tmp2:
        p = Path(tmp2) / "nomes.csv"
        nv.escrever_tabela(tabela, p)
        ida_volta = nv.ler_tabela(p) == tabela
    # verificar_cobertura (antigo 009): passa com tudo coberto; sem um nome, ValueError que lista as células.
    cob_ok = not raises(lambda: nv.verificar_cobertura(tabela, entradas), Exception)
    faltam = [("05_VEHICLE", "B5"), ("10_TIRE", "B86")]
    parcial = [n for n in tabela if (n.aba, n.celula) not in faltam]
    cob_erro = raises(lambda: nv.verificar_cobertura(parcial, entradas), ValueError, "05_VEHICLE!B5") and \
        raises(lambda: nv.verificar_cobertura(parcial, entradas), ValueError, "10_TIRE!B86")
    ref_ok = all(nv.parece_referencia(x) for x in ("A1", "xfd1048576", "R1C1", "r", "C", "RC", "R12", "abc7")) and \
        not any(nv.parece_referencia(x) for x in ("XFE1", "A1048577", "A0", "res_05_x1", "ent_10_b", "R1C1_x"))
    regex_ok = all(nv.NOME_VALIDO.fullmatch(x) for x in ("a", "ent_05_x", "A" * 255)) and \
        not any(nv.NOME_VALIDO.fullmatch(x) for x in ("1a", "_a", "a-b", "a b", "ação", "A" * 256, ""))
    cont = {p: sum(1 for n in tabela if n.papel == p) for p in ("ent", "org", "fte", "res")}
    # Correção 1 (controlador): letras gregas transliteradas; sus_id só do rótulo e da coluna de ID, nunca da Fonte.
    grego_ok = nv.slug("α [rad]") == "alfa_rad" and nv.slug("Δ Fz (μ, γ)") == "delta_fz_mu_gama" and         nv.slug("Kyγ0 ε") == "kygama0_epsilon" and nv.slug("ÿ Ж z") == "y_z"
    sus_de = {(n.aba, n.celula): n.sus_id for n in tabela}
    sus_rotulo_ok = sus_de.get(("05_VEHICLE", "B5")) == "SUS-002" and sus_de.get(("05_VEHICLE", "F5")) == "SUS-002"
    ok_001 = all((grego_ok, sus_rotulo_ok, unicos, not invalidos, not referencias, not inexistentes, not duplicadas, not ent_mal, not ent_sobra,
                  not res_mal, not res_sobra, papeis_ok, not convencao, not sus_mal, ordem_ok, deterministica, csv_igual,
                  csv_lf, csv_cab, ida_volta, cob_ok, cob_erro, ref_ok, regex_ok))
    check("TC-WB40-001: tabela de nomes da V3.0 (tag) com nomes únicos, válidos e sem cara de referência; células "
          "existentes; cada entrada uma vez como ent/org/fte e cada fórmula das abas 05..14 como res; CSV em dia (UTF-8, "
          "LF); verificar_cobertura lista as células sem nome",
          ok_001,
          f"{len(tabela)} nomes {cont}; únicos={unicos} inválidos={invalidos[:3]} ref={referencias[:3]} "
          f"inexistentes={inexistentes[:3]} dup={duplicadas[:3]} ent={ent_mal[:3]}/{ent_sobra[:3]} "
          f"res={res_mal[:3]}/{res_sobra[:3]} papéis={papeis_ok} conv={convencao[:3]} sus={sus_mal[:3]} "
          f"ordem={ordem_ok} det={deterministica} csv={csv_igual}/{csv_lf}/{csv_cab} ida_volta={ida_volta} "
          f"cobertura={cob_ok}/{cob_erro} ref={ref_ok} regex={regex_ok} grego={grego_ok} sus_rótulo={sus_rotulo_ok} | fórmulas={len(formulas)} entradas={len(entradas)}")


# ---------------- Resumo ----------------
print("\n--- Resumo ---")
n_pass = sum(1 for _, s, _ in results if s == "PASS")
n_fail = sum(1 for _, s, _ in results if s == "FAIL")
print(f"{n_pass} PASS, {n_fail} FAIL, de {len(results)} testes")
if n_fail > 0:
    raise SystemExit(1)
