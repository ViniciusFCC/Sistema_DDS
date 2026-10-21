"""
Experimento de viabilidade: a análise de par (equilibrio.analise_par) cabe numa planilha? (V4.0, inc. 1)

EXPERIMENTO FORA DA SUÍTE: não entra em tools/rodar_testes.py e não é critério de aprovação.
Responde a duas perguntas da spec 7.1 (item 4): quanto tempo o motor leva para recalcular a grade
a_y × α em fórmulas, e se as fórmulas reproduzem o Python (capacidade, pico na borda, N/A por faixa
do pneu e saturação interpolada) no Excel e no LibreOffice.

O protótipo NÃO é planilha da equipe: é gerado pelo openpyxl numa pasta temporária, recalculado só
ali (cópias recalculadas pelo motor: EDR-044) e apagado no fim. Só build/exp_viabilidade_par.json
fica (build/ está fora do Git). A planilha viva só é COPIADA para a pasta temporária; a cópia é
recalculada para caso_da_planilha, e o SHA-256 da viva é conferido antes e depois.

Casos (entradas de teste, NÃO são o carro):
- "teste_LMUY045": veículo de tests/test_equilibrio.py (veiculo_base, números copiados) com o pneu
  escalado LMUY = 0,45. Escala só de teste, não é pneu: existe para haver saturação dentro da faixa.
- "teste_LMUY045_toe_camber": o mesmo, com toe dianteiro de 1° e o camber do TC-EQ-003, para que
  γy, o espelhamento da roda esquerda e o Δα do toe não fiquem multiplicados por zero.
- "teste_LMUY045_LKY023_borda": o primeiro com LKY = 0,23 (também escala só de teste): o pico do eixo
  traseiro cai na borda da grade de α a partir de 1,05 g e a busca de saturação tem de parar ali.
- "planilha_viva": caso_da_planilha (valores DEMONSTRATIVOS da V3.0) com o pneu sem escala. Esperado:
  N/A por faixa do pneu ("Risco conhecido" do plano); as fórmulas têm de reproduzir as linhas N/A.

Desenho da pasta: aba COEF com o .tir, o veículo, a grade de α e α no referencial do arquivo por roda;
aba PAR com uma linha por a_y (Fz e γ das quatro rodas como constantes do Python), os termos PAC2002
por roda (dfz, γy, Ky, SHy, μy, Dy, By, SVy, Ey nas três variantes de sinal), uma célula por α e por
eixo com a soma das duas rodas (κ = 0: Gyκ = 1 e SVyκ = 0, então Fy combinado = Fy0), capacidade
(MAX), índice do pico (MATCH), folga, prefixo válido e candidata de saturação; aba RESUMO com MIN.

Uso (Git Bash, depois de export UV_SYSTEM_CERTS=1):
    DDS_RECALC=excel uv run python -X utf8 tests/exp_viabilidade_par.py
    DDS_RECALC=libreoffice uv run python -X utf8 tests/exp_viabilidade_par.py
Sem motor: imprime NÃO EXECUTÁVEL e sai com 3. Saída 0 = critérios atendidos; 1 = divergência.

Regra registrada: este experimento vale para a regra anterior ao EDR-056 (commit a732c39), em que Fz ou
γ fora da faixa do pneu davam N/A e o neutro usava o passo da grade. Ele não é refeito (EDR-056,
Consequences). O protótipo da aba do incremento 4 seguirá o EDR-056: entrada do pneu limitada à faixa
com aviso (help do Adams PAC2002, p. 83), descolamento (Fz <= 0) como N/A e neutro pela resolução
numérica (Milliken, p. 320). Com o equilibrio.py atual, a comparação com o Python pode divergir nas
linhas fora da faixa.
"""

import hashlib
import json
import math
import os
import shutil
import sys
import tempfile
import time
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as COL

import caso_planilha as cp
import equilibrio as eq
import recalculo_planilha as rc
import tire_model as tm
from suspension_model import G, VehicleConfig
from tire_model import AXIS_TRANSFORMS

AY_MAX_G, PASSO_G, N_ALFA = 2.0, 0.05, 25   # grade padrão de analise_par
TOL = 1e-9                                   # mesmo critério da CA-09
RAIZ = Path(__file__).resolve().parent.parent
TIR = RAIZ / "data" / "Hoosier_R25B_13x7_TTC_provisional.tir"
VIVA = RAIZ / "excel" / "FSAE_EV_Suspension_Fase1_V3_0.xlsx"
SAIDA = RAIZ / "build" / "exp_viabilidade_par.json"
RODAS_EIXO = {"F": ("FL", "FR"), "R": ("RL", "RR")}   # (esquerda, direita)
CHAVES_TIR = ("FNOMIN", "LFZO", "LGAY", "PKY1", "PKY2", "PKY3", "LKY", "PHY1", "PHY2", "PHY3", "LHY",
              "LKYG", "PCY1", "LCY", "PDY1", "PDY2", "PDY3", "LMUY", "PEY1", "PEY2", "PEY3", "PEY4",
              "LEY", "PVY1", "PVY2", "PVY3", "PVY4", "LVY", "FZMIN", "FZMAX", "CAMMIN", "CAMMAX",
              "ALPMIN", "ALPMAX")
TERMOS = ("dfz", "gy", "Ky", "SHy", "muy", "Dy", "By", "SVy", "Eyn", "Ey0", "Eyp")


def veiculo_base() -> VehicleConfig:
    """Cópia dos números de veiculo_base() de tests/test_equilibrio.py (teste, NÃO é setup real)."""
    return VehicleConfig(mass_total=250.0, cg_height=0.300, wheelbase=1.550, track_front=1.200,
                         track_rear=1.150, cg_to_front=0.837, cg_to_rear=0.713, unsprung_mass_front=15.0,
                         unsprung_mass_rear=18.0, unsprung_cg_height_front=0.230,
                         unsprung_cg_height_rear=0.230, roll_center_front=0.030, roll_center_rear=0.050,
                         roll_stiffness_front=20826.4, roll_stiffness_rear=26450.0)


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


# ---------------------------------------------------------------------------
# Geração da pasta
# ---------------------------------------------------------------------------

def _aba_coef(wb, caso: eq.CasoEquilibrio) -> dict:
    """Aba COEF: constantes (.tir, eixos, veículo, toe) e a grade de α. Devolve nome -> referência absoluta."""
    ws = wb.create_sheet("COEF")
    ws["A1"] = "PROTÓTIPO DESCARTÁVEL (exp_viabilidade_par.py): não é planilha da equipe; entradas de teste/demonstrativas."
    ref, r = {}, 3
    p = caso.ds.p
    s_a, s_f, s_g = AXIS_TRANSFORMS[caso.ds.axis_system]
    dal = eq.deriva_toe(caso)
    consts = [(k, float(p[k])) for k in CHAVES_TIR] + [
        ("EPS", tm.EPS), ("S_A", s_a), ("S_F", s_f), ("S_G", s_g), ("G", G), ("N_ALFA", N_ALFA),
        ("m", caso.cfg.mass_total), ("a", caso.cfg.cg_to_front), ("b", caso.cfg.cg_to_rear),
        ("L", caso.cfg.wheelbase)]
    consts += [(f"DA_{c}", dal[c]) for c in eq.CANTOS]                       # Δα do toe (deriva_toe)
    consts += [(f"MIR_{c}", -1.0 if c.endswith("L") else 1.0) for c in eq.CANTOS]  # espelhamento
    for nome, v in consts:
        ws.cell(r, 1, nome)
        ws.cell(r, 2, float(v))
        ref[nome] = f"COEF!$B${r}"
        r += 1
    # Amplitude da grade de α (grade_alfa): min(|ALPMIN|, ALPMAX) − max|Δα|
    ws.cell(r, 1, "AMP")
    ws.cell(r, 2, f"=MIN(ABS({ref['ALPMIN']}),{ref['ALPMAX']})-MAX("
                  + ",".join(f"ABS({ref['DA_' + c]})" for c in eq.CANTOS) + ")")
    ref["AMP"] = f"COEF!$B${r}"
    r += 2
    # α_ref_i = −A·i/(n−1) e α no referencial do arquivo por roda: s_α·(espelho·(α_ref + Δα))
    ws.cell(r, 1, "i")
    ws.cell(r + 1, 1, "alfa_ref")
    for i in range(N_ALFA):
        cl = COL(2 + i)
        ws[f"{cl}{r}"] = i
        ws[f"{cl}{r + 1}"] = f"=-{ref['AMP']}*{cl}{r}/({ref['N_ALFA']}-1)"
        ref[("alfa", i)] = f"COEF!${cl}${r + 1}"
    for j, c in enumerate(eq.CANTOS):
        rr = r + 2 + j
        ws.cell(rr, 1, f"alfa_eq_{c}")
        for i in range(N_ALFA):
            cl = COL(2 + i)
            ws[f"{cl}{rr}"] = (f"={ref['S_A']}*({ref['MIR_' + c]}*(COEF!{cl}${r + 1}+{ref['DA_' + c]}))")
            ref[("aeq", c, i)] = f"COEF!${cl}${rr}"
    return ref


def _colunas() -> dict:
    """Mapa nome -> letra da coluna da aba PAR."""
    nomes = ["k", "ay_g", "ay_ms2"] + [f"Fz_{c}" for c in eq.CANTOS] + [f"gam_{c}" for c in eq.CANTOS]
    nomes += [(t, c) for c in eq.CANTOS for t in TERMOS]
    nomes += ["valido_F", "valido_R"]
    nomes += [("fy", e, i) for e in "FR" for i in range(N_ALFA)]
    for e in "FR":
        nomes += [f"cap_{e}", f"idx_{e}", f"ok_{e}", f"req_{e}", f"folga_{e}", f"pref_{e}", f"cand_{e}"]
    return {n: COL(j + 1) for j, n in enumerate(nomes)}


def _aba_par(wb, caso: eq.CasoEquilibrio, res: eq.ResultadoPar, ref: dict) -> dict:
    ws = wb.create_sheet("PAR")
    cl = _colunas()
    for n, letra in cl.items():
        ws[f"{letra}1"] = n if isinstance(n, str) else "_".join(str(x) for x in n)
    P = ref.__getitem__
    fz0 = f"({P('LFZO')}*{P('FNOMIN')})"
    n_lin = len(res.linhas)
    r0 = 2
    for k, ln in enumerate(res.linhas):
        r = r0 + k
        ws[f"{cl['k']}{r}"] = k
        ws[f"{cl['ay_g']}{r}"] = float(k * PASSO_G)                 # a_y = (k·passo)·G, como no Python
        ws[f"{cl['ay_ms2']}{r}"] = f"={cl['ay_g']}{r}*{P('G')}"
        gam = eq.camber(caso, ln.ay_ms2)
        for c in eq.CANTOS:
            ws[f"{cl['Fz_' + c]}{r}"] = float(ln.fz[c])
            ws[f"{cl['gam_' + c]}{r}"] = float(gam[c])
        for c in eq.CANTOS:
            FZ, GAM = f"{cl['Fz_' + c]}{r}", f"{cl['gam_' + c]}{r}"
            T = {t: f"{cl[(t, c)]}{r}" for t in TERMOS}
            guarda = (f"OR({FZ}<{P('FZMIN')},{FZ}>{P('FZMAX')},"
                      f"{P('S_G')}*{GAM}<{P('CAMMIN')},{P('S_G')}*{GAM}>{P('CAMMAX')})")
            expr = {
                "dfz": f"({FZ}-{fz0})/{fz0}",
                "gy": f"{P('S_G')}*{GAM}*{P('LGAY')}",
                "Ky": (f"{P('PKY1')}*{fz0}*SIN(2*ATAN({FZ}/({P('PKY2')}*{fz0})))"
                       f"*(1-{P('PKY3')}*ABS({T['gy']}))*{P('LKY')}"),
                "SHy": f"({P('PHY1')}+{P('PHY2')}*{T['dfz']})*{P('LHY')}+{P('PHY3')}*{T['gy']}*{P('LKYG')}",
                "muy": f"({P('PDY1')}+{P('PDY2')}*{T['dfz']})*(1-{P('PDY3')}*{T['gy']}*{T['gy']})*{P('LMUY')}",
                "Dy": f"{T['muy']}*{FZ}",
                "By": f"{T['Ky']}/({P('PCY1')}*{P('LCY')}*{T['Dy']}+{P('EPS')})",
                "SVy": (f"{FZ}*(({P('PVY1')}+{P('PVY2')}*{T['dfz']})*{P('LVY')}"
                        f"+({P('PVY3')}+{P('PVY4')}*{T['dfz']})*{T['gy']}*{P('LKYG')})*{P('LMUY')}"),
            }
            for t, s in (("Eyn", -1), ("Ey0", 0), ("Eyp", 1)):
                expr[t] = (f"MIN(({P('PEY1')}+{P('PEY2')}*{T['dfz']})"
                           f"*(1-({P('PEY3')}+{P('PEY4')}*{T['gy']})*({s}))*{P('LEY')},1)")
            for t in TERMOS:
                ws[T[t]] = f'=IF({guarda},"N/A",{expr[t]})'
        for e, (cL, cR) in RODAS_EIXO.items():
            ws[f"{cl['valido_' + e]}{r}"] = f"=AND(ISNUMBER({cl[('Dy', cL)]}{r}),ISNUMBER({cl[('Dy', cR)]}{r}))"

            def mf(c, i):
                T = {t: f"{cl[(t, c)]}{r}" for t in TERMOS}
                x = f"({ref[('aeq', c, i)]}+{T['SHy']})"
                Ey = f"CHOOSE(SIGN({x})+2,{T['Eyn']},{T['Ey0']},{T['Eyp']})"
                return (f"{T['Dy']}*SIN({P('PCY1')}*{P('LCY')}*ATAN({T['By']}*{x}-{Ey}"
                        f"*({T['By']}*{x}-ATAN({T['By']}*{x}))))+{T['SVy']}")

            for i in range(N_ALFA):
                ws[f"{cl[('fy', e, i)]}{r}"] = (f"=IF({cl['valido_' + e]}{r},-({P('S_F')}*({mf(cL, i)}))"
                                                f"+{P('S_F')}*({mf(cR, i)}),\"N/A\")")
            faixa = f"{cl[('fy', e, 0)]}{r}:{cl[('fy', e, N_ALFA - 1)]}{r}"
            V, CAP, IDX, OK = (f"{cl[n + '_' + e]}{r}" for n in ("valido", "cap", "idx", "ok"))
            REQ, FOL, PREF, CAND = (f"{cl[n + '_' + e]}{r}" for n in ("req", "folga", "pref", "cand"))
            braco = P("b") if e == "F" else P("a")
            ws[CAP] = f'=IF({V},MAX({faixa}),"N/A")'
            ws[IDX] = f'=IF({V},MATCH({CAP},{faixa},0),"N/A")'
            ws[OK] = f"=IF({V},{IDX}<{P('N_ALFA')},FALSE)"
            ws[REQ] = f"={P('m')}*{cl['ay_ms2']}{r}*{braco}/{P('L')}"
            ws[FOL] = f'=IF({OK},{CAP}-{REQ},"N/A")'
            ws[PREF] = f"={OK}" if k == 0 else f"=AND({cl['pref_' + e]}{r - 1},{OK})"
            if k > 0:   # interpolação da Task 3 (equilibrio._saturacao), só no prefixo válido contínuo
                m0, m1 = f"{cl['folga_' + e]}{r - 1}", FOL
                a0, a1 = f"{cl['ay_ms2']}{r - 1}", f"{cl['ay_ms2']}{r}"
                ws[CAND] = (f'=IF({PREF},IF(AND({m0}>0,{m1}<=0),{a0}+({a1}-{a0})*{m0}/({m0}-{m1}),"-"),"-")')
    wr = wb.create_sheet("RESUMO")
    saida = {"PAR": {}, "linhas": n_lin, "r0": r0, "cl": cl}
    for j, e in enumerate("FR"):
        faixa = f"PAR!{cl['cand_' + e]}{r0 + 1}:{cl['cand_' + e]}{r0 + n_lin - 1}"
        wr.cell(j + 1, 1, f"ay_sat_{e}")
        wr.cell(j + 1, 2, f'=IF(COUNT({faixa})=0,"N/A",MIN({faixa}))')
        saida[f"ay_sat_{e}"] = ("RESUMO", f"B{j + 1}")
    return saida


def gerar(caminho: Path, caso: eq.CasoEquilibrio, res: eq.ResultadoPar | None) -> dict | None:
    """Grava a pasta do protótipo. res = None grava só a aba COEF (pasta de base para t_base_s)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ref = _aba_coef(wb, caso)
    mapa = _aba_par(wb, caso, res, ref) if res is not None else None
    wb.save(caminho)
    return mapa


# ---------------------------------------------------------------------------
# Execução e comparação
# ---------------------------------------------------------------------------

def _num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def _dif_sat(py, sh):
    """Diferença relativa da saturação; 0.0 se os dois são N/A; None (= divergência) se só um é."""
    if py is None and sh == "N/A":
        return 0.0
    if py is None or not _num(sh):
        return None
    return abs(sh - py) / abs(py)


def rodar_caso(nome: str, caso: eq.CasoEquilibrio, motor: str, tmp: Path) -> dict:
    res = eq.analise_par(caso, AY_MAX_G, PASSO_G, N_ALFA)
    alfas = eq.grade_alfa(caso, N_ALFA)
    base, cheia = tmp / f"{nome}_base.xlsx", tmp / f"{nome}.xlsx"
    gerar(base, caso, None)
    mapa = gerar(cheia, caso, res)

    t0 = time.perf_counter()
    rb = rc.recalcular(str(base), motor, timeout=600)
    t_base = time.perf_counter() - t0
    t0 = time.perf_counter()
    rf = rc.recalcular(str(cheia), motor, timeout=600)
    t_rec = time.perf_counter() - t0
    t0 = time.perf_counter()
    rc._resumo(str(cheia.resolve()))   # parte do t_recalculo_s que é leitura pelo openpyxl, não o motor
    t_resumo = time.perf_counter() - t0

    wb = openpyxl.load_workbook(cheia, data_only=True)
    ws, cl, r0 = wb["PAR"], mapa["cl"], mapa["r0"]
    dif_cap, n_cap, dif_validade, dif_borda, dif_idx, borda_sh, borda_py = 0.0, 0, 0, 0, 0, 0, 0
    for k, ln in enumerate(res.linhas):
        r = r0 + k
        for e in "FR":
            py = ln.cap_f if e == "F" else ln.cap_r
            cap_sh, idx_sh = ws[f"{cl['cap_' + e]}{r}"].value, ws[f"{cl['idx_' + e]}{r}"].value
            borda_py += py.motivo == "pico na borda da grade de α"
            borda_sh += _num(idx_sh) and int(idx_sh) == N_ALFA
            if py.fy_max_N is None:
                dif_validade += cap_sh != "N/A"
                continue
            if not (_num(cap_sh) and _num(idx_sh)):
                dif_validade += 1
                continue
            n_cap += 1
            dif_cap = max(dif_cap, abs(cap_sh - py.fy_max_N) / abs(py.fy_max_N))
            dif_borda += (py.motivo == "pico na borda da grade de α") != (int(idx_sh) == N_ALFA)
            dif_idx += alfas.index(py.alfa_pico_rad) != int(idx_sh) - 1
    sat_sh = {e: wb[mapa[f"ay_sat_{e}"][0]][mapa[f"ay_sat_{e}"][1]].value for e in "FR"}
    wb.close()

    # Primeira a_y com capacidade inválida por faixa do pneu e fração usada na última linha válida
    # ("Risco conhecido", item b): só leitura do Python, nada extrapolado.
    prim_na, fracao = None, None
    for k, ln in enumerate(res.linhas):
        if ln.cap_f.fy_max_N is None or ln.cap_r.fy_max_N is None:
            prim_na = ln.ay_ms2 / G
            if k > 0:
                lv = res.linhas[k - 1]
                fracao = {"F": lv.req_f / lv.cap_f.fy_max_N, "R": lv.req_r / lv.cap_r.fy_max_N,
                          "ay_g": lv.ay_ms2 / G}
            break

    d_f, d_r = _dif_sat(res.ay_sat_f, sat_sh["F"]), _dif_sat(res.ay_sat_r, sat_sh["R"])
    erros = rb["total_errors"] + rf["total_errors"]
    aprovado = (erros == 0 and n_cap > 0 and dif_cap < TOL and dif_validade == 0 and dif_borda == 0
                and dif_idx == 0 and d_f is not None and d_f < TOL and d_r is not None and d_r < TOL)
    return {
        "caso": nome, "motor": motor, "linhas": len(res.linhas), "n_alfa": N_ALFA,
        "celulas_formula": rf["total_formulas"], "t_recalculo_s": t_rec, "t_base_s": t_base,
        "t_resumo_openpyxl_s": t_resumo,
        "dif_rel_max_capacidade": dif_cap, "capacidades_comparadas": n_cap,
        "dif_ay_sat_f": d_f, "dif_ay_sat_r": d_r,
        "ay_sat_f_g": {"python": res.ay_sat_f and res.ay_sat_f / G,
                       "planilha": sat_sh["F"] / G if _num(sat_sh["F"]) else sat_sh["F"]},
        "ay_sat_r_g": {"python": res.ay_sat_r and res.ay_sat_r / G,
                       "planilha": sat_sh["R"] / G if _num(sat_sh["R"]) else sat_sh["R"]},
        "pico_na_borda": borda_sh, "pico_na_borda_python": borda_py,
        "divergencias_validade": dif_validade, "divergencias_borda": dif_borda,
        "divergencias_indice_pico": dif_idx,
        "erros_formula": erros, "erros_detalhe": {"base": rb["error_summary"], "cheia": rf["error_summary"]},
        "veredito_python": res.veredito, "motivo_python": res.motivo,
        "primeira_ay_na_g": prim_na, "fracao_usada_ultima_valida": fracao,
        "aprovado": aprovado,
    }


def main() -> int:
    motor, detalhe = rc.escolher()
    if motor is None:
        print(f"NÃO EXECUTÁVEL: nenhum motor de recálculo ({detalhe})")
        return 3
    print(f"Motor: {motor} ({detalhe})")
    ds = tm.load_dataset(str(TIR), source_axis_system="ISO_TYDEX")
    ds_045 = eq.com_fatores(ds, {"LMUY": 0.45})   # escala só de teste, não é pneu
    um_grau = math.radians(1.0)
    casos = [
        ("teste_LMUY045", eq.CasoEquilibrio(cfg=veiculo_base(), ds=ds_045)),
        ("teste_LMUY045_toe_camber", eq.CasoEquilibrio(
            cfg=veiculo_base(), ds=ds_045,
            toe_rad={"FL": um_grau, "FR": um_grau, "RL": 0.0, "RR": 0.0},
            camber_estatico_rad={"FL": -0.01, "FR": -0.01, "RL": -0.02, "RR": -0.02},
            ganho_camber={"FL": 0.5, "FR": -0.5, "RL": 0.3, "RR": -0.3})),
        ("teste_LMUY045_LKY023_borda", eq.CasoEquilibrio(
            cfg=veiculo_base(), ds=eq.com_fatores(ds, {"LMUY": 0.45, "LKY": 0.23}))),
    ]
    saida = []
    with tempfile.TemporaryDirectory(prefix="dds_exp_par_", ignore_cleanup_errors=True) as t:
        tmp = Path(t)
        h0 = _sha(VIVA)
        copia = tmp / "viva_copia.xlsx"
        shutil.copyfile(VIVA, copia)
        rc.recalcular(str(copia), motor, timeout=600)
        casos.append(("planilha_viva", cp.caso_da_planilha(copia, TIR)))
        for nome, caso in casos:
            print(f"\n--- caso {nome} ---")
            d = rodar_caso(nome, caso, motor, tmp)
            saida.append(d)
            for chave in ("linhas", "celulas_formula", "t_base_s", "t_recalculo_s", "t_resumo_openpyxl_s",
                          "erros_formula", "dif_rel_max_capacidade", "capacidades_comparadas",
                          "dif_ay_sat_f", "dif_ay_sat_r", "ay_sat_f_g", "ay_sat_r_g", "pico_na_borda",
                          "pico_na_borda_python", "divergencias_validade", "divergencias_borda",
                          "divergencias_indice_pico", "veredito_python", "motivo_python",
                          "primeira_ay_na_g", "fracao_usada_ultima_valida", "aprovado"):
                print(f"  {chave}: {d[chave]}")
        viva_intacta = _sha(VIVA) == h0
    print(f"\nPlanilha viva intacta (SHA-256): {viva_intacta}")

    # JSON acumulado: uma entrada por (motor, caso); a rodada deste motor substitui a anterior.
    anteriores = []
    if SAIDA.exists():
        try:
            anteriores = json.loads(SAIDA.read_text(encoding="utf-8")).get("execucoes", [])
        except (ValueError, AttributeError):
            anteriores = []
    execucoes = [x for x in anteriores if x.get("motor") != motor] + saida
    SAIDA.parent.mkdir(exist_ok=True)
    SAIDA.write_text(json.dumps({
        "experimento": "viabilidade da análise de par na planilha (V4.0, inc. 1); fora da suíte; EDR-044",
        "grade": {"ay_max_g": AY_MAX_G, "passo_g": PASSO_G, "n_alfa": N_ALFA}, "tolerancia": TOL,
        "execucoes": execucoes}, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")
    ok = viva_intacta and all(d["aprovado"] for d in saida)
    print(f"JSON: {SAIDA}")
    print("RESULTADO:", "critérios atendidos" if ok else "DIVERGÊNCIA (investigar na ordem do CLAUDE.md)")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
