"""
Testes sem motor de recálculo da V4.2 (incremento 5, metas e coerência; projeto T11): estrutura do arquivo gerado e CA-07.
Uso: uv run python -X utf8 tests/test_v4_2.py
Namespace TC-WB42-xxx (IDs autorizados: 001 a 006; este script tem o 001, o 002, o 003 (oráculo metas.py) e o 004 (fluxo da viva);
o 005 e o 006 estão em tests/test_workbook_v4_2.py). A V4.1 de partida é o arquivo da tag v4.1,
extraído para uma pasta temporária; nada é gravado em excel/.

TC-WB42-001 (estrutura da V4.2 gerada da tag v4.1 por gerar_v4_2.gerar; origens do T11 e D16 (Task 2) e metas (Task 3)):
- abas e ordem iguais às da V4.1, mais a 23_LISTAS no fim;
- metas (Task 3): 23_LISTAS com as quatro listas de vocabulário (nome definido, valores exatos de metas.VOCAB) e a tabela de
  relações igual a metas.RELACOES; 03_TARGETS com o cabeçalho H4:Q4 (e as auxiliares ocultas R:U), as linhas 10 e 11 (TARGET-006 e
  007), as notas antigas em A14 e A15 e a validação de dados nas entradas com lista; fórmulas novas só nas famílias G1-G6, cada
  uma na forma aprovada; entradas novas em `celulas_de_entrada_v42()` e nomes de data/nomes_V4_2.csv (o CSV é NOMES_V42);
  toda relação com fonte válida (L na lista conferida `metas.FONTES_L` e a página citada na spec; E existente e APPROVED no EDR;
  M com a entrada como raiz do resultado no grafo e entrada de fato digitada);
- as 117569 fórmulas da V4.1 com texto idêntico fora da 16_EQUILIBRIO B188 e B189, e nelas só o literal
  "● Dado do TX" -> "● Dado do T11"; as fórmulas novas são só as das metas (FORMULAS_ESPERADAS_V42);
- o conjunto de células existentes alteradas é exatamente o aprovado (autorização 5 (a), (b), (c), (d) e (e)), montado aqui de forma
  independente: 32 células de Origem (colunas E e H de 05_VEHICLE e 06_SUSPENSION) com `◐ Dado de outro carro (TX, T09)`,
  a lista Origens (00_README A43:B44 e E43:E44), os textos de regra (00_README A49, 14_DASHBOARD A40, 16_EQUILIBRIO A2,
  20_EQ_TROCA A2, 22_CADEIA A2), as duas fórmulas do D16, a h7, a versão, a linha 22 do changelog e, na 03_TARGETS, A3 e
  as notas A11 e A12 (que descem para A14 e A15); todas as outras células (valor e tipo) iguais fora da região nova das metas;
- nenhuma ocorrência de "Dado do TX" nem "Dado do T09" em valor, fórmula, validação ou formatação condicional fora do
  04_DECISION_REGISTER e do 15_CHANGELOG (o `OUTRO_CARRO` só tem "TX, T09", sem "Dado do");
- 00_README A43:A48 == ORIGENS_V42 (seis posições, só as duas primeiras mudam), significado e categoria de
  SIGNIFICADO_V42; a lista `Origens` continua em '00_README'!$A$43:$A$48; os 1146 nomes definidos da V4.1, as validações, a
  formatação condicional e as colunas ocultas das outras abas ficam iguais aos da V4.1;
- 21_COBERTURA: a emenda da h7 (spec §7.4 linha 7) na coluna "Planilha" (D12); a F12 ("Chave agora") continua a fórmula;
- 00_README A3 com "V4.2" e linha 22 do 15_CHANGELOG no formato da 21.
TC-WB42-002: CA-07 da V4.2: duas regenerações (uma da tag, outra do arquivo extraído) iguais, com estrutura; igual a
excel/FSAE_EV_Suspension_Fase1_V4_2.xlsx se ele existir (ele só vai a excel/ na Task 7); o gerador recusa saída em excel/
(e a viva) e saída = origem antes de abrir qualquer arquivo, sem criar nem mudar arquivo; data/nomes_V4_2.csv igual a
aba_metas.NOMES_V42 e o gerador recusa um CSV diferente.
TC-WB42-003: oráculo metas.py (sem planilha): cada estado de coerência forçado, ordem das regras, 7 METAs ao mesmo tempo (Review
Focus 5), formato do EDR (`EDR-61`, `edr-061`), pior coerência e contador.
TC-WB42-004: fluxo da viva sem motor (EDR-063; python/viva.py, tools/conferir_viva.py, retrato_viva.py, migrar_viva.py), em cópias
numa pasta temporária (a "viva" de teste é a V4.1 da tag com dados de teste injetados; o Drive nem é aberto): localizar_viva;
conferir_viva (viva fiel 0, fórmula trocada 1, EDR citado inexistente ou não aprovado 1, ausente/0 byte/lixo/bloqueada 3); retrato_viva
(JSON, ida e volta, tipos); migrar_viva (dado e rótulo mantidos, conflito listado e não sobrescrito, célula fora dos nomes = 2,
ilegível = 3, recusas de caminho).
"Passou" = estrutura e regeneração conferidas; não é validação física.
"""
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.worksheet.formula import ArrayFormula

import aba_equilibrio as ae
import aba_metas
import dados_planilha as dp
import gerar_v4_2 as g42
import metas
import origens
import viva as viva_mod

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
import comparar_planilhas as cp  # noqa: E402
import conferir_viva as conf_viva  # noqa: E402
import migrar_viva as migr_viva  # noqa: E402
import retrato_viva as retr_viva  # noqa: E402
from nomes_v4_0 import NOME_VALIDO, ler_tabela, parece_referencia  # noqa: E402
from raizes import Grafo  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
CSV42 = RAIZ / "data" / "nomes_V4_2.csv"
VIVA = RAIZ / "excel" / "FSAE_EV_Suspension_Fase1_V3_0.xlsx"
results = []

# Células de Origem com rótulo antigo na V4.1 (varredura de 2026-10-23): 05_VEHICLE e 06_SUSPENSION, colunas E e H.
LINHAS_05 = (5, 6, 7, 8, 30, 32, 33, 34)
LINHAS_06 = (4, 5, 6, 7, 10, 11, 14, 15)
DADO = {("05_VEHICLE", f"{c}{r}") for c in "EH" for r in LINHAS_05} | \
       {("06_SUSPENSION", f"{c}{r}") for c in "EH" for r in LINHAS_06}
LISTA = {("00_README", c) for c in ("A43", "B43", "E43", "A44", "B44", "E44")}
REGRA = {("00_README", "A49"), ("14_DASHBOARD", "A40"), ("16_EQUILIBRIO", "A2"), ("20_EQ_TROCA", "A2"),
         ("22_CADEIA", "A2")}
D16 = {("16_EQUILIBRIO", "B188"), ("16_EQUILIBRIO", "B189")}
OUTRAS = {("21_COBERTURA", "D12"), ("00_README", "A3"), ("15_CHANGELOG", "A22"), ("15_CHANGELOG", "B22"),
          ("15_CHANGELOG", "C22")}
# 03_TARGETS: o título (A3), as duas notas que saem de A11 e A12 (A11 vira o ID TARGET-007) e as duas que descem (A14 e A15)
TARGETS_ALT = {("03_TARGETS", c) for c in ("A3", "A11", "A12", "A14", "A15")}
ESPERADO_ALT = DADO | LISTA | REGRA | D16 | OUTRAS | TARGETS_ALT
ESPERADO_ALT_N = 55
COLS_NOVAS = "HIJKLMNOPQRSTU"


def regiao_nova(k) -> bool:
    """Célula da região nova das metas (vazia na V4.1): 03_TARGETS H4:U11 e A10:G11, 14_DASHBOARD A49:E51, e a 23_LISTAS."""
    aba, cel = k
    col, lin = re.match(r"([A-Z]+)(\d+)", cel).groups()
    lin = int(lin)
    if aba == "23_LISTAS":
        return True
    if aba == "03_TARGETS":
        return (col in COLS_NOVAS and 4 <= lin <= 11) or (col in "ABCDEFG" and lin in (10, 11))
    return aba == "14_DASHBOARD" and col in "ABCDE" and 49 <= lin <= 51


# onde cada coluna das fichas é família de fórmula (G1-G6) e a forma aprovada de cada fórmula (regex)
FAMILIA_COL = {"I": "G1", "K": "G2", "L": "G2", "O": "G5", "R": "G3", "S": "G3", "T": "G3", "U": "G4"}
LINHAS_F = range(5, 12)
CELULAS_FORMULA_NOVAS = {("03_TARGETS", f"{c}{r}") for c in FAMILIA_COL for r in LINHAS_F} |     {("14_DASHBOARD", "B50"), ("14_DASHBOARD", "B51")}
LISTAS_NOMES = {"Lista_Papel": "papel", "Lista_Ciclo_Meta": "ciclo", "Lista_Congelamento": "congelamento",
                "Lista_Coerencia": "coerencia"}
RE_G1 = re.compile(r"=IFERROR\(INDEX\('23_LISTAS'!\$D\$(\d+):\$D\$(\d+),MATCH\(A(\d+),'23_LISTAS'!\$C\$(\d+):\$C\$(\d+),0\)\),"
                   r'"sem relação na 23_LISTAS"\)')
RE_NOME = re.compile(r"=(res_[a-z0-9_]+)")
RE_PIOR_GERADA = re.compile(r"=INDEX\(Origens,MAX\((IFERROR\(MATCH\('[0-9A-Za-z_]+'!\$E\$\d+,Origens,0\),6\)(,|\)\)$))+")
RE_G3 = re.compile(r"=SUMPRODUCT\(--EXACT\(Lista_(Papel|Ciclo_Meta|Congelamento),([HDJ])(\d+)\)\)>0")
RE_G4 = re.compile(r'=IF\(I(\d+)="sem relação na 23_LISTAS",0,SUMPRODUCT\(\(\$I\$5:\$I\$11=I(\d+)\)\*EXACT\(\$H\$5:\$H\$11,"META"\)\)'
                   r'-IF\(EXACT\(H(\d+),"META"\),1,0\)\)')


def edr_status():
    """EDR-ddd -> Status da tabela de índice do registro de decisões."""
    txt = (RAIZ / "docs" / "Engineering_Decision_Register.md").read_text(encoding="utf-8")
    return dict(re.findall(r"^\| (EDR-\d{3}) \|.*\| ([A-Z_]+) \|\s*$", txt, re.M))


def check(name, condition, detail=""):
    status = "PASS" if condition else "FAIL"
    results.append((name, status, detail))
    print(f"[{status}] {name}  {detail}")


def raises(fn, exc=Exception):
    try:
        fn()
        return False
    except exc:
        return True


def formula(v):
    if isinstance(v, ArrayFormula):
        return f"{{{v.ref}}}{v.text}"
    return v if isinstance(v, str) and v.startswith("=") else None


def celulas(wb):
    return {(ws.title, c.coordinate): c.value for ws in wb.worksheets
            for lin in ws.iter_rows() for c in lin if c.value is not None}


def varrer_antigos(wb):
    """Onde sobra "Dado do TX" ou "Dado do T09" (valor, fórmula, validação, formatação condicional), fora do 04 e do 15."""
    achados = []
    for ws in wb.worksheets:
        if ws.title in ("04_DECISION_REGISTER", "15_CHANGELOG"):
            continue
        for lin in ws.iter_rows():
            for c in lin:
                t = c.value.text if isinstance(c.value, ArrayFormula) else c.value
                if isinstance(t, str) and ("Dado do TX" in t or "Dado do T09" in t):
                    achados.append(f"{ws.title}!{c.coordinate}")
        for dv in ws.data_validations.dataValidation:
            if any(x in (s or "") for s in (dv.formula1, dv.formula2) for x in ("Dado do TX", "Dado do T09")):
                achados.append(f"{ws.title}:DV")
        for _, regras in ws.conditional_formatting._cf_rules.items():
            if any(x in f for r in regras for f in (r.formula or []) for x in ("Dado do TX", "Dado do T09")):
                achados.append(f"{ws.title}:CF")
    return achados


def tc_wb42_004(tmp: Path, v41: Path, v42: Path) -> None:
    """TC-WB42-004: fluxo da viva (EDR-063) sobre cópias em pasta temporária. A "viva" de teste é a V4.1 da tag com dados de teste
    injetados. Nada é gravado no Drive nem em excel/; o Drive nem é aberto."""
    d41 = dp.ler_por_nome(v41)
    nums = [n for n in sorted(d41) if n.startswith("ent_") and isinstance(d41[n], (int, float)) and not isinstance(d41[n], bool)]
    orgs_tx = [n for n in sorted(d41) if n.startswith("org_") and d41[n] == origens.TX_V41]
    n_num, n_org_a, n_org_b = nums[0], orgs_tx[0], orgs_tx[1]
    novo_num = d41[n_num] + 1.5
    base = tmp / "viva_t"
    base.mkdir()
    contador = [0]

    def pasta() -> Path:
        contador[0] += 1
        p = base / f"tmp_{contador[0]}"
        p.mkdir()
        return p

    def viva_com(nome: str, dados: dict) -> Path:
        d = base / nome
        d.mkdir()
        arq = d / "FSAE_EV_Suspension_Fase1_V4_1.xlsx"
        dp.injetar(v41, dados, arq)
        return arq

    def editar(origem: Path, nome: str, fn) -> Path:
        d = base / nome
        d.mkdir()
        arq = d / origem.name
        wb = openpyxl.load_workbook(origem)
        try:
            fn(wb)
            wb.save(arq)
        finally:
            wb.close()
        return arq

    # ---- constantes e utilitários de viva.py ----
    ref41, ref42 = viva_mod.REFERENCIAS["V4.1"], viva_mod.REFERENCIAS["V4.2"]
    const_ok = ref41.formulas == cp.FORMULAS_ESPERADAS_V41 and ref42.formulas == cp.FORMULAS_ESPERADAS_V42 \
        and ref41.rotulos == tuple(cp.ROTULOS_ORIGEM) and ref42.rotulos == origens.ORIGENS_V42 \
        and ref41.entradas() == ae.celulas_de_entrada_v41() and ref42.entradas() == aba_metas.celulas_de_entrada_v42() \
        and not ref41.confere_edr and ref42.confere_edr \
        and viva_mod.PASTA_DRIVE.name == "Planilha DDS" and "T11 - Tupã" in str(viva_mod.PASTA_DRIVE)
    ver_ok = viva_mod.versao(v41) == "V4.1" and viva_mod.versao(v42) == "V4.2"
    p_ref = pasta()
    ref_tag = viva_mod.arquivo_de_referencia("V4.1", p_ref)
    ref_tag_ok = ref_tag.read_bytes() == v41.read_bytes() and ref_tag.parent == p_ref \
        and raises(lambda: viva_mod.referencia_da("V9.9"), viva_mod.VivaIlegivel)
    if not viva_mod._tag_existe("v4.2") and not (RAIZ / ref42.caminho_excel).exists():   # antes da Task 7
        ref_tag_ok = ref_tag_ok and raises(lambda: viva_mod.arquivo_de_referencia("V4.2", p_ref), viva_mod.VivaIlegivel)
    edrs = viva_mod.edrs_aprovados(RAIZ / "docs" / "Engineering_Decision_Register.md")
    aprovado = next(e for e in sorted(edrs) if edrs[e] == "APPROVED")
    nao_aprov = next((e for e in sorted(edrs) if edrs[e] != "APPROVED"), None)
    edr_ok = edrs.get("EDR-060") == "APPROVED" and edrs == edr_status() and "EDR-999" not in edrs and nao_aprov is not None
    # localizar_viva: DDS_VIVA primeiro; senão o único arquivo da pasta; mais de um = ValueError; nada = None
    antes_env, antes_pasta = os.environ.pop("DDS_VIVA", None), viva_mod.PASTA_DRIVE
    try:
        pf = base / "pasta_drive"
        pf.mkdir()
        viva_mod.PASTA_DRIVE = pf
        loc_vazia = viva_mod.localizar_viva() is None
        (pf / "FSAE_EV_Suspension_Fase1_V4_1.xlsx").write_bytes(b"x")
        (pf / "~$FSAE_EV_Suspension_Fase1_V4_1.xlsx").write_bytes(b"x")     # arquivo de bloqueio do Excel: ignorado
        (pf / "Hardpoints_Tupa_v4_validado.xlsx").write_bytes(b"x")         # outro arquivo da pasta: ignorado
        loc_um = viva_mod.localizar_viva() == pf / "FSAE_EV_Suspension_Fase1_V4_1.xlsx"
        (pf / "FSAE_EV_Suspension_Fase1_V4_2.xlsx").write_bytes(b"x")
        loc_dois = raises(viva_mod.localizar_viva, ValueError)
        os.environ["DDS_VIVA"] = str(base / "qualquer.xlsx")
        loc_env = viva_mod.localizar_viva() == base / "qualquer.xlsx"
        del os.environ["DDS_VIVA"]
        viva_mod.PASTA_DRIVE = base / "nao_montado"
        loc_nada = viva_mod.localizar_viva() is None
    finally:
        viva_mod.PASTA_DRIVE = antes_pasta
        os.environ.pop("DDS_VIVA", None)
        if antes_env is not None:
            os.environ["DDS_VIVA"] = antes_env
    loc_ok = all((loc_vazia, loc_um, loc_dois, loc_env, loc_nada))

    # ---- conferir_viva ----
    viva_fiel = viva_com("fiel", {n_num: novo_num, n_org_a: origens.T09_V41, n_org_b: origens.EST})
    cod_fiel, rel_fiel, ver_fiel = conf_viva.conferir(viva_fiel, pasta())
    fiel_ok = cod_fiel == 0 and ver_fiel == "V4.1" and "CA-07 (viva) PASS" in rel_fiel and any("SHA-256" in l for l in rel_fiel)

    def trocar_formula(wb):
        wb["16_EQUILIBRIO"]["B188"].value = wb["16_EQUILIBRIO"]["B188"].value + '&""'

    def alterar_estrutura(wb):
        wb["01_SETUP"]["A1"].value = "texto de estrutura alterado"

    viva_formula = editar(viva_fiel, "formula", trocar_formula)
    cod_f, rel_f, _ = conf_viva.conferir(viva_formula, pasta())
    formula_ok = cod_f == 1 and any("16_EQUILIBRIO!B188" in l and "fórmula" in l for l in rel_f)
    viva_estr = editar(viva_fiel, "estrutura", alterar_estrutura)
    cod_e, rel_e, _ = conf_viva.conferir(viva_estr, pasta())
    estr_ok = cod_e == 1 and any("01_SETUP!A1" in l for l in rel_e)
    viva_origem = viva_com("origem_ruim", {n_org_a: origens.T11})       # rótulo da V4.2 numa viva V4.1: fora da lista da V4.1
    cod_o, rel_o, _ = conf_viva.conferir(viva_origem, pasta())
    origem_ok = cod_o == 1 and any("origem fora da lista" in l for l in rel_o)
    nome_viva = "FSAE_EV_Suspension_Fase1_V4_1.xlsx"
    ausente = base / "ausente" / nome_viva
    zero, lixo = base / "zero", base / "lixo"
    zero.mkdir()
    lixo.mkdir()
    (zero / nome_viva).write_bytes(b"")
    (lixo / nome_viva).write_bytes(b"isto nao e um xlsx")
    c_aus, r_aus, _ = conf_viva.conferir(ausente, pasta())
    c_zero, r_zero, _ = conf_viva.conferir(zero / nome_viva, pasta())
    c_lixo, r_lixo, _ = conf_viva.conferir(lixo / nome_viva, pasta())
    copia_orig = shutil.copyfile
    try:     # bloqueio (arquivo aberto por alguém): a cópia falha com PermissionError
        def bloqueada(*a, **k):
            raise PermissionError("arquivo em uso")
        shutil.copyfile = bloqueada
        c_blq, r_blq, _ = conf_viva.conferir(viva_fiel, pasta())
    finally:
        shutil.copyfile = copia_orig
    c_sem_ref, r_sem_ref, _ = conf_viva.conferir(viva_fiel, pasta(), referencia=base / "nao_existe_ref.xlsx")
    comparar_orig = cp.comparar_viva
    try:     # exceção inesperada na leitura (fora de VivaIlegivel): 3, nunca 1
        def estoura(*a, **k):
            raise RuntimeError("leitura quebrou")
        cp.comparar_viva = estoura
        c_inesp, r_inesp, _ = conf_viva.conferir(viva_fiel, pasta())
    finally:
        cp.comparar_viva = comparar_orig
    nao_lida_ok = c_inesp == 3 and "erro inesperado" in r_inesp[0] and all(c == 3 and r[0].startswith("NÃO LIDA") for c, r in
                      ((c_aus, r_aus), (c_zero, r_zero), (c_lixo, r_lixo), (c_blq, r_blq), (c_sem_ref, r_sem_ref))) \
        and "0 byte" in r_zero[0] and "bloqueada" in r_blq[0] and "referência ausente" in r_sem_ref[0]
    # EDR citado (V4.2 em diante): a referência é a V4.2 gerada (a tag v4.2 e excel/ ainda não existem)
    casos = {"ok": ({"N5": aprovado}, 0, None), "inexistente": ({"N5": "EDR-999"}, 1, "03_TARGETS!N5 cita EDR-999"),
             "formato": ({"N5": "EDR-61"}, 1, "03_TARGETS!N5 cita EDR-61"),
             "minusc": ({"N5": "edr-060"}, 1, "03_TARGETS!N5 cita edr-060"),
             "nao_aprovado": ({"N6": nao_aprov}, 1, f"03_TARGETS!N6 cita {nao_aprov}"),
             "decision": ({"F7": "ver EDR-998"}, 1, "03_TARGETS!F7 cita EDR-998"),
             "tbd": ({"N5": "TBD", "M5": "preferência do piloto"}, 0, None)}
    edr_casos = {}
    for k, (cel, esperado, trecho) in casos.items():
        def preencher(wb, cel=cel):
            for c, x in cel.items():
                wb["03_TARGETS"][c].value = x
        arq = editar(v42, f"v42_{k}", preencher)
        cod, rel, ver = conf_viva.conferir(arq, pasta(), referencia=v42)
        edr_casos[k] = cod == esperado and ver == "V4.2" and (trecho is None or any(trecho in l for l in rel))
    edr_conf_ok = all(edr_casos.values()) and viva_mod.edrs_citados(base / "v42_decision" / v42.name) == [("03_TARGETS!F7", "EDR-998")] \
        and viva_mod.edrs_citados(base / "v42_tbd" / v42.name) == [] and viva_mod.edrs_citados(viva_fiel) == []
    # a conferência nunca altera a viva
    sha_fiel = viva_mod.sha256(viva_fiel)
    conf_viva.conferir(viva_fiel, pasta())
    imut_ok = viva_mod.sha256(viva_fiel) == sha_fiel

    # ---- retrato_viva ----
    out_json = base / "retrato" / "retrato-viva.json"
    cod_r = retr_viva.retratar(viva_fiel, out_json, pasta())
    txt = out_json.read_text(encoding="utf-8") if out_json.exists() else ""
    ret = json.loads(txt) if txt else {}
    dados_viva = dp.ler_por_nome(viva_fiel)
    carregado = retr_viva.carregar(out_json) if txt else {"dados": {}}
    forma_ok = cod_r == 0 and txt.startswith('{\n  "versao": "V4.1",\n  "sha256": "') and txt.endswith("}\n") and "\r" not in txt \
        and ret.get("versao") == "V4.1" and ret.get("sha256") == viva_mod.sha256(viva_fiel) \
        and re.fullmatch(r"\d{4}-\d{2}-\d{2}", ret.get("data", "")) and list(ret.get("dados", {})) == sorted(dados_viva) \
        and txt.count("\n") == len(dados_viva) + 7 \
        and ret["dados"][n_num] == novo_num and isinstance(ret["dados"][n_num], type(novo_num)) \
        and ret["dados"][n_org_a] == origens.T09_V41 and ret["dados"][n_org_b] == origens.EST
    volta = base / "retrato" / "volta.xlsx"
    rel_inj = dp.injetar(v41, carregado["dados"], volta)
    dados_volta = dp.ler_por_nome(volta)
    ida_volta_ok = not rel_inj.conflitos and not rel_inj.faltando and set(dados_volta) == set(dados_viva) \
        and all(dp.iguais(dados_volta[n], dados_viva[n]) for n in dados_viva)
    dh, dd, dtm = dt.datetime(2026, 10, 26, 9, 30, 15), dt.date(2026, 10, 26), dt.time(9, 30)
    tipos_ok = retr_viva.decodificar(retr_viva.codificar(dh)) == dh and type(retr_viva.decodificar(retr_viva.codificar(dh))) is dt.datetime \
        and retr_viva.decodificar(retr_viva.codificar(dd)) == dd and type(retr_viva.decodificar(retr_viva.codificar(dd))) is dt.date \
        and retr_viva.decodificar(retr_viva.codificar(dtm)) == dtm \
        and type(retr_viva.codificar(2)) is int and type(retr_viva.codificar(2.0)) is float \
        and raises(lambda: retr_viva.codificar(object()), ValueError) and raises(lambda: retr_viva.decodificar({"x": 1}), ValueError)
    out_ruim = base / "retrato" / "ruim.json"
    cod_ruim = retr_viva.retratar(viva_formula, out_ruim, pasta())
    cod_aus = retr_viva.retratar(ausente, out_ruim, pasta())
    cod_excel = retr_viva.retratar(viva_fiel, RAIZ / "excel" / "_retrato_teste.json", pasta())
    # a guarda de saída vem antes de qualquer escrita: árvore excel/ (não só o pai), Drive (caminho falso), sobre a própria viva
    cod_sub = retr_viva.retratar(viva_fiel, RAIZ / "excel" / "sub" / "_retrato_teste.json", pasta())
    cod_drive = retr_viva.retratar(viva_fiel, viva_mod.PASTA_DRIVE / "sub" / "_retrato_teste.json", pasta())
    pasta_antes = viva_mod.PASTA_DRIVE
    try:     # a mesma recusa com a pasta do Drive apontando para uma pasta temporária (prova que nada foi criado)
        viva_mod.PASTA_DRIVE = base / "drive_falso"
        cod_drive2 = retr_viva.retratar(viva_fiel, base / "drive_falso" / "retrato.json", pasta())
        drive_vazio = not (base / "drive_falso").exists()
    finally:
        viva_mod.PASTA_DRIVE = pasta_antes
    sha_antes = viva_mod.sha256(viva_fiel)
    cod_viva = retr_viva.retratar(viva_fiel, viva_fiel, pasta())
    recusa_ret_ok = cod_ruim == 1 and cod_aus == 3 and cod_excel == 2 and not out_ruim.exists() \
        and not (RAIZ / "excel" / "_retrato_teste.json").exists() and cod_sub == 2 and not (RAIZ / "excel" / "sub").exists() \
        and cod_drive == 2 and cod_drive2 == 2 and drive_vazio and cod_viva == 2 and viva_mod.sha256(viva_fiel) == sha_antes \
        and viva_mod.recusa_de_saida(base / "x.json", [viva_fiel]) is None
    retrato_ok = forma_ok and ida_volta_ok and tipos_ok and recusa_ret_ok

    # ---- migrar_viva ----
    def migrar(viva: Path, para: Path, nome: str):
        s, r = base / f"mig_{nome}" / "saida.xlsx", base / f"mig_{nome}" / "relatorio.md"
        cod = migr_viva.migrar(viva, para, s, r, pasta())
        return cod, s, (r.read_text(encoding="utf-8") if r.exists() else "")

    d42 = dp.ler_por_nome(v42)
    cod_a, s_a, t_a = migrar(viva_fiel, v42, "a")
    da = dp.ler_por_nome(s_a) if s_a.exists() else {}
    # dado mantido; rótulo mapeado (TX/T09 -> outro carro); rótulo da equipe (Estimativa) mantido; o resto igual à V4.2
    resto_a = bool(da) and all(dp.iguais(da[n], d42[n]) for n in d42 if n not in (n_num, n_org_b))
    mig_a_ok = cod_a == 0 and dp.iguais(da.get(n_num), novo_num) and da.get(n_org_a) == origens.OUTRO_CARRO \
        and da.get(n_org_b) == origens.EST and resto_a and f"`{n_num}`" in t_a and f"`{n_org_b}`" in t_a \
        and f"`{n_org_a}`" not in t_a.split("## Escritas")[1] and "Resultado: **PASS**" in t_a and "\r" not in t_a
    # a mesma célula mudada pela equipe e pela versão nova: conflito listado, nunca sobrescrito (Review Focus 4)
    para_c = base / "para_conflito.xlsx"
    dp.injetar(v42, {n_num: d42[n_num] + 100, n_org_b: origens.ARB}, para_c)
    cod_b, s_b, t_b = migrar(viva_fiel, para_c, "b")
    db = dp.ler_por_nome(s_b) if s_b.exists() else {}
    conflitos_b = t_b.split("## Faltando")[0]
    mig_b_ok = cod_b == 1 and dp.iguais(db.get(n_num), d42[n_num] + 100) and db.get(n_org_b) == origens.ARB \
        and f"| `{n_num}` |" in conflitos_b and f"| `{n_org_b}` |" in conflitos_b and "Resultado: **FAIL**" in t_b
    # célula alterada fora dos nomes de dado: código 2, lista a célula, não grava a saída
    cod_c, s_c, t_c = migrar(viva_estr, v42, "c")

    def comentar(wb):
        wb["00_README"]["Z99"].comment = Comment("nota da equipe", "equipe")

    def formula_no_dado(wb):
        aba, cel = dp._destinos(wb)[n_num]
        wb[aba][cel].value = "=1+1"

    cod_d, s_d, t_d = migrar(editar(viva_fiel, "comentario", comentar), v42, "d")        # objeto não migrável
    cod_e0, s_e0, t_e0 = migrar(zero / nome_viva, v42, "e0")                              # ilegível: 0 byte
    cod_e1, s_e1, t_e1 = migrar(ausente, v42, "e1")                                       # ilegível: ausente
    viva_v42 = editar(v42, "viva_v42", lambda wb: None)
    cod_nova, s_nova, t_nova = migrar(viva_v42, v41, "nova")                              # viva mais nova que o destino
    cod_fd, s_fd, t_fd = migrar(editar(viva_fiel, "dado_formula", formula_no_dado), v42, "fd")   # fórmula numa entrada
    erros_ok = cod_c == 2 and not s_c.exists() and "01_SETUP!A1" in t_c and "ERRO" in t_c \
        and cod_d == 2 and not s_d.exists() and "00_README!Z99 (comentário)" in t_d \
        and cod_e0 == 3 and cod_e1 == 3 and not s_e0.exists() and not s_e1.exists() and "ERRO" in t_e0 and "0 byte" in t_e0 \
        and cod_fd == 2 and not s_fd.exists() and "ERRO" in t_fd \
        and cod_nova == 2 and not s_nova.exists() and "mais nova que o destino" in t_nova
    # recusas de caminho: nada é criado nem aberto
    p_rec = pasta()
    recusas = [(viva_fiel, v42, RAIZ / "excel" / "_migrar_teste.xlsx", base / "r.md"),
               (viva_fiel, v42, viva_fiel, base / "r2.md"),
               (viva_fiel, v42, v42, base / "r3.md"),
               (viva_fiel, v42, base / "s4.xlsx", RAIZ / "excel" / "_relatorio_teste.md"),
               (viva_fiel, v42, viva_mod.PASTA_DRIVE / "FSAE_EV_Suspension_Fase1_V4_2.xlsx", base / "r5.md"),
               (viva_fiel, v42, base / "s6.xlsx", viva_mod.PASTA_DRIVE / "relatorio.md"),
               (viva_fiel, v42, base / "s7.txt", base / "r7.md")]
    cods = [migr_viva.migrar(*r, p_rec) for r in recusas]
    recusa_ok = cods == [2] * len(recusas) and not any(p.exists() for p in (
        RAIZ / "excel" / "_migrar_teste.xlsx", RAIZ / "excel" / "_relatorio_teste.md", base / "r.md", base / "s4.xlsx"))
    mig_ok = mig_a_ok and mig_b_ok and erros_ok and recusa_ok

    # ---- CLIs: código de saída e mensagem, sem abrir planilha ----
    py = [sys.executable, "-X", "utf8"]
    r_cli1 = subprocess.run([*py, str(RAIZ / "tools" / "conferir_viva.py"), str(ausente)], capture_output=True, text=True,
                            encoding="utf-8", cwd=RAIZ)
    r_cli2 = subprocess.run([*py, str(RAIZ / "tools" / "migrar_viva.py"), str(ausente), "--para", str(v42), "--saida",
                             str(base / "cli" / "s.xlsx"), "--relatorio", str(base / "cli" / "r.md")], capture_output=True,
                            text=True, encoding="utf-8", cwd=RAIZ)
    cli_ok = r_cli1.returncode == 3 and "NÃO LIDA" in r_cli1.stdout and r_cli2.returncode == 3 \
        and "ERRO" in (base / "cli" / "r.md").read_text(encoding="utf-8")
    check("TC-WB42-004: fluxo da viva sem motor (EDR-063), em cópias: constantes e referência da tag; localizar_viva (DDS_VIVA, pasta, "
          "mais de um, nada); conferir_viva (viva fiel 0, fórmula trocada 1, estrutura 1, rótulo fora da lista 1; ausente, 0 byte, lixo, "
          "bloqueada e sem referência 3; EDR citado inexistente, em formato errado ou não aprovado 1); retrato_viva (forma, ida e volta "
          "pelo JSON, tipos, recusa com conferência diferente de 0); migrar_viva (dado e rótulo mantidos, conflito listado e não "
          "sobrescrito, célula fora dos nomes e objeto = 2, ilegível = 3, recusas de caminho); códigos de saída das CLIs",
          all((const_ok, ver_ok, ref_tag_ok, edr_ok, loc_ok, fiel_ok, formula_ok, estr_ok, origem_ok, nao_lida_ok, edr_conf_ok, imut_ok,
               retrato_ok, mig_ok, cli_ok)),
          f"const={const_ok} versão={ver_ok} tag={ref_tag_ok} edr={edr_ok} localizar={loc_ok} fiel={fiel_ok} fórmula={formula_ok} "
          f"estrutura={estr_ok} origem={origem_ok} não_lida={nao_lida_ok} edr_conferência={edr_casos}/{edr_conf_ok} imutável={imut_ok} "
          f"retrato={forma_ok}/{ida_volta_ok}/{tipos_ok}/{recusa_ret_ok} migrar: a={mig_a_ok} (cód {cod_a}) b={mig_b_ok} (cód {cod_b}) "
          f"erros={erros_ok} ({cod_c}/{cod_d}/{cod_e0}/{cod_e1}/{cod_fd}) recusas={recusa_ok} ({cods}) cli={cli_ok}")


with tempfile.TemporaryDirectory(prefix="dds_v42_") as tmp:
    tmp = Path(tmp)
    v41 = g42.extrair_v41(tmp)
    v42 = tmp / "V4_2.xlsx"
    _, novas = g42.gerar(None, v42)              # da tag v4.1, como o CLI sem --v41; `novas` = fórmulas novas (família)
    wb41, wb42 = openpyxl.load_workbook(v41), openpyxl.load_workbook(v42)

    # ---------------- TC-WB42-001 ----------------
    abas_ok = wb42.sheetnames == wb41.sheetnames + ["23_LISTAS"] and wb41.sheetnames[0] == "00_README" \
        and wb41.sheetnames[-1] == "22_CADEIA"
    c41, c42 = celulas(wb41), celulas(wb42)
    f41 = {k: formula(v) for k, v in c41.items() if formula(v)}
    f42 = {k: formula(v) for k, v in c42.items() if formula(v)}
    n_form = len(f41) == 117569 and set(f41) <= set(f42) and len(f42) == cp.FORMULAS_ESPERADAS_V42 == 117627 \
        and set(f42) - set(f41) == CELULAS_FORMULA_NOVAS
    form_difs = sorted(k for k in f41 if f42.get(k) != f41[k])
    lit_ok = form_difs == sorted(D16) and all(
        f42[k] == f41[k].replace('"● Dado do TX"', '"● Dado do T11"') and f41[k].count('"● Dado do TX"') == 1
        and '"● Dado do TX"' not in f42[k] for k in D16)
    mudou = {k for k in set(c41) | set(c42) if c41.get(k) != c42.get(k) or type(c41.get(k)) is not type(c42.get(k))}
    fora = {k for k in mudou if not regiao_nova(k) or k in ESPERADO_ALT}
    conj_ok = fora == ESPERADO_ALT and len(ESPERADO_ALT) == ESPERADO_ALT_N \
        and all(c41.get(k) is None for k in mudou if regiao_nova(k) and k not in ESPERADO_ALT)
    # dado: o antigo era TX ou T09, o novo é "outro carro" (E e H iguais entre si)
    dado_ok = len(DADO) == 32 and all(c41[k] in (origens.TX_V41, origens.T09_V41) and c42[k] == origens.OUTRO_CARRO
                                      for k in DADO) and all(c42[(a, f"E{r}")] == c42[(a, f"H{r}")]
                                                              for a, r in [(a, int(re.sub(r"[A-Z]", "", c))) for a, c in DADO])
    # varredura independente: toda célula não-fórmula fora do 00_README com o rótulo exato antigo está em DADO
    com_rotulo = {k for k, v in c41.items() if v in (origens.TX_V41, origens.T09_V41) and k[0] != "00_README"}
    varredura_ok = com_rotulo == DADO
    sobra = varrer_antigos(wb42)
    sobra41 = varrer_antigos(wb41)          # a varredura enxerga o que existe na V4.1 (senão "limpo" não prova nada)
    sem_antigo = not sobra and len(sobra41) == 41
    # lista Origens
    r43 = [c42[("00_README", f"A{r}")] for r in range(43, 49)]
    lista_ok = tuple(r43) == origens.ORIGENS_V42 and len(set(r43)) == 6 and r43[2:] == list(origens.ORIGENS_V42[2:]) and all(
        c42[("00_README", f"B{43 + i}")] == origens.SIGNIFICADO_V42[rot][0]
        and c42[("00_README", f"E{43 + i}")] == origens.SIGNIFICADO_V42[rot][1] for i, rot in enumerate(origens.ORIGENS_V42[:2])) \
        and all(c42[("00_README", f"{c}{r}")] == c41[("00_README", f"{c}{r}")] for r in range(45, 49) for c in "ABE")
    # nomes: os 1146 da V4.1 iguais; os novos = NOMES_V42 (CSV) + as quatro listas de vocabulário
    n41 = {k: d.attr_text for k, d in wb41.defined_names.items()}
    n42 = {k: d.attr_text for k, d in wb42.defined_names.items()}
    novos_n42 = {k: v for k, v in n42.items() if k not in n41}
    nomes_csv = ler_tabela(CSV42)
    esperados_n = {n.nome: f"'{n.aba}'!${re.sub(r'[0-9]', '', n.celula)}${re.sub(r'[A-Z]', '', n.celula)}"
                   for n in aba_metas.NOMES_V42}
    esperados_n.update(aba_metas.LISTAS_V42)
    nomes_ok = {k: n42.get(k) for k in n41} == n41 and n42.get("Origens") == "'00_README'!$A$43:$A$48" and len(n41) == 1146 \
        and novos_n42 == esperados_n and nomes_csv == list(aba_metas.NOMES_V42) and len(aba_metas.NOMES_V42) == 75 \
        and len({n.nome for n in nomes_csv}) == len(nomes_csv) \
        and all(NOME_VALIDO.fullmatch(n.nome) and not parece_referencia(n.nome) for n in nomes_csv) \
        and {n.nome for n in nomes_csv}.isdisjoint(n41) \
        and all(n.nome.startswith(("ent_03_", "res_03_", "res_14_")) for n in nomes_csv) \
        and set(aba_metas.LISTAS_V42) == set(LISTAS_NOMES)
    # estrutura: as outras abas iguais às da V4.1; as duas abas novas das metas com o que a Task 3 aprova
    e41, e42 = cp._estrutura(v41), cp._estrutura(v42)
    est = all(e41[a] == e42[a] for a in wb41.sheetnames if a != "03_TARGETS") \
        and e42["03_TARGETS"]["validação"] == sorted([("D5:D11", "Lista_Ciclo_Meta"), ("H5:H11", "Lista_Papel"),
                                                      ("J5:J11", "Lista_Congelamento")]) \
        and e42["03_TARGETS"]["formatação condicional"] == e41["03_TARGETS"]["formatação condicional"] \
        and e42["03_TARGETS"]["colunas ocultas"] == ["R", "S", "T", "U"] and e41["03_TARGETS"]["colunas ocultas"] == [] \
        and e42["23_LISTAS"]["validação"] == [] and e42["23_LISTAS"]["colunas ocultas"] == []
    # 03_TARGETS: cabeçalho, linhas 10 e 11, notas antigas em A14 e A15
    t3 = wb42["03_TARGETS"]
    cab = ("Papel", "Alavanca", "Congelamento", "Valor atual do modelo", "Pior origem da grandeza", "Compromisso declarado",
           "EDR do compromisso", "Coerência", "Ficha", "Objetivo atendido")
    cab_ok = tuple(t3[f"{c}4"].value for c in "HIJKLMNOPQ") == cab and all(t3[f"{c}4"].value for c in "RSTU")
    linhas_ok = (t3["A10"].value, t3["B10"].value, t3["A11"].value, t3["B11"].value) == (
        "TARGET-006", "Gradiente de subesterço K (SUS-1001)", "TARGET-007",
        "Comportamento no limite: push/loose e C_N (SUS-1004)") \
        and all([t3[f"{c}{r}"].value for c in "CDEFG"] == ["TBD", "TBD", "TBD", "-", "-"] for r in (10, 11))
    cond_ini = [(t3[f"D{r}"].value, t3[f"H{r}"].value, t3[f"J{r}"].value, t3[f"M{r}"].value, t3[f"N{r}"].value) for r in LINHAS_F]
    ini_ok = [x[:3] for x in cond_ini] == [("OPEN", "TBD", "F0")] * 5 + [("TBD", "TBD", "F0")] * 2 \
        and all(x[3] is None and x[4] is None for x in cond_ini) \
        and all(t3[f"P{r}"].value == f"docs/metas/TARGET-00{r - 4}.md" for r in LINHAS_F) \
        and all(t3[f"Q{r}"].value == "TBD: lista de objetivos da equipe a definir em EDR" for r in LINHAS_F) \
        and all(t3[f"C{r}"].value == "TBD" and t3[f"E{r}"].value == "TBD" for r in LINHAS_F)
    notas_ok = t3["A14"].value == c41[("03_TARGETS", "A11")] and t3["A15"].value == c41[("03_TARGETS", "A12")] \
        and t3["A12"].value is None and t3["A13"].value is None and c41[("03_TARGETS", "A11")].startswith("'Target' = TBD") \
        and sorted(str(r) for r in t3.merged_cells.ranges) == sorted(["A1:E1", "A3:E3", "A14:G14", "A15:G15"]) \
        and t3["A3"].value.startswith(c41[("03_TARGETS", "A3")]) and t3["A3"].value != c41[("03_TARGETS", "A3")]
    # 23_LISTAS: as quatro listas (nome definido, valores exatos) e a tabela de relações
    ws23 = wb42["23_LISTAS"]

    def faixa(nome):
        aba, ref = n42[nome].split("!")
        a, b = ref.replace("$", "").split(":")
        col = re.sub(r"[0-9]", "", a)
        return aba.strip("'"), [ws23[f"{col}{r}"].value for r in range(int(re.sub(r"[A-Z]", "", a)), int(re.sub(r"[A-Z]", "", b)) + 1)]

    listas_ok = all(faixa(n)[0] == "23_LISTAS" and tuple(faixa(n)[1]) == metas.VOCAB[v] for n, v in LISTAS_NOMES.items())
    rel_cab = tuple(ws23.cell(aba_metas.LINHA_REL0 - 1, c).value for c in range(1, 10))
    rel_lin = [tuple(ws23.cell(aba_metas.LINHA_REL0 + i, c).value for c in range(1, 10)) for i in range(len(metas.RELACOES))]
    rel_cab_ok = rel_cab == ("Nº", "Tipo", "Meta", "Alavanca", "Fonte (L/E/M)", "Fonte", "Resultado", "Entrada da alavanca", "Nota")
    rel_tab_ok = rel_lin == [(r.n, r.tipo, r.meta, r.alavanca, r.fonte_tipo, r.fonte, r.resultado, r.entrada,
                              metas.NOTAS_RELACAO[r.n]) for r in metas.RELACOES] \
        and ws23.cell(aba_metas.LINHA_REL0 + len(metas.RELACOES), 1).value is None
    textos23 = " ".join(str(c.value) for lin in ws23.iter_rows() for c in lin if c.value is not None)
    nota_ok = "linha sem fonte não existe; ausência de linha = não avaliado" in textos23 and "Guia G.1" in textos23 \
        and "§3.2" in textos23
    # significados da 23_LISTAS iguais aos do oráculo (uma fonte só)
    sig_ok = True
    for chave, nome in (("papel", "Lista_Papel"), ("ciclo", "Lista_Ciclo_Meta"), ("congelamento", "Lista_Congelamento"),
                        ("coerencia", "Lista_Coerencia")):
        aba_, ref = n42[nome].split("!")
        a, b = ref.replace("$", "").split(":")
        r0, r1 = int(re.sub(r"[A-Z]", "", a)), int(re.sub(r"[A-Z]", "", b))
        for r in range(r0, r1 + 1):
            sig = metas.SIGNIFICADO[chave][ws23[f"A{r}"].value]
            esperado = list(sig) if isinstance(sig, tuple) else [sig]
            sig_ok &= [ws23[f"{c}{r}"].value for c in "BC"[:len(esperado)]] == esperado
    # fontes das relações: L conferida e na spec; E existente e aprovada; M com a entrada como raiz do resultado
    spec = (RAIZ / "docs" / "superpowers" / "specs" / "2026-10-20-v4-0-metas-equilibrio-design.md").read_text(encoding="utf-8")
    status = edr_status()
    grafo = Grafo(wb42)
    entradas_v42 = aba_metas.celulas_de_entrada_v42()

    spec_linha = " ".join(spec.split())   # as citações atravessam quebras de linha da spec

    def fonte_l_ok(fonte):
        """Fonte L: uma das conferidas (`metas.FONTES_L`) e citada na spec com a obra e a página juntas, como está escrita
        (só o número da página, em qualquer lugar da spec, não basta)."""
        return fonte in metas.FONTES_L and fonte in spec_linha

    def fonte_e_ok(fonte):
        return bool(re.fullmatch(r"EDR-\d{3}", fonte)) and status.get(fonte) in ("APPROVED", "IMPLEMENTED", "VERIFIED")

    fontes_ok, fontes_det = True, []
    for r in metas.RELACOES:
        if r.fonte_tipo == "L":
            ok = fonte_l_ok(r.fonte)
        elif r.fonte_tipo == "E":
            ok = fonte_e_ok(r.fonte)
        elif r.fonte_tipo == "M":
            dest_r, dest_e = (ae._destino(wb42, r.resultado), ae._destino(wb42, r.entrada))
            ok = (dest_e in grafo.raizes_preenchidas(dest_r)
                  and dest_e in entradas_v42
                  and r.entrada.startswith("ent_")
                  and dest_r in grafo.formulas
                  and r.fonte == f"{r.resultado} ← {r.entrada}")
        else:
            ok = False
        fontes_ok &= bool(ok)
        if not ok:
            fontes_det.append(r.n)
    # os verificadores pegam fonte ruim (senão "ok" não prova nada): página que não é a conferida, livro sem página,
    # EDR inexistente ou fora do formato; e a entrada que não é raiz do resultado
    pega_l = (not fonte_l_ok("Gillespie p. 999") and not fonte_l_ok("Milliken") and not fonte_l_ok("Gillespie p. 215")
              and not fonte_l_ok("p. 214") and not fonte_l_ok("214"))   # só a página, sem a obra, não passa
    pega_e = not fonte_e_ok("EDR-999") and not fonte_e_ok("EDR-61") and fonte_e_ok("EDR-060")
    pega_m = ae._destino(wb42, "ent_06_k_arb_traseiro") not in grafo.raizes_preenchidas(
        ae._destino(wb42, "res_06_frequencia_natural_dianteira"))
    # entradas novas
    entradas_novas = {k for k in aba_metas.celulas_de_entrada_v42() - ae.celulas_de_entrada_v41()}
    ent_esperadas = {("03_TARGETS", f"{c}{r}") for r in (10, 11) for c in "CDEFG"} | \
        {("03_TARGETS", f"{c}{r}") for r in LINHAS_F for c in "HJMNQ"}
    ent_ok = entradas_novas == ent_esperadas and len(ent_esperadas) == 45 \
        and aba_metas.celulas_de_entrada_v42() >= ae.celulas_de_entrada_v41() \
        and {n.celula for n in aba_metas.NOMES_V42 if n.papel == "ent"} == {c for _, c in ent_esperadas} \
        and all((n.aba, n.celula) in ent_esperadas for n in aba_metas.NOMES_V42 if n.papel == "ent") \
        and all(n.nome.startswith("ent_03_") for n in aba_metas.NOMES_V42 if n.papel == "ent") \
        and {(n.aba, n.celula) for n in aba_metas.NOMES_V42 if n.papel == "res"} == CELULAS_FORMULA_NOVAS - \
        {("03_TARGETS", f"{c}{r}") for c in "RSTU" for r in LINHAS_F} \
        and aba_metas.LINHA_TARGET == {f"TARGET-00{i}": 4 + i for i in range(1, 8)}
    # famílias G1-G6: cada fórmula nova está no registro da geração, na família e na forma aprovadas
    reg_novas = {(n["aba"], n["cel"]): n for n in novas}
    fam_ok = set(reg_novas) == CELULAS_FORMULA_NOVAS and all(reg_novas[k]["formula"] == f42[k] for k in reg_novas)
    rel_alav = {r.meta: r for r in metas.RELACOES if r.tipo == "alavanca" and r.fonte_tipo == "M"}
    for r in LINHAS_F:
        tgt = f"TARGET-00{r - 4}"
        fam = lambda col: reg_novas[("03_TARGETS", f"{col}{r}")]["familia"] == FAMILIA_COL[col]
        fam_ok &= all(fam(c) for c in FAMILIA_COL)
        m1 = RE_G1.fullmatch(f42[("03_TARGETS", f"I{r}")])
        fam_ok &= bool(m1) and m1.group(1) == m1.group(4) and m1.group(2) == m1.group(5) and int(m1.group(3)) == r \
            and int(m1.group(2)) - int(m1.group(1)) + 1 >= len(metas.RELACOES) \
            and [ws23[f"C{k}"].value for k in range(int(m1.group(1)), int(m1.group(2)) + 1)][:len(metas.RELACOES)] == \
            [x.meta for x in metas.RELACOES]
        fam_ok &= f42[("03_TARGETS", f"K{r}")] == "=" + rel_alav[tgt].resultado
        fam_ok &= bool(RE_NOME.fullmatch(f42[("03_TARGETS", f"L{r}")]) or RE_PIOR_GERADA.fullmatch(f42[("03_TARGETS", f"L{r}")]))
        for col, lista, campo in (("R", "Papel", "H"), ("S", "Ciclo_Meta", "D"), ("T", "Congelamento", "J")):
            m3 = RE_G3.fullmatch(f42[("03_TARGETS", f"{col}{r}")])
            fam_ok &= bool(m3) and (m3.group(1), m3.group(2), int(m3.group(3))) == (lista, campo, r)
        m4 = RE_G4.fullmatch(f42[("03_TARGETS", f"U{r}")])
        fam_ok &= bool(m4) and all(int(g) == r for g in m4.groups())
        o = f42[("03_TARGETS", f"O{r}")]
        fam_ok &= o.startswith(f'=IF(NOT(R{r}),"CONFLITO: Papel fora do vocabulário"') and o.endswith(f'&" · "&L{r}') \
            and f'EXACT(LEFT(N{r}&"",4),"EDR-")' in o and "SUMPRODUCT(--ISNUMBER(FIND(MID(N" in o
    # L das metas 001 a 003, 006 e 007 aponta para a pior origem que já existe (14_DASHBOARD coluna E); 004 e 005, gerada
    esperado_l = {5: "=res_14_frequencia_natural_dianteira_e15", 6: "=res_14_frequencia_natural_traseira_e16",
                  7: "=res_14_distribuicao_de_rigidez_dianteira_e17", 10: "=res_14_gradiente_de_subesterco_k_linear_e41",
                  11: "=res_14_equilibrio_no_limite_push_loose_e42"}
    l_ok = all(f42[("03_TARGETS", f"L{r}")] == v for r, v in esperado_l.items()) \
        and all(RE_PIOR_GERADA.fullmatch(f42[("03_TARGETS", f"L{r}")]) for r in (8, 9)) \
        and all(n42[v[1:]].endswith(c) for v, c in ((esperado_l[5], "$E$15"), (esperado_l[6], "$E$16"), (esperado_l[7], "$E$17"),
                                                    (esperado_l[10], "$E$41"), (esperado_l[11], "$E$42")))
    # 14_DASHBOARD: bloco das metas
    d14 = wb42["14_DASHBOARD"]
    dash_ok = d14["A49"].value == "Metas (V4.2; 03_TARGETS)" and d14["A50"].value == "Pior coerência das fichas" \
        and d14["A51"].value == "Compromissos declarados" and all(d14[f"A{r}"].value is None for r in (48, 52)) \
        and f42[("14_DASHBOARD", "B50")].startswith('=IF(COUNTIF(\'03_TARGETS\'!$O$5:$O$11,"CONFLITO*")>0,"CONFLITO",') \
        and f42[("14_DASHBOARD", "B51")] == "=COUNTIF('03_TARGETS'!$O$5:$O$11,\"COMPROMISSO DECLARADO*\")" \
        and reg_novas[("14_DASHBOARD", "B50")]["familia"] == reg_novas[("14_DASHBOARD", "B51")]["familia"] == "G6"
    # h7
    h7 = c42[("21_COBERTURA", "D12")]
    h7_ok = h7.startswith("Sim (entrada, 0 até medir; um sentido por vez, pela entrada \"Sentido da curva\": o outro se vê trocando o sentido;") \
        and h7.endswith("EDR-060, 2026-10-22)") and c41[("21_COBERTURA", "D12")] == "Sim (entrada, 0 até medir)" \
        and c42[("21_COBERTURA", "F12")] == "=ent_16_chave_h7_cg_lateral"
    a3 = c42[("00_README", "A3")]
    a3_ok = a3.startswith("Versão do workbook: V4.2 (incremento 5: metas e coerência; projeto T11 (Origens e D16))") \
        and "V4.1" not in a3 and a3.endswith("| Status: DRAFT — cross-check, ainda NÃO REFERENCE-validated") \
        and c41[("00_README", "A3")].replace("V4.1", "V4.2").endswith(a3[a3.index("| Status"):])
    ch = wb42["15_CHANGELOG"]
    ch_ok = ch["A22"].value == "2026-10" and ch["B22"].value.startswith("V4.2, ") and "EDR-062" in ch["C22"].value \
        and "EDR-061" in ch["C22"].value and "23_LISTAS" in ch["B22"].value \
        and "docs/registro-de-versoes-V4_2-celulas.md" in ch["C22"].value and all(ch[f"{c}23"].value is None for c in "ABC") \
        and all(list(ch[f"{c}22"]._style) == list(ch[f"{c}21"]._style) for c in "ABC")
    # registro do gerador: antes = V4.1, novo = V4.2, mesmo conjunto
    reg = g42.alteracoes_v4_2(wb41)
    reg_set = {(a["aba"], a["cel"]) for a in reg}
    reg_ok = reg_set == ESPERADO_ALT and len(reg) == ESPERADO_ALT_N and all(a["ref"] for a in reg) \
        and all(c41.get((a["aba"], a["cel"])) == a["antes"] and c42.get((a["aba"], a["cel"])) == a["novo"] for a in reg)
    metas_ok = all((cab_ok, linhas_ok, ini_ok, notas_ok, listas_ok, rel_cab_ok, rel_tab_ok, nota_ok, sig_ok, fontes_ok,
                    pega_l, pega_e, pega_m, ent_ok, fam_ok, l_ok, dash_ok))
    ok_001 = abas_ok and n_form and lit_ok and conj_ok and dado_ok and varredura_ok and sem_antigo and lista_ok \
        and nomes_ok and est and h7_ok and a3_ok and ch_ok and reg_ok and metas_ok
    check("TC-WB42-001: estrutura da V4.2 gerada da tag v4.1: abas e ordem (+23_LISTAS); 117569 fórmulas da V4.1 iguais fora da "
          "16_EQUILIBRIO B188/B189 (só o literal \"● Dado do T11\") e as 58 novas só nas famílias G1-G6 (cada uma na forma aprovada); "
          "células existentes alteradas = o conjunto aprovado; 32 células de dado com \"outro carro\"; nenhum \"Dado do TX\"/\"Dado do "
          "T09\" fora do 04 e do 15; lista Origens inalterada de nome; 1146 nomes da V4.1 iguais e os novos = data/nomes_V4_2.csv "
          "(NOMES_V42) + 4 listas; validações nas entradas com lista; 03_TARGETS (H:Q, linhas 10-11, A14/A15), 23_LISTAS (listas, "
          "tabela de relações, fontes L/E/M conferidas) e bloco das metas no 14_DASHBOARD; emenda da h7; A3; changelog linha 22",
          ok_001,
          f"abas={abas_ok} fórmulas={n_form} ({len(f41)}/{len(f42)}) literal={lit_ok} difs_fórmula={form_difs} "
          f"conjunto={conj_ok} (fora da região nova: {len(fora)}; fora do esperado {sorted(fora - ESPERADO_ALT)[:3]}, "
          f"faltou {sorted(ESPERADO_ALT - fora)[:3]}) dado={dado_ok} varredura={varredura_ok} sobra={sobra[:3]} "
          f"(V4.1 tinha {len(sobra41)}) lista={lista_ok} nomes={nomes_ok} ({len(n42)}) estrutura={est} h7={h7_ok} A3={a3_ok} "
          f"changelog={ch_ok} registro={reg_ok} | metas: cab={cab_ok} linhas={linhas_ok} ini={ini_ok} notas={notas_ok} "
          f"listas={listas_ok} relações={rel_cab_ok}/{rel_tab_ok} nota23={nota_ok} sig={sig_ok} fontes={fontes_ok}{fontes_det} "
          f"pega={pega_l}/{pega_e}/{pega_m} entradas={ent_ok} famílias={fam_ok} L={l_ok} dashboard={dash_ok}")
    wb41.close()
    wb42.close()

    # ---------------- TC-WB42-002 ----------------
    regen1, regen2 = tmp / "V4_2_regen1.xlsx", tmp / "V4_2_regen2.xlsx"
    g42.gerar(None, regen1)
    g42.gerar(v41, regen2, tmp / "registro.md")
    ok_dup, rel_dup = cp.comparar_xlsx(regen1, regen2, cp.FORMULAS_ESPERADAS_V42, estrutura=True)
    ok_dup2, _ = cp.comparar_xlsx(v42, regen1, cp.FORMULAS_ESPERADAS_V42, estrutura=True)
    entregue = cp.ENTREGUE_XLSX_V42
    if entregue.exists():
        ok_ent, rel_ent = cp.comparar_xlsx(regen1, entregue, cp.FORMULAS_ESPERADAS_V42, estrutura=True)
    else:
        ok_ent, rel_ent = True, [f"excel/…V4_2.xlsx ainda não existe (vai a excel/ na Task 7): comparação adiada"]
    reg_md = (tmp / "registro.md").read_text(encoding="utf-8")
    reg_md_ok = reg_md.startswith("# Registro de células da V4.2") and reg_md.count("\n| ") >= 50 and "\r" not in reg_md
    fantasma = tmp / "nao_existe.xlsx"
    saida_excel = RAIZ / "excel" / "_teste_recusa_V4_2.xlsx"
    rec_excel = raises(lambda: g42.gerar(fantasma, saida_excel), SystemExit) and not saida_excel.exists()
    viva_antes = VIVA.stat().st_mtime_ns if VIVA.exists() else None
    rec_viva = raises(lambda: g42.gerar(fantasma, VIVA), SystemExit) and \
        (VIVA.stat().st_mtime_ns if VIVA.exists() else None) == viva_antes
    b41 = v41.read_bytes()
    rec_origem = raises(lambda: g42.gerar(v41, v41), SystemExit) and v41.read_bytes() == b41
    csv_ok = ler_tabela(CSV42) == list(aba_metas.NOMES_V42) and CSV42.read_bytes().endswith(b"\n") \
        and raises(lambda: g42.gerar(v41, tmp / "csv_ruim.xlsx", nomes=list(aba_metas.NOMES_V42)[:-1]), ValueError) \
        and not (tmp / "csv_ruim.xlsx").exists()
    check("TC-WB42-002: CA-07 da V4.2: duas regenerações (da tag e do arquivo extraído) iguais, com estrutura "
          f"({cp.FORMULAS_ESPERADAS_V42} fórmulas); igual a excel/…V4_2.xlsx se existir; registro .md gravado; recusa saída em "
          "excel/ (e a viva) e saída = origem sem abrir nem criar arquivo; data/nomes_V4_2.csv = aba_metas.NOMES_V42 e o gerador recusa um CSV diferente",
          ok_dup and ok_dup2 and ok_ent and reg_md_ok and rec_excel and rec_viva and rec_origem and csv_ok,
          f"regen={ok_dup}/{ok_dup2} entregue={ok_ent} registro={reg_md_ok} recusas={rec_excel}/{rec_viva}/{rec_origem} csv={csv_ok} | "
          + " ; ".join(r for r in rel_ent if not r.startswith("  DIF"))[:300]
          + ("" if ok_dup else " | " + " ; ".join(rel_dup[:6])))

    # ---------------- TC-WB42-004 ----------------
    tc_wb42_004(tmp, v41, v42)


# ---------------- TC-WB42-003 (oráculo metas.py; sem planilha) ----------------
PO = "○ Estimativa"


def F(target, **kw):
    """Ficha de teste: valores iniciais da decisão 6 (Papel TBD, Target TBD, F0), com a pior origem de exemplo."""
    base = dict(target=target, valor="TBD", papel="TBD", ciclo="OPEN", congelamento="F0", compromisso=None, edr=None,
                pior_origem=PO)
    base.update(kw)
    return metas.Ficha(**base)


ALVOS = tuple(f"TARGET-00{i}" for i in range(1, 8))
ALAV_LAMBDA = "distribuição de rolagem λ_f (barras)"


def sete(**kw):
    return [F(t, **kw) for t in ALVOS]


def estado(f, fichas=None):
    return metas.estado_coerencia(f, fichas if fichas is not None else [f])


# vocabulários exatos e relações
vocab_ok = metas.VOCAB == {
    "papel": ("META", "CONSEQUÊNCIA", "TBD"),
    "ciclo": ("TBD", "OPEN", "PROPOSED", "APPROVED", "FROZEN", "STALE"),
    "congelamento": ("F0", "F1", "F2", "F3", "F4"),
    "coerencia": ("COERENTE", "COMPROMISSO DECLARADO", "CONFLITO", "NÃO AVALIÁVEL"),
    "fonte": ("L", "E", "M")}
alav = {t: metas.alavanca(t) for t in ALVOS}
alav_ok = alav == {"TARGET-001": "mola dianteira", "TARGET-002": "mola traseira", "TARGET-003": ALAV_LAMBDA,
                   "TARGET-004": "geometria dianteira (anti-dive)", "TARGET-005": "geometria traseira (anti-squat)",
                   "TARGET-006": ALAV_LAMBDA, "TARGET-007": ALAV_LAMBDA} and metas.alavanca("TARGET-009") is None \
    and metas.alavanca("TARGET-006 × TARGET-007") is None
rel_m = [r for r in metas.RELACOES if r.tipo == "alavanca" and r.fonte_tipo == "M"]
rel_ok = sorted(r.meta for r in rel_m) == list(ALVOS) and all(r.resultado and r.entrada for r in rel_m) \
    and [r.n for r in metas.RELACOES] == list(range(1, len(metas.RELACOES) + 1)) \
    and all(r.fonte_tipo in metas.VOCAB["fonte"] and r.fonte for r in metas.RELACOES) \
    and sum(r.tipo == "regime" for r in metas.RELACOES) == 1 \
    and all(r.resultado is None and r.entrada is None for r in metas.RELACOES if r.fonte_tipo != "M") \
    and set(metas.FONTES_L) >= {r.fonte for r in metas.RELACOES if r.fonte_tipo == "L"}
# campo_valido (exato: maiúsculas e minúsculas importam, como a fórmula EXACT da planilha)
cv = metas.campo_valido
campo_ok = cv("papel", "META") and cv("papel", "CONSEQUÊNCIA") and cv("papel", "TBD") and not cv("papel", "meta") \
    and not cv("papel", "META ") and not cv("papel", None) and not cv("papel", "") and not cv("papel", 5) \
    and cv("ciclo", "OPEN") and not cv("ciclo", "Open") and cv("ciclo", "STALE") and not cv("ciclo", "DRAFT") \
    and cv("congelamento", "F0") and cv("congelamento", "F4") and not cv("congelamento", "F5") \
    and not cv("congelamento", "f1") and raises(lambda: cv("campo_que_nao_existe", "x"), KeyError)
# formato do EDR
ef = metas.edr_no_formato
edr_ok = ef("EDR-061") and ef("EDR-001") and ef("EDR-999") and not any(
    ef(v) for v in ("EDR-61", "edr-061", "Edr-061", "EDR-0611", "EDR-06a", "EDR-1.5", "EDR-1e2", "EDR- 61", " EDR-061",
                    "EDR-061 ", "EDR-061\n", "EDR061", "", None, 61, 61.0))
# cada estado, forçado (a primeira regra que bate decide)
est_voc_p = estado(F("TARGET-001", papel="meta")) == ("CONFLITO", "Papel fora do vocabulário")
est_voc_c = estado(F("TARGET-001", ciclo="Open")) == ("CONFLITO", "Ciclo de vida fora do vocabulário")
est_voc_f = estado(F("TARGET-001", congelamento="F9")) == ("CONFLITO", "Congelamento fora do vocabulário")
est_ordem_voc = estado(F("TARGET-001", papel="x", ciclo="y", congelamento="z")) == ("CONFLITO", "Papel fora do vocabulário") \
    and estado(F("TARGET-001", ciclo="y", congelamento="z")) == ("CONFLITO", "Ciclo de vida fora do vocabulário")
duas = [F("TARGET-003", papel="META", valor=0.55), F("TARGET-006", papel="META", valor=2.0)]
est_sobre = estado(duas[0], duas) == ("CONFLITO", f"sobredeterminação ({ALAV_LAMBDA}; escolha qual manda)") \
    and estado(duas[1], duas) == estado(duas[0], duas)
# a regra 1 vem antes da 2: vocabulário inválido na ficha ganha da sobredeterminação (e "meta" minúsculo não é META)
f_voc = F("TARGET-003", papel="META", ciclo="Open")
est_voc_antes = estado(f_voc, [f_voc, F("TARGET-006", papel="META")])[1] == "Ciclo de vida fora do vocabulário"
meta_minuscula = metas.sobredeterminadas([F("TARGET-003", papel="META"), F("TARGET-006", papel="meta")]) == set()
# a regra 2 vem antes da 3: sobredeterminação mesmo com Target TBD
f_sob = F("TARGET-003", papel="META")
est_sobre_antes = estado(f_sob, [f_sob, F("TARGET-007", papel="META")])[0] == "CONFLITO"
# regra 3: NÃO AVALIÁVEL
nao_av = ("NÃO AVALIÁVEL", "meta sem valor ou papel (TBD)")
est_nao_av = estado(F("TARGET-001")) == nao_av and estado(F("TARGET-001", papel="META", valor="TBD")) == nao_av \
    and estado(F("TARGET-001", papel="TBD", valor=2.5)) == nao_av \
    and all(estado(F("TARGET-001", papel="META", valor=v)) == nao_av for v in (None, "", "   ", "tbd", " TBD "))
# regra 4: compromisso com e sem EDR; a regra 3 vem antes
ok4 = F("TARGET-007", papel="META", valor="PUSH", compromisso="push por preferência do piloto", edr="EDR-061")
est_comp = estado(ok4) == ("COMPROMISSO DECLARADO", "") and metas.texto_coerencia(ok4, [ok4]) == f"COMPROMISSO DECLARADO · {PO}"
est_comp_sem = all(estado(F("TARGET-007", papel="META", valor="PUSH", compromisso="x", edr=e)) ==
                   ("CONFLITO", "compromisso sem EDR") for e in (None, "", "EDR-61", "edr-061", "EDR-999999", 61))
est_comp_tbd = estado(F("TARGET-007", papel="META", valor="TBD", compromisso="x", edr="EDR-061")) == nao_av
# compromisso vazio ou "TBD" não conta como declarado
est_comp_vazio = all(estado(F("TARGET-007", papel="META", valor="PUSH", compromisso=c))[0] == "COERENTE"
                     for c in (None, "", "  ", "TBD", "tbd"))
# COERENTE e os textos finais
coer = F("TARGET-001", papel="CONSEQUÊNCIA", valor=2.5)
est_coer = estado(coer) == ("COERENTE", "regras avaliadas: vocabulário e sobredeterminação") and \
    metas.texto_coerencia(coer, [coer]) == f"COERENTE (regras avaliadas: vocabulário e sobredeterminação) · {PO}" and \
    estado(F("TARGET-001", papel="META", valor=0))[0] == "COERENTE"
f_inv = F("TARGET-001", papel="meta")
textos_ok = metas.texto_coerencia(F("TARGET-001"), [F("TARGET-001")]) == f"NÃO AVALIÁVEL: meta sem valor ou papel (TBD) · {PO}" \
    and metas.texto_coerencia(f_inv, [f_inv]) == f"CONFLITO: Papel fora do vocabulário · {PO}" \
    and metas.texto_coerencia(F("TARGET-001", pior_origem="? A confirmar"), [F("TARGET-001")]).endswith(" · ? A confirmar")
# os quatro estados aparecem, todos do vocabulário
todos4 = {estado(f)[0] for f in (F("TARGET-001", papel="meta"), F("TARGET-001"), ok4, coer)}
quatro_ok = todos4 == set(metas.VOCAB["coerencia"])
# Review Focus 5: Papel META nas 7 fichas de uma vez
sete_meta = sete(papel="META")
sobre7 = metas.sobredeterminadas(sete_meta)
est7 = {f.target: metas.estado_coerencia(f, sete_meta)[0] for f in sete_meta}
txt7 = [metas.texto_coerencia(f, sete_meta) for f in sete_meta]
rf5_ok = sobre7 == {"TARGET-003", "TARGET-006", "TARGET-007"} and est7 == {
    "TARGET-001": "NÃO AVALIÁVEL", "TARGET-002": "NÃO AVALIÁVEL", "TARGET-003": "CONFLITO", "TARGET-004": "NÃO AVALIÁVEL",
    "TARGET-005": "NÃO AVALIÁVEL", "TARGET-006": "CONFLITO", "TARGET-007": "CONFLITO"} \
    and metas.pior_coerencia(txt7) == "CONFLITO" and metas.compromissos(txt7) == 0
# com valores preenchidos, as 4 sem sobredeterminação ficam COERENTE
sete_val = [F(t, papel="META", valor=1.0) for t in ALVOS]
rf5b_ok = {f.target: metas.estado_coerencia(f, sete_val)[0] for f in sete_val} == {
    "TARGET-001": "COERENTE", "TARGET-002": "COERENTE", "TARGET-003": "CONFLITO", "TARGET-004": "COERENTE",
    "TARGET-005": "COERENTE", "TARGET-006": "CONFLITO", "TARGET-007": "CONFLITO"}
# uma só META em λ_f: sem sobredeterminação
uma = [F("TARGET-003", papel="META"), F("TARGET-006", papel="CONSEQUÊNCIA"), F("TARGET-007", papel="CONSEQUÊNCIA")]
uma_ok = metas.sobredeterminadas(uma) == set()
# pior coerência e contador
C_, D_, X_, N_ = (f"COERENTE (regras avaliadas: vocabulário e sobredeterminação) · {PO}", f"COMPROMISSO DECLARADO · {PO}",
                  f"CONFLITO: Papel fora do vocabulário · {PO}", f"NÃO AVALIÁVEL: meta sem valor ou papel (TBD) · {PO}")
pior_ok = metas.pior_coerencia([C_, C_]) == "COERENTE" and metas.pior_coerencia([C_, D_]) == "COMPROMISSO DECLARADO" \
    and metas.pior_coerencia([C_, D_, N_]) == "NÃO AVALIÁVEL" and metas.pior_coerencia([C_, D_, N_, X_]) == "CONFLITO" \
    and metas.pior_coerencia([X_, C_]) == "CONFLITO" and metas.pior_coerencia([]) == "NÃO AVALIÁVEL" \
    and raises(lambda: metas.pior_coerencia(["OK"]), ValueError)
cont_ok = metas.compromissos([C_, D_, D_, N_]) == 2 and metas.compromissos([]) == 0 and metas.compromissos([C_, X_]) == 0
check("TC-WB42-003: oráculo metas.py: vocabulários e relações; campo_valido exato; formato do EDR; cada estado de coerência forçado "
      "(vocabulário, sobredeterminação, NÃO AVALIÁVEL, compromisso com e sem EDR, COERENTE) na ordem das regras; textos finais; "
      "7 METAs ao mesmo tempo (003, 006 e 007 em CONFLITO; Review Focus 5); pior coerência e contador",
      all((vocab_ok, alav_ok, rel_ok, campo_ok, edr_ok, est_voc_p, est_voc_c, est_voc_f, est_ordem_voc, est_sobre, est_voc_antes,
           meta_minuscula, est_sobre_antes, est_nao_av, est_comp, est_comp_sem, est_comp_tbd, est_comp_vazio, est_coer,
           textos_ok, quatro_ok, rf5_ok, rf5b_ok, uma_ok, pior_ok, cont_ok)),
      f"vocab={vocab_ok} alavanca={alav_ok} relações={rel_ok} campo={campo_ok} edr={edr_ok} voc={est_voc_p}/{est_voc_c}/{est_voc_f}/"
      f"{est_ordem_voc}/{est_voc_antes} sobre={est_sobre}/{meta_minuscula}/{est_sobre_antes} nãoav={est_nao_av} comp={est_comp}/"
      f"{est_comp_sem}/{est_comp_tbd}/{est_comp_vazio} coer={est_coer} textos={textos_ok} quatro={quatro_ok} rf5={rf5_ok}/{rf5b_ok}/"
      f"{uma_ok} pior={pior_ok} cont={cont_ok}")

print("\n--- Resumo ---")
n_pass = sum(1 for _, s, _ in results if s == "PASS")
n_fail = sum(1 for _, s, _ in results if s == "FAIL")
print(f"{n_pass} PASS, {n_fail} FAIL, de {len(results)} testes")
if n_fail > 0:
    raise SystemExit(1)
