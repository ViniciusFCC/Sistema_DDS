"""
Testes do módulo equilibrio (análise de par, Milliken §7.1; V4.0, incremento 1).
TC-EQ-001..003 (IDs autorizados pelo usuário): caso, cargas e camber por roda.
TC-EQ-004..008 (idem): curva do eixo, capacidade, borda da grade, faixa do pneu e sinal do toe.
TC-EQ-009..015 (idem): análise de par (força exigida, saturação, veredito e C_N, Milliken §7.1 e §8.4).
TC-EQ-016..018 (idem): robustez do veredito ao pneu, nomes de fator e texto do veredito (spec 7.3 e 7.6).
TC-EQ-019..020 (idem): caso lido da planilha recalculada e pior origem (leitura apenas, sem salvar a planilha).
TC-EQ-039..043 (idem; V4.0, inc. 2): Ackermann e cos δ; CG lateral, dois sentidos de curva e aero com a velocidade.
TC-EQ-044..046 (idem; V4.0, inc. 2): tração na traseira, elipse de atrito e diferencial aberto (spec 7.4, h9 e h10).
Roda como script: uv run python -X utf8 tests/test_equilibrio.py
O veículo e os números de camber abaixo são entradas de teste, não do carro.
"""

import copy
import dataclasses
import math
import os

import equilibrio as eq
import origens
import suspension_model as sm
import tire_model as tm
from suspension_model import VehicleConfig, G

TOL = 1e-9
HERE = os.path.dirname(os.path.abspath(__file__))
TIR = os.path.join(HERE, "..", "data", "Hoosier_R25B_13x7_TTC_provisional.tir")
results = []


def raises(fn, exc=ValueError, needle=None):
    try:
        fn()
        return False
    except exc as e:
        return needle is None or needle in str(e)


def check(name, condition, detail=""):
    status = "PASS" if condition else "FAIL"
    results.append((name, status, detail))
    print(f"[{status}] {name}  {detail}")


def veiculo_base(**over) -> VehicleConfig:
    """Veículo de teste (números plausíveis, NÃO é setup real)."""
    base = dict(
        mass_total=250.0,
        cg_height=0.300,
        wheelbase=1.550,
        track_front=1.200,
        track_rear=1.150,
        cg_to_front=0.837,
        cg_to_rear=0.713,
        unsprung_mass_front=15.0,
        unsprung_mass_rear=18.0,
        unsprung_cg_height_front=0.230,
        unsprung_cg_height_rear=0.230,
        roll_center_front=0.030,
        roll_center_rear=0.050,
        roll_stiffness_front=20826.4,
        roll_stiffness_rear=26450.0,
    )
    base.update(over)
    return VehicleConfig(**base)


cfg = veiculo_base()
ds = tm.load_dataset(TIR, source_axis_system="ISO_TYDEX")

# ---------------- TC-EQ-001 : a_y = 0 ----------------
caso = eq.CasoEquilibrio(cfg=cfg, ds=ds, downforce_N=400.0, aero_balance_front=0.45)
c0 = eq.cargas(caso, 0.0)
est = sm.static_fz_corner(cfg)
aer = sm.aero_fz_corner(400.0, 0.45)
dif = max(abs(c0[c] - (est[c] + aer[c])) for c in eq.CANTOS)
check("TC-EQ-001: a_y=0 -> Fz = estático + aero em cada canto", dif < TOL, f"max dif={dif:.2e} N")

# ---------------- TC-EQ-002 : a_y = 1,2 g ----------------
c12 = eq.cargas(caso, 1.2 * G)
soma = sum(c12.values())
check("TC-EQ-002: a_y=1,2g -> soma = m·G + DF; FR>FL e RR>RL",
      abs(soma - (cfg.mass_total * G + 400.0)) < TOL and c12["FR"] > c12["FL"] and c12["RR"] > c12["RL"],
      f"soma={soma:.6f} FL={c12['FL']:.1f} FR={c12['FR']:.1f} RL={c12['RL']:.1f} RR={c12['RR']:.1f}")

# ---------------- TC-EQ-003 : camber por roda e validação ----------------
estatico = {"FL": -0.01, "FR": -0.01, "RL": -0.02, "RR": -0.02}
ganho = {"FL": 0.5, "FR": -0.5, "RL": 0.3, "RR": -0.3}
caso_c = eq.CasoEquilibrio(cfg=cfg, ds=ds, camber_estatico_rad=estatico, ganho_camber=ganho)
phi = sm.lateral_dfz(cfg, 1.0 * G)["roll_angle_rad"]
gam = eq.camber(caso_c, 1.0 * G)
dif = max(abs(gam[c] - (estatico[c] + ganho[c] * phi)) for c in eq.CANTOS)
erro_toe = raises(lambda: eq.CasoEquilibrio(cfg=cfg, ds=ds, toe_rad={"FL": 0.0}), ValueError, "toe_rad")
check("TC-EQ-003: camber = estático + ganho·φ (φ de lateral_dfz); toe_rad com cantos faltando -> ValueError",
      dif < 1e-15 and phi != 0.0 and erro_toe, f"phi={phi:.6f} rad, max dif={dif:.2e}")

# ---------------- TC-EQ-004 : fy_eixo = soma das duas rodas ----------------
caso0 = eq.CasoEquilibrio(cfg=cfg, ds=ds)
fz0 = eq.cargas(caso0, 0.0)
dif = 0.0
for a in eq.grade_alfa(caso0, 25):
    esperado = (tm.wheel_forces(ds, "FL", fz0["FL"], a, 0.0, 0.0)[1]
                + tm.wheel_forces(ds, "FR", fz0["FR"], a, 0.0, 0.0)[1])
    dif = max(dif, abs(eq.fy_eixo(caso0, "F", 0.0, a) - esperado))
check("TC-EQ-004: fy_eixo = Fy(FL) + Fy(FR) em cada alfa da grade (a_y=0, toe e camber 0)",
      dif < TOL, f"max dif={dif:.2e} N")

# ---------------- TC-EQ-005 : máximo numérico = forma fechada |Dy| + SVy ----------------
fnom = ds.p["FNOMIN"]
info = tm.pure_fy0(ds.p, fnom, 0.0, 0.0)[1]
fechada = abs(info["Dy"]) + info["SVy"]
N = 4001
alfas = [i * ds.p["ALPMIN"] / (N - 1) for i in range(N)]
fys = [tm.wheel_forces(ds, "FR", fnom, a, 0.0, 0.0)[1] for a in alfas]
fmax = max(fys)
i_pico = fys.index(fmax)
rel = abs(fmax - fechada) / fechada
check("TC-EQ-005: max numérico de Fy (FR, FNOMIN, gamma=0) = |Dy|+SVy (erro rel < 1e-4); pico fora da borda",
      rel < 1e-4 and i_pico < N - 1,
      f"num={fmax:.3f} N, fechada={fechada:.3f} N, rel={rel:.2e}, idx={i_pico}, "
      f"alfa={math.degrees(alfas[i_pico]):.2f} deg")

# ---------------- TC-EQ-006 : pico na borda da grade ----------------
# Dataset de teste (PKY1 x 0,05): pneu pouco rígido, pico além da faixa de alfa. Não é pneu.
ds_mole = dataclasses.replace(ds, p={**ds.p, "PKY1": ds.p["PKY1"] * 0.05})
caso_mole = eq.CasoEquilibrio(cfg=cfg, ds=ds_mole)
cap_mole = eq.capacidade_eixo(caso_mole, "F", 0.0, 200)
check("TC-EQ-006: pneu pouco rígido -> motivo 'pico na borda da grade de α' (valor mantido)",
      cap_mole.motivo == "pico na borda da grade de α" and cap_mole.fy_max_N is not None,
      f"motivo={cap_mole.motivo!r}, fy_max={cap_mole.fy_max_N}")

# ---------------- TC-EQ-007 : faixa do pneu (EDR-056; help do Adams PAC2002, p. 83) ----------------
# Fz > 0 e γ fora da faixa entram limitados ao FZMIN/FZMAX ou CAMMIN/CAMMAX, com aviso, e a capacidade
# é válida; Fz <= 0 é descolamento (fora do modelo de transferência de carga, SUS-104) e fica None;
# toe que consome a faixa de α continua None. Nenhum caso pode lançar exceção.
def _cap_sem_excecao(caso_x, ay, eixo="F"):
    try:
        return eq.capacidade_eixo(caso_x, eixo, ay, 25)
    except Exception:  # o check precisa registrar qualquer exceção como falha
        return None


def _cap_ref(caso_x, eixo, fz_usado, gam_usado):
    """Capacidade recomposta: máximo sobre grade_alfa com Fz e γ dados por roda (κ = 0)."""
    dal = eq.deriva_toe(caso_x)
    rodas = ("FL", "FR") if eixo == "F" else ("RL", "RR")
    return max(sum(tm.wheel_forces(ds, c, fz_usado[c], a + dal[c], gam_usado[c], 0.0)[1] for c in rodas)
               for a in eq.grade_alfa(caso_x, 25))


def _limites(cap):
    return getattr(cap, "limites", ()) or ()


def _v(x, fmt):
    """Número com vírgula decimal, como nos motivos e avisos de equilibrio (os mesmos números do formato)."""
    return format(x, fmt).replace(".", ",")


# 2,0 g no veículo de teste: 0 < Fz_FL < FZMIN (limitado) e Fz_RL <= 0 (descolamento)
fz2 = eq.cargas(caso0, 2.0 * G)
cap_f2 = _cap_sem_excecao(caso0, 2.0 * G, "F")
cap_r2 = _cap_sem_excecao(caso0, 2.0 * G, "R")
ref_f2 = _cap_ref(caso0, "F", {"FL": ds.p["FZMIN"], "FR": fz2["FR"]}, {"FL": 0.0, "FR": 0.0})
aviso_fl = f"Fz_FL = {_v(fz2['FL'], '.1f')} N → {_v(ds.p['FZMIN'], '.1f')} N (FZMIN)"
ok_fzmin = (0.0 < fz2["FL"] < ds.p["FZMIN"] and ds.p["FZMIN"] <= fz2["FR"] <= ds.p["FZMAX"]
            and cap_f2 is not None and cap_f2.fy_max_N is not None and cap_f2.motivo is None
            and aviso_fl in _limites(cap_f2) and abs(cap_f2.fy_max_N - ref_f2) < TOL)
ok_desc = (fz2["RL"] <= 0.0 and cap_r2 is not None and cap_r2.fy_max_N is None and cap_r2.alfa_pico_rad is None
           and (cap_r2.motivo or "").startswith("Fz_RL = ") and "descolamento" in cap_r2.motivo)
# camber estático além de CAMMAX numa roda dianteira (a_y = 0): avaliado com γ = CAMMAX
cam_fora = ds.p["CAMMAX"] + math.radians(1.0)
caso_cam = eq.CasoEquilibrio(cfg=cfg, ds=ds, camber_estatico_rad={"FL": 0.0, "FR": cam_fora, "RL": 0.0, "RR": 0.0})
cap_cam = _cap_sem_excecao(caso_cam, 0.0)
ref_cam = _cap_ref(caso_cam, "F", fz0, {"FL": 0.0, "FR": ds.p["CAMMAX"]})
ok_cam = (cap_cam is not None and cap_cam.fy_max_N is not None and cap_cam.motivo is None
          and any(s.startswith("γ_FR = ") and "(CAMMAX)" in s for s in _limites(cap_cam))
          and abs(cap_cam.fy_max_N - ref_cam) < TOL)
# downforce de teste (não é do carro) que leva Fz das rodas dianteiras acima de FZMAX em a_y = 0
caso_df = eq.CasoEquilibrio(cfg=cfg, ds=ds, downforce_N=4400.0, aero_balance_front=0.5)
fz_df = eq.cargas(caso_df, 0.0)
cap_df = _cap_sem_excecao(caso_df, 0.0)
ref_df = _cap_ref(caso_df, "F", {"FL": ds.p["FZMAX"], "FR": ds.p["FZMAX"]}, {"FL": 0.0, "FR": 0.0})
ok_fzmax = (min(fz_df["FL"], fz_df["FR"]) > ds.p["FZMAX"] and cap_df is not None and cap_df.fy_max_N is not None
            and cap_df.motivo is None and sum("(FZMAX)" in s for s in _limites(cap_df)) == 2
            and abs(cap_df.fy_max_N - ref_df) < TOL)
# toe maior que a faixa de α: continua None
toe_grande = 2.0 * min(abs(ds.p["ALPMIN"]), ds.p["ALPMAX"])
cap_toe = _cap_sem_excecao(eq.CasoEquilibrio(
    cfg=cfg, ds=ds, toe_rad={"FL": toe_grande, "FR": toe_grande, "RL": 0.0, "RR": 0.0}), 0.0)
ok_toe = cap_toe is not None and cap_toe.fy_max_N is None and cap_toe.motivo == "toe consome a faixa de α"
check("TC-EQ-007: faixa do pneu (EDR-056): Fz < FZMIN, Fz > FZMAX e γ > CAMMAX -> limitados, válidos, com aviso "
      "e iguais à recomposição; Fz <= 0 -> None 'descolamento'; toe que consome a faixa -> None; sem exceção",
      ok_fzmin and ok_desc and ok_cam and ok_fzmax and ok_toe,
      f"{ok_fzmin}/{ok_desc}/{ok_cam}/{ok_fzmax}/{ok_toe}; Fz_FL={fz2['FL']:.1f} N, Fz_RL={fz2['RL']:.1f} N, "
      f"Fz_F(DF)={fz_df['FL']:.1f} N; limites={_limites(cap_f2)}, {_limites(cap_cam)}, {_limites(cap_df)}; "
      f"motivos={getattr(cap_r2, 'motivo', None)!r}; {getattr(cap_toe, 'motivo', None)!r}")

# ---------------- TC-EQ-008 : sinal do toe (EDR-002) ----------------
um_grau = math.radians(1.0)
caso_toe = eq.CasoEquilibrio(cfg=cfg, ds=ds, toe_rad={"FL": um_grau, "FR": um_grau, "RL": 0.0, "RR": 0.0})
fzt = eq.cargas(caso_toe, 0.0)
dal = eq.deriva_toe(caso_toe)
fy_fl = tm.wheel_forces(ds, "FL", fzt["FL"], dal["FL"], 0.0, 0.0)[1]
fy_fr = tm.wheel_forces(ds, "FR", fzt["FR"], dal["FR"], 0.0, 0.0)[1]
fy_ax = eq.fy_eixo(caso_toe, "F", 0.0, 0.0)
check("TC-EQ-008: toe-out simétrico de 1 grau -> Fy(FL) > 0, Fy(FR) < 0, soma ~ 0",
      dal["FL"] < 0 < dal["FR"] and fy_fl > 0 and fy_fr < 0 and abs(fy_ax) < TOL,
      f"FL={fy_fl:.3f} N, FR={fy_fr:.3f} N, soma={fy_ax:.2e}")


# ---------------- Análise de par (TC-EQ-009..015) ----------------
def _par_sem_excecao(caso_x, **kw):
    try:
        return eq.analise_par(caso_x, **kw), None
    except Exception as e:  # o check precisa registrar qualquer exceção como falha
        return None, f"{type(e).__name__}: {e}"


def _folga_em(linhas, eixo, ay):
    """Folga cap − req do eixo ('f' ou 'r') interpolada linearmente em ay, entre as linhas vizinhas."""
    for la, lb in zip(linhas, linhas[1:]):
        if la.ay_ms2 <= ay <= lb.ay_ms2:
            ma = getattr(la, f"cap_{eixo}").fy_max_N - getattr(la, f"req_{eixo}")
            mb = getattr(lb, f"cap_{eixo}").fy_max_N - getattr(lb, f"req_{eixo}")
            return ma + (mb - ma) * (ay - la.ay_ms2) / (lb.ay_ms2 - la.ay_ms2)
    raise AssertionError(f"a_y={ay} fora das linhas")


def _resumo(r):
    if r is None:
        return "sem resultado"
    return (f"veredito={r.veredito}, sat_f={r.ay_sat_f}, sat_r={r.ay_sat_r}, "
            f"lim={r.ay_lim}, c_n={r.c_n}, motivo={r.motivo!r}")


# Dataset de teste: LMUY = 0,5. Escala só de teste, não é pneu ("Risco conhecido" do plano: com o
# .tir sem escala, a roda interna sai da faixa de Fz antes de o eixo saturar).
ds_teste = dataclasses.replace(ds, p={**ds.p, "LMUY": 0.5})
W = cfg.mass_total * G

# ---------------- TC-EQ-009 : força lateral exigida por eixo ----------------
cfg_ab = veiculo_base(cg_to_front=0.85, cg_to_rear=0.70)
caso_ab = eq.CasoEquilibrio(cfg=cfg_ab, ds=ds)
try:
    req = eq.forca_exigida(caso_ab, 9.81)
except Exception as e:  # o check precisa registrar qualquer exceção como falha
    req = (math.nan, math.nan)
esp = (cfg_ab.mass_total * 9.81 * 0.70 / cfg_ab.wheelbase, cfg_ab.mass_total * 9.81 * 0.85 / cfg_ab.wheelbase)
check("TC-EQ-009: forca_exigida = (m·a_y·b/L, m·a_y·a/L) com a != b",
      abs(req[0] - esp[0]) < TOL and abs(req[1] - esp[1]) < TOL,
      f"obtido={req}, esperado={esp}")

# ---------------- TC-EQ-010 : carro simétrico -> NEUTRO ----------------
k_soma = cfg.roll_stiffness_front + cfg.roll_stiffness_rear


def veiculo_simetrico(lambda_f=0.5, **over):
    L = cfg.wheelbase
    return veiculo_base(**{**dict(cg_to_front=L / 2, cg_to_rear=L / 2, track_rear=cfg.track_front,
                                  roll_stiffness_front=lambda_f * k_soma, roll_stiffness_rear=(1 - lambda_f) * k_soma,
                                  roll_center_rear=cfg.roll_center_front, unsprung_mass_rear=cfg.unsprung_mass_front,
                                  unsprung_cg_height_rear=cfg.unsprung_cg_height_front), **over})


r_sim, err_sim = _par_sem_excecao(eq.CasoEquilibrio(cfg=veiculo_simetrico(0.5), ds=ds_teste))
_res_sim = getattr(r_sim, "resolucao_ms2", None)
check("TC-EQ-010: carro simétrico (LMUY 0,5) -> as duas saturações existem, NEUTRO, |S_f − S_r| <= resolução "
      "numérica (EDR-056), c_n None",
      r_sim is not None and r_sim.ay_sat_f is not None and r_sim.ay_sat_r is not None
      and isinstance(_res_sim, float) and math.isfinite(_res_sim) and _res_sim > 0.0
      and r_sim.veredito == eq.NEUTRO and abs(r_sim.ay_sat_f - r_sim.ay_sat_r) <= _res_sim
      and r_sim.c_n is None,
      (err_sim or _resumo(r_sim)) + f", resolucao={_res_sim}")

# ---------------- TC-EQ-011 : distribuição de rolagem (Gillespie p. 214) ----------------
# LMUY = 0,45: escala só de teste, não é pneu. Com este pneu, o efeito de lambda_f 0,5 -> 0,7 é
# ~0,03-0,04 g, abaixo do passo padrão de 0,05 g; pelo EDR-056 (Milliken p. 320) o veredito não pode
# depender do passo de exibição: o mesmo veredito sai na grade padrão (0,05 g) e na de 0,01 g, e a
# resolução numérica medida fica abaixo de 0,01 g.
ds_teste_011 = dataclasses.replace(ds, p={**ds.p, "LMUY": 0.45})
PASSO_011 = 0.01
r_push, err_push = _par_sem_excecao(eq.CasoEquilibrio(cfg=veiculo_simetrico(0.7), ds=ds_teste_011),
                                    passo_g=PASSO_011)
r_loose, err_loose = _par_sem_excecao(eq.CasoEquilibrio(cfg=veiculo_simetrico(0.3), ds=ds_teste_011),
                                      passo_g=PASSO_011)
r_push_p, err_push_p = _par_sem_excecao(eq.CasoEquilibrio(cfg=veiculo_simetrico(0.7), ds=ds_teste_011))
r_loose_p, err_loose_p = _par_sem_excecao(eq.CasoEquilibrio(cfg=veiculo_simetrico(0.3), ds=ds_teste_011))


def _pre_saturacoes(r):
    """As duas saturações existem e nenhuma roda sai da faixa (nos dois eixos) antes delas."""
    if r is None or r.ay_sat_f is None or r.ay_sat_r is None:
        return False
    ate = max(r.ay_sat_f, r.ay_sat_r)
    return all(ln.cap_f.motivo is None and ln.cap_r.motivo is None for ln in r.linhas if ln.ay_ms2 <= ate)


def _sat_g(r):
    if r is None or r.ay_sat_f is None or r.ay_sat_r is None:
        return "sat indisponível"
    return f"sat_f={r.ay_sat_f / G:.4f} g, sat_r={r.ay_sat_r / G:.4f} g"


def _res_ok(r):
    res = getattr(r, "resolucao_ms2", None)
    return isinstance(res, float) and 0.0 < res < 0.01 * G


pre_011 = _pre_saturacoes(r_push) and _pre_saturacoes(r_loose)
pre_011p = _pre_saturacoes(r_push_p) and _pre_saturacoes(r_loose_p)
check("TC-EQ-011: lambda_f = 0,7 -> PUSH; lambda_f = 0,3 -> LOOSE, o mesmo na grade padrão (0,05 g) e na de "
      "0,01 g, com resolução numérica < 0,01 g (LMUY 0,45; pré-condição: duas saturações, nenhuma capacidade "
      "inválida antes delas)",
      pre_011 and pre_011p and all(_res_ok(r) for r in (r_push, r_loose, r_push_p, r_loose_p))
      and r_push.veredito == eq.PUSH and r_push_p.veredito == eq.PUSH
      and r_loose.veredito == eq.LOOSE and r_loose_p.veredito == eq.LOOSE,
      f"0,7 @0,01: {_sat_g(r_push)}; {err_push or _resumo(r_push)}; res={getattr(r_push, 'resolucao_ms2', None)} | "
      f"0,7 @0,05: {_sat_g(r_push_p)}; {err_push_p or _resumo(r_push_p)}; "
      f"res={getattr(r_push_p, 'resolucao_ms2', None)} | "
      f"0,3 @0,01: {_sat_g(r_loose)}; {err_loose or _resumo(r_loose)}; "
      f"res={getattr(r_loose, 'resolucao_ms2', None)} | "
      f"0,3 @0,05: {_sat_g(r_loose_p)}; {err_loose_p or _resumo(r_loose_p)}; "
      f"res={getattr(r_loose_p, 'resolucao_ms2', None)}")

# ---------------- TC-EQ-012 : sinal e recomposição de C_N ----------------
ok_012, det_012 = False, "pré-condição do TC-EQ-011 falhou"
if pre_011 and r_push.veredito == eq.PUSH and r_loose.veredito == eq.LOOSE:
    cs = veiculo_simetrico(0.7)
    a_s, b_s, L_s = cs.cg_to_front, cs.cg_to_rear, cs.wheelbase
    cn_push = _folga_em(r_push.linhas, "r", r_push.ay_lim) * b_s / (W * L_s)
    cn_loose = -_folga_em(r_loose.linhas, "f", r_loose.ay_lim) * a_s / (W * L_s)
    ok_012 = (r_push.c_n > 0 and r_loose.c_n < 0
              and abs(r_push.c_n - cn_push) < 1e-12 and abs(r_loose.c_n - cn_loose) < 1e-12)
    det_012 = (f"PUSH c_n={r_push.c_n:.6f} (recomp. {cn_push:.6f}); "
               f"LOOSE c_n={r_loose.c_n:.6f} (recomp. {cn_loose:.6f})")
check("TC-EQ-012: C_N > 0 em PUSH, < 0 em LOOSE e igual à recomposição pelas linhas", ok_012, det_012)

# ---------------- TC-EQ-013 : ramos de validade do veredito (EDR-056) ----------------
# Pneu sem escala: a roda interna descola (Fz <= 0) antes de os eixos saturarem -> N/A por descolamento,
# ou, se um eixo saturar, o outro segue válido além de S + resolução.
r_base, err_base = _par_sem_excecao(caso0)
ok_013, det_013 = False, err_base or _resumo(r_base)
if r_base is not None:
    if r_base.veredito == eq.NA:
        # sem a_y_lim, limitacoes junta as capacidades válidas de todas as linhas: um aviso por roda e
        # limite, o de maior excesso (aqui, o menor Fz > 0 abaixo de FZMIN), recomposto pelas linhas
        esp_lim = []
        for roda, eixo in (("FL", "f"), ("FR", "f"), ("RL", "r"), ("RR", "r")):
            fzs = [ln.fz[roda] for ln in r_base.linhas
                   if getattr(ln, f"cap_{eixo}").motivo is None and ln.fz[roda] < ds.p["FZMIN"]]
            if fzs:
                esp_lim.append(f"Fz_{roda} = {_v(min(fzs), '.1f')} N → {_v(ds.p['FZMIN'], '.1f')} N (FZMIN)")
        # limitacoes_detalhe (API pública): o mesmo pior caso de cada roda e limite, na mesma ordem
        det_lim = eq.limitacoes_detalhe(r_base, caso0)
        ok_013 = ("descolamento" in (r_base.motivo or "") and r_base.c_n is None
                  and len(esp_lim) >= 2 and r_base.limitacoes == tuple(esp_lim)
                  and [d["pior"].texto for d in det_lim] == esp_lim
                  and all(d["nome"] == "FZMIN" and d["primeira_ay"] <= d["pior_ay"] for d in det_lim))
        ramo = "N/A"
    elif r_base.veredito in (eq.PUSH, eq.LOOSE):
        nao_sat = "r" if r_base.veredito == eq.PUSH else "f"
        validas = [ln.ay_ms2 for ln in r_base.linhas if getattr(ln, f"cap_{nao_sat}").motivo is None]
        ok_013 = (r_base.ay_lim is not None and r_base.resolucao_ms2 is not None
                  and any(a >= r_base.ay_lim + r_base.resolucao_ms2 for a in validas))
        ramo = r_base.veredito
    else:
        ramo = r_base.veredito
    # nos dois ramos: as saturações conhecidas caem dentro da parte válida da grade de cada eixo
    for eixo, sat in (("f", r_base.ay_sat_f), ("r", r_base.ay_sat_r)):
        if sat is not None:
            ate = [ln for ln in r_base.linhas if ln.ay_ms2 <= sat + 1e-12]
            seguinte = [ln for ln in r_base.linhas if ln.ay_ms2 > sat][:1]
            ok_013 = ok_013 and all(getattr(ln, f"cap_{eixo}").motivo is None for ln in ate + seguinte)
    det_013 = f"ramo={ramo}; {_resumo(r_base)}; limitacoes={r_base.limitacoes}"

# Sub-casos determinísticos dos ramos de validade do veredito (A: eixo não saturado segue válido além de
# S + resolução -> PUSH/LOOSE com c_n; B: a roda interna do eixo não saturado descola (Fz <= 0) antes de S
# do outro -> N/A com o motivo dele). Pelo EDR-056, Fz e γ fora da faixa do pneu não invalidam mais a
# linha (entram limitados); o descolamento continua invalidando. Escala LMUY 0,45 só de teste, não é
# pneu; passo 0,01 g.
ds_013 = dataclasses.replace(ds, p={**ds.p, "LMUY": 0.45})
L_013 = cfg.wheelbase


def _cn_recomp(r, eixo_folga):
    folga = _folga_em(r.linhas, eixo_folga, r.ay_lim)
    # carro simétrico: a = b = L/2, então o braço é L/2 nos dois sinais
    return folga * (L_013 / 2) / (W * L_013) * (1 if eixo_folga == "r" else -1)


def _caso_013(lf, **over):
    return eq.CasoEquilibrio(cfg=veiculo_simetrico(lf, **over), ds=ds_013)


# Ramo A: o eixo traseiro (PUSH) ou dianteiro (LOOSE) não satura, mas a grade o mantém válido além de S.
rA_p, eA_p = _par_sem_excecao(_caso_013(0.7), passo_g=0.01, ay_max_g=1.17)
rA_l, eA_l = _par_sem_excecao(_caso_013(0.3), passo_g=0.01, ay_max_g=1.17)
okA = False
if rA_p is not None and rA_l is not None:
    pre_p = (rA_p.ay_sat_f is not None and rA_p.ay_sat_r is None and rA_p.resolucao_ms2 is not None
             and all(ln.cap_r.motivo is None for ln in rA_p.linhas)
             and rA_p.linhas[-1].ay_ms2 >= rA_p.ay_sat_f + rA_p.resolucao_ms2)
    pre_l = (rA_l.ay_sat_r is not None and rA_l.ay_sat_f is None and rA_l.resolucao_ms2 is not None
             and all(ln.cap_f.motivo is None for ln in rA_l.linhas)
             and rA_l.linhas[-1].ay_ms2 >= rA_l.ay_sat_r + rA_l.resolucao_ms2)
    okA = (pre_p and pre_l
           and rA_p.veredito == eq.PUSH and rA_p.motivo is None and rA_p.c_n is not None and rA_p.c_n > 0
           and abs(rA_p.c_n - _cn_recomp(rA_p, "r")) < 1e-12
           and rA_l.veredito == eq.LOOSE and rA_l.motivo is None and rA_l.c_n is not None and rA_l.c_n < 0
           and abs(rA_l.c_n - _cn_recomp(rA_l, "f")) < 1e-12)

# Ramo B: CG mais alto (h = 0,45 m, só de teste, não é do carro). Com lambda_f = 0,3 a traseira transfere
# mais carga e a RL descola antes de a frente saturar (o PUSH não pode ser confirmado); com lambda_f = 0,7,
# o espelho: a FL descola antes de a traseira saturar.
_H_013B = 0.45
rB_p, eB_p = _par_sem_excecao(_caso_013(0.3, cg_height=_H_013B), passo_g=0.01)
rB_l, eB_l = _par_sem_excecao(_caso_013(0.7, cg_height=_H_013B), passo_g=0.01)


def _pre_b(r, sat, eixo_inv, roda):
    """Pré-condições do ramo B pelas linhas: só o outro eixo satura; o eixo alvo fica inválido por
    descolamento da roda interna numa linha com a_y <= S do outro; antes disso ele é válido."""
    if r is None or sat(r) is None:
        return False
    inval = [ln for ln in r.linhas if getattr(ln, f"cap_{eixo_inv}").motivo is not None]
    if not inval:
        return False
    prim = inval[0]
    validas_antes = all(getattr(ln, f"cap_{eixo_inv}").motivo is None for ln in r.linhas if ln.ay_ms2 < prim.ay_ms2)
    return (prim.ay_ms2 <= sat(r) and validas_antes and prim.fz[roda] <= 0.0
            and "descolamento" in getattr(prim, f"cap_{eixo_inv}").motivo)


okB = False
if rB_p is not None and rB_l is not None:
    pre_bp = rB_p.ay_sat_r is None and _pre_b(rB_p, lambda r: r.ay_sat_f, "r", "RL")
    pre_bl = rB_l.ay_sat_f is None and _pre_b(rB_l, lambda r: r.ay_sat_r, "f", "FL")
    okB = (pre_bp and pre_bl
           and rB_p.veredito == eq.NA and rB_p.c_n is None
           and (rB_p.motivo or "").startswith("Fz_RL = ") and "descolamento" in rB_p.motivo
           and rB_l.veredito == eq.NA and rB_l.c_n is None
           and (rB_l.motivo or "").startswith("Fz_FL = ") and "descolamento" in rB_l.motivo)
ok_013 = ok_013 and okA and okB
det_013 += (f" | A push: {eA_p or _resumo(rA_p)} | A loose: {eA_l or _resumo(rA_l)}"
            f" | B push: {eB_p or _resumo(rB_p)} | B loose: {eB_l or _resumo(rB_l)}")
check("TC-EQ-013: ramos de validade do veredito, determinísticos (ds sem escala: N/A por descolamento ou "
      "PUSH/LOOSE válido; A: eixo não saturado válido além de S + resolução -> PUSH/LOOSE com c_n e recomposição; "
      "B: roda interna do eixo não saturado descola antes de S -> N/A com o motivo dele); "
      "nenhuma capacidade inválida usada",
      ok_013, det_013)

# ---------------- TC-EQ-014 : grade curta -> não satura ----------------
r_curta, err_curta = _par_sem_excecao(caso0, ay_max_g=0.3)
# 0,33 g com passo 0,05: a grade para em 0,30 g (k = 6), e o motivo cita a última linha varrida, não ay_max_g
r_curta33, err_curta33 = _par_sem_excecao(caso0, ay_max_g=0.33, passo_g=0.05)
check("TC-EQ-014: ay_max_g = 0,3 -> N/A 'não satura até 0,3 g'; ay_max_g = 0,33 e passo 0,05 -> o motivo cita a "
      "última linha varrida (0,3 g), não 0,33",
      r_curta is not None and r_curta.veredito == eq.NA and "não satura até 0,3 g" in (r_curta.motivo or "")
      and r_curta.c_n is None and r_curta.ay_lim is None
      and r_curta33 is not None and r_curta33.veredito == eq.NA and len(r_curta33.linhas) == 7
      and "não satura até 0,3 g" in (r_curta33.motivo or "") and "0,33" not in (r_curta33.motivo or ""),
      (err_curta or _resumo(r_curta)) + " | " + (err_curta33 or _resumo(r_curta33)))

# ---------------- TC-EQ-015 : parâmetros inválidos ----------------
erros = [raises(lambda: eq.analise_par(caso0, ay_max_g=0)),
         raises(lambda: eq.analise_par(caso0, passo_g=-0.05)),
         raises(lambda: eq.analise_par(caso0, n_alfa=2))]
check("TC-EQ-015: ay_max_g=0, passo_g=-0,05 e n_alfa=2 -> ValueError cada um", all(erros), f"{erros}")

# ---------------- TC-EQ-016 : robustez ao pneu (spec 7.6) ----------------
_M = "mesmo sinal na varredura exploratória do pneu"
_dep = (eq.NAO_CONCLUSIVO, "o sinal depende do pneu")
_parte = (eq.NAO_CONCLUSIVO, "N/A em parte da varredura (sem inversão de sinal)")
ok_016 = (eq.veredito_robusto([eq.PUSH, eq.PUSH]) == (eq.PUSH, _M)
          and eq.veredito_robusto([eq.LOOSE, eq.LOOSE]) == (eq.LOOSE, _M)
          and eq.veredito_robusto([eq.PUSH, eq.LOOSE]) == _dep
          and eq.veredito_robusto([eq.PUSH, eq.NEUTRO]) == _dep
          and eq.veredito_robusto([eq.PUSH, eq.NA]) == _parte
          and eq.veredito_robusto([eq.LOOSE, eq.NA, eq.LOOSE]) == _parte
          and eq.veredito_robusto([eq.PUSH, eq.LOOSE, eq.NA]) == _dep
          and eq.veredito_robusto([eq.LOOSE, eq.NEUTRO, eq.NA]) == _dep
          and eq.veredito_robusto([eq.NA, eq.NA]) == (eq.NA, "N/A em toda a varredura")
          and eq.veredito_robusto([]) == (eq.NA, "robustez ao pneu não avaliada"))
check("TC-EQ-016: veredito_robusto (PUSH/PUSH, LOOSE/LOOSE, inversão, NEUTRO, N/A misto com e sem inversão, N/A, vazio)", ok_016)

# ---------------- TC-EQ-017 : varredura do pneu ----------------
# Grade reduzida e fatores abaixo são só de teste, não são amplitudes do EDR-055.
_kw = dict(ay_max_g=1.0, passo_g=0.1, n_alfa=9)
v1 = eq.varredura_pneu(caso0, {"LMUY": [1.0]}, **_kw)
ok_017a = (len(v1) == 1 and v1[0][0] == {"LMUY": 1.0} and v1[0][1] == eq.analise_par(caso0, **_kw))
p_antes = copy.deepcopy(ds.p)
v4 = eq.varredura_pneu(caso0, {"LFZO": [0.9, 1.1], "LKY": [0.9, 1.1]}, **_kw)
ok_017b = len(v4) == 4 and ds.p == p_antes and caso0.ds is ds
ok_017c = eq.varredura_pneu(caso0, {}, **_kw) == []
# LMUY 0,5 (só de teste): igual a analise_par com o dataset escalado à mão (cópia de ds.p com LMUY·0,5)
_p_mao = copy.deepcopy(ds.p)
_p_mao["LMUY"] = ds.p["LMUY"] * 0.5
v05 = eq.varredura_pneu(caso0, {"LMUY": [0.5]}, **_kw)
ok_017d = (len(v05) == 1 and v05[0][0] == {"LMUY": 0.5}
           and v05[0][1] == eq.analise_par(dataclasses.replace(caso0, ds=dataclasses.replace(ds, p=_p_mao)), **_kw))
check("TC-EQ-017: fator 1,0 reproduz analise_par; LMUY 0,5 = analise_par com o .tir escalado à mão; 4 combinações; "
      "ds.p intacto; dict vazio -> []",
      ok_017a and ok_017b and ok_017c and ok_017d, f"{ok_017a}/{ok_017b}/{ok_017c}/{ok_017d}")

# ---------------- TC-EQ-018 : nomes de fator e texto do veredito ----------------
erros_018 = [raises(lambda: eq.com_fatores(ds, {"PDY2": 1.1})),
             raises(lambda: eq.com_fatores(ds, {"LMUX ": 1.0})),
             raises(lambda: eq.com_fatores(ds, {"LXYZ": 1.0})),
             # LONGVL existe no .tir e começa com "L", mas não é fator de escala (help do Adams, pp. 23-25)
             raises(lambda: eq.com_fatores(ds, {"LONGVL": 1.0})),
             raises(lambda: eq.com_fatores(ds, {"LMUY": 0})),
             raises(lambda: eq.com_fatores(ds, {"LMUY": float("nan")}))]
ds_f = eq.com_fatores(ds, {"LMUY": 0.5})
ok_018a = all(erros_018) and ds_f.p["LMUY"] == ds.p["LMUY"] * 0.5 and ds.p == p_antes
_r = lambda v, m=None: eq.ResultadoPar(linhas=(), passo_ms2=0.0, ay_sat_f=None, ay_sat_r=None,
                                       ay_lim=None, veredito=v, motivo=m, c_n=None)
TX, T09 = origens.TX_V41, origens.T09_V41
T11 = origens.T11
OUTRO = origens.OUTRO_CARRO
_nc = "DEMONSTRATIVO · ? NÃO CONCLUSIVO: o sinal depende do pneu"
ok_018b = (eq.texto_veredito(_r(eq.PUSH), T09) == "DEMONSTRATIVO · ▲ PUSH: sai de frente (final understeer)"
           and eq.texto_veredito(_r(eq.LOOSE), T11) == "▼ LOOSE: sai de traseira (final oversteer)"
           and eq.texto_veredito(_r(eq.NEUTRO), T11) == "● NEUTRO no limite"
           and eq.texto_veredito(_r(eq.NA, "motivo x"), T11) == "N/A: motivo x"
           and eq.texto_veredito(_r(eq.PUSH), T11, (eq.PUSH, "mesmo")) == "▲ PUSH: sai de frente (final understeer)"
           and eq.texto_veredito(_r(eq.PUSH), T09, _dep) == _nc
           # V4.2: o rótulo do TX da V4.1 já não é o carro do projeto (prefixo); com origem_do_carro=TX_V41 volta o comportamento antigo
           and eq.texto_veredito(_r(eq.LOOSE), TX) == "DEMONSTRATIVO · ▼ LOOSE: sai de traseira (final oversteer)"
           and eq.texto_veredito(_r(eq.LOOSE), TX, origem_do_carro=TX) == "▼ LOOSE: sai de traseira (final oversteer)"
           and eq.texto_veredito(_r(eq.LOOSE), T11, origem_do_carro=TX) == "DEMONSTRATIVO · ▼ LOOSE: sai de traseira (final oversteer)"
           # NÃO CONCLUSIVO mostra o motivo, com e sem o prefixo D16
           and eq.texto_veredito(_r(eq.PUSH), T11, _dep) == "? NÃO CONCLUSIVO: o sinal depende do pneu"
           and eq.texto_veredito(_r(eq.PUSH), T11, _parte)
           == "? NÃO CONCLUSIVO: N/A em parte da varredura (sem inversão de sinal)"
           and eq.texto_veredito(_r(eq.PUSH), OUTRO, _parte)
           == "DEMONSTRATIVO · ? NÃO CONCLUSIVO: N/A em parte da varredura (sem inversão de sinal)")
check("TC-EQ-018: nome de fator inválido (inclusive LONGVL, que não é fator de escala) e valor 0 ou NaN -> "
      "ValueError; textos exatos de 7.3 e prefixo DEMONSTRATIVO",
      ok_018a and ok_018b, f"{ok_018a}/{ok_018b}")

# ---------------- TC-EQ-019 : planilha sem valores calculados ----------------
import hashlib
import tempfile

import openpyxl

import caso_planilha as cp
import raizes

with tempfile.TemporaryDirectory(prefix="dds_eq019_") as _tmp:
    _f = os.path.join(_tmp, "sem_valores.xlsx")
    _wb = openpyxl.Workbook()
    _wb.remove(_wb.active)
    _end = cp.enderecos_das_entradas(_wb)   # planilha sem nomes definidos: cai nos endereços da tabela
    for _aba in sorted({a for a, _ in _end.values()}):
        _wb.create_sheet(_aba)
    for _aba, _cel in _end.values():
        _wb[_aba][_cel].value = "=1+1"
    _wb.save(_f)
    ok_019a = raises(lambda: cp.caso_da_planilha(_f, TIR), ValueError, "precisa ser recalculada")
ok_019b = (len(cp.ENTRADAS_CASO) == 23 and cp.ENTRADAS_CASO["DF"] == "ent_08_downforce_total"
           and cp.ENTRADAS_CASO["bal"] == "ent_08_aero_balance_fracao_dianteira"
           and _end["DF"] == ("08_VEHICLE_STATE", "B8") and _end["bal"] == ("08_VEHICLE_STATE", "B9"))
_V40 = os.path.join(HERE, "..", "excel", "FSAE_EV_Suspension_Fase1_V4_0.xlsx")
if os.path.exists(_V40):   # só leitura: os 23 nomes da V4.0 apontam para os mesmos endereços da tabela (V3.0)
    _wb40 = openpyxl.load_workbook(_V40, read_only=False)
    ok_019b = ok_019b and cp.enderecos_das_entradas(_wb40) == _end
    _wb40.close()
check("TC-EQ-019: planilha sem valor calculado -> ValueError 'precisa ser recalculada'; mapa de entradas",
      ok_019a and ok_019b, f"{ok_019a}/{ok_019b}")

# ---------------- TC-EQ-020 : pior origem da planilha viva (só leitura) ----------------
VIVA = os.path.join(HERE, "..", "excel", "FSAE_EV_Suspension_Fase1_V3_0.xlsx")


def _sha(p):
    with open(p, "rb") as fh:
        return hashlib.sha256(fh.read()).hexdigest()


_h0 = _sha(VIVA)
_pior = cp.pior_origem_do_caso(VIVA)
# sem o pneu: valor conhecido em 2026-10-21 (SUS-007/008 arbitrários na viva)
_pior_ent = cp.pior_origem_entradas(VIVA)
_h1 = _sha(VIVA)
ok_020 = (_pior in raizes.PIOR_PRIMEIRO and _pior != "● Dado do TX" and _h0 == _h1
          and _pior_ent == "○ Arbitrário / demonstrativo")
check("TC-EQ-020: pior origem da viva é rótulo conhecido e não é TX; pior origem das entradas (sem o pneu) = "
      "'○ Arbitrário / demonstrativo'; arquivo intacto (SHA-256)",
      ok_020, f"pior = {_pior}; entradas = {_pior_ent}")


# ---------------- TC-EQ-039..041 : Ackermann e cos δ (V4.0, inc. 2; spec 7.1 item 3 e 7.4 hipóteses 2 e 3) ----------------
import kinematics_model as km

R_TESTE = 8.0  # entrada de teste rotulada (R do skidpad "a confirmar"); não é dado do carro
ok_039a = raises(lambda: eq.CasoEquilibrio(cfg=cfg, ds=ds, ackermann_pct=100.0), ValueError, "raio_m")
ok_039b = not raises(lambda: eq.CasoEquilibrio(cfg=cfg, ds=ds, raio_m=R_TESTE, ackermann_pct=100.0))
ok_039c = raises(lambda: eq.CasoEquilibrio(cfg=cfg, ds=ds, raio_m=-1.0), ValueError, "raio_m")
check("TC-EQ-039: ackermann_pct sem raio_m -> ValueError citando raio_m; com raio_m é aceito",
      ok_039a and ok_039b and ok_039c, f"{ok_039a}/{ok_039b}/{ok_039c}")

# 040: só o cos δ_ref (sem Ackermann) multiplica a capacidade dianteira
caso_r = eq.CasoEquilibrio(cfg=cfg, ds=ds, raio_m=R_TESTE)
cap0 = eq.capacidade_eixo(caso0, "F", 0.0, 25)
cap_r = eq.capacidade_eixo(caso_r, "F", 0.0, 25)
cosd = math.cos(cfg.wheelbase / R_TESTE)
erro_f = abs(cap_r.fy_max_N / (cap0.fy_max_N * cosd) - 1.0)
cap0_t = eq.capacidade_eixo(caso0, "R", 0.0, 25)
cap_r_t = eq.capacidade_eixo(caso_r, "R", 0.0, 25)
check("TC-EQ-040: R = 8 m sem Ackermann -> capacidade dianteira × cos(L/R) (erro < 1e-12); traseira igual",
      erro_f < 1e-12 and cap_r_t.fy_max_N == cap0_t.fy_max_N and cap_r_t.alfa_pico_rad == cap0_t.alfa_pico_rad,
      f"erro rel F = {erro_f:.2e}, cos = {cosd:.6f}")

# 041: δ das rodas e Δα extra
d_ref = cfg.wheelbase / R_TESTE
d0 = eq.esterco_rodas(eq.CasoEquilibrio(cfg=cfg, ds=ds, raio_m=R_TESTE, ackermann_pct=0.0))
ok_041a = abs(d0["FL"] - d_ref) < 1e-12 and abs(d0["FR"] - d_ref) < 1e-12
da0 = eq.deriva_ackermann(eq.CasoEquilibrio(cfg=cfg, ds=ds, raio_m=R_TESTE, ackermann_pct=0.0))
ok_041b = all(abs(v) < 1e-12 for v in da0.values())
caso_a = eq.CasoEquilibrio(cfg=cfg, ds=ds, raio_m=R_TESTE, ackermann_pct=100.0)
d100 = eq.esterco_rodas(caso_a)
d_o, d_i = d100["FR"], d100["FL"]
ideal_i = math.radians(km.ackermann_inner_ideal_deg(math.degrees(d_o), cfg.wheelbase, cfg.track_front))
ok_041c = (abs(d_i - ideal_i) < 1e-9 and abs((d_o + d_i) / 2.0 - d_ref) < 1e-9 and d_i > d_o)
da = eq.deriva_ackermann(caso_a)
ok_041d = da["FL"] < 0.0 < da["FR"] and da["RL"] == 0.0 and da["RR"] == 0.0
ok_041e = (abs(da["FL"] + (d_i - d_ref)) < 1e-12 and abs(da["FR"] + (d_o - d_ref)) < 1e-12)
check("TC-EQ-041: Ackermann 0 % -> δ = δ_ref e Δα extra nulo; 100 % -> δ_i = ideal (1e-9), média δ_ref, "
      "Δα de FL e FR com sinais opostos",
      ok_041a and ok_041b and ok_041c and ok_041d and ok_041e,
      f"δ_ref = {math.degrees(d_ref):.4f}°, δ_o = {math.degrees(d_o):.4f}°, δ_i = {math.degrees(d_i):.4f}° "
      f"({ok_041a}/{ok_041b}/{ok_041c}/{ok_041d}/{ok_041e})")


# ---------------- TC-EQ-042..043 : CG lateral, dois sentidos e aero com a velocidade (V4.0, inc. 2; spec 7.4, hipóteses 7 e 8) ----------------
def _campos_veredito(r):
    return (r.veredito, r.ay_sat_f, r.ay_sat_r, r.ay_lim, r.c_n, r.motivo, r.resolucao_ms2, r.limitacoes)


def _par2_sem_excecao(caso_x, **kw):
    try:
        return eq.analise_par_dois_sentidos(caso_x, **kw), None
    except Exception as e:  # o check precisa registrar qualquer exceção como falha
        return None, f"{type(e).__name__}: {e}"


# 042a: carro simétrico, y_cg = 0, com toe, camber estático e ganho de camber simétricos (só de teste; o ganho
# antissimétrico FL = −g, FR = +g é o de braços simétricos, gradiente de eixo +g) -> os dois sentidos iguais.
_toe_sim = {"FL": math.radians(0.2), "FR": math.radians(0.2), "RL": 0.0, "RR": 0.0}
_cam_sim = {"FL": math.radians(-1.0), "FR": math.radians(-1.0), "RL": math.radians(-0.5), "RR": math.radians(-0.5)}
_gan_sim = {"FL": -0.6, "FR": 0.6, "RL": -0.4, "RR": 0.4}
caso_42s = eq.CasoEquilibrio(cfg=veiculo_simetrico(0.7), ds=ds_teste_011, toe_rad=_toe_sim,
                             camber_estatico_rad=_cam_sim, ganho_camber=_gan_sim)
d42s, e42s = _par2_sem_excecao(caso_42s)
ok_042a = (d42s is not None and d42s["esquerda"] == d42s["direita"]
           and _campos_veredito(d42s["esquerda"]) == _campos_veredito(d42s["direita"])
           and d42s["esquerda"].ay_sat_f is not None and d42s["pior"] == "empate")
ok_042a2 = eq.espelhar(caso_42s) == caso_42s

# 042b: y_cg = +0,02 m (só de teste) -> estática: soma igual e o lado esquerdo ganha W_eixo·y/t por eixo.
Y_TESTE = 0.02
cfg_42 = veiculo_simetrico(0.7)
st = sm.static_fz_corner(cfg_42)
stl = sm.static_fz_corner_cg_lateral(cfg_42, Y_TESTE)
W_f42, W_r42 = st["FL"] + st["FR"], st["RL"] + st["RR"]
dzf, dzr = W_f42 * Y_TESTE / cfg_42.track_front, W_r42 * Y_TESTE / cfg_42.track_rear
ok_042b = (abs(sum(stl.values()) - sum(st.values())) < TOL
           and abs(stl["FL"] - (st["FL"] + dzf)) < TOL and abs(stl["FR"] - (st["FR"] - dzf)) < TOL
           and abs(stl["RL"] - (st["RL"] + dzr)) < TOL and abs(stl["RR"] - (st["RR"] - dzr)) < TOL)
caso_42y = eq.CasoEquilibrio(cfg=cfg_42, ds=ds_teste_011, cg_lateral_m=Y_TESTE)
c42 = eq.cargas(caso_42y, 0.0)
ok_042c = all(abs(c42[c] - stl[c]) < TOL for c in eq.CANTOS)

# 042c: dois sentidos com y_cg = +0,02 m (LMUY 0,45 do TC-EQ-011): a direita (rodas externas = as esquerdas, já mais
# carregadas) é a pior. Esperado físico; se der o contrário, o check falha e os números ficam no detalhe.
d42, e42 = _par2_sem_excecao(caso_42y)
ok_042d = (d42 is not None and d42["pior"] == "direita"
           and d42["direita"].ay_lim is not None and d42["esquerda"].ay_lim is not None
           and d42["direita"].ay_lim <= d42["esquerda"].ay_lim)
_det42 = e42 or (f"esquerda: {_resumo(d42['esquerda'])}, res={d42['esquerda'].resolucao_ms2} | "
                 f"direita: {_resumo(d42['direita'])}, res={d42['direita'].resolucao_ms2} | pior={d42['pior']}")

# 042d: tabela_linear (1ª linha) com o carro assimétrico (fix round 1 da Task 6). Veículo base e .tir real.
# y_cg = −0,02 m (só de teste): a Fy do eixo em α = 0 já passa da exigência em a_y = 0, o ângulo necessário é > 0 e
# angulos_necessarios precisa achá-lo; a 1ª linha tem k_inclinacao numérico. y_cg = 0: o ângulo em a_y = 0 é (0, 0)
# e a 1ª linha é a de antes da correção: igual à Richardson das secantes sem desconto (forma anterior, recalculada
# aqui) e ao valor de e2e3130 neste .tir (−6.1009286892091e-05 rad/(m/s²)).
import equilibrio_linear as el

tab_m = el.tabela_linear(eq.CasoEquilibrio(cfg=cfg, ds=ds, cg_lateral_m=-Y_TESTE), ay_max_g=0.05)
ok_042e = tab_m[0]["k_inclinacao"] is not None and math.isfinite(tab_m[0]["k_inclinacao"])
caso_42z = eq.CasoEquilibrio(cfg=cfg, ds=ds)
tab_z = el.tabela_linear(caso_42z, ay_max_g=0.05)
_sec = [-(a[0] - a[1]) / h for h, a in ((h, el.angulos_necessarios(caso_42z, h)) for h in (0.005 * G, 0.01 * G))]
ok_042f = (el.angulos_necessarios(caso_42z, 0.0) == (0.0, 0.0)
           and tab_z[0]["k_inclinacao"] == 2.0 * _sec[0] - _sec[1]
           and tab_z[0]["k_inclinacao"] == -6.1009286892091e-05)
check("TC-EQ-042: simétrico com y_cg = 0 -> esquerda == direita (ResultadoPar inteiro) e espelhar(caso) == caso; "
      "y_cg = +0,02 m (só de teste) -> soma estática igual, esquerda ganha W_eixo·y/t por eixo, cargas(a_y=0) = "
      "estática deslocada; dois sentidos (LMUY 0,45) -> pior = direita; tabela_linear: y_cg = −0,02 m -> 1ª linha "
      "com k_inclinacao numérico, y_cg = 0 -> 1ª linha igual à de antes",
      ok_042a and ok_042a2 and ok_042b and ok_042c and ok_042d and ok_042e and ok_042f,
      f"({ok_042a}/{ok_042a2}/{ok_042b}/{ok_042c}/{ok_042d}/{ok_042e}/{ok_042f}) sim: "
      f"{e42s or _resumo(d42s['esquerda'])} | ΔFz_F = {dzf:.3f} N, ΔFz_R = {dzr:.3f} N | {_det42} | "
      f"k_incl(y=−0,02) = {tab_m[0]['k_inclinacao']!r} ({tab_m[0]['motivo_inclinacao']}), "
      f"k_incl(y=0) = {tab_z[0]['k_inclinacao']!r}")

# ---------------- TC-EQ-043 : aero com V = √(a_y·R) ----------------
CZA_TESTE, RHO_TESTE, BAL_TESTE = 3.0, 1.2, 0.45   # entradas de teste rotuladas (CzA, ρ, balanço "a confirmar")
caso_43 = eq.CasoEquilibrio(cfg=cfg, ds=ds, aero_balance_front=BAL_TESTE, raio_m=R_TESTE, cza_m2=CZA_TESTE,
                            rho_kg_m3=RHO_TESTE)
erros_43, det_43 = [], []
for ay_g in (0.0, 0.5, 1.0, 1.5):
    ay = ay_g * G
    f43 = eq.cargas(caso_43, ay)
    df_esp = 0.5 * RHO_TESTE * CZA_TESTE * (ay * R_TESTE)
    st43 = sm.static_fz_corner(cfg)
    e_soma = abs(sum(f43.values()) - (W + df_esp))
    e_f = abs(f43["FL"] + f43["FR"] - (st43["FL"] + st43["FR"] + BAL_TESTE * df_esp))
    e_r = abs(f43["RL"] + f43["RR"] - (st43["RL"] + st43["RR"] + (1.0 - BAL_TESTE) * df_esp))
    erros_43.append(max(e_soma, e_f, e_r))
    det_43.append(f"{ay_g:g} g: DF = {df_esp:.2f} N, erro = {erros_43[-1]:.1e} N")
ok_043a = max(erros_43) < 1e-9
ok_043b = abs(sm.downforce_por_velocidade(1.2, 3.0, 10.0) - 180.0) < 1e-12
ok_043c = (raises(lambda: eq.CasoEquilibrio(cfg=cfg, ds=ds, cza_m2=3.0, rho_kg_m3=1.2), ValueError, "raio_m")
           and raises(lambda: eq.CasoEquilibrio(cfg=cfg, ds=ds, cza_m2=3.0, raio_m=R_TESTE), ValueError, "rho_kg_m3")
           and raises(lambda: eq.CasoEquilibrio(cfg=cfg, ds=ds, cza_m2=3.0, raio_m=R_TESTE, rho_kg_m3=1.2,
                                                downforce_N=400.0), ValueError, "downforce_N"))
check("TC-EQ-043: CzA = 3,0 m², ρ = 1,2 kg/m³, R = 8 m (só de teste) -> Σ Fz = m·G + 0,5·ρ·CzA·a_y·R e divisão "
      "F/R pelo balanço (erro < 1e-9 N); cza_m2 sem raio_m, sem rho_kg_m3 ou junto com downforce_N -> ValueError",
      ok_043a and ok_043b and ok_043c, f"({ok_043a}/{ok_043b}/{ok_043c}) " + "; ".join(det_43))


# ---------------- TC-EQ-044..046 : tração na traseira, elipse de atrito, diferencial aberto (spec 7.4, h9 e h10) ----------------
def _fx_max_ref(fz_c, gam_c):
    """Pico de Fx puro do PAC2002 (κ > 0, tração) pela forma fechada |Dx| + SVx (PCX1 ≥ 1 neste .tir)."""
    _, info = tm.pure_fx0(ds.p, fz_c, gam_c, 0.0)
    return abs(info["Dx"]) + info["SVx"]


def _fx_exigido_ref(caso_x, ay, cda, crr):
    """Fx exigido no eixo traseiro, recomposto: 0,5·ρ·CdA·a_y·R + Crr·m·G + F_yf·tan(L/R).

    F_yf é a componente lateral (no carro) da força do pneu dianteiro, F_pneu·cos δ; a componente para trás é
    F_pneu·sen δ = F_yf·tan δ (EDR-057)."""
    cfg_x = caso_x.cfg
    f_yf = cfg_x.mass_total * ay * cfg_x.cg_to_rear / cfg_x.wheelbase
    return (0.5 * RHO_TESTE * cda * ay * R_TESTE + crr * cfg_x.mass_total * G
            + f_yf * math.tan(cfg_x.wheelbase / R_TESTE))


# 044: CdA enorme (200 m², só de teste) a 1 g -> a roda interna (RL, menos carregada) passa do atrito longitudinal
CDA_ENORME = 200.0
caso_44 = eq.CasoEquilibrio(cfg=cfg, ds=ds, raio_m=R_TESTE, rho_kg_m3=RHO_TESTE, cda_m2=CDA_ENORME)
cap_44 = _cap_sem_excecao(caso_44, 1.0 * G, "R")
fz44 = eq.cargas(caso_44, 1.0 * G)
fx44 = _fx_exigido_ref(caso_44, 1.0 * G, CDA_ENORME, 0.0) / 2.0
ok_044a = (fx44 > _fx_max_ref(fz44["RL"], 0.0) and cap_44 is not None and cap_44.fy_max_N is None
           and cap_44.alfa_pico_rad is None and cap_44.motivo == "tração excede o atrito longitudinal em RL")
ok_044b = (raises(lambda: eq.CasoEquilibrio(cfg=cfg, ds=ds, rho_kg_m3=RHO_TESTE, cda_m2=1.0), ValueError, "raio_m")
           and raises(lambda: eq.CasoEquilibrio(cfg=cfg, ds=ds, raio_m=R_TESTE, cda_m2=1.0), ValueError, "rho_kg_m3")
           and not raises(lambda: eq.CasoEquilibrio(cfg=cfg, ds=ds, crr=0.015)))
check("TC-EQ-044: CdA = 200 m² (só de teste), a_y = 1 g -> capacidade traseira None com 'tração excede o atrito "
      "longitudinal em RL', sem exceção; cda_m2 sem raio_m ou sem rho_kg_m3 -> ValueError; crr sozinho é aceito",
      ok_044a and ok_044b,
      f"({ok_044a}/{ok_044b}) Fx/roda = {fx44:.1f} N, Fx_max RL = {_fx_max_ref(fz44['RL'], 0.0):.1f} N, "
      f"cap = {cap_44}")

# 045: Crr = 0,015 e CdA = 1,0 m² (só de teste), R = 8 m, ρ = 1,2 -> capacidade traseira = máximo sobre a grade de α de
# Σ Fy_c·√(1 − (Fx/Fx_max,c)²) (escala por roda antes do máximo do eixo); a frente igual ao caso sem tração.
CDA_045, CRR_045 = 1.0, 0.015
caso_45_sem = eq.CasoEquilibrio(cfg=cfg, ds=ds, raio_m=R_TESTE, rho_kg_m3=RHO_TESTE)
caso_45 = dataclasses.replace(caso_45_sem, cda_m2=CDA_045, crr=CRR_045)
erros_45, det_45, ok_045f = [], [], True
for ay_g in (0.0, 0.5, 1.0):
    ay = ay_g * G
    fz45 = eq.cargas(caso_45, ay)
    gam45 = eq.camber(caso_45, ay)
    fx_roda = _fx_exigido_ref(caso_45, ay, CDA_045, CRR_045) / 2.0
    fat = {c: math.sqrt(1.0 - (fx_roda / _fx_max_ref(fz45[c], gam45[c])) ** 2) for c in ("RL", "RR")}
    dal45 = eq.deriva_total(caso_45)
    esp45 = max(sum(fat[c] * tm.wheel_forces(ds, c, fz45[c], a + dal45[c], gam45[c], 0.0)[1] for c in ("RL", "RR"))
                for a in eq.grade_alfa(caso_45, 25))
    cap45 = _cap_sem_excecao(caso_45, ay, "R")
    ok_lim = cap45 is not None and cap45.motivo is None and not _limites(cap45)
    erros_45.append(abs(cap45.fy_max_N / esp45 - 1.0) if ok_lim else math.inf)
    cf_com, cf_sem = _cap_sem_excecao(caso_45, ay, "F"), _cap_sem_excecao(caso_45_sem, ay, "F")
    ok_045f = ok_045f and cf_com is not None and cf_com == cf_sem
    det_45.append(f"{ay_g:g} g: Fx/roda = {fx_roda:.2f} N, fatores RL/RR = {fat['RL']:.6f}/{fat['RR']:.6f}, "
                  f"erro rel = {erros_45[-1]:.1e}")
# Fx_max: a forma fechada |Dx| + SVx só vale com o κ do pico dentro da faixa. Na carga de RR a 1 g do caso acima, a
# forma fechada (fx_max_puro) bate com o máximo numérico de pure_fx0 em κ ∈ [0, KPUMAX] (20001 pontos). Com KPUMAX
# reduzido a 0,03 (só de teste; o pico do .tir fica entre 0,065 e 0,092), o pico sai da faixa e fx_max_puro vira o
# máximo numérico na faixa, abaixo da forma fechada.
fz_fx = eq.cargas(caso_45, 1.0 * G)["RR"]


def _fx_num(p_x, kmax, n=20000):
    return max(tm.pure_fx0(p_x, fz_fx, 0.0, kmax * i / n)[0] for i in range(n + 1))


fx_fech = eq.fx_max_puro(ds, fz_fx, 0.0)
fx_num = _fx_num(ds.p, ds.p["KPUMAX"])
ds_kcurto = dataclasses.replace(ds, p={**ds.p, "KPUMAX": 0.03})
fx_kcurto = eq.fx_max_puro(ds_kcurto, fz_fx, 0.0)
fx_kcurto_num = _fx_num(ds_kcurto.p, 0.03)
ok_045x = (abs(fx_fech / fx_num - 1.0) < 1e-6 and abs(fx_fech - _fx_max_ref(fz_fx, 0.0)) < 1e-9
           and abs(fx_kcurto / fx_kcurto_num - 1.0) < 1e-6 and fx_kcurto < fx_fech * (1.0 - 1e-3))
det_45.append(f"Fx_max(RR, 1 g): fechada = {fx_fech:.3f} N, numérica = {fx_num:.3f} N; KPUMAX 0,03: "
              f"{fx_kcurto:.3f} N (numérica {fx_kcurto_num:.3f} N)")
ok_045 = max(erros_45) < 1e-9 and ok_045f and ok_045x
check("TC-EQ-045: Crr = 0,015, CdA = 1,0 m² (só de teste) -> capacidade traseira = máx_α Σ Fy_c·√(1 − (Fx/Fx_max,c)²) "
      "(erro rel < 1e-9), Fx = 0,5·ρ·CdA·a_y·R + Crr·m·G + F_yf·tan δ_ref dividido igual entre RL e RR; frente igual; "
      "Fx_max fechado = máximo numérico em [0, KPUMAX] e, com o pico fora da faixa, o numérico",
      ok_045, f"(frente igual: {ok_045f}; Fx_max: {ok_045x}) " + "; ".join(det_45))

# 046: teste de direção (spec 7.4, h9: ignorar a tração superestima a traseira e puxa para push). Veículo simétrico,
# λ_f = 0,5, LMUY 0,45 (do TC-EQ-011, só de teste), R = 8 m, ρ = 1,2, Crr = 0,015, CdA = 1,0 m² (só de teste).
# Com R, o cos δ_ref já deixa o caso sem tração em PUSH; a tração reduz S_r e C_N (rumo a loose) sem mexer na
# frente, e o veredito continua PUSH neste caso. Só com Crr (sem R, caso NEUTRO), a queda de S_r fica abaixo da
# resolução numérica: o veredito não muda (sem inversão falsa).
caso_46_sem = eq.CasoEquilibrio(cfg=veiculo_simetrico(0.5), ds=ds_teste_011, raio_m=R_TESTE, rho_kg_m3=RHO_TESTE)
caso_46_com = dataclasses.replace(caso_46_sem, crr=CRR_045, cda_m2=CDA_045)
r46_sem, e46_sem = _par_sem_excecao(caso_46_sem)
r46_com, e46_com = _par_sem_excecao(caso_46_com)
ok_046a = (r46_sem is not None and r46_com is not None
           and None not in (r46_sem.ay_sat_f, r46_sem.ay_sat_r, r46_com.ay_sat_f, r46_com.ay_sat_r)
           and r46_sem.veredito == eq.PUSH and r46_com.veredito == eq.PUSH
           and r46_com.ay_sat_r < r46_sem.ay_sat_r and r46_com.c_n < r46_sem.c_n
           and r46_com.ay_sat_f == r46_sem.ay_sat_f
           and all(a.cap_f == b.cap_f for a, b in zip(r46_sem.linhas, r46_com.linhas)))
caso_46n = eq.CasoEquilibrio(cfg=veiculo_simetrico(0.5), ds=ds_teste_011)
r46n_sem, e46n_sem = _par_sem_excecao(caso_46n)
r46n_com, e46n_com = _par_sem_excecao(dataclasses.replace(caso_46n, crr=CRR_045))
ok_046b = (r46n_sem is not None and r46n_com is not None and r46n_sem.veredito == eq.NEUTRO
           and r46n_com.veredito == eq.NEUTRO and r46n_com.ay_sat_r < r46n_sem.ay_sat_r
           and r46n_sem.ay_sat_r - r46n_com.ay_sat_r < r46n_com.resolucao_ms2)


def _g(x):
    return f"{x / G:.6f}" if isinstance(x, float) else repr(x)


check("TC-EQ-046: tração (Crr 0,015, CdA 1,0, R 8 m; só de teste) -> S_r e C_N caem (rumo a loose), S_f e capacidades "
      "dianteiras iguais, veredito segue PUSH neste caso; só Crr (sem R) -> queda de S_r abaixo da resolução, segue NEUTRO",
      ok_046a and ok_046b,
      f"({ok_046a}/{ok_046b}) R 8 sem: {e46_sem or ''} {r46_sem and r46_sem.veredito}, S_f={_g(r46_sem and r46_sem.ay_sat_f)} g, "
      f"S_r={_g(r46_sem and r46_sem.ay_sat_r)} g, C_N={r46_sem and r46_sem.c_n} | R 8 com: {e46_com or ''} "
      f"{r46_com and r46_com.veredito}, S_f={_g(r46_com and r46_com.ay_sat_f)} g, S_r={_g(r46_com and r46_com.ay_sat_r)} g, "
      f"C_N={r46_com and r46_com.c_n} | só Crr: ΔS_r={_g(r46n_sem and r46n_com and r46n_sem.ay_sat_r - r46n_com.ay_sat_r)} g, "
      f"res={_g(r46n_com and r46n_com.resolucao_ms2)} g {e46n_sem or ''}{e46n_com or ''}")


# ---------------- Resumo ----------------
print("\n--- Resumo ---")
n_pass = sum(1 for _, s, _ in results if s == "PASS")
n_fail = sum(1 for _, s, _ in results if s == "FAIL")
print(f"{n_pass} PASS, {n_fail} FAIL, de {len(results)} testes")
if n_fail > 0:
    raise SystemExit(1)
