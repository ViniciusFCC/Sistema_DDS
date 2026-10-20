"""Extrato das entradas da planilha viva (EDR-051, fechamento da V3.0): uma linha por célula de entrada, com o rótulo
da linha, para o commit de dados ter um diff legível ("quem mudou o quê"). É saída, não fonte: a planilha continua
sendo o arquivo da equipe. Só lê a planilha.

Uso:
    uv run python -X utf8 tools/extrato_entradas.py [PLANILHA.xlsx] [--saida data/entradas-V3_0.csv]
"""
import argparse
import csv
import re
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "python"))
from atualizar_planilha import celulas_de_entrada  # noqa: E402


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("planilha", nargs="?", default=str(RAIZ / "excel" / "FSAE_EV_Suspension_Fase1_V3_0.xlsx"))
    ap.add_argument("--saida", default=str(RAIZ / "build" / "entradas-V3_0.csv"))
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.planilha)
    linha = lambda cel: int(re.sub(r"[A-Z]", "", cel))
    chaves = sorted(celulas_de_entrada(), key=lambda k: (k[0], linha(k[1]), k[1]))
    Path(a.saida).parent.mkdir(parents=True, exist_ok=True)
    with open(a.saida, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["aba", "celula", "rotulo_da_linha", "valor"])
        for aba, cel in chaves:
            ws = wb[aba]
            v = ws[cel].value
            w.writerow([aba, cel, ws[f"A{linha(cel)}"].value or "", "" if v is None else v])
    print(len(chaves), "entradas ->", a.saida)


if __name__ == "__main__":
    main()
