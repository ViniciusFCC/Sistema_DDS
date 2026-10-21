"""Relatório de viabilidade e robustez do incremento 1 da V4.0 (evidência registrada).

Uso (Git Bash, depois de export UV_SYSTEM_CERTS=1):
    uv run python -X utf8 tools/relatorio_inc1.py --saida docs/v4-0/inc1-viabilidade-e-robustez.md

Regras (mesmo cuidado de `tools/rodar_testes.py --saida`):
- `--saida` é obrigatório; o arquivo não pode existir (o relatório nunca é regravado).
- A árvore precisa estar limpa (`git status --porcelain` vazio): o commit registrado é o que gerou o relatório.
- A planilha viva só é COPIADA para uma pasta temporária; só a cópia é recalculada (EDR-044). O SHA-256 da viva
  é conferido antes e depois; se mudar, o relatório não é gravado.
- Nada é inventado: todo número sai do Python (equilibrio.py, caso_planilha.py), da cópia recalculada, de
  `data/varredura_pneu_exploratoria.json` (EDR-055), da tabela 7.4 da spec ou de `build/exp_viabilidade_par.json`
  (Task 6). Não há limiar de "vai ou não vai" (spec §9, item 10): a seção 6 mostra só números, e a recomendação vai
  separada e marcada como recomendação.

Saída: 0 = relatório gravado; 2 = recusa (arquivo existe, árvore suja, JSON da Task 6 ausente ou incompleto);
3 = NÃO EXECUTÁVEL (sem motor de recálculo); 1 = a planilha viva mudou durante a execução.
"""
from __future__ import annotations

import argparse
import dataclasses
import datetime as dt
import hashlib
import json
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import time
from importlib.metadata import version
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "tests"))   # recalculo_planilha mora em tests/ (como em tests/exp_viabilidade_par.py)

import openpyxl  # noqa: E402

import caso_planilha as cp  # noqa: E402
import equilibrio as eq  # noqa: E402
import recalculo_planilha as rc  # noqa: E402
import tire_model as tm  # noqa: E402
from atualizar_planilha import CONF, TX  # noqa: E402
from raizes import PIOR_PRIMEIRO, SEM_ROTULO, Grafo, origem_da_raiz  # noqa: E402
from suspension_model import G, VehicleConfig  # noqa: E402

VIVA = RAIZ / "excel" / "FSAE_EV_Suspension_Fase1_V3_0.xlsx"
TIR = RAIZ / "data" / "Hoosier_R25B_13x7_TTC_provisional.tir"
VARREDURA = RAIZ / "data" / "varredura_pneu_exploratoria.json"
EXP_JSON = RAIZ / "build" / "exp_viabilidade_par.json"
SPEC = RAIZ / "docs" / "superpowers" / "specs" / "2026-10-20-v4-0-metas-equilibrio-design.md"
MOTORES = ("excel", "libreoffice")
CASOS_EXP = ("teste_LMUY045", "teste_LMUY045_toe_camber", "teste_LMUY045_LKY023_borda", "planilha_viva")
PREFIXO = "DEMONSTRATIVO · "
GRADE = {"ay_max_g": 2.0, "passo_g": 0.05, "n_alfa": 25}   # grade passada a analise_par e varredura_pneu
ORIGEM_TESTE = "○ Arbitrário / demonstrativo"           # dataset de teste (LMUY 0,45): nunca é dado do TX

# Hipóteses da tabela 7.4 da spec que o incremento 1 implementa (número -> o que o código faz).
# 2 e 8 entram em parte: o resto delas vai para a lista das não incluídas.
INCLUIDAS = {
    "1": "α_ref comum por eixo mais o Δα de toe de cada roda (`equilibrio.fy_eixo`).",
    "2": "Só o toe estático por roda; o Ackermann fica para o incremento 2.",
    "4": "Camber estático mais ganho linear dγ/dφ por roda (`equilibrio.camber`).",
    "5": "Três massas, RC fixo e rolagem linear (`suspension_model.lateral_dfz`).",
    "6": "Termo gravitacional da rolagem, m_s·g·h2 (EDR-046, `suspension_model.lateral_dfz`).",
}
PARCIAIS = {
    "2": "Ackermann por roda (SUS-414, precisa do raio R).",
    "8": "Dependência da downforce com a velocidade (V = √(a_y·R)); o incremento 1 usa downforce e balanço fixos "
         "do Setup (EDR-013).",
}


# ----------------------------------------------------------------------------------------------------------------
# utilitários
# ----------------------------------------------------------------------------------------------------------------

def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for bloco in iter(lambda: f.read(1 << 20), b""):
            h.update(bloco)
    return h.hexdigest()


def git(*args: str) -> str:
    r = subprocess.run(["git", *args], cwd=RAIZ, capture_output=True, text=True, encoding="utf-8")
    if r.returncode != 0:
        raise RuntimeError(f"git {' '.join(args)} falhou: {r.stderr.strip()}")
    return r.stdout.strip()


def n(x: float | None, casas: int = 3) -> str:
    """Número com vírgula decimal; None vira '—'."""
    if x is None:
        return "—"
    return f"{x:.{casas}f}".replace(".", ",")


def g_de(ay_ms2: float | None, casas: int = 3) -> str:
    return "—" if ay_ms2 is None else f"{n(ay_ms2 / G, casas)} g"


def pct(x: float | None) -> str:
    return "—" if x is None else f"{n(100.0 * x, 1)} %"


def pior(*rotulos: str) -> str:
    for r in rotulos:
        if r not in PIOR_PRIMEIRO:
            raise ValueError(f"rótulo de origem fora do vocabulário: {r!r}")
    return min(rotulos, key=PIOR_PRIMEIRO.index)


def recusa(msg: str, codigo: int = 2) -> int:
    print(f"RECUSADO: {msg}", file=sys.stderr)
    return codigo


# ----------------------------------------------------------------------------------------------------------------
# leitura de fontes
# ----------------------------------------------------------------------------------------------------------------

def tabela_74() -> list[dict]:
    """Linhas da tabela 7.4 da spec (Hipóteses), lidas do arquivo, sem cópia à mão."""
    texto = SPEC.read_text(encoding="utf-8")
    m = re.search(r"^### 7\.4 Hipóteses\s*$(.*?)^### ", texto, flags=re.M | re.S)
    if not m:
        raise RuntimeError(f"seção '### 7.4 Hipóteses' não encontrada em {SPEC}")
    linhas = []
    for ln in m.group(1).splitlines():
        cel = [c.strip() for c in ln.strip().strip("|").split("|")]
        if len(cel) == 6 and cel[0].isdigit():
            linhas.append(dict(zip(("num", "hipotese", "base", "planilha", "python", "efeito"), cel)))
    if [int(x["num"]) for x in linhas] != list(range(1, 17)):
        raise RuntimeError("a tabela 7.4 da spec não tem as hipóteses 1 a 16 na ordem; confira o arquivo")
    return linhas


def entradas_e_origens(copia_recalc: Path, copia_formulas: Path) -> list[dict]:
    """Valor, rótulo, unidade, ID e pior origem das raízes de cada entrada de ENTRADAS_CASO."""
    wv = openpyxl.load_workbook(copia_recalc, data_only=True, read_only=True)
    wf = openpyxl.load_workbook(copia_formulas, data_only=False)
    grafo = Grafo(wf)
    saida = []
    for nome, (aba, cel) in cp.ENTRADAS_CASO.items():
        linha = int(re.sub(r"[A-Z]", "", cel))
        ws = wv[aba]
        rot = []
        for ra, rc_ in grafo.raizes_preenchidas((aba, cel)):
            r = origem_da_raiz(wf, ra, rc_)
            rot.append(CONF if r == SEM_ROTULO else r)
        saida.append({"nome": nome, "aba": aba, "cel": cel, "valor": ws[cel].value,
                      "rotulo": ws.cell(linha, 1).value, "unidade": ws.cell(linha, 3).value,
                      "id": ws.cell(linha, 4).value, "origem": pior(*rot) if rot else CONF,
                      "formula": isinstance(wf[aba][cel].value, str) and wf[aba][cel].value.startswith("=")})
    wv.close()
    return saida


def ler_exp_json() -> tuple[dict | None, str | None]:
    """JSON da Task 6. Devolve (dados, None) ou (None, motivo da recusa)."""
    instr = ("rode `DDS_RECALC=excel uv run python -X utf8 tests/exp_viabilidade_par.py` e "
             "`DDS_RECALC=libreoffice uv run python -X utf8 tests/exp_viabilidade_par.py` (Git Bash, depois de "
             "`export UV_SYSTEM_CERTS=1`) e rode este relatório de novo")
    if not EXP_JSON.exists():
        return None, f"{EXP_JSON.relative_to(RAIZ).as_posix()} não existe; {instr}"
    try:
        d = json.loads(EXP_JSON.read_text(encoding="utf-8"))
        execs = d["execucoes"]
    except (ValueError, KeyError, TypeError) as e:
        return None, f"{EXP_JSON.relative_to(RAIZ).as_posix()} ilegível ({e}); {instr}"
    faltam = [f"{m}/{c}" for m in MOTORES for c in CASOS_EXP
              if not any(x.get("motor") == m and x.get("caso") == c for x in execs)]
    if faltam:
        return None, f"faltam execuções no JSON da Task 6: {', '.join(faltam)}; {instr}"
    return d, None


# ----------------------------------------------------------------------------------------------------------------
# análises
# ----------------------------------------------------------------------------------------------------------------

def faixa(res: eq.ResultadoPar) -> dict | None:
    """Primeira linha da grade com capacidade inválida por faixa do pneu (Fz ou camber) e a última válida antes dela."""
    for k, ln in enumerate(res.linhas):
        motivos = [c.motivo for c in (ln.cap_f, ln.cap_r) if c.motivo and "fora da faixa do pneu" in c.motivo]
        if motivos:
            lv = res.linhas[k - 1] if k > 0 else None
            fr = None
            if lv is not None and lv.cap_f.fy_max_N and lv.cap_r.fy_max_N:
                fr = {"F": lv.req_f / lv.cap_f.fy_max_N, "R": lv.req_r / lv.cap_r.fy_max_N}
            return {"k": k, "linha": ln, "motivos": motivos, "ultima_valida": lv, "fracao": fr}
    return None


def rodas_fora(ln: eq.LinhaPar, p: dict) -> list[tuple[str, float, str]]:
    """(roda, Fz, limite violado) para cada roda fora de [FZMIN, FZMAX] na linha."""
    fora = []
    for c in eq.CANTOS:
        if ln.fz[c] < p["FZMIN"]:
            fora.append((c, ln.fz[c], "FZMIN"))
        elif ln.fz[c] > p["FZMAX"]:
            fora.append((c, ln.fz[c], "FZMAX"))
    return fora


def veiculo_teste_011(lambda_f: float) -> VehicleConfig:
    """veiculo_simetrico(lambda_f) de tests/test_equilibrio.py (TC-EQ-011), números copiados. Teste, NÃO é setup."""
    base = dict(mass_total=250.0, cg_height=0.300, wheelbase=1.550, track_front=1.200, track_rear=1.150,
                cg_to_front=0.837, cg_to_rear=0.713, unsprung_mass_front=15.0, unsprung_mass_rear=18.0,
                unsprung_cg_height_front=0.230, unsprung_cg_height_rear=0.230, roll_center_front=0.030,
                roll_center_rear=0.050, roll_stiffness_front=20826.4, roll_stiffness_rear=26450.0)
    k = base["roll_stiffness_front"] + base["roll_stiffness_rear"]
    L = base["wheelbase"]
    base.update(cg_to_front=L / 2, cg_to_rear=L / 2, track_rear=base["track_front"],
                roll_stiffness_front=lambda_f * k, roll_stiffness_rear=(1 - lambda_f) * k,
                roll_center_rear=base["roll_center_front"], unsprung_mass_rear=base["unsprung_mass_front"],
                unsprung_cg_height_rear=base["unsprung_cg_height_front"])
    return VehicleConfig(**base)


# ----------------------------------------------------------------------------------------------------------------
# relatório
# ----------------------------------------------------------------------------------------------------------------

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--saida", type=Path, required=True, help="arquivo .md do relatório (não pode existir)")
    a = ap.parse_args()
    saida = a.saida if a.saida.is_absolute() else Path.cwd() / a.saida

    if saida.exists():
        return recusa(f"{saida} já existe; o relatório registrado nunca é regravado")
    sujo = git("status", "--porcelain")
    if sujo:
        return recusa("a árvore tem mudanças pendentes (git status --porcelain):\n" + sujo)
    exp, erro = ler_exp_json()
    if exp is None:
        return recusa(erro)

    motor, detalhe_motor = rc.escolher()
    if motor is None:
        print(f"NÃO EXECUTÁVEL: nenhum motor de recálculo ({detalhe_motor})", file=sys.stderr)
        return 3
    print(f"Motor: {motor} ({detalhe_motor})")

    data = dt.datetime.now().isoformat(timespec="seconds")
    commit = git("rev-parse", "HEAD")
    ramo = git("rev-parse", "--abbrev-ref", "HEAD")
    h_antes = sha256(VIVA)

    ds = tm.load_dataset(str(TIR), source_axis_system="ISO_TYDEX")
    p = ds.p
    var = json.loads(VARREDURA.read_text(encoding="utf-8"))
    fatores = var["fatores"]

    with tempfile.TemporaryDirectory(prefix="dds_rel_inc1_", ignore_cleanup_errors=True) as t:
        tmp = Path(t)
        copia_f = tmp / "viva_formulas.xlsx"     # só leitura de fórmulas e raízes (nunca recalculada)
        copia_r = tmp / "viva_recalculada.xlsx"  # recalculada pelo motor (EDR-044)
        shutil.copyfile(VIVA, copia_f)
        shutil.copyfile(VIVA, copia_r)
        t0 = time.perf_counter()
        resumo_rec = rc.recalcular(str(copia_r), motor, timeout=600)
        t_rec_viva = time.perf_counter() - t0
        caso = cp.caso_da_planilha(copia_r, TIR)
        pior_caso = cp.pior_origem_do_caso(copia_f)
        entradas = entradas_e_origens(copia_r, copia_f)
    h_depois = sha256(VIVA)
    if h_depois != h_antes:
        print(f"ERRO: a planilha viva mudou durante a execução (SHA-256 {h_antes} -> {h_depois}); "
              "relatório não gravado", file=sys.stderr)
        return 1

    # --- nominal ---
    res = eq.analise_par(caso, **GRADE)
    txt_nominal = eq.texto_veredito(res, pior_caso)
    fx = faixa(res)

    # --- varredura (EDR-055) ---
    pior_var = pior(pior_caso, var["origem"])
    t0 = time.perf_counter()
    varr = eq.varredura_pneu(caso, fatores, **GRADE)
    t_varr = time.perf_counter() - t0
    robusto = eq.veredito_robusto([r.veredito for _, r in varr])
    txt_robusto = eq.texto_veredito(res, pior_var, robusto)

    # --- nota de resolução: cenário do TC-EQ-011 (veículo de teste, LMUY 0,45 só de teste) ---
    ds_011 = dataclasses.replace(ds, p={**ds.p, "LMUY": 0.45})
    r011 = {}
    for lf in (0.5, 0.7):
        for passo in (0.05, 0.01):
            r011[(lf, passo)] = eq.analise_par(eq.CasoEquilibrio(cfg=veiculo_teste_011(lf), ds=ds_011),
                                               **{**GRADE, "passo_g": passo})

    textos = [txt_nominal, txt_robusto] + [eq.texto_veredito(r, pior_var) for _, r in varr]
    textos += [eq.texto_veredito(r, pior(ORIGEM_TESTE)) for r in r011.values()]

    md: list[str] = []
    w = md.append

    w("# V4.0, incremento 1: viabilidade da análise de par na planilha e robustez ao pneu")
    w("")
    w("Gerado por `tools/relatorio_inc1.py`. Não edite à mão: para atualizar, gere um relatório novo com outro nome.")
    w("")
    w("Todos os números deste relatório são **demonstrativos**: o caso vem da planilha V3.0 com valores do T09, "
      "estimativas e entradas arbitrárias, e o pneu é o `.tir` provisório do TTC (R25B), não o pneu do carro. "
      "Nada aqui é resultado do TX nem validação física (Validation = NOT_VALIDATED).")
    w("")

    # 1. Origem
    w("## 1. Origem")
    w("")
    w(f"- Data: {data}")
    w(f"- Commit: `{commit}` (ramo `{ramo}`; árvore sem mudanças pendentes ao iniciar)")
    w(f"- Planilha viva: `{VIVA.relative_to(RAIZ).as_posix()}`")
    w(f"  - SHA-256 antes: `{h_antes}`")
    w(f"  - SHA-256 depois: `{h_depois}` (igual: a viva não foi alterada; só uma cópia temporária foi recalculada)")
    w(f"- Motor de recálculo da cópia: {motor} ({detalhe_motor}); {resumo_rec['total_formulas']} fórmulas, "
      f"{resumo_rec['total_errors']} erros, {n(t_rec_viva, 1)} s")
    w(f"- Pneu: `{TIR.relative_to(RAIZ).as_posix()}` (SHA-256 `{sha256(TIR)}`), carregado em ISO_TYDEX (EDR-039)")
    w(f"- Varredura: `{VARREDURA.relative_to(RAIZ).as_posix()}` ({var['edr']}, origem \"{var['origem']}\")")
    w(f"- Experimento de viabilidade: `{EXP_JSON.relative_to(RAIZ).as_posix()}` (Task 6, `tests/exp_viabilidade_par.py`; "
      f"fora do Git; SHA-256 `{sha256(EXP_JSON)}`, modificado em "
      f"{dt.datetime.fromtimestamp(EXP_JSON.stat().st_mtime).isoformat(timespec='seconds')})")
    w(f"- Python {platform.python_version()}, openpyxl {version('openpyxl')}, numpy {version('numpy')}")
    w("")

    # 2. Caso
    w("## 2. Caso")
    w("")
    w(f"Entradas lidas da cópia recalculada (`caso_planilha.ENTRADAS_CASO`). Pior origem do caso: **{pior_caso}** "
      "(pior rótulo entre as raízes das entradas e o pneu; `caso_planilha.pior_origem_do_caso`).")
    w("")
    w("| Entrada | Célula | ID | Valor | Unidade | Pior origem das raízes |")
    w("|---|---|---|---|---|---|")
    for e in entradas:
        cal = " (fórmula)" if e["formula"] else ""
        w(f"| {e['rotulo']} | `{e['aba']}!{e['cel']}`{cal} | {e['id']} | {n(float(e['valor']), 4)} | "
          f"{e['unidade']} | {e['origem']} |")
    w("")
    cfg = caso.cfg
    w("Derivados usados pelo Python (`caso_planilha.caso_da_planilha`):")
    w("")
    w(f"- K_φ dianteiro = {n(cfg.roll_stiffness_front, 1)} N·m/rad; K_φ traseiro = {n(cfg.roll_stiffness_rear, 1)} N·m/rad "
      f"(λ_f = {n(cfg.roll_stiffness_front / (cfg.roll_stiffness_front + cfg.roll_stiffness_rear), 3)})")
    w(f"- Downforce = {n(caso.downforce_N, 1)} N e balanço dianteiro = {n(caso.aero_balance_front, 3)}, fixos (EDR-013)")
    for nome, d in (("Toe estático", caso.toe_rad), ("Camber estático", caso.camber_estatico_rad),
                    ("Ganho de camber", caso.ganho_camber)):
        estado = "sem dado (EDR-036); usado 0 nas quatro rodas" if all(v == 0.0 for v in d.values()) else \
            ", ".join(f"{c} = {n(v, 4)}" for c, v in d.items())
        w(f"- {nome}: {estado}")
    w(f"- Pneu: FZMIN = {n(p['FZMIN'], 0)} N, FZMAX = {n(p['FZMAX'], 0)} N, ALPMIN/ALPMAX = "
      f"{n(p['ALPMIN'], 4)}/{n(p['ALPMAX'], 4)} rad, CAMMIN/CAMMAX = {n(p['CAMMIN'], 4)}/{n(p['CAMMAX'], 4)} rad, "
      f"FNOMIN = {n(p['FNOMIN'], 0)} N, LFZO = {n(p['LFZO'], 2)}, LMUY = {n(p['LMUY'], 2)}, LKY = {n(p['LKY'], 2)}")
    w(f"- Grade da análise de par (valores passados a `analise_par`): a_y de 0 a {n(res.linhas[-1].ay_ms2 / G, 2)} g em "
      f"passos de {n(res.passo_ms2 / G, 2)} g ({len(res.linhas)} linhas); {GRADE['n_alfa']} valores de α por eixo; "
      "κ = 0 (sem tração)")
    w("")
    t74 = tabela_74()
    w("### Hipóteses incluídas no incremento 1 (spec, tabela 7.4)")
    w("")
    w("| # | Hipótese | Base | Como entra no incremento 1 |")
    w("|---|---|---|---|")
    for h in t74:
        if h["num"] in INCLUIDAS:
            extra = ""
            if h["num"] == "2":
                extra = " Toe: " + ("sem dado (EDR-036)." if all(v == 0.0 for v in caso.toe_rad.values())
                                   else "digitado.")
            if h["num"] == "4":
                zero = all(v == 0.0 for v in caso.camber_estatico_rad.values()) and \
                    all(v == 0.0 for v in caso.ganho_camber.values())
                extra = " Camber: " + ("sem dado (EDR-036)." if zero else "digitado.")
            w(f"| {h['num']} | {h['hipotese']} | {h['base']} | {INCLUIDAS[h['num']]}{extra} |")
    w(f"| 8 (parte) | Aero fixo por Setup | EDR-013 | Downforce e balanço da aba 08_VEHICLE_STATE, constantes em a_y. |")
    w("")
    w("### Hipóteses não incluídas (incremento 2 ou depois)")
    w("")
    w("O efeito é o da coluna \"Efeito se ignorada\" da tabela 7.4 da spec, copiado do arquivo por este script.")
    w("")
    w("| # | Hipótese | Base | Efeito se ignorada |")
    w("|---|---|---|---|")
    for h in t74:
        if h["num"] in INCLUIDAS and h["num"] not in PARCIAIS:
            continue
        nome = PARCIAIS[h["num"]] if h["num"] in PARCIAIS else h["hipotese"]
        rot = f"{h['num']} (parte)" if h["num"] in PARCIAIS else h["num"]
        w(f"| {rot} | {nome} | {h['base']} | {h['efeito']} |")
    w("")

    # 3. Resultado nominal
    w("## 3. Resultado nominal (pneu sem escala)")
    w("")
    w(f"- Veredito: **{txt_nominal}**")
    w(f"- a_y de saturação do eixo dianteiro: {g_de(res.ay_sat_f, 4)}")
    w(f"- a_y de saturação do eixo traseiro: {g_de(res.ay_sat_r, 4)}")
    w(f"- a_y_lim: {g_de(res.ay_lim, 4)}")
    w(f"- C_N: {n(res.c_n, 5) if res.c_n is not None else '— (só existe em PUSH ou LOOSE)'}")
    if fx is None:
        w("- Nenhuma roda sai da faixa do pneu na grade.")
        w("")
    else:
        ln = fx["linha"]
        w(f"- Primeira a_y em que alguma roda sai da faixa do pneu: **{g_de(ln.ay_ms2, 2)}** "
          f"({'; '.join(fx['motivos'])})")
        w("")
        fora = rodas_fora(ln, p)
        lv = fx["ultima_valida"]
        w("Cargas verticais por roda na transição (curva à esquerda: FR e RR são as externas):")
        w("")
        w("| a_y | Fz_FL [N] | Fz_FR [N] | Fz_RL [N] | Fz_RR [N] | Faixa do pneu |")
        w("|---|---|---|---|---|---|")
        for lx, rotulo in ((lv, "última linha válida"), (ln, "primeira linha fora")):
            if lx is None:
                continue
            w(f"| {g_de(lx.ay_ms2, 2)} ({rotulo}) | " + " | ".join(n(lx.fz[c], 1) for c in eq.CANTOS)
              + f" | [{n(p['FZMIN'], 0)}, {n(p['FZMAX'], 0)}] N |")
        w("")
        limites = {lim for _, _, lim in fora}
        if fora:
            desc = "; ".join(f"Fz_{c} = {n(fz, 1)} N passa de {lim} = {n(p[lim], 0)} N" for c, fz, lim in fora)
            w(f"Causa: {desc}.")
            if limites == {"FZMAX"}:
                interna_min = min(ln.fz["FL"], ln.fz["RL"])
                w(f"A roda que sai da faixa é uma roda **externa**, por excesso de carga. A roda interna mais "
                  f"leve tem {n(interna_min, 1)} N nessa linha, acima de FZMIN = {n(p['FZMIN'], 0)} N.")
                w("")
                w("**Correção do plano.** Segundo o plano (`docs/superpowers/plans/2026-10-20-v4-0-inc1-viabilidade-"
                  "robustez.md`, seção \"Risco conhecido\"), a roda RL do veículo de teste fica abaixo de 100 N perto de "
                  "1,55 g e o resultado mais provável é N/A por faixa do pneu, também no caso da planilha; e a "
                  "decisão (c) do plano é sobre a roda interna abaixo de FZMIN (texto do plano, não reverificado aqui). "
                  f"No caso da planilha viva a causa é outra: {', '.join(c for c, _, _ in fora)} "
                  f"({'externa' if all(c in ('FR', 'RR') for c, _, _ in fora) else 'ver tabela'}) passa de FZMAX.")
            w("")
        if res.veredito == eq.NA and fx["fracao"] is not None:
            fr = fx["fracao"]
            w(f"Fração da capacidade usada por eixo na última linha válida ({g_de(lv.ay_ms2, 2)}), "
              "req/cap. Indica qual eixo está mais perto do limite, sem extrapolar o pneu. **Não é veredito.**")
            w("")
            w("| Eixo | Força exigida [N] | Capacidade [N] | req/cap |")
            w("|---|---|---|---|")
            w(f"| Dianteiro | {n(lv.req_f, 1)} | {n(lv.cap_f.fy_max_N, 1)} | {pct(fr['F'])} |")
            w(f"| Traseiro | {n(lv.req_r, 1)} | {n(lv.cap_r.fy_max_N, 1)} | {pct(fr['R'])} |")
            w("")
    w("Linhas da grade até a primeira linha fora da faixa (capacidade = máximo de Fy do eixo na grade de α):")
    w("")
    w("| a_y [g] | Fz_FL | Fz_FR | Fz_RL | Fz_RR | req_F [N] | cap_F [N] | req_R [N] | cap_R [N] | Observação |")
    w("|---|---|---|---|---|---|---|---|---|---|")
    ate = fx["k"] if fx else len(res.linhas) - 1
    for lx in res.linhas[:ate + 1]:
        obs = "; ".join(c.motivo for c in (lx.cap_f, lx.cap_r) if c.motivo) or ""
        w(f"| {n(lx.ay_ms2 / G, 2)} | " + " | ".join(n(lx.fz[c], 0) for c in eq.CANTOS)
          + f" | {n(lx.req_f, 0)} | {n(lx.cap_f.fy_max_N, 0)} | {n(lx.req_r, 0)} | {n(lx.cap_r.fy_max_N, 0)} | {obs} |")
    w("")
    w("### Decisão pendente do usuário: política da faixa do pneu (\"Risco conhecido\", item c)")
    w("")
    w("O plano deixou para o usuário como tratar uma roda fora da faixa de carga do pneu. As opções escritas no "
      "plano são: (1) Fy = 0 só quando Fz ≤ 0 (descolamento, que é física e não extrapolação) e N/A entre 0 e "
      "FZMIN; (2) manter o N/A. Esta decisão não foi tomada neste incremento.")
    w("")
    if fx is not None and {lim for _, _, lim in rodas_fora(fx["linha"], p)} == {"FZMAX"}:
        w("Com os números acima, a opção (1) não muda o resultado do caso da planilha: ela trata a roda interna "
          "com pouca carga, e aqui a análise para porque a roda externa passa de FZMAX. Qualquer política que "
          "permita seguir além desse ponto teria de dizer o que fazer acima de FZMAX, o que seria extrapolar o pneu; "
          "a decisão continua sendo do usuário.")
        w("")
    w("### Nota sobre a resolução da grade de a_y")
    w("")
    w("Cenário do TC-EQ-011 de `tests/test_equilibrio.py`, recalculado aqui: veículo de teste simétrico "
      "(`veiculo_simetrico`, NÃO é setup real) com LMUY = 0,45 (escala só de teste, não é pneu).")
    w("")
    w("| λ_f | Passo [g] | a_y sat. F | a_y sat. R | Diferença | Veredito |")
    w("|---|---|---|---|---|---|")
    for (lf, passo), r in r011.items():
        dif = (abs(r.ay_sat_f - r.ay_sat_r) / G) if (r.ay_sat_f is not None and r.ay_sat_r is not None) else None
        w(f"| {n(lf, 1)} | {n(passo, 2)} | {g_de(r.ay_sat_f, 4)} | {g_de(r.ay_sat_r, 4)} | "
          f"{'—' if dif is None else n(dif, 4) + ' g'} | {eq.texto_veredito(r, pior(ORIGEM_TESTE))} |")
    r5, r7 = r011[(0.5, 0.01)], r011[(0.7, 0.01)]
    if None not in (r5.ay_sat_f, r5.ay_sat_r, r7.ay_sat_f, r7.ay_sat_r):
        w("")
        w(f"De λ_f = 0,5 para 0,7 (passo 0,01 g), a saturação dianteira muda {n((r7.ay_sat_f - r5.ay_sat_f) / G, 4)} g "
          f"e a traseira {n((r7.ay_sat_r - r5.ay_sat_r) / G, 4)} g; a diferença entre os eixos fica em "
          f"{n(abs(r7.ay_sat_f - r7.ay_sat_r) / G, 4)} g. Com o passo padrão de 0,05 g, uma diferença menor que o "
          "passo dá NEUTRO (critério |S_f − S_r| < passo de `analise_par`). O passo padrão não separa um efeito "
          "desse tamanho.")
    w("")

    # 4. Varredura
    w("## 4. Robustez ao pneu: varredura exploratória (EDR-055)")
    w("")
    w(f"Fatores λ do PAC2002 multiplicados pelos valores de `{VARREDURA.relative_to(RAIZ).as_posix()}`: "
      + "; ".join(f"{k} × {', '.join(n(v, 1) for v in vs)}" for k, vs in fatores.items())
      + f" ({len(varr)} combinações). Rótulo: \"{var['origem']}\" ({var['nota']}). Os fatores não alteram "
      f"FZMIN/FZMAX. Tempo do Python para as {len(varr)} análises: {n(t_varr, 1)} s.")
    w("")
    w("| " + " | ".join(fatores) + " | Veredito | a_y_lim | Primeira a_y fora da faixa | req/cap F | req/cap R |")
    w("|" + "---|" * (len(fatores) + 5))
    for comb, r in varr:
        f2 = faixa(r)
        frc = f2["fracao"] if (f2 and r.veredito == eq.NA) else None
        w("| " + " | ".join(n(comb[k], 1) for k in fatores)
          + f" | {eq.texto_veredito(r, pior_var)} | {g_de(r.ay_lim, 4)} | "
          f"{g_de(f2['linha'].ay_ms2, 2) if f2 else '—'} | "
          f"{pct(frc['F']) if frc else '—'} | {pct(frc['R']) if frc else '—'} |")
    w("")
    contagem = {}
    for _, r in varr:
        contagem[r.veredito] = contagem.get(r.veredito, 0) + 1
    w("Contagem: " + ", ".join(f"{k} = {v}" for k, v in sorted(contagem.items())) + ".")
    w("")
    w(f"Veredito robusto (`veredito_robusto`): **{robusto[0]}** ({robusto[1]}).")
    w("")
    w(f"Texto de exibição com a robustez (`texto_veredito`, pior origem {pior_var}): **{txt_robusto}**")
    w("")

    # 5. Viabilidade
    w("## 5. Viabilidade da análise de par em fórmulas (Task 6)")
    w("")
    w(f"Dados de `{EXP_JSON.relative_to(RAIZ).as_posix()}`. O protótipo é uma pasta de trabalho gerada pelo openpyxl numa pasta "
      "temporária (não é planilha da equipe), com uma grade "
      f"de {exp['grade']['ay_max_g']} g / {exp['grade']['passo_g']} g × {exp['grade']['n_alfa']} valores de α. "
      "\"Base\" é o recálculo da pasta de trabalho só com a aba de coeficientes; \"recálculo\" é o da cheia. "
      f"Tolerância da comparação: {exp['tolerancia']:g} (a mesma da CA-09). Casos `teste_*` usam o veículo de teste "
      "e escalas só de teste (não são o carro).")
    w("")
    w("| Motor | Caso | Linhas | Células de fórmula | Base [s] | Recálculo [s] | Leitura openpyxl [s] | Erros | "
      "Dif. rel. máx. capacidade | Dif. a_y sat. F | Dif. a_y sat. R | Picos na borda (planilha/Python) | Aprovado |")
    w("|---|---|---|---|---|---|---|---|---|---|---|---|---|")
    execs = exp["execucoes"]
    for m in MOTORES:
        for c in CASOS_EXP:
            x = next(e for e in execs if e["motor"] == m and e["caso"] == c)

            def ex(v):
                return "—" if v is None else f"{v:.1e}"
            w(f"| {m} | `{c}` | {x['linhas']} | {x['celulas_formula']} | {n(x['t_base_s'], 3)} | "
              f"{n(x['t_recalculo_s'], 3)} | {n(x['t_resumo_openpyxl_s'], 3)} | {x['erros_formula']} | "
              f"{ex(x['dif_rel_max_capacidade'])} | {ex(x['dif_ay_sat_f'])} | {ex(x['dif_ay_sat_r'])} | "
              f"{x['pico_na_borda']}/{x['pico_na_borda_python']} | {'sim' if x['aprovado'] else 'NÃO'} |")
    w("")
    viva_exp = [next(e for e in execs if e["motor"] == m and e["caso"] == "planilha_viva") for m in MOTORES]
    coincide = all(e["veredito_python"] == res.veredito and e["motivo_python"] == res.motivo for e in viva_exp)
    w("Conferência: o veredito e o motivo do caso `planilha_viva` no JSON "
      + ("**coincidem**" if coincide else "**NÃO coincidem**")
      + f" com a análise deste relatório ({res.veredito}; {res.motivo}).")
    w("")
    t_rec = {m: [e["t_recalculo_s"] for e in execs if e["motor"] == m] for m in MOTORES}

    # 6. Vai ou não vai
    w("## 6. Vai ou não vai: os números")
    w("")
    w("Não há limiar registrado para decidir (spec §9, item 10: limite de desempenho só por decisão registrada). "
      "Os números abaixo não são veredito de viabilidade.")
    w("")
    w(f"- Veredito nominal do caso da planilha: {txt_nominal}")
    w(f"- Veredito com a varredura ({len(varr)} combinações): {txt_robusto} (veredito robusto: {robusto[0]}; nota: {robusto[1]})")
    if fx is not None:
        w(f"- Primeira a_y fora da faixa do pneu: {g_de(fx['linha'].ay_ms2, 2)}; "
          + "; ".join(f"Fz_{c} = {n(fz, 1)} N > {lim}" if lim == "FZMAX" else f"Fz_{c} = {n(fz, 1)} N < {lim}"
                      for c, fz, lim in rodas_fora(fx["linha"], p)))
        if fx["fracao"]:
            w(f"- req/cap em {g_de(fx['ultima_valida'].ay_ms2, 2)}: dianteiro {pct(fx['fracao']['F'])}, "
              f"traseiro {pct(fx['fracao']['R'])}")
    for m in MOTORES:
        w(f"- Recálculo do protótipo ({m}): {n(min(t_rec[m]), 2)} a {n(max(t_rec[m]), 2)} s por caso; "
          f"{max(e['celulas_formula'] for e in execs if e['motor'] == m)} células de fórmula; "
          f"{sum(e['erros_formula'] for e in execs if e['motor'] == m)} erros; diferença relativa máxima da "
          f"capacidade {max(e['dif_rel_max_capacidade'] for e in execs if e['motor'] == m):.1e}")
    w(f"- Casos aprovados no experimento: {sum(e['aprovado'] for e in execs)} de {len(execs)}")
    celulas = max(e["celulas_formula"] for e in execs)
    w(f"- Estimativa linear, **não medida**: a varredura de {len(varr)} combinações em fórmulas teria da ordem de "
      f"{len(varr)} × {celulas} = {len(varr) * celulas} células de fórmula.")
    w(f"- Varredura no Python: {n(t_varr, 1)} s para {len(varr)} análises")
    w("")
    w("### Recomendação (não é decisão)")
    w("")
    w("O que os números sugerem, para o usuário decidir:")
    w("")
    tudo_ok = all(e["aprovado"] and e["erros_formula"] == 0 for e in execs)
    if tudo_ok:
        w(f"1. A análise de par cabe em fórmulas nos dois motores: as fórmulas reproduziram o Python em todos os "
          f"casos, com 0 erros. O recálculo do protótipo levou cerca de {n(max(t_rec['excel']), 1)} s no Excel e "
          f"{n(max(t_rec['libreoffice']), 1)} s no LibreOffice. Se esse tempo é aceitável é decisão da equipe.")
    else:
        w("1. O experimento não reproduziu o Python em todos os casos (ver a seção 5); antes de levar a análise "
          "para a planilha, a divergência precisa ser investigada na ordem do CLAUDE.md.")
    w(f"2. A varredura em fórmulas multiplicaria o protótipo por {len(varr)} (estimativa acima, não medida). A spec "
      "(7.6, item 4) prevê a alternativa de o Python gerar o relatório da varredura e a planilha mostrar só o "
      "veredito nominal com o aviso \"robustez ao pneu: ver relatório\".")
    if res.veredito == eq.NA and robusto[0] == eq.NA and fx is not None:
        lims = sorted({lim for _, _, lim in rodas_fora(fx["linha"], p)})
        rodas = ", ".join(f"Fz_{c} ({lim})" for c, _, lim in rodas_fora(fx["linha"], p))
        sem_sat = res.ay_sat_f is None and res.ay_sat_r is None
        w(f"3. Com os dados demonstrativos atuais e o `.tir` provisório, o incremento 1 não produz veredito de "
          f"equilíbrio no limite para o caso da planilha: o resultado é N/A no nominal e nas {len(varr)} combinações. "
          f"No nominal, a primeira a_y fora da faixa é {g_de(fx['linha'].ay_ms2, 2)}, por {rodas}"
          + (", e nenhum eixo saturou nas linhas válidas anteriores. " if sem_sat else ". ")
          + ("A fração req/cap na última linha válida mostra o quanto falta, sem dar o sinal de push ou loose. "
             if fx["fracao"] else "")
          + ("A causa é o limite FZMAX; a opção (1) do plano, que só trata Fz ≤ 0, não muda esse resultado. "
             if lims == ["FZMAX"] else
             "A opção (1) do plano só trata Fz ≤ 0; confira na tabela da seção 3 se ela cobre essa causa. ")
          + "Antes de usar o veredito, o usuário precisa decidir a política da faixa do pneu, ou usar o pneu do "
            "carro quando houver dado dele (hoje o .tir é o R25B provisório, EDR-022).")
    w("")
    w("A decisão de seguir, ajustar ou parar é do usuário.")
    w("")

    # conferências finais antes de gravar
    sem_prefixo = [t_ for t_ in textos if pior_var != TX and not t_.startswith(PREFIXO)]
    if pior_caso != TX and not txt_nominal.startswith(PREFIXO):
        sem_prefixo.append(txt_nominal)
    if sem_prefixo:
        print(f"ERRO: veredito sem o prefixo {PREFIXO!r}: {sem_prefixo[:3]}", file=sys.stderr)
        return 1

    saida.parent.mkdir(parents=True, exist_ok=True)
    with open(saida, "x", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(md) + "\n")
    print(f"Veredito nominal: {txt_nominal}")
    print(f"Veredito robusto: {robusto} -> {txt_robusto}")
    print(f"Planilha viva intacta (SHA-256): {h_antes == h_depois}")
    print(f"Relatório: {saida}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
