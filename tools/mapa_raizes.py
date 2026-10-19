"""Mapa de raízes da planilha V3.0: para cada fórmula, as entradas de que ela depende e a pior origem entre elas.

Refaz, a partir das fórmulas e da coluna Origem atuais (EDR-035), o mapa da auditoria V2.1
(data/placeholder-dependency-map-V2.1.csv), que não conhece as fórmulas da V3.0. É a base da fase 2 do FB-03
(pior origem ao vivo no 14_DASHBOARD), que só entra na planilha com aprovação.

Uso:
    uv run python -X utf8 tools/mapa_raizes.py [V3_0.xlsx] [--saida data/placeholder-dependency-map-V3.0.csv]

Regras (nada é inferido):
- Raiz = célula referenciada que não é fórmula. Constantes dentro da fórmula (9.81, 0.0254) não são raízes.
- Origem da raiz: coluna E da linha nas abas 05/06/07/08/11; na 10_TIRE, coluna I da linha (casos 23-25) e I84 para
  os coeficientes B86:B175. Qualquer outra raiz fica "sem rótulo" e é listada como tal.
- "Pior origem": a mais distante de "dado do carro", na ordem de PIOR_PRIMEIRO. A ordem dentro do mesmo tier da
  auditoria V2.1 (T1: Arbitrário e Estimativa; T2: T09 e TTC) é PROPOSTA, não aprovada; por isso o CSV guarda
  também o tier e a contagem por rótulo.
"""
import argparse
import csv
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "python"))
from atualizar_planilha import ABAS_ORIGEM, ORIGENS_LISTA, TX, T09, TTC, EST, ARB, CONF  # noqa: E402

SEM_ROTULO = "sem rótulo de origem"
PIOR_PRIMEIRO = (CONF, ARB, EST, TTC, T09, TX)   # "? A confirmar" primeiro: origem desconhecida não conta como boa
TIER = {TX: "T3", T09: "T2", TTC: "T2", EST: "T1", ARB: "T1", CONF: "?"}
COEF_TIRE = range(86, 176)                       # 10_TIRE: coeficientes do .tir; origem em I84
COLUNA = re.compile(r"^(?:(?P<aba>'[^']+'|[^!']+)!)?\$?(?P<c1>[A-Z]{1,3}):\$?(?P<c2>[A-Z]{1,3})$")
REF = re.compile(r"^(?:(?P<aba>'[^']+'|[^!']+)!)?\$?(?P<c1>[A-Z]{1,3})\$?(?P<r1>\d+)(?::\$?(?P<c2>[A-Z]{1,3})\$?(?P<r2>\d+))?$")


def referencias(formula: str, aba_atual: str, colunas: list | None = None) -> list[tuple[str, str]]:
    """Células referenciadas. Coluna inteira (ex.: COUNTIF de IDs) não vira raiz: vai para `colunas`."""
    out = []
    for tok in Tokenizer(formula).items:
        if tok.type != "OPERAND" or tok.subtype != "RANGE":
            continue
        m = REF.match(tok.value)
        if not m:
            mc = COLUNA.match(tok.value)
            if mc is None:
                raise ValueError(f"referência não reconhecida {tok.value!r} em {formula!r}")
            if colunas is not None:
                colunas.append(f"{(mc['aba'] or aba_atual).strip(chr(39))}!{mc['c1']}:{mc['c2']}")
            continue
        aba = (m["aba"] or aba_atual).strip("'")
        c1, r1 = column_index_from_string(m["c1"]), int(m["r1"])
        c2, r2 = (column_index_from_string(m["c2"]), int(m["r2"])) if m["c2"] else (c1, r1)
        for c in range(min(c1, c2), max(c1, c2) + 1):
            for r in range(min(r1, r2), max(r1, r2) + 1):
                out.append((aba, f"{get_column_letter(c)}{r}"))
    return out


def origem_da_raiz(wb, aba: str, cel: str) -> str:
    r = int(re.sub(r"[A-Z]", "", cel))
    if aba in ABAS_ORIGEM:
        v = wb[aba][f"E{r}"].value
    elif aba == "10_TIRE":
        v = wb[aba][f"I{84 if r in COEF_TIRE else r}"].value
    else:
        v = None
    return v if v in ORIGENS_LISTA else SEM_ROTULO


def mapa(caminho: Path) -> list[dict]:
    wb = openpyxl.load_workbook(caminho)
    formulas = {(ws.title, c.coordinate): c.value for ws in wb.worksheets for linha in ws.iter_rows() for c in linha
                if isinstance(c.value, str) and c.value.startswith("=")}
    memo: dict = {}
    colunas_por_formula: dict = {}

    def raizes(no, pilha=()):
        if no in memo:
            return memo[no]
        if no not in formulas:
            return {no}
        if no in pilha:
            raise RuntimeError(f"referência circular em {no}")
        acc = set()
        cols = colunas_por_formula.setdefault(no, [])
        for ref in referencias(formulas[no], no[0], cols):
            acc |= raizes(ref, pilha + (no,))
        memo[no] = acc
        return acc

    linhas = []
    for (aba, cel), f in sorted(formulas.items(), key=lambda kv: (kv[0][0], int(re.sub(r"[A-Z]", "", kv[0][1])), kv[0][1])):
        rz = sorted(raizes((aba, cel)), key=lambda x: (x[0], int(re.sub(r"[A-Z]", "", x[1])), x[1]))
        vazias = [x for x in rz if wb[x[0]][x[1]].value is None]
        rz = [x for x in rz if wb[x[0]][x[1]].value is not None]
        org = {x: origem_da_raiz(wb, *x) for x in rz}
        rotuladas = [o for o in org.values() if o != SEM_ROTULO]
        pior = next((o for o in PIOR_PRIMEIRO if o in rotuladas), "")
        linhas.append({
            "sheet": aba, "cell": cel, "label": wb[aba][f"A{int(re.sub(r'[A-Z]', '', cel))}"].value or "",
            "pior_origem": pior, "tier_pior": TIER.get(pior, ""), "n_raizes": len(rz),
            **{f"n_{o}": rotuladas.count(o) for o in PIOR_PRIMEIRO},
            "n_sem_rotulo": len(rz) - len(rotuladas),
            "raizes_nao_TX": "; ".join(f"{a}!{c}[{org[(a, c)]}]" for a, c in rz if org[(a, c)] not in (TX, SEM_ROTULO)),
            "raizes_sem_rotulo": "; ".join(f"{a}!{c}" for a, c in rz if org[(a, c)] == SEM_ROTULO),
            "refs_vazias": "; ".join(f"{a}!{c}" for a, c in vazias),
            "refs_coluna_inteira": "; ".join(colunas_por_formula.get((aba, cel), [])),
            "formula": f,
        })
    return linhas


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("planilha", nargs="?", default=str(RAIZ / "excel" / "FSAE_EV_Suspension_Fase1_V3_0.xlsx"))
    ap.add_argument("--saida", default=str(RAIZ / "build" / "placeholder-dependency-map-V3.0.csv"))
    a = ap.parse_args()
    linhas = mapa(Path(a.planilha))
    Path(a.saida).parent.mkdir(parents=True, exist_ok=True)
    with open(a.saida, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(linhas[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(linhas)
    print("fórmulas", len(linhas), "->", a.saida)


if __name__ == "__main__":
    main()
