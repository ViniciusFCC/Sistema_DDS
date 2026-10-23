"""Relatório do incremento 4 da V4.1 (equilíbrio na planilha): caso da V4.1, resultado nominal, limitações, robustez ao
pneu (EDR-055) e tempos de recálculo.

É o documento que a célula de robustez da 16_EQUILIBRIO cita (aba_equilibrio.AVISO_ROBUSTEZ): a planilha mostra o
veredito nominal, e a varredura do pneu (27 combinações, EDR-055) roda aqui, no Python, sobre o caso que a planilha lê
(caso_planilha.caso_da_planilha_v41). Nenhum número do Python é copiado para a planilha (spec 7.5).

Uso (Git Bash, depois de export UV_SYSTEM_CERTS=1):
    uv run python -X utf8 tools/relatorio_inc4.py [--saida docs/v4-1/inc4-equilibrio-na-planilha.md] [--sobrescrever]

Regras:
- `excel/FSAE_EV_Suspension_Fase1_V4_1.xlsx` só é COPIADA para uma pasta temporária; só a cópia é recalculada (EDR-044).
  O SHA-256 da referência é conferido antes e depois; se mudar, o relatório não é gravado.
- Todo número sai da cópia recalculada, do Python (equilibrio.py, equilibrio_linear.py, caso_planilha.py), de
  `data/varredura_pneu_exploratoria.json` (EDR-055) ou das medições do portão da Task 4 (docs/v4-1/). Valores
  demonstrativos (T09/E03, pneu provisório do TTC) saem sempre rotulados como tais; nada aqui é resultado do carro.
- O arquivo não é regravado sem --sobrescrever. Gravado em UTF-8 com LF.

Saída: 0 = gravado; 2 = recusa; 3 = NÃO EXECUTÁVEL (sem motor de recálculo); 1 = a referência mudou durante a execução.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import platform
import shutil
import subprocess
import sys
import tempfile
import time
from importlib.metadata import version
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "tests"))   # recalculo_planilha mora em tests/

import openpyxl  # noqa: E402

import aba_equilibrio as ae  # noqa: E402
import caso_planilha as cp  # noqa: E402
import equilibrio as eq  # noqa: E402
import equilibrio_linear as el  # noqa: E402
import origens  # noqa: E402
import recalculo_planilha as rc  # noqa: E402
from raizes import PIOR_PRIMEIRO  # noqa: E402
from suspension_model import G  # noqa: E402

V41 = RAIZ / "excel" / "FSAE_EV_Suspension_Fase1_V4_1.xlsx"
TIR = RAIZ / "data" / "Hoosier_R25B_13x7_TTC_provisional.tir"
VARREDURA = RAIZ / "data" / "varredura_pneu_exploratoria.json"
PORTAO = RAIZ / "docs" / "v4-1" / "inc4-viabilidade-tabela-de-troca.md"
SAIDA_PADRAO = RAIZ / "docs" / "v4-1" / "inc4-equilibrio-na-planilha.md"
GRADE = {"ay_max_g": 2.0, "passo_g": 0.05, "n_alfa": 25}   # a grade da 17 (analise_par e varredura_pneu)


def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for bloco in iter(lambda: f.read(1 << 20), b""):
            h.update(bloco)
    return h.hexdigest()


def git(*args: str) -> str:
    r = subprocess.run(["git", *args], cwd=RAIZ, capture_output=True, text=True, encoding="utf-8")
    return r.stdout.strip() if r.returncode == 0 else f"indisponível ({r.stderr.strip()})"


def n(x, casas: int = 3) -> str:
    """Número com vírgula decimal; None vira '—'; texto passa como está."""
    if x is None:
        return "—"
    if isinstance(x, str):
        return x
    return f"{x:.{casas}f}".replace(".", ",")


def sci(x) -> str:
    return "—" if x is None else f"{x:.2e}".replace(".", ",")


def g_txt(v, casas: int = 4) -> str:
    """Valor já em g da planilha: número com unidade, ou o texto (N/A) como está."""
    return v if isinstance(v, str) else f"{n(v, casas)} g"


def g_de(ay_ms2, casas: int = 3) -> str:
    return "—" if ay_ms2 is None else f"{n(ay_ms2 / G, casas)} g"


def pior(*rotulos: str) -> str:
    for r in rotulos:
        if r not in PIOR_PRIMEIRO:
            raise ValueError(f"rótulo de origem fora do vocabulário: {r!r}")
    return min(rotulos, key=PIOR_PRIMEIRO.index)


def valor(wb, nome: str):
    aba, cel = ae._destino(wb, nome)
    return wb[aba][cel].value


def caso_com_karb(copia: Path, fator: float):
    """caso_da_planilha_v41 com o K_ARB dianteiro da 06 × fator (mesma troca do TC-WB41-006); a cópia não muda."""
    orig = cp._cfg
    cp._cfg = lambda x: orig({**x, "KarbF": x["KarbF"] * fator})
    try:
        return cp.caso_da_planilha_v41(copia, TIR)
    finally:
        cp._cfg = orig


def linha_troca(copia: Path, fator: float) -> dict:
    """Esperado do Python para uma linha da tabela de troca (chaves como em ae.TAB20; None = N/A)."""
    try:
        caso = caso_com_karb(copia, fator)
    except ValueError as e:
        return {"invalida": str(e)}
    r = eq.analise_par(caso, **GRADE)
    p = el.parcelas_k(caso, 0.0)
    kf, kr = caso.cfg.roll_stiffness_front, caso.cfg.roll_stiffness_rear
    return {"kphi": kf, "lam": kf / (kf + kr), "k": p.k_total, "sf": r.ay_sat_f, "sr": r.ay_sat_r, "aylim": r.ay_lim,
            "ver": r.veredito, "cn": r.c_n, "mot": r.motivo or "-"}


def neutro(linhas: list[dict]):
    """a_y do acerto neutro pela regra da 20_EQ_TROCA (spec 5, item 4): 1º par de linhas vizinhas em que S_r − S_f troca
    de sinal (interpolação linear), ou vale 0; senão a linha NEUTRO de menor |S_r − S_f|; senão None."""
    d = [ln["sr"] - ln["sf"] if ln.get("sf") is not None and ln.get("sr") is not None else None for ln in linhas]
    for j, dj in enumerate(d):
        if dj is None:
            continue
        if dj == 0:
            return linhas[j]["sf"], f"S_r = S_f na linha {j}"
        if j + 1 < len(d) and d[j + 1] is not None and dj * d[j + 1] < 0:
            sf, sf1 = linhas[j]["sf"], linhas[j + 1]["sf"]
            return sf + (sf1 - sf) * dj / (dj - d[j + 1]), f"interpolação entre as linhas {j} e {j + 1}"
    neu = [(abs(dj), j) for j, dj in enumerate(d) if dj is not None and linhas[j]["ver"] == eq.NEUTRO]
    if neu:
        j = min(neu)[1]
        return linhas[j]["aylim"], f"linha NEUTRO {j}"
    return None, "neutro fora da faixa da tabela"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--saida", type=Path, default=SAIDA_PADRAO)
    ap.add_argument("--sobrescrever", action="store_true", help="regrava um relatório que já existe")
    a = ap.parse_args()
    saida = a.saida if a.saida.is_absolute() else Path.cwd() / a.saida
    if saida.exists() and not a.sobrescrever:
        print(f"RECUSADO: {saida} já existe (use --sobrescrever)", file=sys.stderr)
        return 2
    if not V41.exists():
        print(f"RECUSADO: falta {V41}", file=sys.stderr)
        return 2
    motor, detalhe_motor = rc.escolher()
    if motor is None:
        print(f"NÃO EXECUTÁVEL: nenhum motor de recálculo ({detalhe_motor})", file=sys.stderr)
        return 3
    print(f"Motor: {motor} ({detalhe_motor})")

    data = dt.datetime.now().isoformat(timespec="seconds")
    commit = git("rev-parse", "HEAD")
    ramo = git("rev-parse", "--abbrev-ref", "HEAD")
    sujo = bool(git("status", "--porcelain", "--untracked-files=no"))
    h_antes = sha256(V41)
    var = json.loads(VARREDURA.read_text(encoding="utf-8"))
    fatores = var["fatores"]

    with tempfile.TemporaryDirectory(prefix="dds_rel_inc4_", ignore_cleanup_errors=True) as t:
        tmp = Path(t)
        copia_f = tmp / "v41_formulas.xlsx"       # só leitura de fórmulas e raízes (nunca recalculada)
        copia_r = tmp / "v41_recalculada.xlsx"    # recalculada pelo motor (EDR-044)
        shutil.copyfile(V41, copia_f)
        shutil.copyfile(V41, copia_r)
        t0 = time.perf_counter()
        resumo_rec = rc.recalcular(str(copia_r), motor, timeout=900)
        t_rec = time.perf_counter() - t0
        wb = openpyxl.load_workbook(copia_r, data_only=True)
        x = cp.valores_v41(wb)
        pl = {k: valor(wb, nome) for k, nome in ae.NOME16.items()
              if k in ("status", "ver", "mot", "aylim", "S_F", "S_R", "resol", "cn", "lims", "tvolta", "pior", "kTotg",
                       "kmot", "vchar", "vcrit", "txt_k", "txt_ver", "apoio", "txt_sent")}
        tab = [{q: valor(wb, ae.NOME20[(q, j)]) for q, _, _ in ae.TAB20 if q != "fator"} for j in range(len(ae.FATORES20))]
        custo = {q: valor(wb, nm) for q, nm in ae.NOME20_CUSTO.items()}
        wb.close()
        pior_caso = cp.pior_origem_do_caso_v41(copia_f)
        pior_k = cp.pior_origem_do_caso_v41(copia_f, linear=True)
        erro_caso = None
        try:
            caso = cp.caso_da_planilha_v41(copia_r, TIR)
        except ValueError as e:
            caso, erro_caso = None, str(e)
        py_troca = [linha_troca(copia_r, f) for f in ae.FATORES20]
        formulas_por_aba = {}
        wf = openpyxl.load_workbook(copia_f)
        for ws in wf.worksheets:
            c = sum(1 for lin in ws.iter_rows() for cel in lin
                    if isinstance(cel.value, str) and cel.value.startswith("="))
            if c:
                formulas_por_aba[ws.title] = c
        wf.close()
    if sha256(V41) != h_antes:
        print("ERRO: a V4.1 de referência mudou durante a execução; relatório não gravado", file=sys.stderr)
        return 1

    # --- Python: nominal, K, varredura ---
    res = eq.analise_par(caso, **GRADE) if caso is not None else None
    txt_py = eq.texto_veredito(res, pior_caso, origem_do_carro=origens.TX_V41) if res is not None else None
    par = el.parcelas_k(caso, 0.0) if caso is not None else None
    txt_k_py = el.texto_k(par, pior_k, origem_do_carro=origens.TX_V41) if par is not None else None
    t0 = time.perf_counter()
    varr = eq.varredura_pneu(caso, fatores, **GRADE) if caso is not None else []
    t_varr = time.perf_counter() - t0
    robusto = eq.veredito_robusto([r.veredito for _, r in varr])
    pior_var = pior(pior_caso, var["origem"])
    ds = caso.ds if caso is not None else None

    md: list[str] = []
    w = md.append
    w("# V4.1, incremento 4: equilíbrio na planilha, caso nominal, robustez ao pneu e tempos")
    w("")
    w("Gerado por `tools/relatorio_inc4.py`. Não edite à mão: rode o script de novo (`--sobrescrever`).")
    w("")
    w("Todos os números deste relatório são **demonstrativos**: o caso vem da planilha V4.1 de referência, com valores do "
      "T09/E03, estimativas e entradas arbitrárias, e o pneu é o `.tir` provisório do TTC (R25B), não o pneu do carro. "
      "Nada aqui é resultado do TX nem validação física (Validation = NOT_VALIDATED, EDR-022). Este é o documento que a "
      "célula de robustez da `16_EQUILIBRIO` cita; a decisão sobre ele está no EDR-060.")
    w("")

    w("## 1. Origem")
    w("")
    w(f"- Data: {data}")
    w(f"- Commit: `{commit}` (ramo `{ramo}`" + ("; com mudanças não commitadas em arquivos versionados" if sujo else
                                                  "; árvore sem mudanças pendentes") + ")")
    w(f"- Planilha: `{V41.relative_to(RAIZ).as_posix()}` (SHA-256 antes `{h_antes}`; depois igual: só uma cópia temporária "
      "foi recalculada)")
    w(f"- Motor de recálculo da cópia: {motor} ({detalhe_motor}); {resumo_rec['total_formulas']} fórmulas, "
      f"{resumo_rec['total_errors']} erros, {n(t_rec, 1)} s (recálculo pelo motor, incluindo a partida do programa, mais a leitura do `_resumo` pelo openpyxl; ver a seção 7)")
    w(f"- Pneu: `{TIR.relative_to(RAIZ).as_posix()}` (SHA-256 `{sha256(TIR)}`), em ISO_TYDEX (EDR-039)")
    w(f"- Varredura: `{VARREDURA.relative_to(RAIZ).as_posix()}` ({var['edr']}, origem \"{var['origem']}\")")
    w(f"- Python {platform.python_version()}, openpyxl {version('openpyxl')}, numpy {version('numpy')}")
    w("")

    w("## 2. Caso da V4.1")
    w("")
    w(f"Entradas lidas da cópia recalculada pelos nomes definidos (`caso_planilha.caso_da_planilha_v41`). Pior origem do "
      f"limite: **{pior_caso}**; pior origem do K (inclui o roll steer): **{pior_k}** (regra de `aba_equilibrio`; chaves e "
      "sentido são escolhas de modelagem e ficam fora dela). Por isso o texto de exibição leva o prefixo "
      "\"DEMONSTRATIVO · \" (D16).")
    w("")
    w(f"Situação das entradas na planilha: **{pl['status']}**." if erro_caso is None else
      f"Situação das entradas na planilha: **{pl['status']}**. O Python não montou o caso: {erro_caso}")
    w("")
    w("Entradas novas da `16_EQUILIBRIO` (valor digitado; vazio = \"a confirmar\", EDR-036):")
    w("")
    w("| Entrada | Valor | Unidade |")
    w("|---|---|---|")
    for e in ae.ENTRADAS16:
        v = x.get(e.chave)
        w(f"| {e.rotulo} | {'vazio (a confirmar)' if v is None else v} | {e.unidade} |")
    w("")
    w("Chaves h2, h3, h7, h8 e h9 desligadas: o limite é o da análise de par básica (rolagem de três massas, SUS-103, "
      "downforce fixa do Setup, EDR-013; sem tração).")
    w("")

    w("## 3. Resultado nominal")
    w("")
    if res is None:
        w(f"Sem resultado: o caso é inválido para o Python ({erro_caso}). A planilha mostra: {pl['txt_ver']}.")
    else:
        w(f"Grade da análise de par: a_y de 0 a {n(res.linhas[-1].ay_ms2 / G, 2)} g em passos de "
          f"{n(res.passo_ms2 / G, 2)} g ({len(res.linhas)} linhas), {GRADE['n_alfa']} valores de α por eixo, κ = 0.")
        w("")
        w("| Grandeza | Python (`analise_par`) | Planilha (16_EQUILIBRIO) |")
        w("|---|---|---|")
        w(f"| Texto de exibição | {txt_py} | {pl['txt_ver']} |")
        w(f"| S_f | {g_de(res.ay_sat_f, 4)} | {n(pl['S_F'], 4) if isinstance(pl['S_F'], str) else g_de(pl['S_F'], 4)} |")
        w(f"| S_r | {g_de(res.ay_sat_r, 4)} | {n(pl['S_R'], 4) if isinstance(pl['S_R'], str) else g_de(pl['S_R'], 4)} |")
        w(f"| a_y limite | {g_de(res.ay_lim, 4)} | {n(pl['aylim'], 4) if isinstance(pl['aylim'], str) else g_de(pl['aylim'], 4)} |")
        w(f"| Resolução numérica | {g_de(res.resolucao_ms2, 5)} | {n(pl['resol'], 5) if isinstance(pl['resol'], str) else g_de(pl['resol'], 5)} |")
        w(f"| C_N | {n(res.c_n, 5)} | {n(pl['cn'], 5)} |")
        w(f"| Veredito / motivo | {res.veredito} / {res.motivo or '-'} | {pl['ver']} / {pl['mot']} |")
        w(f"| Avisos de limitação do pneu | {'; '.join(res.limitacoes) if res.limitacoes else 'nenhum'} | {pl['lims']} |")
        w("")
        w(f"K do passo 1 (gradiente de subesterço, parcelas do Gillespie cap. 6): Python `{txt_k_py}`; planilha "
          f"`{pl['txt_k']}`. Velocidades característica e crítica (Gillespie pp. 203-204): planilha "
          f"{n(pl['vchar'])} / {n(pl['vcrit'])} m/s (V_char só existe com K > 0 e V_crit só com K < 0; o outro fica N/A).")
        if par is not None:
            w("")
            w(f"Parcelas de K em a_y = 0, em °/g: pneus {n(el.para_graus_por_g(par.k_pneus) if par.k_pneus is not None else None, 4)}, "
              f"transferência de carga {n(el.para_graus_por_g(par.k_transferencia) if par.k_transferencia is not None else None, 4)}, "
              f"deslocamento com a carga {n(el.para_graus_por_g(par.k_deslocamento) if par.k_deslocamento is not None else None, 4)}, "
              f"camber {n(el.para_graus_por_g(par.k_camber) if par.k_camber is not None else None, 4)}, "
              f"roll steer {n(el.para_graus_por_g(par.k_roll_steer) if par.k_roll_steer is not None else None, 4)}, "
              f"total {n(el.para_graus_por_g(par.k_total) if par.k_total is not None else None, 4)} (Python).")
        w("")
        w("O caso nominal dá N/A (a planilha e o Python concordam): nenhum eixo satura dentro da grade de 2 g. A causa não foi "
          "investigada. Há duas hipóteses, nenhuma verificada: o teto de 2 g da grade estar perto da capacidade do pneu do TTC "
          "sem escala (na varredura da seção 5, com LMUY = 0,8 o caso satura em LOOSE, e com LMUY maior não satura), ou o "
          "camber em curva estar no limite do pneu pelo ganho demonstrativo (seção 4).")
        w("")
        if par is not None and par.k_camber is not None and par.k_total is not None:
            w(f"O texto \"Sobresterço\" do K ({n(el.para_graus_por_g(par.k_total), 2)} °/g) vem quase todo da parcela de camber "
              f"({n(el.para_graus_por_g(par.k_camber), 2)} °/g), que sai do ganho de camber demonstrativo implausível "
              "(∂γ/∂φ ≈ −5,7 rad/rad, do SUS-404 = −0,6 °/mm; seção 4). Não é resultado do carro.")
            w("")
        w(f"Tempo de volta no skidpad (SUS-1007): planilha `{pl['tvolta']}` (precisa do raio R, que está vazio).")
    w("")

    w("## 4. Limitações do modelo e deste relatório")
    w("")
    w("- A planilha mostra o veredito nominal de **um** sentido de curva por vez (h7); troque o sentido para ver o outro "
      "(decisão 2 do EDR-060).")
    w("- A robustez ao pneu **não é recalculada na planilha**: a seção 5 abaixo é o único lugar em que ela existe, e vale "
      "para as entradas deste momento (decisão 1 do EDR-060).")
    w("- Pneu provisório do TTC (EDR-038); geometria sintética (não há hardpoints do CAD); mola, MR, barra e RC são "
      "valores demonstrativos do T09/E03, não do TX.")
    w("- Diferencial aberto é o único implementado (h10); as hipóteses 11 a 16 da spec ficam fora (aba `21_COBERTURA`).")
    w("- Pior origem do K conservadora e \"Neutro\" só com K = 0 exato; outras limitações conhecidas estão no EDR-060.")
    w("- O ganho de camber em rolagem sai das entradas SUS-404 da 07_KINEMATICS (autorização 7 do plano); o valor demonstrativo "
      "atual (−0,6 °/mm) dá ∂γ/∂φ ≈ −5,7 rad/rad, implausível, e por isso o camber em curva fica limitado ao CAMMIN/CAMMAX "
      "do pneu, com o aviso de limitação da seção 3. A causa é o dado demonstrativo, não o método.")
    w("- No caso demonstrativo nenhuma linha da tabela de troca satura até 2 g (seção 6), então o a_y do neutro e o custo ficam "
      "N/A. O cruzamento por interpolação é exercitado nos testes (TR1 do TC-WB41-006, com valores de teste) e passa a "
      "aparecer na planilha quando as entradas forem reais.")
    w("- Passou nos testes = a planilha reproduz o oráculo em Python; não é validação física.")
    w("")

    w("## 5. Robustez ao pneu (EDR-055)")
    w("")
    w(f"`equilibrio.varredura_pneu` sobre `caso_planilha.caso_da_planilha_v41`: LFZO, LMUY e LKY × "
      f"{{{'; '.join(n(v, 1) for v in fatores['LFZO'])}}} em produto cartesiano ({len(varr)} combinações). Os "
      f"multiplicadores são exploratórios, origem \"{var['origem']}\", e nunca uma faixa do pneu (EDR-022). Tempo da varredura "
      f"no Python: {n(t_varr, 1)} s.")
    w("")
    if varr:
        w("| LFZO | LMUY | LKY | Veredito | a_y limite | C_N | Motivo do N/A |")
        w("|---|---|---|---|---|---|---|")
        for comb, r in varr:
            w(f"| {n(comb['LFZO'], 1)} | {n(comb['LMUY'], 1)} | {n(comb['LKY'], 1)} | {r.veredito} | {g_de(r.ay_lim, 3)} | "
              f"{n(r.c_n, 4)} | {r.motivo or '-'} |")
        w("")
        w(f"Resumo (`veredito_robusto`, spec 7.6 item 3): **{robusto[0]}**, {robusto[1]}. Texto de exibição com a "
          f"robustez: `{eq.texto_veredito(res, pior_var, robusto, origem_do_carro=origens.TX_V41)}`.")
        vs = {}
        for _, r in varr:
            vs[r.veredito] = vs.get(r.veredito, 0) + 1
        w("")
        w("Contagem por veredito: " + ", ".join(f"{k} {v}" for k, v in vs.items()) + ".")
        w("")
        n_na = vs.get(eq.NA, 0)
        n_pu = vs.get(eq.PUSH, 0)
        n_lo = vs.get(eq.LOOSE, 0)
        w(f"Leitura do resumo: houve {n_lo} {eq.LOOSE}, {n_na} {eq.NA} e {n_pu} {eq.PUSH}. Se não há {eq.PUSH}, o sinal não se "
          f"inverte em nenhum ponto. O {robusto[0]} vem da regra de `equilibrio.veredito_robusto`, que trata qualquer mistura de "
          f"vereditos, inclusive de {eq.LOOSE} com {eq.NA}, como não conclusivo; a spec §7.6 item 3 só fala em inversão de "
          "push/loose, e essa regra mais estrita é pendência de decisão do usuário (EDR-060, item 8, PENDING_CONFIRMATION). "
          f"O que de fato depende do pneu aqui é o carro saturar antes de 2 g ou não ({eq.NA} = não satura até o teto da grade). "
          "Não se deve ler a frase \"o sinal depende do pneu\" como uma inversão entre PUSH e LOOSE.")
    else:
        w("Sem varredura: o caso é inválido para o Python.")
    w("")
    w("Robusto aqui significa apenas \"o sinal não mudou com esta perturbação demonstrativa\"; não valida o pneu.")
    w("")

    w("## 6. Tabela de troca da barra dianteira (20_EQ_TROCA)")
    w("")
    w("K_ARB dianteiro × {0,5; 0,75; 1; 1,5; 2} do valor da `06_SUSPENSION` (grade de exibição rotulada, **não é meta** nem "
      "faixa de projeto; decisão do usuário de 2026-10-22). Valores da planilha recalculada; a coluna \"Python\" confere o "
      "veredito e o a_y limite de cada linha.")
    w("")
    w("| Fator | K_φF [N·m/rad] | λ_f | K [°/g] | S_f [g] | S_r [g] | a_y limite [g] | Veredito (planilha) | Python |")
    w("|---|---|---|---|---|---|---|---|---|")
    for f, s, p in zip(ae.FATORES20, tab, py_troca):
        if "invalida" in p:
            py_txt = f"N/A ({p['invalida']})"
        else:
            py_txt = f"{p['ver']}, a_y_lim {g_de(p['aylim'], 4)}"
        w(f"| ×{n(f, 2)} | {n(s['kphi'], 1)} | {n(s['lam'], 4)} | {n(s['kg'], 4)} | {n(s['sfg'], 4)} | {n(s['srg'], 4)} | "
          f"{n(s['aylimg'], 4)} | {s['ver']} | {py_txt} |")
    w("")
    ayn_py, ramo_n = neutro([p for p in py_troca if "invalida" not in p]) if all("invalida" not in p for p in py_troca) \
        else (None, "linhas inválidas")
    w(f"a_y do acerto neutro (cruzamento das saturações, spec 5 item 4): planilha {g_txt(custo['ayng'])} "
      f"(motivo `{custo['aynm']}`); Python {g_de(ayn_py, 4)} ({ramo_n}). Custo do comportamento = a_y do neutro − a_y limite "
      f"atual: planilha {g_txt(custo['perdag'])}. Veredito do custo: `{custo['limiar']}` (a spec §5 item 5 não define "
      "limiar; nenhum número aqui diz se o custo é grande ou pequeno).")
    w("")

    w("## 7. Tempos de recálculo")
    w("")
    w(f"- Cópia da V4.1 inteira ({resumo_rec['total_formulas']} fórmulas): {n(t_rec, 1)} s no motor {motor} "
      "(`t_rec`: tempo de `rc.recalcular`, que soma o recálculo pelo motor (abrir, recalcular tudo, gravar e fechar), a partida do "
      "programa, por ser a primeira chamada de uma execução, e a leitura do `_resumo` pelo openpyxl, que carrega o arquivo mais de uma vez "
      "e varre as células à procura de erros; uma execução). **Não é comparável** aos 8,8 s (Excel) e 10,6 s (LibreOffice) das Tasks 5 e 6 "
      "(V4.1 da época, 117532 a 117568 fórmulas): essas medidas cronometram só o recálculo do motor (`_excel_recalc` no Excel), "
      "sem a partida nem o `_resumo`. Em 2026-10-22, quatro execuções deste mesmo relatório deram 22,9 s, 66,5 s, 67,9 s e 66,7 s no "
      "mesmo computador. Não verificado: a diferença entre as execuções pode vir de um processo do Excel residual de rodadas anteriores "
      "ou da inspeção do antivírus ao gravar o arquivo temporário; nenhuma das duas hipóteses foi testada. O tempo medido aqui não é "
      "um limite nem uma garantia.")
    w("- Fórmulas por aba nova: " + ", ".join(f"{k} {v}" for k, v in formulas_por_aba.items()
                                              if k[:2] in ("16", "17", "18", "19", "20", "21", "22")) + ".")
    w(f"- Medição do portão da Task 4, com o Excel recalculando só a troca do K_ARB (recálculo automático): "
      f"`{PORTAO.relative_to(RAIZ).as_posix()}`. Lá, o recálculo completo (CalculateFull) com a tabela de 5 linhas e a "
      "grade refinada ficou em cerca de 0,3 s e a troca de uma entrada em cerca de 0,4 s; no LibreOffice, o recálculo com "
      "abrir e gravar ficou entre 11 e 12 s. Sem limiar registrado, o tempo não decide nada sozinho (spec §9, item 10).")
    w("")

    saida.parent.mkdir(parents=True, exist_ok=True)
    saida.write_text("\n".join(md) + "\n", encoding="utf-8", newline="\n")
    print(f"relatório: {saida}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
