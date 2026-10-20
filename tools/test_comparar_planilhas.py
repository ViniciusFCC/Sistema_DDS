"""Testes do tools/comparar_planilhas.py, a guarda do CA-07 e do CA-08.

Uso:
    uv run python -X utf8 tools/test_comparar_planilhas.py

O modo de falha mais perigoso do comparador é aceitar tudo. Cada caso altera
uma cópia temporária da planilha (ou do layout) entregue e exige FAIL, com a
célula ou a chave apontada no relatório. Os controles exigem PASS: a cópia
idêntica e a cópia salva pelo openpyxl sem mudança (o build também salva pelo
openpyxl), para que os FAIL venham da alteração, e não do ato de salvar.
"""
import json
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import comparar_planilhas as cp  # noqa: E402

results = []


def check(name, cond, detail=""):
    results.append((name, "PASS" if cond else "FAIL"))
    print(f"[{'PASS' if cond else 'FAIL'}] {name}  {detail}")


def primeira(ws, formula: bool):
    for linha in ws.iter_rows():
        for c in linha:
            f = cp._formula(c.value)
            if c.value is not None and (f is not None) == formula and (formula or isinstance(c.value, (int, float))):
                return c
    raise LookupError(ws.title)


with tempfile.TemporaryDirectory(prefix="dds_cmp_") as tmp:
    tmp = Path(tmp)

    def variante(nome, alterar):
        """Cópia salva pelo openpyxl depois de `alterar(wb)`; devolve (caminho, descrição)."""
        wb = openpyxl.load_workbook(cp.ENTREGUE_XLSX)
        desc = alterar(wb)
        caminho = tmp / f"{nome}.xlsx"
        wb.save(caminho)
        return caminho, desc

    copia = tmp / "copia.xlsx"
    shutil.copyfile(cp.ENTREGUE_XLSX, copia)
    ok, rel = cp.comparar_xlsx(copia)
    check("comparador 01: cópia byte a byte da entregue -> PASS", ok, rel[0])

    salva, _ = variante("salva", lambda wb: None)
    ok, rel = cp.comparar_xlsx(salva)
    check("comparador 02: entregue salva pelo openpyxl sem mudança -> PASS (controle)", ok, rel[0])

    def muda_formula(wb):
        c = primeira(wb["10_TIRE"], formula=True)
        c.value = c.value + "+0"
        return f"10_TIRE!{c.coordinate}"
    p, alvo = variante("formula", muda_formula)
    ok, rel = cp.comparar_xlsx(p)
    check("comparador 03: texto de uma fórmula alterado (+0) -> FAIL apontando a célula",
          not ok and any(alvo in r and "fórmula" in r for r in rel), alvo)

    def formula_vira_valor(wb):
        c = primeira(wb["10_TIRE"], formula=True)
        c.value = 0
        return f"10_TIRE!{c.coordinate}"
    p, alvo = variante("formula_valor", formula_vira_valor)
    ok, rel = cp.comparar_xlsx(p)
    check("comparador 04: fórmula trocada por valor -> FAIL (célula e contagem 346)",
          not ok and any(alvo in r for r in rel) and any("gerada 346" in r for r in rel), alvo)

    def muda_valor(wb):
        c = primeira(wb["10_TIRE"], formula=False)
        c.value = c.value + 1e-9
        return f"10_TIRE!{c.coordinate}"
    p, alvo = variante("valor", muda_valor)
    ok, rel = cp.comparar_xlsx(p)
    check("comparador 05: valor numérico sem fórmula alterado em 1e-9 -> FAIL apontando a célula",
          not ok and any(alvo in r and "valor" in r for r in rel), alvo)

    def apaga_valor(wb):
        c = primeira(wb["10_TIRE"], formula=False)
        c.value = None
        return f"10_TIRE!{c.coordinate}"
    p, alvo = variante("apagado", apaga_valor)
    ok, rel = cp.comparar_xlsx(p)
    check("comparador 06: célula apagada -> FAIL apontando a célula", not ok and any(alvo in r for r in rel), alvo)

    def troca_ordem(wb):
        wb.move_sheet(wb.worksheets[-1], offset=-1)
        return "ordem das abas"
    p, alvo = variante("ordem", troca_ordem)
    ok, rel = cp.comparar_xlsx(p)
    check("comparador 07: ordem das abas trocada -> FAIL", not ok and any("abas diferem" in r for r in rel))

    layout = json.loads(cp.ENTREGUE_JSON.read_text(encoding="utf-8"))
    j = tmp / "layout_igual.json"
    j.write_text(json.dumps(layout), encoding="utf-8")
    ok, rel = cp.comparar_json(j)
    check("comparador 08: layout igual ao entregue -> PASS", ok, rel[0])

    l2 = json.loads(json.dumps(layout)); l2["ref"]["fx"] += 1e-9
    j.write_text(json.dumps(l2), encoding="utf-8")
    ok, rel = cp.comparar_json(j)
    check("comparador 09: valor do layout alterado em 1e-9 -> FAIL apontando $.ref.fx",
          not ok and any("$.ref.fx" in r for r in rel))

    l3 = json.loads(json.dumps(layout)); l3["rows"].pop(next(iter(l3["rows"])))
    j.write_text(json.dumps(l3), encoding="utf-8")
    ok, rel = cp.comparar_json(j)
    check("comparador 10: chave removida do layout -> FAIL", not ok and any("só em entregue" in r for r in rel))

    # ---- CA-07 da planilha viva (EDR-051, fechamento da V3.0): só as entradas podem mudar ----
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "python"))
    import atualizar_planilha as ap  # noqa: E402
    ENT = ap.celulas_de_entrada()
    REF = cp.ENTREGUE_XLSX_V30

    def viva(nome, alterar):
        wb = openpyxl.load_workbook(REF)
        alterar(wb)
        caminho = tmp / f"viva_{nome}.xlsx"
        wb.save(caminho)
        return caminho

    def preenche(wb):
        v = wb["05_VEHICLE"]
        v["B5"].value, v["E5"].value, v["F5"].value = 0.300, "○ Estimativa", "Teste: estimativa da equipe"
        wb["03_TARGETS"]["C5"].value = 2.5
        wb["01_SETUP"]["B9"].value = "G-REV-B"
        wb["10_TIRE"]["B23"].value = -4
    ok, rel = cp.comparar_viva(viva("preenchida", preenche), REF, ENT)
    check("comparador 11 (viva): entradas preenchidas (valor, origem, fonte, meta, revisão, caso de carga) -> PASS", ok, rel[-1])

    def formula_valor(wb):
        wb["05_VEHICLE"]["B23"].value = 290.0
    ok, rel = cp.comparar_viva(viva("formula_valor", formula_valor), REF, ENT)
    check("comparador 12 (viva): fórmula trocada por valor -> FAIL apontando a célula",
          not ok and any("05_VEHICLE!B23" in r for r in rel))

    def rotulo(wb):
        wb["05_VEHICLE"]["C5"].value = "mm"
    ok, rel = cp.comparar_viva(viva("rotulo", rotulo), REF, ENT)
    check("comparador 13 (viva): unidade (estrutura fora das entradas) alterada -> FAIL", not ok and any("05_VEHICLE!C5" in r for r in rel))

    def sem_lista(wb):
        wb["06_SUSPENSION"].data_validations.dataValidation = []
    ok, rel = cp.comparar_viva(viva("sem_lista", sem_lista), REF, ENT)
    check("comparador 14 (viva): lista de Origem removida de uma aba -> FAIL", not ok and any("06_SUSPENSION: validação" in r for r in rel))

    def origem_ruim(wb):
        wb["05_VEHICLE"]["E6"].value = "medido"
    ok, rel = cp.comparar_viva(viva("origem_ruim", origem_ruim), REF, ENT)
    check("comparador 15 (viva): Origem fora dos 6 rótulos (colada por cima da lista) -> FAIL", not ok and any("origem fora da lista" in r for r in rel))

print("\n--- Resumo ---")
p = sum(1 for _, s in results if s == "PASS"); f = len(results) - p
print(f"{p} PASS, {f} FAIL, de {len(results)} testes")
if f:
    raise SystemExit(1)
