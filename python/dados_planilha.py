"""Dados por chave entre a V3.0 viva e a V4.0 (incremento 3, fundação; EDR-058).

Os dados da equipe (células `ent`, `org` e `fte` da tabela `data/nomes_V4_0.csv`) passam da V3.0 para a V4.0 pelo
nome, não pelo endereço:
- `ler_por_endereco`: lê a V3.0 (a viva ou a da tag) pelos endereços da tabela;
- `ler_por_nome`: lê a V4.0 resolvendo os nomes definidos da própria planilha;
- `injetar`: escreve os dados numa cópia da V4.0, com detecção de conflito contra uma base;
- `celulas_alteradas`: células em que a viva difere da V3.0 da tag (para `verificar_cobertura`).
Os nomes `res` (fórmulas) nunca são lidos nem escritos como dado.

Os valores são comparados crus, como o openpyxl os devolve (sem recálculo): None é igual a None; números (int e
float, fora bool) comparam pelo valor exato; o resto exige o mesmo tipo e igualdade.
O injetor grava só a cópia `saida`, nunca em excel/ nem por cima da origem (recusa antes de abrir qualquer arquivo),
e nunca recalcula.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

from gerar_v4_0 import _recusar
from nomes_v4_0 import Nome

PAPEIS_DADOS = ("ent", "org", "fte")


def eh_dado(nome: str) -> bool:
    """True se o nome é de dado (`ent_`, `org_` ou `fte_`); `res_` e outros nomes (por exemplo `Origens`) não são."""
    return nome.split("_", 1)[0] in PAPEIS_DADOS


def iguais(a, b) -> bool:
    """Igualdade crua de dois valores de célula: números pelo valor exato, o resto pelo tipo e pelo valor."""
    num = (int, float)
    if isinstance(a, num) and isinstance(b, num) and not isinstance(a, bool) and not isinstance(b, bool):
        return a == b
    return type(a) is type(b) and a == b


def _eh_formula(v) -> bool:
    return isinstance(v, ArrayFormula) or (isinstance(v, str) and v.startswith("="))


def ler_por_endereco(xlsx_v30: Path, tabela: Iterable[Nome]) -> dict[str, object]:
    """Valores das células `ent`, `org` e `fte` da V3.0 em `xlsx_v30`, por nome, lidos pelos endereços da tabela."""
    dados = [n for n in tabela if n.papel in PAPEIS_DADOS]
    wb = openpyxl.load_workbook(xlsx_v30)
    try:
        faltam = sorted({n.aba for n in dados} - set(wb.sheetnames))
        if faltam:
            raise ValueError(f"abas da tabela que não existem em {xlsx_v30}: {faltam}")
        return {n.nome: wb[n.aba][n.celula].value for n in dados}
    finally:
        wb.close()


def _destinos(wb) -> dict[str, tuple[str, str]]:
    """nome de dado -> (aba, célula), dos nomes definidos no escopo do workbook (um destino por nome)."""
    out = {}
    for nome, dn in wb.defined_names.items():
        if not eh_dado(nome):
            continue
        dest = list(dn.destinations)
        if len(dest) != 1:
            raise ValueError(f"nome {nome} com {len(dest)} destinos ({dn.attr_text}); esperado 1")
        aba, ref = dest[0]
        cel = ref.replace("$", "")
        if ":" in cel:
            raise ValueError(f"nome {nome} aponta para um intervalo ({dn.attr_text}); esperado uma célula")
        out[nome] = (aba, cel)
    return out


def ler_por_nome(xlsx_v40: Path) -> dict[str, object]:
    """Valores das células `ent`, `org` e `fte` da V4.0 em `xlsx_v40`, por nome, resolvendo os nomes definidos dela."""
    wb = openpyxl.load_workbook(xlsx_v40)
    try:
        return {nome: wb[aba][cel].value for nome, (aba, cel) in _destinos(wb).items()}
    finally:
        wb.close()


@dataclass
class RelatorioInjecao:
    """Resultado de `injetar`: nomes escritos, já iguais ao novo, em conflito (nome, base, destino, novo) e sem célula
    na V4.0."""
    escritas: list[str] = field(default_factory=list)
    iguais: list[str] = field(default_factory=list)
    conflitos: list[tuple[str, object, object, object]] = field(default_factory=list)
    faltando: list[str] = field(default_factory=list)


def injetar(xlsx_v40: Path, dados: dict, saida: Path, base: dict | None = None) -> RelatorioInjecao:
    """Escreve `dados` (nome -> valor) numa cópia da V4.0 gravada em `saida`.

    Para cada nome: destino igual ao novo vai para `iguais`; destino igual à base (padrão: `ler_por_nome` da própria
    `xlsx_v40`) é escrito; qualquer outro caso é conflito, e a célula não é escrita. Nome sem célula na V4.0 vai para
    `faltando`. Nome que não é de dado (`res_`…) ou valor com fórmula levantam ValueError antes de gravar.
    Recusa `saida` em excel/ ou igual a `xlsx_v40` antes de abrir qualquer arquivo. Nunca recalcula."""
    xlsx_v40, saida = Path(xlsx_v40), Path(saida)
    _recusar(xlsx_v40, saida)
    nao_dados = sorted(k for k in dados if not eh_dado(k))
    if nao_dados:
        raise ValueError(f"nomes que não são de dado (só ent/org/fte são injetados): {nao_dados[:10]}")
    com_formula = sorted(k for k, v in dados.items() if _eh_formula(v))
    if com_formula:
        raise ValueError(f"valores com fórmula não são migrados como dado: {com_formula[:10]}")
    if base is None:
        base = ler_por_nome(xlsx_v40)
    rel = RelatorioInjecao()
    wb = openpyxl.load_workbook(xlsx_v40)
    try:
        destinos = _destinos(wb)
        for nome in sorted(dados):
            novo = dados[nome]
            if nome not in destinos:
                rel.faltando.append(nome)
                continue
            aba, cel = destinos[nome]
            c = wb[aba][cel]
            atual = c.value
            if iguais(atual, novo):
                rel.iguais.append(nome)
            elif nome in base and iguais(atual, base[nome]):
                c.value = novo
                rel.escritas.append(nome)
            else:
                rel.conflitos.append((nome, base.get(nome), atual, novo))
        saida.parent.mkdir(parents=True, exist_ok=True)
        wb.save(saida)
    finally:
        wb.close()
    return rel


def _valores(caminho: Path) -> tuple[list[str], dict[tuple[str, str], object]]:
    """Abas e valores crus não vazios (fórmula como texto; fórmula matricial como ("array", ref, texto))."""
    wb = openpyxl.load_workbook(caminho)
    try:
        out = {}
        for ws in wb.worksheets:
            for lin in ws.iter_rows():
                for c in lin:
                    v = c.value
                    if v is None:
                        continue
                    out[(ws.title, c.coordinate)] = ("array", v.ref, v.text) if isinstance(v, ArrayFormula) else v
        return wb.sheetnames, out
    finally:
        wb.close()


class AbasDiferentes(ValueError):
    """A viva e a V3.0 da tag não têm as mesmas abas."""


def celulas_alteradas(viva: Path, v30: Path) -> set[tuple[str, str]]:
    """Células (aba, célula) em que `viva` difere de `v30` (a V3.0 da tag), comparando os valores crus. Aba que só
    existe numa das duas levanta AbasDiferentes. Para a migração, estas células têm de ter nome de dado na tabela."""
    abas_v, vv = _valores(viva)
    abas_t, vt = _valores(v30)
    if set(abas_v) != set(abas_t):
        raise AbasDiferentes(f"só na viva {sorted(set(abas_v) - set(abas_t))}, "
                             f"só na V3.0 {sorted(set(abas_t) - set(abas_v))}")
    return {k for k in set(vv) | set(vt) if not iguais(vv.get(k), vt.get(k))}
