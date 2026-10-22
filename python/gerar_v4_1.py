"""Gera a planilha V4.1 (incremento 4: equilíbrio na planilha) a partir da V4.0 da tag v4.0.

Uso:
    uv run python -X utf8 python/gerar_v4_1.py [--v40 ARQ] [--saida build/FSAE_EV_Suspension_Fase1_V4_1.xlsx]
                                               [--registro docs/registro-de-versoes-V4_1-celulas.md]
                                               [--nomes-csv data/nomes_V4_1.csv]

Sem --v40, o arquivo da tag v4.0 é extraído (git show, binário) para uma pasta temporária: nunca se usa a V4.0 da
árvore. A tabela dos nomes que a V4.1 acrescenta vem de data/nomes_V4_1.csv (a tabela completa da V4.1 é a de
data/nomes_V4_0.csv mais essa); o gerador recusa um CSV diferente de `aba_equilibrio.NOMES_V41`, porque as fórmulas
das abas novas usam esses nomes. Com --nomes-csv, o CSV é regravado a partir de NOMES_V41 antes de gerar.

O que muda em relação à V4.0 (EDR-060; plano docs/superpowers/plans/2026-10-21-v4-1-inc4-equilibrio-na-planilha.md):
- abas novas 16_EQUILIBRIO, 17_EQ_PAR, 18_EQ_PAR_FINA, 19_EQ_LINEAR, 20_EQ_TROCA, 20a…20d_TROCA, 21_COBERTURA e
  22_CADEIA, e o bloco novo do 14_DASHBOARD (`aba_equilibrio.montar`, fórmulas por família do catálogo F1-F19);
- os nomes definidos de `aba_equilibrio.NOMES_V41` (os 698 da V4.0 ficam iguais);
- as células existentes de `alteracoes_v4_1()`, aprovadas pelo usuário em 2026-10-21 (item 6 das autorizações do plano):
  (a) a 10_TIRE alinhada ao EDR-056: Fz > 0, α e γ fora da faixa do .tir entram no cálculo limitados ao valor mínimo
      ou máximo da faixa (convenção do PAC2002, help do Adams 2024.1, p. 83), e a linha de status da coluna mostra o
      aviso (entrada, limite, valor real e usado); Fz <= 0 continua N/A (descolamento) e κ fora da faixa também (o
      EDR-056 não trata do κ). Dentro da faixa nenhum valor muda;
  (b) os textos da spec §12.1 (13_VALIDATION A25, 00_README A8, 14_DASHBOARD A19);
  (c) a versão na 00_README A3 e a linha nova do 15_CHANGELOG, como na V4.0.
A saída nunca fica em excel/ nem por cima da origem: a recusa vem antes de abrir qualquer arquivo. A V4.1 de
referência é gerada em build/ e copiada para excel/ num commit próprio.
"""
from __future__ import annotations

import argparse
import copy
import re
import tempfile
from collections import defaultdict
from pathlib import Path

import openpyxl

import aba_equilibrio as ae
from atualizar_planilha import escrever_registro
from nomes_v4_0 import Nome, escrever_tabela, ler_tabela

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
PASTA_EXCEL = RAIZ / "excel"
SAIDA_PADRAO = RAIZ / "build" / "FSAE_EV_Suspension_Fase1_V4_1.xlsx"
CSV_NOMES = RAIZ / "data" / "nomes_V4_1.csv"
PLANO = "docs/superpowers/plans/2026-10-21-v4-1-inc4-equilibrio-na-planilha.md"
REGISTRO = "docs/registro-de-versoes-V4_1-celulas.md"
REF_6A = ("EDR-056, item 1 (help do Adams 2024.1 PAC2002, p. 83: entrada fora da faixa avaliada no limite, com aviso); "
          "plano, autorização 6a (usuário, 2026-10-21); dependentes fora da 10_TIRE: 13_VALIDATION B40:B43 e "
          "14_DASHBOARD B34:B37 (o B43 conta só status exatamente \"OK\", sem os limitados); teste: TC-WB41-003")
REF_6B = "spec §12.1; plano, autorização 6b (usuário, 2026-10-21); teste: TC-WB41-001 (só o trecho muda)"
REF_6C = "EDR-060; plano, autorização 6c (padrão da V4.0); teste: TC-WB41-001"
REF_NOVAS = "EDR-060; catálogo de fórmulas do plano (famílias F1-F19 e REF)"

ABAS_NOVAS = (ae.ABA16, ae.ABA17, ae.ABA18, ae.ABA19, ae.ABA20, *ae.AUX20.values(), ae.ABA21, ae.ABA22)
FAMILIAS = tuple(f"F{i}" for i in range(1, 20)) + ("REF",)

TEXTO_A3 = ("Versão do workbook: V4.1 (incremento 4: equilíbrio no limite e linear em fórmulas, abas 16 a 22; 10_TIRE "
            "com a entrada do pneu limitada à faixa do .tir, EDR-056) | Status: DRAFT — cross-check, ainda NÃO "
            "REFERENCE-validated")
LINHA_CHANGELOG = 21   # primeira linha livre do 15_CHANGELOG na V4.0 da tag (a 20 é a da V4.0)

# ---------------------------------------------------------------------------------------------------------------------
# 6a: 10_TIRE pelo EDR-056
# ---------------------------------------------------------------------------------------------------------------------
ABA_TIRE = "10_TIRE"
COLS_TIRE = "BCDEF"            # FL, FR, RL, RR e VERIF
L_CANTO, L_FZ, L_ALFA, L_GAMA, L_KAPPA = 20, 22, 23, 24, 25
L_STATUS = 26
L_ALFA_ARQ, L_GAMA_ARQ, L_ESPELHO = 38, 39, 35
# Faixas do .tir na 10_TIRE (FZMIN B87, FZMAX B88, ALPMIN B89, ALPMAX B90, CAMMIN B91, CAMMAX B92, KPUMIN B93,
# KPUMAX B94) e os sinais do sistema de eixos (s_α B12, s_γ B14). α e γ são limitados no referencial do arquivo
# (linhas 38 e 39), onde a faixa vale, como `equilibrio._entrada_limitada` faz com s_g·γ.
P = {"FZMIN": "$B$87", "FZMAX": "$B$88", "ALPMIN": "$B$89", "ALPMAX": "$B$90", "CAMMIN": "$B$91", "CAMMAX": "$B$92",
     "KPUMIN": "$B$93", "KPUMAX": "$B$94", "s_a": "$B$12", "s_g": "$B$14", "valido": "$B$17"}
NOTA_STATUS = ("Sem extrapolação: fora da faixa do dataset, Fz (> 0), α e γ entram no cálculo limitados ao valor mínimo "
               "ou máximo da faixa (FZMIN/FZMAX, ALPMIN/ALPMAX, CAMMIN/CAMMAX; convenção do PAC2002, help do Adams "
               "2024.1, p. 83; EDR-056), e o status mostra \"OK - limitado …\" com a entrada, o valor real, o usado e "
               "o limite. κ fora da faixa continua N/A. Fz≤0 = perda de contato (N/A). Com limitação, o μ efetivo "
               "(linha 31) divide pelo Fz usado (limitado à faixa), não pelo Fz real da linha 22.")


def _rel(col: str, lin: int) -> re.Pattern:
    """Referência relativa à coluna `col` e absoluta à linha `lin` (`B$22`), sem pegar `$B$22` nem `B$220`."""
    return re.compile(rf"(?<![A-Za-z$]){col}\${lin}(?!\d)")


def _limitado(col: str, lin: int, lo: str, hi: str) -> str:
    return f"MIN(MAX({col}${lin},{P[lo]}),{P[hi]})"


def status_tire(col: str) -> str:
    """Status do ponto de operação da coluna (linha 26) pela regra do EDR-056.

    Ordem: entrada ausente → dataset inválido → Fz <= 0 (descolamento) → κ fora da faixa → "OK" (tudo dentro) ou
    "OK - limitado à faixa do .tir (EDR-056): <avisos>". Cada aviso segue o texto de `equilibrio._entrada_limitada`
    (vírgula decimal; Fz com 1 casa, ângulos em graus com 2), com a roda da linha 20: "Fz_FR = 1526,3 N → 1521,0 N
    (FZMAX)", "γ_FL = -4,50° → -4,00° (CAMMIN)"; o α no mesmo formato, no referencial do projeto (ISO 8855). O valor
    usado de α e de γ volta do referencial do arquivo: α = s_α·espelho·α_arq e γ = s_γ·γ_arq (fatores ±1)."""
    c = col
    roda = f"{c}${L_CANTO}"
    fz, al, ga = f"{c}${L_FZ}", f"{c}${L_ALFA_ARQ}", f"{c}${L_GAMA_ARQ}"

    def fz_txt(lim):
        return (f'"; Fz_"&{roda}&" = "&{ae._f(fz, 1)}&" N → "&{ae._f(P[lim], 1)}&" N ({lim})"')

    def ang_txt(simb, real, usado, lim):
        return (f'"; {simb}_"&{roda}&" = "&{ae._fs(real, 2)}&"° → "&{ae._fs(usado, 2)}&"° ({lim})"')

    a_usado = lambda lim: f"DEGREES({P[lim]}*{P['s_a']}*{c}${L_ESPELHO})"
    g_usado = lambda lim: f"DEGREES({P[lim]}*{P['s_g']})"
    avisos = (f'IF({fz}<{P["FZMIN"]},{fz_txt("FZMIN")},IF({fz}>{P["FZMAX"]},{fz_txt("FZMAX")},""))'
              f'&IF({al}<{P["ALPMIN"]},{ang_txt("α", f"{c}${L_ALFA}", a_usado("ALPMIN"), "ALPMIN")},'
              f'IF({al}>{P["ALPMAX"]},{ang_txt("α", f"{c}${L_ALFA}", a_usado("ALPMAX"), "ALPMAX")},""))'
              f'&IF({ga}<{P["CAMMIN"]},{ang_txt("γ", f"{c}${L_GAMA}", g_usado("CAMMIN"), "CAMMIN")},'
              f'IF({ga}>{P["CAMMAX"]},{ang_txt("γ", f"{c}${L_GAMA}", g_usado("CAMMAX"), "CAMMAX")},""))')
    dentro = (f"AND({fz}>={P['FZMIN']},{fz}<={P['FZMAX']},{al}>={P['ALPMIN']},{al}<={P['ALPMAX']},"
              f"{ga}>={P['CAMMIN']},{ga}<={P['CAMMAX']})")
    k = f"{c}${L_KAPPA}"
    return (f'=IF(NOT(AND(ISNUMBER({c}${L_FZ}),ISNUMBER({c}${L_ALFA}),ISNUMBER({c}${L_GAMA}),ISNUMBER({k}))),'
            f'"N/A - entrada ausente",IF({P["valido"]}<>"OK","N/A - dataset inválido",'
            f'IF({fz}<=0,"N/A - Fz<=0 (liftoff)",IF(OR({k}<{P["KPUMIN"]},{k}>{P["KPUMAX"]}),"N/A - κ fora da faixa",'
            f'IF({dentro},"OK","OK - limitado à faixa do .tir (EDR-056): "&MID({avisos},3,2000))))))')


def formula_tire_limitada(antes: str, col: str) -> str:
    """Fórmula de cálculo da 10_TIRE (linhas 29-83) com a regra do EDR-056: a guarda `X$26<>"OK"` passa a aceitar o
    status "OK - limitado …" (LEFT(X$26,2)), e Fz (linha 22), α do arquivo (38) e γ do arquivo (39) entram limitados à
    faixa. O resto do texto fica igual. Recusa fórmula sem a guarda ou com referência relativa a outra coluna."""
    guarda = f'{col}${L_STATUS}<>"OK"'
    if guarda not in antes:
        raise ValueError(f"10_TIRE {col}: fórmula sem a guarda {guarda}: {antes!r}")
    outras = [x for x in COLS_TIRE if x != col and re.search(rf"(?<![A-Za-z$]){x}\$\d", antes)]
    if outras:
        raise ValueError(f"10_TIRE {col}: referência relativa a outra coluna ({outras}): {antes!r}")
    novo = antes.replace(guarda, f'LEFT({col}${L_STATUS},2)<>"OK"')
    for lin, lo, hi in ((L_FZ, "FZMIN", "FZMAX"), (L_ALFA_ARQ, "ALPMIN", "ALPMAX"), (L_GAMA_ARQ, "CAMMIN", "CAMMAX")):
        novo = _rel(col, lin).sub(_limitado(col, lin, lo, hi), novo)
    return novo


def _alteracoes_tire(ws) -> list[dict]:
    out = []
    for col in COLS_TIRE:
        c = ws[f"{col}{L_STATUS}"]
        out.append({"aba": ABA_TIRE, "cel": c.coordinate, "antes": c.value, "novo": status_tire(col), "ref": REF_6A})
    for r in range(L_STATUS + 1, ws.max_row + 1):
        for col in COLS_TIRE:
            v = ws[f"{col}{r}"].value
            if isinstance(v, str) and v.startswith("=") and f'{col}${L_STATUS}<>"OK"' in v:
                out.append({"aba": ABA_TIRE, "cel": f"{col}{r}", "antes": v, "novo": formula_tire_limitada(v, col),
                            "ref": REF_6A})
    i26 = ws[f"I{L_STATUS}"].value
    out.append({"aba": ABA_TIRE, "cel": f"I{L_STATUS}", "antes": i26, "novo": NOTA_STATUS,
                "ref": REF_6A + "; nota do status (texto)"})
    return out


# Textos que a 6a deixaria falsos (ruling do controlador na Task 7): o resumo do pneu e o item 7 do guia na 00_README e o
# rótulo do μ efetivo na 10_TIRE. O 04_DECISION_REGISTER C17 fica: é o registro histórico do DR-013.
# (aba, célula, trecho de hoje, trecho novo): só o trecho muda, o resto da frase fica
TEXTOS_6A = (
    ("00_README", "A37", "fora da faixa do dataset o resultado é N/A",
     "fora da faixa do dataset, Fz (> 0), α e γ entram limitados ao valor mínimo ou máximo da faixa, com aviso no "
     "status (EDR-056); Fz ≤ 0 e κ fora da faixa dão N/A"),
    ("00_README", "A53", "valor fora da faixa do pneu",
     "Fz ≤ 0 ou κ fora da faixa do pneu; Fz, α e γ fora da faixa entram limitados ao limite, com aviso, EDR-056"),
    ("10_TIRE", "A31", "√(Fx²+Fy²)/Fz", "√(Fx²+Fy²)/Fz usado (limitado à faixa com aviso, EDR-056)"),
)
REF_6A_TEXTO = REF_6A + "; texto que a 6a deixaria falso (ruling do controlador, Task 7)"


# ---------------------------------------------------------------------------------------------------------------------
# 6b e 6c: textos
# ---------------------------------------------------------------------------------------------------------------------
# (aba, célula, trecho de hoje, trecho novo): só o trecho muda, o resto da frase fica (spec §12.1)
TEXTOS_6B = (
    ("13_VALIDATION", "A25", "16/16 Test Cases PASS em suspension_model_V2_0.py / test_suspension_model_V2_0.py",
     "54/54 Test Cases PASS em suspension_model.py / test_suspension_model.py"),
    ("00_README", "A8", "(suspension_model_V2_1.py, tire_model_V2_2.py e seus testes",
     "(suspension_model.py, tire_model.py e seus testes"),
    ("14_DASHBOARD", "A19", "Status de Validação", "Status de verificação (identidades algébricas, EDR-007)"),
)


def _changelog() -> list[tuple[str, str]]:
    lin = LINHA_CHANGELOG
    return [
        (f"A{lin}", "2026-10"),
        (f"B{lin}", "V4.1, equilibrio na planilha (incremento 4): abas 16_EQUILIBRIO a 22_CADEIA (analise de par do "
                    "Milliken com as regras do EDR-056, gradiente de subesterco K pelas parcelas, tabela de troca da "
                    "barra dianteira, cobertura e cadeia de decisao) e bloco novo do 14_DASHBOARD; nomes de "
                    "data/nomes_V4_1.csv; 10_TIRE com Fz, alfa e gama limitados a faixa do .tir com aviso (EDR-056); "
                    "textos da spec 12.1. Gerada da tag v4.0 por python/gerar_v4_1.py."),
        (f"C{lin}", f"EDR-060 (equilibrio na planilha V4.1), EDR-056 (10_TIRE), EDR-059 (faixa SUS-10XX). Plano: {PLANO}. "
                    f"Registro celula a celula: {REGISTRO}. NOT_VALIDATED fisicamente."),
    ]


def alteracoes_v4_1(origem: Path | None = None) -> list[dict]:
    """Células existentes que a V4.1 muda em relação à V4.0 (item 6 das autorizações), cada uma com aba, cel, antes
    (lido da V4.0), novo e ref. Sem `origem`, lê o arquivo da tag v4.0 (só leitura, numa pasta temporária)."""
    with tempfile.TemporaryDirectory(prefix="dds_v41_") as tmp:
        caminho = Path(origem) if origem is not None else ae.extrair_v40(Path(tmp))
        wb = openpyxl.load_workbook(caminho)
        try:
            out = _alteracoes_tire(wb[ABA_TIRE])
            for textos, ref in ((TEXTOS_6A, REF_6A_TEXTO), (TEXTOS_6B, REF_6B)):
                for aba, cel, trecho, novo in textos:
                    antes = wb[aba][cel].value
                    if not isinstance(antes, str) or antes.count(trecho) != 1:
                        raise ValueError(f"{aba}!{cel}: o texto de hoje não tem o trecho {trecho!r}: {antes!r}")
                    out.append({"aba": aba, "cel": cel, "antes": antes, "novo": antes.replace(trecho, novo), "ref": ref})
            out.append({"aba": "00_README", "cel": "A3", "antes": wb["00_README"]["A3"].value, "novo": TEXTO_A3,
                        "ref": REF_6C})
            for cel, novo in _changelog():
                antes = wb["15_CHANGELOG"][cel].value
                if antes is not None:
                    raise ValueError(f"15_CHANGELOG!{cel} já preenchida na origem: {antes!r}")
                out.append({"aba": "15_CHANGELOG", "cel": cel, "antes": None, "novo": novo, "ref": REF_6C})
        finally:
            wb.close()
    return out


# ---------------------------------------------------------------------------------------------------------------------
# Geração
# ---------------------------------------------------------------------------------------------------------------------
def _recusar(caminho_v40: Path, saida: Path) -> None:
    """Recusa saída dentro de excel/ ou igual à origem. Só compara caminhos: não abre nenhum arquivo."""
    s = saida.resolve()
    if s == Path(caminho_v40).resolve():
        raise SystemExit(f"saída igual à origem: recusado ({saida})")
    if s.parent == PASTA_EXCEL.resolve() or PASTA_EXCEL.resolve() in s.parents:
        raise SystemExit(f"saída em excel/ recusada ({saida}): gere em build/ ou numa pasta temporária; a cópia "
                         "para excel/ é um passo à parte, num commit próprio")


def conferir_tabela(tabela: list[Nome]) -> None:
    """A tabela (data/nomes_V4_1.csv) tem de ser exatamente `aba_equilibrio.NOMES_V41`: as fórmulas usam esses nomes."""
    if list(tabela) != list(ae.NOMES_V41):
        dif = [n.nome for n in set(tabela) ^ set(ae.NOMES_V41)][:5]
        raise ValueError(f"data/nomes_V4_1.csv difere de aba_equilibrio.NOMES_V41 (regrave com --nomes-csv): {dif}")


def montar_v41(wb, alteracoes: list[dict]) -> tuple[list[dict], list[dict]]:
    """Workbook da V4.0 (aberto, com os nomes da V4.0) → V4.1 em memória: abas novas e bloco do 14 (`montar`), nomes
    da V4.1 e `alteracoes` (cada "antes" tem de bater com a célula). Devolve (registro das alterações, registro das
    fórmulas novas [{aba, cel, familia, formula}]). Não salva. Usado por `gerar` e por `aba_equilibrio.copia_v41`."""
    novas = ae.montar(wb)
    ae.acrescentar_nomes(wb)
    changelog = wb["15_CHANGELOG"]
    for col in "ABCDE":   # mesma formatação da linha anterior, como na V4.0
        changelog[f"{col}{LINHA_CHANGELOG}"]._style = copy.copy(changelog[f"{col}{LINHA_CHANGELOG - 1}"]._style)
    registro = []
    for alt in alteracoes:
        c = wb[alt["aba"]][alt["cel"]]
        if c.value != alt["antes"]:
            raise ValueError(f"{alt['aba']}!{alt['cel']}: a origem não é a V4.0 esperada ({c.value!r} != "
                             f"{alt['antes']!r})")
        c.value = alt["novo"]
        registro.append(dict(alt))
    wb.calculation.fullCalcOnLoad = True
    return registro, novas


def resumo_novas(novas: list[dict]) -> list[dict]:
    """Linhas do registro para as fórmulas novas, uma por aba e família: quantas células, a primeira e a última (ordem
    de escrita) e o texto da fórmula da primeira. O texto de cada célula é reproduzível pelo gerador (CA-07, TC-WB41-002)."""
    grupos: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for r in novas:
        grupos[(r["aba"], r["familia"])].append(r)
    out = []
    for (aba, fam), rs in grupos.items():
        out.append({"aba": aba, "cel": f"{len(rs)} células ({rs[0]['cel']} … {rs[-1]['cel']})", "antes": None,
                    "novo": f"família {fam}; exemplo {rs[0]['cel']}: {rs[0]['formula']}", "ref": REF_NOVAS})
    return out


def gerar(caminho_v40: Path, saida: Path, tabela: list[Nome] | None = None) -> list[dict]:
    """Copia a V4.0, escreve as abas novas e o bloco do 14 (`aba_equilibrio.montar`), acrescenta os nomes da V4.1
    (escopo do workbook), aplica `alteracoes_v4_1()` e grava `saida`. Devolve o registro: as alterações (aba, cel, antes,
    novo, ref), uma linha por aba e família de fórmula nova e uma linha final de resumo dos nomes (aba "(workbook)")."""
    caminho_v40, saida = Path(caminho_v40), Path(saida)
    _recusar(caminho_v40, saida)
    tabela = ler_tabela(CSV_NOMES) if tabela is None else tabela
    conferir_tabela(tabela)
    alts = alteracoes_v4_1(caminho_v40)
    wb = openpyxl.load_workbook(caminho_v40)
    try:
        existentes = list(wb.defined_names)
        registro, novas = montar_v41(wb, alts)
        registro += resumo_novas(novas)
        registro.append({"aba": "(workbook)", "cel": "nomes definidos", "antes": f"{len(existentes) - 1} nomes da V4.0 e a lista Origens",
                         "novo": f"{len(tabela)} nomes de data/nomes_V4_1.csv acrescentados (escopo do workbook); "
                                 f"{len(novas)} fórmulas novas nas abas {ABAS_NOVAS[0]} a {ABAS_NOVAS[-1]} e no "
                                 "14_DASHBOARD", "ref": REF_NOVAS})
        saida.parent.mkdir(parents=True, exist_ok=True)
        wb.save(saida)
    finally:
        wb.close()
    return registro


INTRO_REGISTRO = (
    "# Registro de células da V4.1 (gerado por python/gerar_v4_1.py)\n\n"
    "Gerado da tag v4.0; não edite à mão (regere com `uv run python -X utf8 python/gerar_v4_1.py --registro "
    f"{REGISTRO}`). Primeiro, cada célula existente que a V4.1 muda (autorizações 6a, 6b e 6c do plano {PLANO}), "
    "com o texto de antes e o de depois. A 6a (10_TIRE, EDR-056) troca a guarda `X$26<>\"OK\"` por "
    "`LEFT(X$26,2)<>\"OK\"` e usa Fz (linha 22), α do arquivo (38) e γ do arquivo (39) limitados à faixa do .tir "
    "(`MIN(MAX(…))`); dentro da faixa nenhum valor muda (TC-WB41-003). Depois, as fórmulas novas, uma linha por aba e "
    "família do catálogo (F1-F19, REF), com o número de células, a primeira e a última e o texto da primeira: as 117181 "
    "fórmulas novas somam cerca de 31 MB de texto, e o de cada célula é reproduzível pelo gerador (CA-07, TC-WB41-002). "
    "Por fim, o resumo dos nomes definidos.\n\n")


def escrever_registro_v41(registro: list[dict], caminho: Path) -> None:
    """Registro em Markdown: a introdução (INTRO_REGISTRO) e a tabela de `atualizar_planilha.escrever_registro`, LF."""
    caminho = Path(caminho)
    escrever_registro(registro, caminho)
    caminho.write_text(INTRO_REGISTRO + caminho.read_text(encoding="utf-8"), encoding="utf-8", newline="\n")


def main() -> None:
    p = argparse.ArgumentParser(description="Gera a V4.1 (equilíbrio na planilha) a partir da V4.0 da tag v4.0.")
    p.add_argument("--v40", type=Path, default=None, help="arquivo da V4.0 (padrão: extraído da tag v4.0)")
    p.add_argument("--saida", type=Path, default=SAIDA_PADRAO)
    p.add_argument("--registro", type=Path, default=None, help="registro célula a célula (.md)")
    p.add_argument("--nomes-csv", type=Path, default=None,
                   help="regrava a tabela de nomes da V4.1 (CSV) a partir de aba_equilibrio.NOMES_V41 antes de gerar")
    a = p.parse_args()
    if a.nomes_csv is not None:
        escrever_tabela(ae.NOMES_V41, a.nomes_csv)
    tabela = ler_tabela(a.nomes_csv or CSV_NOMES)
    with tempfile.TemporaryDirectory() as tmp:
        if a.v40 is None:
            _recusar(Path(tmp) / "V4_0_tag.xlsx", a.saida)
            v40 = ae.extrair_v40(Path(tmp))
        else:
            v40 = a.v40
        reg = gerar(v40, a.saida, tabela)
    if a.registro:
        escrever_registro_v41(reg, a.registro)
    n_alt = sum(1 for r in reg if r["ref"] != REF_NOVAS)
    print(f"salvo {a.saida}: {len(tabela)} nomes novos, {n_alt} células existentes alteradas")


if __name__ == "__main__":
    main()
