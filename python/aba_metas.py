"""Fichas de meta na planilha da V4.2 (incremento 5; EDR-061): colunas H a Q da 03_TARGETS, as linhas TARGET-006 e 007, a aba
23_LISTAS e o bloco de metas do 14_DASHBOARD.

O oráculo é `metas.py`; cada família de fórmula do catálogo G1-G6 (plano docs/superpowers/plans/2026-10-23-v4-2-inc5-metas-e-
coerencia.md) reproduz a função citada, e o TC-WB42-006 confere as duas contra as mesmas fichas:
- G1  03_TARGETS I5:I11 (Alavanca): `INDEX`/`MATCH` na tabela de relações da 23_LISTAS (uma fonte só, spec §5 item 8), com
      "sem relação na 23_LISTAS" quando a meta não tem relação (`metas.alavanca`);
- G2  03_TARGETS K5:K11 (Valor atual do modelo): `=nome` do resultado da relação de fonte M da meta; L5:L11 (Pior origem da
      grandeza): `=nome` da célula de pior origem que já existe (14_DASHBOARD coluna E ou 16_EQUILIBRIO coluna E, por nome);
      para TARGET-004 e 005 o 14_DASHBOARD não tem linha de resultado, e L é a fórmula de pior origem gerada pelo grafo, a mesma
      de `atualizar_planilha._formula_pior_origem` que preenche a coluna E do dashboard (uma célula de pior origem a mais, nunca
      um texto digitado);
- G3  R:T (auxiliares ocultas): vocabulário de Papel, Ciclo de vida e Congelamento (`metas.campo_valido`). `EXACT` na lista, e
      não `ISNUMBER(MATCH(...))`: o `MATCH` do Excel não distingue "meta" de "META", e o oráculo distingue (a planilha nunca
      mostra um estado que o oráculo não reproduz);
- G4  U (auxiliar oculta): quantas outras METAs há na mesma alavanca (`metas.sobredeterminadas`); `EXACT` pelo mesmo motivo;
- G5  O5:O11 (Coerência): a cadeia da decisão 4, `<estado> · <pior origem>` (`metas.estado_coerencia`, `metas.texto_coerencia`);
      o formato `EDR-ddd` é `LEN = 7`, `EXACT(LEFT(...), "EDR-")` e três dígitos de fato (um `--MID` sozinho aceitaria "EDR-1e2");
- G6  14_DASHBOARD B50 e B51: pior coerência e quantos compromissos declarados (`metas.pior_coerencia`, `metas.compromissos`).
Fora das famílias: os textos fixos, o cabeçalho, as linhas 10 e 11, as listas e a tabela de relações da 23_LISTAS e as alterações
de células existentes (`alteracoes_03`: A3, A11 e A12, e as notas que descem para A14 e A15).

Entradas novas (nome `ent_03_*`): C:G das linhas 10 e 11 e H, J, M, N e Q das linhas 5 a 11 (`celulas_de_entrada_v42`). Resultados
novos (`res_03_*`, `res_14_*`): I, K, L e O das linhas 5 a 11 e B50 e B51; as auxiliares R:U são cálculo auxiliar, sem nome.
Nomes das listas (fora do CSV, como `Origens`): Lista_Papel, Lista_Ciclo_Meta, Lista_Congelamento e Lista_Coerencia.
Nenhuma função salva a planilha viva nem a de referência.
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import column_index_from_string, get_column_letter as COL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import metas
from aba_equilibrio import ABA14, _destino, _estilo, _estilos
from atualizar_planilha import _formula_pior_origem
from nomes_v4_0 import NOME_VALIDO, Nome, parece_referencia, slug
from raizes import Grafo

ABA03, ABA23 = "03_TARGETS", "23_LISTAS"
ALVOS = tuple(f"TARGET-00{i}" for i in range(1, 8))
LINHA_TARGET = {t: 5 + i for i, t in enumerate(ALVOS)}            # TARGET-001 na linha 5 … TARGET-007 na linha 11
LINHAS = tuple(LINHA_TARGET.values())
SUS_ALVO = {"TARGET-001": "SUS-205", "TARGET-002": "SUS-205", "TARGET-003": "SUS-208", "TARGET-004": "SUS-301",
            "TARGET-005": "SUS-302", "TARGET-006": "SUS-1001", "TARGET-007": "SUS-1004"}
NOVAS_LINHAS = {"TARGET-006": "Gradiente de subesterço K (SUS-1001)",
                "TARGET-007": "Comportamento no limite: push/loose e C_N (SUS-1004)"}
COLS_ENT = "HJMNQ"                    # entradas novas das fichas (todas as linhas 5 a 11)
COLS_FORMULA = {"I": "G1", "K": "G2", "L": "G2", "O": "G5", "R": "G3", "S": "G3", "T": "G3", "U": "G4"}
CAB_H = {"H": "Papel", "I": "Alavanca", "J": "Congelamento", "K": "Valor atual do modelo", "L": "Pior origem da grandeza",
         "M": "Compromisso declarado", "N": "EDR do compromisso", "O": "Coerência", "P": "Ficha", "Q": "Objetivo atendido",
         "R": "Papel válido (G3)", "S": "Ciclo de vida válido (G3)", "T": "Congelamento válido (G3)",
         "U": "Outras METAs na alavanca (G4)"}
OBJETIVO_TBD = "TBD: lista de objetivos da equipe a definir em EDR"
FICHA_TXT = "docs/metas/{t}.md"
NOTA_SEM_FONTE = "linha sem fonte não existe; ausência de linha = não avaliado"

# 03_TARGETS: valores iniciais (decisão 6; nenhum inventado)
INICIAL = {"papel": "TBD", "congelamento": "F0", "valor": "TBD", "owner": "TBD"}

# ---------------------------------------------------------------------------------------------------------------------
# 23_LISTAS: posições (uma fonte para as fórmulas, os nomes das listas e o teste)
# ---------------------------------------------------------------------------------------------------------------------
_SECOES = (   # chave de VOCAB, título, cabeçalho, nome da lista (ou None)
    ("papel", "Papel (spec §5, item 2)", ("Valor", "Significado"), "Lista_Papel"),
    ("ciclo", "Ciclo de vida da meta (Guia G.1)", ("Valor", "Significado"), "Lista_Ciclo_Meta"),
    ("congelamento", "Congelamento (Guia §3.2)", ("Valor", "Significado", "Uso"), "Lista_Congelamento"),
    ("coerencia", "Coerência da ficha (spec §4.4)", ("Valor", "Significado"), "Lista_Coerencia"),
    ("fonte", "Fonte da relação", ("Valor", "Significado"), None),
)


def _posicoes() -> tuple[dict[str, tuple[int, int, int]], int]:
    """(chave -> (linha do título, 1ª linha de valor, última linha de valor), linha do título das relações). A aba tem o título
    em A1, a explicação em A2 e começa as seções na linha 4; cada seção: título, cabeçalho, valores e uma linha em branco."""
    pos, r = {}, 4
    for chave, *_ in _SECOES:
        n = len(metas.VOCAB[chave])
        pos[chave] = (r, r + 2, r + 1 + n)
        r += n + 3
    return pos, r


POS, LINHA_TITULO_REL = _posicoes()
LINHA_REL0 = LINHA_TITULO_REL + 3         # 1ª linha de relação (título, nota e cabeçalho antes)
LINHA_REL1 = LINHA_REL0 + len(metas.RELACOES) - 1
CAB_REL = ("Nº", "Tipo", "Meta", "Alavanca", "Fonte (L/E/M)", "Fonte", "Resultado", "Entrada da alavanca", "Nota")
LISTAS_V42: dict[str, str] = {nome: f"'{ABA23}'!$A${POS[ch][1]}:$A${POS[ch][2]}" for ch, _, _, nome in _SECOES if nome}


# ---------------------------------------------------------------------------------------------------------------------
# Nomes e entradas
# ---------------------------------------------------------------------------------------------------------------------
def _ref(celula: str) -> tuple[str, int]:
    m = re.fullmatch(r"([A-Z]+)(\d+)", celula)
    return m.group(1), int(m.group(2))


def _nomes_v42() -> list[Nome]:
    """Nomes da V4.2 (data/nomes_V4_2.csv): `ent_03_<id>_<célula>` nas entradas novas (rótulo = o ID da meta, como os da V4.0);
    `res_03_<coluna>_<id>` em I, K, L e O das linhas 5 a 11; e os dois do dashboard (`res_14_<rótulo>_b<linha>`)."""
    out = []
    for t, r in LINHA_TARGET.items():
        cols = ("C", "D", "E", "F", "G") if r in (10, 11) else ()
        for c in (*cols, *COLS_ENT):
            out.append(Nome(f"ent_03_{slug(t)}_{c.lower()}{r}", ABA03, f"{c}{r}", "ent", SUS_ALVO[t], t))
        for c in "IKLO":
            rot = f"{CAB_H[c]} {t}"
            out.append(Nome(f"res_03_{slug(rot)}", ABA03, f"{c}{r}", "res", SUS_ALVO[t], rot))
    for r, rot in ((50, "Pior coerência das fichas"), (51, "Compromissos declarados")):
        out.append(Nome(f"res_14_{slug(rot)}_b{r}", ABA14, f"B{r}", "res", "", rot))
    out.sort(key=lambda n: (n.aba, _ref(n.celula)[1], column_index_from_string(_ref(n.celula)[0])))
    vistos = set()
    for n in out:
        if n.nome in vistos or not NOME_VALIDO.fullmatch(n.nome) or parece_referencia(n.nome):
            raise ValueError(f"nome inválido ou repetido na V4.2: {n.nome}")
        vistos.add(n.nome)
    return out


NOMES_V42: list[Nome] = _nomes_v42()
_NOME = {(n.aba, n.celula): n.nome for n in NOMES_V42}


def celulas_de_entrada_v42() -> set[tuple[str, str]]:
    """Células de entrada da V4.2: as da V4.1 (`aba_equilibrio.celulas_de_entrada_v41()`), C:G das linhas 10 e 11 da 03_TARGETS e
    H, J, M, N e Q das linhas 5 a 11."""
    from aba_equilibrio import celulas_de_entrada_v41
    return celulas_de_entrada_v41() | {(ABA03, f"{c}{r}") for r in (10, 11) for c in "CDEFG"} \
        | {(ABA03, f"{c}{r}") for r in LINHAS for c in COLS_ENT}


def acrescentar_nomes(wb) -> None:
    """Nomes definidos da V4.2 (NOMES_V42) e as quatro listas de vocabulário, no escopo do workbook; recusa nome já existente."""
    existentes = set(wb.defined_names)
    for n in NOMES_V42:
        _definir(wb, existentes, n.nome, f"'{n.aba}'!${_ref(n.celula)[0]}${_ref(n.celula)[1]}")
    for nome, ref in LISTAS_V42.items():
        _definir(wb, existentes, nome, ref)


def _definir(wb, existentes: set, nome: str, ref: str) -> None:
    if nome in existentes:
        raise ValueError(f"nome já definido: {nome}")
    existentes.add(nome)
    wb.defined_names[nome] = DefinedName(nome, attr_text=ref)


# ---------------------------------------------------------------------------------------------------------------------
# Alterações em células existentes da 03_TARGETS (autorização 5 (c)) para o registro e para `gerar_v4_2.aplicar`
# ---------------------------------------------------------------------------------------------------------------------
REF_03 = ("EDR-061, decisões 4 a 6 (fichas de meta); plano, autorização 5 (c) (03_TARGETS: A3, A11 e A12 descem para A14 e A15, "
          "A11 vira o ID TARGET-007); teste: TC-WB42-001")


def alteracoes_03(wb) -> list[dict]:
    """Células existentes da 03_TARGETS que a V4.2 muda (antes lido de `wb`, a V4.1 aberta): o título A3, a nota A11 (que vira o
    ID TARGET-007), a nota A12 (que sai) e as duas notas que descem para A14 e A15. Não grava nada."""
    ws = wb[ABA03]
    a3, a11, a12 = ws["A3"].value, ws["A11"].value, ws["A12"].value
    if not (isinstance(a11, str) and a11.startswith("'Target' = TBD") and isinstance(a12, str)
            and a12.startswith("Parametros como Motion Ratio") and ws["A14"].value is None and ws["A15"].value is None):
        raise ValueError(f"{ABA03}: notas A11/A12 inesperadas na V4.1: {a11!r} / {a12!r}")
    return [
        {"aba": ABA03, "cel": "A3", "antes": a3, "novo": f"{a3}; V4.2: fichas de meta nas colunas H a Q (EDR-061)", "ref": REF_03},
        {"aba": ABA03, "cel": "A11", "antes": a11, "novo": "TARGET-007", "ref": REF_03},
        {"aba": ABA03, "cel": "A12", "antes": a12, "novo": None, "ref": REF_03},
        {"aba": ABA03, "cel": "A14", "antes": None, "novo": a11, "ref": REF_03},
        {"aba": ABA03, "cel": "A15", "antes": None, "novo": a12, "ref": REF_03},
    ]


# ---------------------------------------------------------------------------------------------------------------------
# Fórmulas (G1-G6)
# ---------------------------------------------------------------------------------------------------------------------
def formula_g1(r: int) -> str:
    c, d = f"'{ABA23}'!$C${LINHA_REL0}:$C${LINHA_REL1}", f"'{ABA23}'!$D${LINHA_REL0}:$D${LINHA_REL1}"
    return f'=IFERROR(INDEX({d},MATCH(A{r},{c},0)),"{metas.SEM_RELACAO}")'


def formula_g3(lista: str, campo: str, r: int) -> str:
    return f"=SUMPRODUCT(--EXACT({lista},{campo}{r}))>0"


def formula_g4(r: int) -> str:
    return (f'=IF(I{r}="{metas.SEM_RELACAO}",0,SUMPRODUCT(($I$5:$I$11=I{r})*EXACT($H$5:$H$11,"META"))'
            f'-IF(EXACT(H{r},"META"),1,0))')


def formula_g5(r: int) -> str:
    """Cadeia IF da decisão 4 (a primeira regra que bate decide) e `· pior origem`; espelha `metas.estado_coerencia`."""
    n = f'N{r}&""'
    edr = (f'AND(LEN({n})=7,EXACT(LEFT({n},4),"EDR-"),'
           f'SUMPRODUCT(--ISNUMBER(FIND(MID({n},{{5,6,7}},1),"0123456789")))=3)')
    preenchido = f'AND(TRIM(M{r}&"")<>"",UPPER(TRIM(M{r}&""))<>"TBD")'
    tbd = f'OR(H{r}="TBD",TRIM(C{r}&"")="",UPPER(TRIM(C{r}&""))="TBD")'
    cadeia = (
        f'IF(NOT(R{r}),"CONFLITO: Papel fora do vocabulário",'
        f'IF(NOT(S{r}),"CONFLITO: Ciclo de vida fora do vocabulário",'
        f'IF(NOT(T{r}),"CONFLITO: Congelamento fora do vocabulário",'
        f'IF(AND(EXACT(H{r},"META"),U{r}>0),"CONFLITO: sobredeterminação ("&I{r}&"; escolha qual manda)",'
        f'IF({tbd},"NÃO AVALIÁVEL: {metas.MOTIVO_NAO_AVALIAVEL}",'
        f'IF({preenchido},IF({edr},"COMPROMISSO DECLARADO","CONFLITO: {metas.MOTIVO_SEM_EDR}"),'
        f'"COERENTE ({metas.MOTIVO_COERENTE})"))))))')
    return f'={cadeia}&"{metas.SEPARADOR}"&L{r}'


def formula_g6_pior() -> str:
    faixa = f"'{ABA03}'!$O$5:$O$11"
    return ("=" + "".join(f'IF(COUNTIF({faixa},"{e}*")>0,"{e}",' for e in metas.PIOR_PRIMEIRO[:-1])
            + f'"{metas.PIOR_PRIMEIRO[-1]}"' + ")" * (len(metas.PIOR_PRIMEIRO) - 1))


def formula_g6_conta() -> str:
    return f"=COUNTIF('{ABA03}'!$O$5:$O$11,\"COMPROMISSO DECLARADO*\")"


# ---------------------------------------------------------------------------------------------------------------------
# Montagem
# ---------------------------------------------------------------------------------------------------------------------
def _vazia(ws, cel: str) -> None:
    if ws[cel].value is not None:
        raise ValueError(f"{ws.title}!{cel} já preenchida; as metas só usam células vazias (as existentes passam por alteracoes_03)")


def _put(registro: list, ws, cel: str, familia: str, formula: str) -> None:
    _vazia(ws, cel)
    ws[cel] = formula
    registro.append({"aba": ws.title, "cel": cel, "familia": familia, "formula": formula})


def _alvo_m() -> dict[str, metas.Relacao]:
    out = {}
    for r in metas.RELACOES:
        if r.tipo == "alavanca" and r.fonte_tipo == "M":
            if r.meta in out:
                raise ValueError(f"mais de uma relação de fonte M para {r.meta}")
            out[r.meta] = r
    if set(out) != set(ALVOS):
        raise ValueError(f"toda meta precisa de uma relação de fonte M: faltam {sorted(set(ALVOS) - set(out))}")
    return out


# pior origem ao lado (G2, coluna L): células que já existem, por nome
PIOR_EXISTENTE = {"TARGET-001": "res_14_frequencia_natural_dianteira_e15", "TARGET-002": "res_14_frequencia_natural_traseira_e16",
                  "TARGET-003": "res_14_distribuicao_de_rigidez_dianteira_e17",
                  "TARGET-006": "res_14_gradiente_de_subesterco_k_linear_e41",
                  "TARGET-007": "res_14_equilibrio_no_limite_push_loose_e42"}


def _linha_alvo(ws) -> None:
    """Linhas 10 e 11 (TARGET-006 e 007): A:G com o estilo da linha 9, valores como as linhas existentes."""
    for t, r in (("TARGET-006", 10), ("TARGET-007", 11)):
        for c in "ABCDEFG":
            ws[f"{c}{r}"]._style = copy.copy(ws[f"{c}9"]._style)
        if r == 10:
            _vazia(ws, "A10")
            ws["A10"] = t                      # A11 (que era a nota) é escrita por `alteracoes_03`
        ws[f"B{r}"] = NOVAS_LINHAS[t]
        ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"], ws[f"F{r}"], ws[f"G{r}"] = "TBD", "TBD", "TBD", "-", "-"


def _aba03(wb, registro: list) -> None:
    ws = wb[ABA03]
    m = _alvo_m()
    # as notas A11 e A12 descem para A14 e A15 (os valores vêm de `alteracoes_03`); só o estilo e as mesclagens aqui
    for celula in ("A11", "A12"):
        ws.unmerge_cells(f"{celula}:G{celula[1:]}")
    estilo_nota = {"A14": copy.copy(ws["A11"]._style), "A15": copy.copy(ws["A12"]._style)}
    for celula in ("A11", "A12"):
        for c in "BCDEFG":
            _vazia(ws, f"{c}{celula[1:]}")
    ws["A12"]._style = copy.copy(ws["A13"]._style)
    for c in "ABCDEFG":
        ws[f"{c}11"]._style = copy.copy(ws[f"{c}9"]._style)
    ws["A14"]._style, ws["A15"]._style = estilo_nota["A14"], estilo_nota["A15"]
    ws.merge_cells("A14:G14")
    ws.merge_cells("A15:G15")
    _linha_alvo(ws)
    # cabeçalho
    for c, rot in CAB_H.items():
        _vazia(ws, f"{c}4")
        ws[f"{c}4"] = rot
        ws[f"{c}4"]._style = copy.copy(ws["G4"]._style)
    g = None
    for t, r in LINHA_TARGET.items():
        for c in CAB_H:
            ws[f"{c}{r}"]._style = copy.copy(ws["G5"]._style)
            ws[f"{c}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        # entradas (valores iniciais, decisão 6)
        for c, v in (("H", INICIAL["papel"]), ("J", INICIAL["congelamento"]), ("P", FICHA_TXT.format(t=t)), ("Q", OBJETIVO_TBD)):
            _vazia(ws, f"{c}{r}")
            ws[f"{c}{r}"] = v
        # G1, G2
        _put(registro, ws, f"I{r}", "G1", formula_g1(r))
        _put(registro, ws, f"K{r}", "G2", "=" + m[t].resultado)
        if t in PIOR_EXISTENTE:
            _put(registro, ws, f"L{r}", "G2", "=" + PIOR_EXISTENTE[t])
        else:
            g = g or Grafo(wb)
            res = _destino(wb, m[t].resultado)
            celulas = g.celulas_origem(res)
            if not celulas:
                raise RuntimeError(f"{t}: nenhuma raiz com coluna Origem em {m[t].resultado}")
            _put(registro, ws, f"L{r}", "G2", _formula_pior_origem(celulas))
        # G3, G4, G5
        for col, lista, campo in (("R", "Lista_Papel", "H"), ("S", "Lista_Ciclo_Meta", "D"), ("T", "Lista_Congelamento", "J")):
            _put(registro, ws, f"{col}{r}", "G3", formula_g3(lista, campo, r))
        _put(registro, ws, f"U{r}", "G4", formula_g4(r))
        _put(registro, ws, f"O{r}", "G5", formula_g5(r))
    for c in "RSTU":
        ws.column_dimensions[c].hidden = True
    for c, w in zip("HIJKLMNOPQ", (14, 30, 16, 22, 30, 30, 18, 60, 26, 36)):
        ws.column_dimensions[c].width = w
    # validação de dados nas entradas com lista (Papel, Ciclo de vida da meta = coluna Status, Congelamento)
    for faixa, lista in (("H5:H11", "Lista_Papel"), ("D5:D11", "Lista_Ciclo_Meta"), ("J5:J11", "Lista_Congelamento")):
        dv = DataValidation(type="list", formula1=f"={lista}", allow_blank=True, showErrorMessage=True,
                            errorTitle="Valor fora do vocabulário",
                            error=f"Escolha um valor da lista ({lista}, aba {ABA23}).")
        ws.add_data_validation(dv)
        dv.add(faixa)


def _aba23(wb) -> None:
    ws = wb.create_sheet(ABA23)
    est = _estilos(wb)
    ws["A1"] = "23_LISTAS — vocabulários e relações das metas (V4.2; EDR-061)"
    _estilo(ws["A1"], est, "titulo")
    ws["A2"] = ("Fonte única das listas e das relações: python/metas.py (oráculo) e esta aba dizem o mesmo (TC-WB42-001). A coluna "
                "Alavanca da 03_TARGETS lê a tabela de relações abaixo. Relação sem fonte não existe: " + NOTA_SEM_FONTE
                + ", nunca \"ok\". Fonte L = livro e página conferidos na imagem do PDF; E = EDR existente e aprovado; "
                "M = nome do resultado e nome da entrada da alavanca (a entrada é raiz do resultado no grafo de fórmulas).")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:I2")
    ws.row_dimensions[2].height = 62
    for chave, titulo, cab, _ in _SECOES:
        r_t, r0, r1 = POS[chave]
        ws[f"A{r_t}"] = titulo
        _estilo(ws[f"A{r_t}"], est, "secao")
        for j, t in enumerate(cab):
            ws[f"{COL(1 + j)}{r_t + 1}"] = t
            _estilo(ws[f"{COL(1 + j)}{r_t + 1}"], est, "cab")
        for i, valor in enumerate(metas.VOCAB[chave]):
            r = r0 + i
            sig = metas.SIGNIFICADO[chave][valor]
            ws[f"A{r}"] = valor
            for j, t in enumerate(sig if isinstance(sig, tuple) else (sig,)):
                ws[f"{COL(2 + j)}{r}"] = t
            for c in "ABC":
                ws[f"{c}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws[f"A{LINHA_TITULO_REL}"] = "Relações meta → alavanca (spec §5, item 8)"
    _estilo(ws[f"A{LINHA_TITULO_REL}"], est, "secao")
    ws[f"A{LINHA_TITULO_REL + 1}"] = (f"Tabela esparsa: {NOTA_SEM_FONTE}. A linha do tipo regime é informativa e fica fora das "
                                      "fórmulas.")
    for j, t in enumerate(CAB_REL):
        ws[f"{COL(1 + j)}{LINHA_REL0 - 1}"] = t
        _estilo(ws[f"{COL(1 + j)}{LINHA_REL0 - 1}"], est, "cab")
    for i, r in enumerate(metas.RELACOES):
        lin = LINHA_REL0 + i
        for j, v in enumerate((r.n, r.tipo, r.meta, r.alavanca, r.fonte_tipo, r.fonte, r.resultado, r.entrada,
                               metas.NOTAS_RELACAO[r.n])):
            ws[f"{COL(1 + j)}{lin}"] = v
            ws[f"{COL(1 + j)}{lin}"].alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHI", (24, 58, 38, 38, 14, 46, 46, 42, 80)):
        ws.column_dimensions[col].width = w


def _aba14(wb, registro: list) -> None:
    ws = wb[ABA14]
    for r in range(48, 53):
        for c in "ABCDE":
            if r in (49, 50, 51) and ws[f"{c}{r}"].value is not None:
                raise ValueError(f"{ABA14}!{c}{r} já preenchida; o bloco das metas só usa linhas vazias")
    est = {k: ws[c]._style for k, c in (("secao", "A32"), ("rotulo", "A33"), ("valor", "B33"), ("un", "C33"), ("ref", "D33"),
                                        ("pior", "E33"))}
    ws["A49"] = "Metas (V4.2; 03_TARGETS)"
    ws["A49"]._style = copy.copy(est["secao"])
    ws.merge_cells("A49:E49")
    for r, rot, formula in ((50, "Pior coerência das fichas", formula_g6_pior()), (51, "Compromissos declarados", formula_g6_conta())):
        ws[f"A{r}"], ws[f"C{r}"], ws[f"D{r}"] = rot, "-", "EDR-061"
        _put(registro, ws, f"B{r}", "G6", formula)
        for col, e in (("A", "rotulo"), ("B", "valor"), ("C", "un"), ("D", "ref"), ("E", "pior")):
            ws[f"{col}{r}"]._style = copy.copy(est[e])
        ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")


def montar(wb) -> list[dict]:
    """Escreve as fichas de meta no workbook da V4.1 (aberto pelo openpyxl, com os nomes da V4.1 e depois de `alteracoes_03`
    aplicada): colunas H:U e linhas 10 e 11 da 03_TARGETS, a aba 23_LISTAS no fim e B50 e B51 do 14_DASHBOARD. Devolve o
    registro [{aba, cel, familia, formula}] de toda célula com fórmula escrita. Não acrescenta os nomes (`acrescentar_nomes`) nem
    salva."""
    if ABA23 in wb.sheetnames:
        raise ValueError(f"a aba {ABA23} já existe")
    if wb.sheetnames[-1] != "22_CADEIA":
        raise ValueError(f"a última aba deveria ser 22_CADEIA; é {wb.sheetnames[-1]}")
    registro: list[dict] = []
    _aba03(wb, registro)      # antes da 23: a fórmula G1 só aponta para ela
    _aba23(wb)
    _aba14(wb, registro)
    return registro


# ---------------------------------------------------------------------------------------------------------------------
# Leitura (teste da planilha)
# ---------------------------------------------------------------------------------------------------------------------
def fichas_da_planilha(xlsx_recalculado: Path) -> list[metas.Ficha]:
    """As sete fichas da 03_TARGETS de uma cópia recalculada (valores, `data_only`), para `metas.estado_coerencia`. A pior origem
    é a da coluna L (a célula que a fórmula G5 usa)."""
    wb = openpyxl.load_workbook(xlsx_recalculado, data_only=True)
    try:
        ws = wb[ABA03]
        out = []
        for t, r in LINHA_TARGET.items():
            if ws[f"A{r}"].value != t:
                raise ValueError(f"{ABA03}!A{r} deveria ser {t}: {ws[f'A{r}'].value!r}")
            out.append(metas.Ficha(t, ws[f"C{r}"].value, ws[f"H{r}"].value, ws[f"D{r}"].value, ws[f"J{r}"].value,
                                   ws[f"M{r}"].value, ws[f"N{r}"].value, ws[f"L{r}"].value))
        return out
    finally:
        wb.close()
