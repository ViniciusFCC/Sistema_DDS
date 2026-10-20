"""Gera a planilha V3.0 a partir da V2.4: revisão metodológica contra a literatura (EDR-045 a EDR-049), 1ª rodada de feedback
(feedback da equipe: origem de cada entrada, valores derivados como fórmula, guia de preenchimento; EDR-035, EDR-050, EDR-051)
e 2ª rodada da revisão metodológica (hipótese do CG não suspenso como entrada, centro da roda separado, N/A e TBD no guia; EDR-050, EDR-035).

Uso:
    uv run python -X utf8 python/atualizar_planilha.py [ORIGEM_V2_4.xlsx] [SAIDA_V3_0.xlsx] [--registro ARQ.md]

Padrão: origem excel/FSAE_EV_Suspension_Fase1_V2_4.xlsx, saída build/FSAE_EV_Suspension_Fase1_V3_0.xlsx
(nunca por cima da entregue). Lê a origem e grava uma cópia nova pelo openpyxl; nenhuma planilha é
recalculada nem salva pelo Excel ou pelo LibreOffice.

Toda célula tocada está em ALTERACOES, com a fonte e o EDR. O registro (--registro) lista, para cada
uma, aba, célula, conteúdo anterior (lido da origem) e novo: é a rastreabilidade que o CLAUDE.md exige
para mudança de fórmula. Referências: docs/revisao-metodologica.md.
"""
import argparse
import copy
import re
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment

AQUI = Path(__file__).resolve().parent
V = "'05_VEHICLE'!"
S = "'06_SUSPENSION'!"
K = "'07_KINEMATICS'!"
E = "'08_VEHICLE_STATE'!"
L = "'09_LOADS'!"
D = "'11_DESIGN'!"
T = "'10_TIRE'!"
AUTOR = "Sistema_DDS V3.0"

# Rigidez vertical do pneu: único valor com fonte no projeto é o do .tir TTC provisório (EDR-038).
K_T_PROVISORIO = 120312.0   # N/m, VERTICAL_STIFFNESS de data/Hoosier_R25B_13x7_TTC_provisional.tir
NOTA_KT = ("PROVISÓRIO (EDR-045): VERTICAL_STIFFNESS do .tir TTC provisório (Hoosier R25B, EDR-038). "
           "Não é o pneu do carro nem medição da equipe: substituir pela rigidez vertical do pneu real. "
           "Milliken p. 591, Gillespie eq. 5-8 e Seward eq. 4.1 põem o pneu em série com a mola.")


def _a(aba, cel, novo, ref, estilo=None, comentario=None):
    return {"aba": aba, "cel": cel, "novo": novo, "ref": ref, "estilo": estilo, "comentario": comentario}


def alteracoes(ky_row: int) -> list[dict]:
    """Lista completa das células alteradas ou criadas na V3.0. ky_row = linha de Ky em 10_TIRE (layout)."""
    a = []
    # ---------------- 05_VEHICLE: massa suspensa por eixo e a_s (EDR-047; Milliken p. 601-603) ----------------
    for r, rot, f, un in (
        (25, "Massa suspensa - eixo dianteiro (m·b/L − m_uF)", f"={V}$B$4*{V}$B$10/{V}$B$6-{V}$B$14", "kg"),
        (26, "Massa suspensa - eixo traseiro (m·a/L − m_uR)", f"={V}$B$4*{V}$B$9/{V}$B$6-{V}$B$15", "kg"),
        (27, "CG suspenso -> eixo dianteiro (a_s = L·m_s,R/m_s)", f"={V}$B$6*{V}$B$26/{V}$B$23", "m"),
    ):
        a += [_a("05_VEHICLE", f"A{r}", rot, "EDR-047", ("05_VEHICLE", 23, "A")),
              _a("05_VEHICLE", f"B{r}", f, "EDR-047; Milliken p. 601-603; Seward eq. 5.14", ("05_VEHICLE", 23, "B")),
              _a("05_VEHICLE", f"C{r}", un, "EDR-047", ("05_VEHICLE", 23, "C")),
              _a("05_VEHICLE", f"D{r}", "SUS-009", "EDR-047", ("05_VEHICLE", 23, "D"))]

    # ---------------- 06_SUSPENSION: pneu em série (EDR-045) ----------------
    st = lambda r, c: ("06_SUSPENSION", r, c)
    a.append(_a("06_SUSPENSION", "A2", "Mola, MR, barra e roll center (B4:B15): valores DEMONSTRATIVOS herdados do T09/E03, para referência e "
                                       "teste das fórmulas; não são o TX/E05 (EDR-012/022).", "EDR-012/022", st(35, "A")))
    a.append(_a("06_SUSPENSION", "A46","Pneu em série e ride rate (EDR-045) — K_T PROVISÓRIO: rigidez vertical do .tir TTC, não é o pneu do carro",
                "EDR-045", st(17, "A")))
    for r, rot, val, un, sid, com in (
        (47, "Rigidez vertical do pneu - dianteiro (K_T)", K_T_PROVISORIO, "N/m", "SUS-211", NOTA_KT),
        (48, "Rigidez vertical do pneu - traseiro (K_T)", K_T_PROVISORIO, "N/m", "SUS-211", NOTA_KT),
    ):
        a += [_a("06_SUSPENSION", f"A{r}", rot, "EDR-045", st(4, "A")),
              _a("06_SUSPENSION", f"B{r}", val, "EDR-045 (valor: .tir provisório, EDR-038)", st(4, "B"), com),
              _a("06_SUSPENSION", f"C{r}", un, "EDR-045", st(4, "C")),
              _a("06_SUSPENSION", f"D{r}", sid, "EDR-045", st(4, "D"))]
    for r, rot, f, un, sid in (
        (49, "Ride rate dianteiro K_R = K_W·K_T/(K_W+K_T)", f"={S}$B$18*{S}$B$47/({S}$B$18+{S}$B$47)", "N/m", "SUS-212"),
        (50, "Ride rate traseiro K_R = K_W·K_T/(K_W+K_T)", f"={S}$B$19*{S}$B$48/({S}$B$19+{S}$B$48)", "N/m", "SUS-212"),
        (51, "Rigidez de rolagem do par de pneus - dianteira (K_T·t_F²/2)", f"={S}$B$47*{V}$B$7^2/2", "N*m/rad", "SUS-206"),
        (52, "Rigidez de rolagem do par de pneus - traseira (K_T·t_R²/2)", f"={S}$B$48*{V}$B$8^2/2", "N*m/rad", "SUS-207"),
    ):
        a += [_a("06_SUSPENSION", f"A{r}", rot, "EDR-045", st(18, "A")),
              _a("06_SUSPENSION", f"B{r}", f, "EDR-045; Milliken p. 589-592; Gillespie eq. 5-8; Seward eq. 4.1/4.5", st(18, "B")),
              _a("06_SUSPENSION", f"C{r}", un, "EDR-045", st(18, "C")),
              _a("06_SUSPENSION", f"D{r}", sid, "EDR-045", st(18, "D"))]
    a += [
        _a("06_SUSPENSION", "A20", "Roll Stiffness sem ARB e sem pneu (K_W·t_F²/2) - dianteira", "EDR-045"),
        _a("06_SUSPENSION", "A21", "Roll Stiffness sem ARB e sem pneu (K_W·t_R²/2) - traseira", "EDR-045"),
        _a("06_SUSPENSION", "A22", "Roll Stiffness dianteira K_phiF = (K_W·t²/2 + K_ARB) em série com K_T·t²/2", "EDR-045"),
        _a("06_SUSPENSION", "B22", f"=1/(1/({S}$B$20+{S}$B$10)+1/{S}$B$51)", "EDR-045; Milliken p. 592 e p. 680"),
        _a("06_SUSPENSION", "A23", "Roll Stiffness traseira K_phiR = (K_W·t²/2 + K_ARB) em série com K_T·t²/2", "EDR-045"),
        _a("06_SUSPENSION", "B23", f"=1/(1/({S}$B$21+{S}$B$11)+1/{S}$B$52)", "EDR-045; Milliken p. 592 e p. 680"),
        # massa suspensa por canto (EDR-047)
        _a("06_SUSPENSION", "A27", "Massa suspensa por canto (SUS-009, por eixo: (m·b/L − m_uF)/2 e (m·a/L − m_uR)/2)", "EDR-047"),
        _a("06_SUSPENSION", "B28", f"={V}$B$25/2", "EDR-047; Milliken p. 601, Tab. 16.3"),
        _a("06_SUSPENSION", "B29", f"={V}$B$26/2", "EDR-047; Milliken p. 601, Tab. 16.3"),
        # frequência com ride rate (EDR-045)
        _a("06_SUSPENSION", "A31", "Frequência Natural (SUS-205) — com a ride rate (pneu em série)", "EDR-045"),
        _a("06_SUSPENSION", "B32", f"=(1/(2*PI()))*SQRT({S}$B$49/{S}$B$28)", "EDR-045; Gillespie eq. 5-9; Seward eq. 4.2"),
        _a("06_SUSPENSION", "B33", f"=(1/(2*PI()))*SQRT({S}$B$50/{S}$B$29)", "EDR-045; Gillespie eq. 5-9; Seward eq. 4.2"),
        # anti com sinal e ponto de reação do carro (EDR-048)
        _a("06_SUSPENSION", "A34", "Anti-Dive / Anti-Squat (SUS-301/302) — com sinal (EDR-026/048)", "EDR-048"),
        _a("06_SUSPENSION", "A35", "Milliken p. 618-619: freio no cubo (outboard) -> linha pelo contact patch; motor no chassi com semieixos -> "
                                   "linha pelo centro da roda (Fig. 17.15b). + = anti, − = pro (p. 620). IC de 07_KINEMATICS ainda é PLACEHOLDER.",
           "EDR-048"),
        _a("06_SUSPENSION", "A38", "Inclinação dz/dx da linha contact patch -> IC lateral dianteiro (com sinal; freio no cubo)", "EDR-048"),
        _a("06_SUSPENSION", "A39", "Inclinação dz/dx da linha centro da roda -> IC lateral traseiro (com sinal; motor no chassi)", "EDR-048"),
        _a("06_SUSPENSION", "B39", f"=({K}$B$12-{K}$B$14)/{K}$B$13", "EDR-048; Milliken Fig. 17.15(b); Gillespie eq. 7-19"),
        _a("06_SUSPENSION", "A40", "Anti-Dive (%) = FB_f·(−dz/dx)·L/h·100 (+ = anti, − = pro-dive)", "EDR-048"),
        _a("06_SUSPENSION", "B40", f"={S}$B$36*(-{S}$B$38)*({V}$B$6/{V}$B$5)*100", "EDR-048; Milliken p. 618"),
        _a("06_SUSPENSION", "A41", "Anti-Squat (%) = TF_r·(dz/dx)·L/h·100 (+ = anti, − = pro-squat)", "EDR-048"),
    ]
    # ---------------- 07_KINEMATICS: IC lateral com sinal e altura do centro da roda (EDR-048) ----------------
    a += [
        _a("07_KINEMATICS", "A11", "Δx do IC lateral dianteiro ao contact patch (x_IC − x_CP; ISO: + = à frente)", "EDR-048"),
        _a("07_KINEMATICS", "B11", -2.5, "EDR-048: mesmo PLACEHOLDER (2,5 m), agora com sinal: IC atrás do contact patch (Milliken p. 617)"),
        _a("07_KINEMATICS", "A13", "Δx do IC lateral traseiro ao centro da roda (x_IC − x_WC; ISO: + = à frente)", "EDR-048"),
        _a("07_KINEMATICS", "A14", "Altura do centro da roda traseira (ponto de reação do anti-squat)", "EDR-048", ("07_KINEMATICS", 12, "A")),
        _a("07_KINEMATICS", "B14", f"={V}$B$17", "EDR-048: PLACEHOLDER = SUS-011 (aproximado pela altura do centro da roda)", ("07_KINEMATICS", 12, "B"),
           "PLACEHOLDER: usa a altura do CG não suspenso traseiro (SUS-011), que a ficha aproxima pela altura do centro da roda. "
           "Substituir pela altura do centro da roda do CAD."),
        _a("07_KINEMATICS", "C14", "m", "EDR-048", ("07_KINEMATICS", 12, "C")),
        _a("07_KINEMATICS", "D14", "SUS-406-R", "EDR-048", ("07_KINEMATICS", 12, "D")),
    ]
    # ---------------- 09_LOADS: SUS-103 completa (EDR-046/047) ----------------
    h2 = f"({V}$B$24-{L}$B$18)"
    a += [
        _a("09_LOADS", "A17", "Transferência Lateral — Milliken, três massas, com termo gravitacional (SUS-103)", "EDR-046"),
        _a("09_LOADS", "A18", "Altura do roll axis no CG suspenso (z_s, em a_s)", "EDR-047"),
        _a("09_LOADS", "B18", f"={S}$B$14+({S}$B$15-{S}$B$14)*{V}$B$27/{V}$B$6", "EDR-047; Milliken p. 603"),
        _a("09_LOADS", "B19", f"=IF(ISNUMBER({L}$B$21),({S}$B$22*{L}$B$21+{E}$B$5*({V}$B$25*{S}$B$14+{V}$B$14*{V}$B$16))/{V}$B$7,\"N/A\")",
           "EDR-046/047; Milliken p. 604 e p. 682"),
        _a("09_LOADS", "B20", f"=IF(ISNUMBER({L}$B$21),({S}$B$23*{L}$B$21+{E}$B$5*({V}$B$26*{S}$B$15+{V}$B$15*{V}$B$17))/{V}$B$8,\"N/A\")",
           "EDR-046/047; Milliken p. 604 e p. 682"),
        _a("09_LOADS", "A21", "Ângulo de rolagem φ = m_s·a_y·h2/(K_φF+K_φR − m_s·g·h2) (+ = para fora da curva)", "EDR-046", ("09_LOADS", 19, "A")),
        _a("09_LOADS", "B21", (f"=IF({S}$B$22+{S}$B$23-{V}$B$23*9.81*{h2}<=0,\"N/A - instabilidade de rolagem (m_s·g·h2 >= K_φ, Milliken p. 682)\","
            f"{V}$B$23*{E}$B$5*{h2}/({S}$B$22+{S}$B$23-{V}$B$23*9.81*{h2}))"),
           "EDR-046; Milliken p. 682; Gillespie eq. 6-31", ("09_LOADS", 19, "B")),
        _a("09_LOADS", "C21", "rad", "EDR-046", ("09_LOADS", 19, "C")),
        _a("09_LOADS", "D21", "SUS-103", "EDR-046", ("09_LOADS", 19, "D")),
    ]
    # ---------------- 11_DESIGN: modo reverso com pneu (EDR-045) ----------------
    w = lambda c, m: f"(2*PI()*{D}$B${c})^2*{S}$B${m}"
    a += [
        _a("11_DESIGN", "A7", "Rigidez necessária (equação reversa, SUS-205/212/202; pneu em série)", "EDR-045"),
        _a("11_DESIGN", "A8", "Wheel Rate necessário - dianteira (K_R alvo -> K_W = K_R·K_T/(K_T − K_R))", "EDR-045"),
        _a("11_DESIGN", "B8", f"={w(4, 28)}*{S}$B$47/({S}$B$47-{w(4, 28)})", "EDR-045; Milliken p. 591; Seward eq. 4.4"),
        _a("11_DESIGN", "A9", "Wheel Rate necessário - traseira (K_R alvo -> K_W = K_R·K_T/(K_T − K_R))", "EDR-045"),
        _a("11_DESIGN", "B9", f"={w(5, 29)}*{S}$B$48/({S}$B$48-{w(5, 29)})", "EDR-045; Milliken p. 591; Seward eq. 4.4"),
        _a("11_DESIGN", "A20", "K_ARB dianteiro necessário (com o pneu em série)", "EDR-045"),
        _a("11_DESIGN", "B20", f"={D}$B$19*{S}$B$51/({S}$B$51-{D}$B$19)-{S}$B$20", "EDR-045; Milliken p. 592"),
    ]
    # ---------------- 13_VALIDATION: TC-005 com o termo gravitacional (EDR-046) ----------------
    a += [
        _a("13_VALIDATION", "A6", "TC-005: t_F*dFzF + t_R*dFzR - m*ay*h_CG - m_s*g*h2*φ [deve ser ~0]", "EDR-046"),
        _a("13_VALIDATION", "B6", f"={V}$B$7*{L}$B$19+{V}$B$8*{L}$B$20-{V}$B$4*{E}$B$5*{V}$B$5-{V}$B$23*9.81*{h2}*{L}$B$21",
           "EDR-046; Milliken p. 682-683"),
        _a("13_VALIDATION", "A26", "Piloto de referência: MR=0.8, Ks=100 N/mm, ms=50 kg, pneu rígido -> fn=8.897 Hz (verificado)", "EDR-045"),
    ]
    # ---------------- 10_TIRE: Ky com um único λFz0 (EDR-049) ----------------
    a.append(_a("10_TIRE", f"A{ky_row}", "Ky = PKY1·F'z0·sin(2·atan(Fz/(PKY2·F'z0)))·(1−PKY3·|γy|)·LKY  (um λFz0, EDR-049)", "EDR-049"))
    for c in "BCDEF":
        a.append(_a("10_TIRE", f"{c}{ky_row}", None, "EDR-049; Pacejka (2006) eq. 4.E25; Seward eq. 5.9"))
    return a + _rodada2() + _feedback2() + _eixo_iso() + _direcao()


# =====================================================================================================================
# 1ª rodada de feedback da V3.0 (feedback da equipe, 2026-10-14; decisões do usuário em 2026-10-19): EDR-035, EDR-050, EDR-051.
# =====================================================================================================================
# Rótulos da coluna "Origem" (EDR-035). O símbolo dá a leitura em P&B; a ordem é a da lista suspensa.
TX, T09, TTC, EST, ARB, CONF = ("● Dado do TX", "◐ Dado do T09", "◐ Provisório (.tir TTC)", "○ Estimativa",
                                "○ Arbitrário / demonstrativo", "? A confirmar")
ORIGENS_LISTA = (TX, T09, TTC, EST, ARB, CONF)
SIGNIFICADO = {   # rótulo -> (significado para quem preenche, categorias da auditoria V2.1 e tier)
    TX: ("Medido no carro atual ou tirado de documento do TX/E05 (ex.: Key Parameters).", "T3_FONTE_TX_REGISTRADA (T3)"),
    T09: ("Valor do carro anterior (T09/E03): serve para testar as fórmulas, não é o TX.", "REAL_OUTRO_VEICULO_T09/E03 (T2)"),
    TTC: ("Vem do arquivo de pneu provisório do TTC; não é o pneu do carro.", "REAL_OUTRO_VEICULO_PROVISORIO (T2)"),
    EST: ("Estimado ou suposto: inclui hipótese de modelo e caso de carga sem fonte.",
          "ESTIMATIVA_NAO_MEDIDA (T1), HIPOTESE_CONGELADA_V1 e CASO_DE_CARGA_SEM_PROVENIENCIA (T2); vale a pior, T1"),
    ARB: ("Placeholder, valor arbitrário ou meta de exemplo: não usar para decisão de projeto.",
          "ARBITRARIO, PLACEHOLDER, ALVO_DEMONSTRATIVO (T1)"),
    CONF: ("A origem do valor não está registrada.", "sem classificação"),
}
_T09 = ("Demonstrativo (EDR-012/022): valor do T09/E03 para referência, comparação e teste das fórmulas "
        "(usuário, 2026-10-19); não é o TX/E05. ")
_KIN = "Placeholder geométrico (07_KINEMATICS inteira é placeholder até a Fase 2); não é hardpoint real."
_CARGA = "Caso de carga sem fonte registrada. "
# Entradas das abas 05/06/07/08/11: (aba, linha, origem; None = célula calculada, sem lista), fonte (coluna F).
# Fonte: comentários das células na V2.4/V3.0 (migrados para a coluna F) e mapa da auditoria V2.1
# (data/placeholder-dependency-map-V2.1.csv). O RC (06 B14/B15) é T09/E03 por confirmação do usuário (2026-10-19).
ORIGENS = [
    ("05_VEHICLE", 4, None, "Calculada (EDR-050): massa medida sem piloto (B30) + massa do piloto (B31)."),
    ("05_VEHICLE", 5, TX, "E05, Key Parameters (315 mm). Medida a partir do plano do solo (SUS-000), não da origem do CAD."),
    ("05_VEHICLE", 6, TX, "E05, Key Parameters (1594 mm). Diverge do E03 (1550 mm, chassi anterior)."),
    ("05_VEHICLE", 7, TX, "E05, Key Parameters (1277 mm)."),
    ("05_VEHICLE", 8, TX, "E05, Key Parameters (1271 mm)."),
    ("05_VEHICLE", 9, None, "Calculada (EDR-050): L·(1 − f), com f = fração de peso no eixo dianteiro (B32)."),
    ("05_VEHICLE", 10, None, "Calculada (EDR-050): L·f, com f = fração de peso no eixo dianteiro (B32)."),
    ("05_VEHICLE", 14, ARB, "Não disponível para o TX/E05: a pesagem em plataforma mede a carga total no solo e não separa massa "
                            "suspensa de não suspensa (confirmado pela equipe). Valor arbitrário: não usar para decisão de projeto."),
    ("05_VEHICLE", 15, ARB, "Não disponível para o TX/E05 (mesma razão do SUS-007). Valor arbitrário: não usar para decisão de projeto."),
    ("05_VEHICLE", 16, None, "Calculada (EDR-050): diâmetro do pneu (B33) × 0,0254 / 2. Estimativa: o raio do pneu aproxima a altura "
                             "do CG não suspenso; não é medição da equipe."),
    ("05_VEHICLE", 17, None, "Calculada (EDR-050): diâmetro do pneu (B34) × 0,0254 / 2. Estimativa: o raio do pneu aproxima a altura "
                             "do CG não suspenso; não é medição da equipe."),
    ("05_VEHICLE", 20, CONF, "SUS-000: coordenada Z do solo no sistema de coordenadas da equipe. As alturas desta planilha já são "
                             "medidas a partir do solo (ficha SUS-000 do Engineering Standard). A origem do valor −0,05 m não está registrada."),
    ("05_VEHICLE", 30, TX, "Adesivo SAE Brasil 2026: medição real, sem piloto (248,5 kg)."),
    ("05_VEHICLE", 31, EST, "Piloto de referência (68 kg, padrão FSAE); não é a massa medida do piloto."),
    ("05_VEHICLE", 32, TX, "Distribuição medida no adesivo do TX/E05: 48,43 % no eixo dianteiro, com piloto (coerente com a massa total, "
                           "que inclui o piloto)."),
    ("05_VEHICLE", 33, T09, "Pneu 20.0-13 do E03 (diâmetro 20,0 in)."),
    ("05_VEHICLE", 34, T09, "Pneu 20.0-13 do E03 (mesmo pneu na frente e atrás)."),
    ("06_SUSPENSION", 4, T09, _T09 + "Calculado a partir da planilha detalhada do T09: ride rate dianteiro 16 112,23 N/m (que é o wheel "
                                     "rate) × MR² (MR = 1,76). Corrige o valor anterior, do spec sheet resumido, que arredondava F e R como iguais."),
    ("06_SUSPENSION", 5, T09, _T09 + "Calculado a partir da planilha detalhada do T09: ride rate traseiro 18 418,2 N/m × MR² (MR = 1,76)."),
    ("06_SUSPENSION", 6, T09, _T09 + "E03 (1,76:1). Convenção wheel/spring assumida (igual à do projeto); o E03 não diz a direção."),
    ("06_SUSPENSION", 7, T09, _T09 + "E03 (1,76:1, igual à dianteira)."),
    ("06_SUSPENSION", 10, T09, _T09 + "Planilha detalhada do T09 (AntiRoll Bar Rate dianteiro), valor direto."),
    ("06_SUSPENSION", 11, T09, _T09 + "Planilha detalhada do T09 (AntiRoll Bar Rate traseiro, ≈ 0, coerente com 'sem ARB')."),
    ("06_SUSPENSION", 14, T09, _T09 + "Planilha detalhada do T09 (0,0675 m), mais precisa que o spec sheet do E03 (60 mm, arredondado). "
                                      "Provisório até a Fase 2 (OptimumKinematics)."),
    ("06_SUSPENSION", 15, T09, _T09 + "E03 (110 mm, roll center estático acima do solo). Provisório até a Fase 2 (OptimumKinematics)."),
    ("06_SUSPENSION", 36, ARB, "Placeholder: fração da força de frenagem no eixo dianteiro (brake bias). Substituir pelo balanço de freio "
                               "real do Setup."),
    ("06_SUSPENSION", 37, EST, "Hipótese congelada da V1: TF_r = 1 (tração só traseira). Revisitar se a equipe mudar para AWD."),
    ("06_SUSPENSION", 44, ARB, "Placeholder: razão entre o ângulo do volante e o ângulo médio das rodas dianteiras. Substituir pela "
                               "medição da cremalheira ou pela geometria do rack."),
    ("06_SUSPENSION", 47, TTC, "VERTICAL_STIFFNESS do .tir TTC provisório (Hoosier R25B; EDR-038, EDR-045). Não é o pneu do carro nem "
                               "medição da equipe: substituir pela rigidez vertical do pneu real."),
    ("06_SUSPENSION", 48, TTC, "VERTICAL_STIFFNESS do .tir TTC provisório (Hoosier R25B; EDR-038, EDR-045). Não é o pneu do carro nem "
                               "medição da equipe: substituir pela rigidez vertical do pneu real."),
    ("07_KINEMATICS", 10, ARB, "Placeholder geométrico para testar a fórmula; não é hardpoint real. Fase 2: REFERENCE do OptimumKinematics."),
    ("07_KINEMATICS", 11, ARB, _KIN), ("07_KINEMATICS", 12, ARB, _KIN), ("07_KINEMATICS", 13, ARB, _KIN),
    ("07_KINEMATICS", 14, None, "Calculada: = SUS-011 (05_VEHICLE B17). Placeholder: substituir pela altura do centro da roda do CAD."),
    ("07_KINEMATICS", 17, ARB, _KIN), ("07_KINEMATICS", 18, ARB, _KIN), ("07_KINEMATICS", 19, ARB, _KIN),
    ("07_KINEMATICS", 20, ARB, _KIN), ("07_KINEMATICS", 23, ARB, _KIN), ("07_KINEMATICS", 24, ARB, _KIN),
    ("07_KINEMATICS", 27, ARB, _KIN), ("07_KINEMATICS", 28, ARB, _KIN),
    ("07_KINEMATICS", 31, ARB, "Placeholder geométrico; não é a medição do curso instalado. Substituir pelo curso real (bump + rebound) "
                               "com o piloto sentado, conforme V.3.1.1."),
    ("07_KINEMATICS", 32, ARB, "Placeholder geométrico; não é a medição do curso instalado (ver B31)."),
    ("08_VEHICLE_STATE", 4, EST, _CARGA + "ISO 8855: a_x + = aceleração para frente (transfere carga para o eixo traseiro); "
                                          "negativo = frenagem."),
    ("08_VEHICLE_STATE", 5, None, "Calculada (EDR-050): a_y em g (B11) × 9,81. ISO 8855: a_y + = para a esquerda (o lado direito ganha carga)."),
    ("08_VEHICLE_STATE", 8, EST, _CARGA + "Aero constante por Setup (SUS-701, EDR-013)."),
    ("08_VEHICLE_STATE", 9, EST, _CARGA + "Fração dianteira da downforce (SUS-702, EDR-013)."),
    ("08_VEHICLE_STATE", 11, EST, _CARGA + "Curva à esquerda de 1,2 g."),
    ("11_DESIGN", 4, ARB, "Alvo demonstrativo: a meta da equipe está TBD em 03_TARGETS (TARGET-001)."),
    ("11_DESIGN", 5, ARB, "Alvo demonstrativo: a meta da equipe está TBD em 03_TARGETS (TARGET-002)."),
    ("11_DESIGN", 18, ARB, "Alvo demonstrativo: a meta da equipe está TBD em 03_TARGETS (TARGET-003)."),
]
# 10_TIRE: as colunas E-H estão ocupadas (cantos, unidade, ID); a origem vai em I e a fonte em J, por linha de caso e
# para o bloco de coeficientes (linha 84, logo acima do cabeçalho mesclado A85:I85). Sem o aviso de desatualização (cinco valores por linha).
TIRE_ORIGENS = [
    (23, ARB, "Demonstrativos (EDR-022): α, γ e κ não vêm de cálculo do carro (não há modelo de slip). Valores arbitrários coerentes "
              "em sinal com 08_VEHICLE_STATE (curva à esquerda, frenagem). Substituir por casos de carga da equipe. VERIF (F): caso fixo de teste."),
    (24, ARB, "Camber de entrada manual (EDR-036 sem decisão: ainda não há ID para camber estático). Convenção do projeto: "
              "+ = topo para fora, nos dois lados."),
    (25, ARB, "Demonstrativo (EDR-022), como α e γ."),
    (84, TTC, "Todos os coeficientes da tabela abaixo (B86:B175) foram transcritos do .tir provisório (ver os metadados acima). REFERENCE "
              "por natureza (DR-011), mas o dataset é provisório. Mudou o dataset: mudar o Tire REV (01_SETUP) e reexecutar "
              "13_VALIDATION; os resultados antigos viram STALE."),
]
ABAS_ORIGEM = ("05_VEHICLE", "06_SUSPENSION", "07_KINEMATICS", "08_VEHICLE_STATE", "11_DESIGN")
LINHA_COMO_PREENCHER = 39          # 00_README: início da seção "Como preencher"
LINHAS_ORIGENS = range(43, 49)     # 00_README: A = rótulos (nome definido "Origens"), B:D = significado, E = auditoria V2.1
NOME_LISTA = "Origens"
G_MS2 = "9.81"                     # mesmo g literal das fórmulas existentes (09_LOADS, 13_VALIDATION)


def _rodada2() -> list[dict]:
    a = []
    ref35, ref50 = "EDR-035 (feedback da equipe, 2026-10-14)", "EDR-050 (feedback da equipe, 2026-10-14)"
    inp = lambda aba, r: (aba, r, "B")                     # estilo de uma entrada existente
    # ---- FB-05 / EDR-050: parcelas externas viram entrada; a célula vira a conta ----
    a.append(_a("05_VEHICLE", "A29", "Entradas de onde saem B4, B9/B10 e B16/B17 (EDR-050: valor que sai de conta não é digitado)",
                ref50, ("05_VEHICLE", 22, "A")))
    for r, rot, val, un, sid in (
        (30, "Massa medida sem piloto", 248.5, "kg", "SUS-001"),
        (31, "Massa do piloto", 68, "kg", "SUS-001"),
        (32, "Fração de peso no eixo dianteiro (f)", 0.4843, "-", "SUS-006"),
        (33, "Diâmetro do pneu dianteiro", 20.0, "in", "SUS-010"),
        (34, "Diâmetro do pneu traseiro", 20.0, "in", "SUS-011"),
    ):
        a += [_a("05_VEHICLE", f"A{r}", rot, ref50, ("05_VEHICLE", 14, "A")),
              _a("05_VEHICLE", f"B{r}", val, ref50 + "; valor: comentário da V2.4 (ver coluna F)", inp("05_VEHICLE", 14)),
              _a("05_VEHICLE", f"C{r}", un, ref50, ("05_VEHICLE", 14, "C")),
              _a("05_VEHICLE", f"D{r}", sid, ref50, ("05_VEHICLE", 14, "D"))]
    calc = ("05_VEHICLE", 23, "B")
    a += [
        _a("05_VEHICLE", "B4", f"={V}$B$30+{V}$B$31", ref50, calc),
        _a("05_VEHICLE", "B9", f"={V}$B$6*(1-{V}$B$32)", ref50, calc),
        _a("05_VEHICLE", "B10", f"={V}$B$6*{V}$B$32", ref50, calc),
        _a("05_VEHICLE", "B16", f"={V}$B$33*0.0254/2", ref50, calc),
        _a("05_VEHICLE", "B17", f"={V}$B$34*0.0254/2", ref50, calc),
        _a("08_VEHICLE_STATE", "A10", "Entrada de onde sai a_y (B5) — EDR-050", ref50, ("08_VEHICLE_STATE", 3, "A")),
        _a("08_VEHICLE_STATE", "A11", "Aceleração lateral em g (entrada de a_y)", ref50, ("08_VEHICLE_STATE", 5, "A")),
        _a("08_VEHICLE_STATE", "B11", 1.2, ref50 + "; valor: 11,772 m/s² / 9,81", inp("08_VEHICLE_STATE", 5)),
        _a("08_VEHICLE_STATE", "C11", "g", ref50, ("08_VEHICLE_STATE", 5, "C")),
        _a("08_VEHICLE_STATE", "D11", "-", ref50, ("08_VEHICLE_STATE", 5, "D")),
        _a("08_VEHICLE_STATE", "B5", f"={E}$B$11*{G_MS2}", ref50, ("05_VEHICLE", 23, "B")),
    ]
    # ---- FB-03 / EDR-035: origem (E), fonte (F) e, ocultos, valor e origem da geração (G, H) ----
    for aba in ABAS_ORIGEM:
        a += [_a(aba, "E2", "Origem do valor (lista)", ref35, (aba, 3, "A")),
              _a(aba, "F2", "Fonte / observação", ref35, (aba, 3, "A"))]
    for aba, r, org, fonte in ORIGENS:
        st = (aba, r, "C")
        if org is not None:
            a.append(_a(aba, f"E{r}", org, ref35, st))
        a.append(_a(aba, f"F{r}", fonte, ref35, st))
        if org is not None:
            a += [_a(aba, f"G{r}", ("COPIA", f"B{r}"), ref35 + ": valor da geração (oculto; aviso de desatualização)", st),
                  _a(aba, f"H{r}", org, ref35 + ": origem da geração (oculto; aviso de desatualização)", st)]
    a += [_a("10_TIRE", "I20", "Origem do valor (lista)", ref35, ("10_TIRE", 20, "A")),
          _a("10_TIRE", "J20", "Fonte / observação", ref35, ("10_TIRE", 20, "A"))]
    for r, org, fonte in TIRE_ORIGENS:
        a += [_a("10_TIRE", f"I{r}", org, ref35, ("10_TIRE", 86, "H")),
              _a("10_TIRE", f"J{r}", fonte, ref35, ("10_TIRE", 86, "H"))]
    # ---- FB-02: legenda com o estilo real das células; a classe "amarelo" sai (o fundo amarelo é removido em aplicar) ----
    a += [
        _a("00_README", "A16", "■ Azul", "FB-02; estilo = 05_VEHICLE B14", ("05_VEHICLE", 14, "B")),
        _a("00_README", "A17", "■ Verde-teal", "FB-02; estilo = 05_VEHICLE B6", ("05_VEHICLE", 6, "B")),
        _a("00_README", "A18", "■ Laranja + fundo", "FB-02; estilo = 06_SUSPENSION B14", ("06_SUSPENSION", 14, "B")),
        _a("00_README", "A19", "■ Preto + fundo cinza", "FB-02; estilo = 05_VEHICLE B23", ("05_VEHICLE", 23, "B")),
        _a("00_README", "A20", "Origem do valor", "FB-02/EDR-035", ("00_README", 19, "B")),
        _a("00_README", "B20", "Não é cor da célula: fica na coluna E de cada entrada (● ◐ ○ ?). Ver \"Como preencher\", abaixo.",
           "FB-02/EDR-035"),
    ]
    # ---- FB-06: guia de preenchimento na própria planilha ----
    h, t = ("00_README", 22, "A"), ("00_README", 23, "A")
    r0 = LINHA_COMO_PREENCHER
    guia = [
        (r0, "Como preencher esta planilha (EDR-035, EDR-050, EDR-051)", h),
        (r0 + 1, "1. Edite só as entradas (fonte azul; verde-teal e laranja quando houver o dado real) das abas 05, 06, 07, 08, "
                 "10 (casos de carga) e 11 (metas). Célula com fundo cinza é calculada: não edite.", t),
        (r0 + 2, "2. Cada entrada tem ao lado a Origem (coluna E; na 10_TIRE, coluna I) e a Fonte (coluna F; na 10_TIRE, J). "
                 "Ao trocar um valor, escolha a Origem na lista e escreva na Fonte de onde veio o número. Origens possíveis:", t),
    ]
    a += [_a("00_README", f"A{r}", txt, "FB-06/EDR-035", st) for r, txt, st in guia]
    a += [_a("00_README", f"A{r0 + 3}", "Origem", "FB-06/EDR-035", t),
          _a("00_README", f"B{r0 + 3}", "Significado", "FB-06/EDR-035", t),
          _a("00_README", f"E{r0 + 3}", "Categorias da auditoria V2.1 (data/placeholder-dependency-map-V2.1.csv)", "FB-06/EDR-035", t)]
    for r, org in zip(LINHAS_ORIGENS, ORIGENS_LISTA):
        sig, aud = SIGNIFICADO[org]
        a += [_a("00_README", f"A{r}", org, "FB-06/EDR-035; nome definido 'Origens'", t),
              _a("00_README", f"B{r}", sig, "FB-06/EDR-035", t),
              _a("00_README", f"E{r}", aud, "FB-06/EDR-035", t)]
    fim = LINHAS_ORIGENS.stop
    resto = [
        "3. Avisos: Origem com fundo vermelho-alaranjado = o valor mudou e a Origem continuou a mesma (revise a Origem). "
        "Fonte com fundo vermelho = a Origem diz \"Dado do TX\" e falta a Fonte. Na 10_TIRE não há esses avisos: ao trocar um "
        "caso de carga, troque a Origem à mão.",
        "4. Valor que sai de conta não é digitado (EDR-050): digite as parcelas, não o resultado. Exemplo: massa total = massa "
        "medida + piloto (05_VEHICLE B30 e B31).",
        "5. Até a V3.0 ser fechada, esta planilha é gerada pelo script do repositório e o que for digitado nela se perde na próxima "
        "geração (EDR-051). Depois do fechamento, ela passa a ser o arquivo da equipe.",
        "6. Convenções de sinais e eixos: seção acima. Decisões e fichas: pasta docs/ do repositório "
        "(Engineering_Decision_Register.md e Engineering Standard).",
    ]
    a += [_a("00_README", f"A{fim + i}", txt, "FB-06/EDR-035/050/051", t) for i, txt in enumerate(resto)]
    return a


# =====================================================================================================================
# 2ª rodada da revisão metodológica da V3.0 (decisões do usuário em 2026-10-19): EDR-050 estendido.
# =====================================================================================================================
# A hipótese "CG não suspenso na altura do centro da roda" vira entrada (razão h_ns / raio), e a altura do centro da roda
# traseira (07 B14, ponto de reação do anti-squat) deixa de passar pelo CG não suspenso: são grandezas diferentes que a V3.0
# fazia coincidir na mesma célula (05 B17).
_RAZAO_HNS = ("Estimativa: Milliken p. 672 (cap. 18, altura do CG suspenso) toma a altura do centro da roda como \"boa primeira "
              "aproximação\" da altura do CG não suspenso e usa o raio carregado (R_LF, R_LR); razão 1 = essa aproximação. Aqui o "
              "raio é o sem carga (diâmetro / 2), sem a deflexão estática do pneu. Trocar pela medição do conjunto não suspenso quando houver.")
ORIGENS_FB2 = [   # mesmo formato de ORIGENS; processadas junto com elas na formatação (lista, cores, avisos)
    ("05_VEHICLE", 35, EST, _RAZAO_HNS),
    ("05_VEHICLE", 36, EST, _RAZAO_HNS),
    ("05_VEHICLE", 16, None, "Calculada (EDR-050): raio do pneu (B33 × 0,0254 / 2) × razão h_ns / raio (B35). Estimativa: a hipótese "
                             "está na razão (B35), não é medição da equipe."),
    ("05_VEHICLE", 17, None, "Calculada (EDR-050): raio do pneu (B34 × 0,0254 / 2) × razão h_ns / raio (B36). Estimativa: a hipótese "
                             "está na razão (B36), não é medição da equipe."),
    ("07_KINEMATICS", 14, None, "Calculada: raio sem carga do pneu traseiro (05_VEHICLE B34 × 0,0254 / 2), sem a deflexão estática. "
                                "Placeholder: substituir pela altura do centro da roda do CAD."),
]


def _todas_origens() -> list[tuple]:
    return ORIGENS + ORIGENS_FB2 + ORIGENS_DIR


def _feedback2() -> list[dict]:
    a = []
    ref = "EDR-050 (2ª rodada da revisão metodológica, usuário 2026-10-19)"
    for r, rot, sid in ((35, "Razão h_ns / raio do pneu dianteiro", "SUS-010"), (36, "Razão h_ns / raio do pneu traseiro", "SUS-011")):
        a += [_a("05_VEHICLE", f"A{r}", rot, ref, ("05_VEHICLE", 14, "A")),
              _a("05_VEHICLE", f"B{r}", 1.0, ref + "; Milliken p. 672", ("05_VEHICLE", 14, "B")),
              _a("05_VEHICLE", f"C{r}", "-", ref, ("05_VEHICLE", 14, "C")),
              _a("05_VEHICLE", f"D{r}", sid, ref, ("05_VEHICLE", 14, "D"))]
    calc = ("05_VEHICLE", 23, "B")
    a += [_a("05_VEHICLE", "B16", f"={V}$B$33*0.0254/2*{V}$B$35", ref, calc),
          _a("05_VEHICLE", "B17", f"={V}$B$34*0.0254/2*{V}$B$36", ref, calc),
          _a("07_KINEMATICS", "B14", f"={V}$B$34*0.0254/2", ref + ": centro da roda separado do CG não suspenso (EDR-048)",
             ("07_KINEMATICS", 12, "B"))]
    for aba, r, org, fonte in ORIGENS_FB2:
        st = (aba, r, "C")
        if org is not None:
            a += [_a(aba, f"E{r}", org, "EDR-035", st),
                  _a(aba, f"G{r}", ("COPIA", f"B{r}"), "EDR-035: valor da geração (oculto; aviso de desatualização)", st),
                  _a(aba, f"H{r}", org, "EDR-035: origem da geração (oculto; aviso de desatualização)", st)]
        a.append(_a(aba, f"F{r}", fonte, ref, st))
    # Item 7 do guia: significado de N/A e TBD (EDR-035, parte de N/A e TBD aprovada pelo usuário em 2026-10-19).
    a.append(_a("00_README", f"A{LINHA_NA_TBD}",
                "7. N/A = não se aplica, ou não dá para calcular na configuração atual (falta dado, valor fora da faixa do pneu, "
                "instabilidade): a célula nunca mostra número inventado. TBD = falta um dado, uma meta ou uma decisão (equivale a "
                "OPEN). \"? A confirmar\" (coluna Origem) é outra coisa: há valor, mas a origem dele não está registrada.",
                "EDR-035 (N/A e TBD; usuário, 2026-10-19)", ("00_README", 23, "A")))
    # Kinematics REV ainda não definida = pendência (TBD), como a Geometry REV em 07 B68 (auditoria dos usos; usuário, 2026-10-19).
    a.append(_a("01_SETUP", "B10", "TBD (Fase 2)", "EDR-035 (N/A e TBD; troca aprovada pelo usuário, 2026-10-19)"))
    return a


LINHA_NA_TBD = LINHAS_ORIGENS.stop + 4   # 00_README: logo depois do item 6 do guia

# Direção (cap. 19 da Milliken; Gillespie cap. 6 e 8; Seward 6.3): EDR-052, fichas SUS-414 e SUS-415 (usuário, 2026-10-19:
# "crie as fichas, implemente os cálculos e metodologias de acordo com a literatura"). Bloco novo na 07_KINEMATICS.
_REF52 = "EDR-052 (direção, cap. 19 da Milliken; usuário 2026-10-19)"
DIR_CABECALHOS = {71: "Direção — Ackermann (SUS-414) — entradas PLACEHOLDER até haver a medição ou a geometria do TX",
                  77: "Direção — relação de direção pela cremalheira (SUS-601, Milliken p. 718)",
                  81: "Direção — roll steer e gradiente de subesterço (SUS-415, Gillespie eq. 6-57)"}
_ACK = f"{K}$B$72"
DIR_LINHAS = [   # (linha, rótulo, valor ou fórmula, unidade, ID, origem ou None, fonte)
    (72, "Ângulo da roda dianteira externa (δo)", 20.0, "°", "SUS-414", EST,
     "Estimativa: carros de FSAE precisam de pouco mais de 20° de esterço na roda externa (Seward p. 162-163, Tab. 6.2). "
     "Não é o batente do TX."),
    (73, "Ângulo da roda dianteira interna (δi) no mesmo esterço", 20.0, "°", "SUS-414", ARB,
     "Placeholder: igual a δo (direção paralela, 0 % Ackermann) até haver a medição do TX ou a geometria (Python: "
     "kinematics_model.ackermann_from_kinematics)."),
    (74, "δi para Ackermann exato", f"=DEGREES(ATAN(1/(1/TAN(RADIANS({_ACK}))-{V}$B$7/{V}$B$6)))", "°", "SUS-414", None,
     "Calculada: cot δo − cot δi = t/L (Gillespie eq. 8-1 e 8-2, p. 277-278), com t = 05 B7 e L = 05 B6. Reproduz a Tab. 6.2 do Seward."),
    (75, "% Ackermann", f"=({K}$B$73-{_ACK})/({K}$B$74-{_ACK})*100", "%", "SUS-414", None,
     "Calculada: 100 % = Ackermann exato; 0 % = paralelo; < 0 = reverso (Milliken p. 714-716; Seward, glossário). A razão "
     "entre os dois extremos é a forma adotada pelo projeto (EDR-052)."),
    (76, "Raio da trajetória da roda dianteira externa", f"={V}$B$6/SIN(RADIANS({_ACK}))", "m", "SUS-414", None,
     "Calculada: L / sin δo (Seward, Tab. 6.2)."),
    (78, "Curso da cremalheira por volta do pinhão (c-factor)", 0.0508, "m/volta", "SUS-601", ARB,
     "Placeholder: o exemplo de \"2-inch rack\" do Milliken p. 718; não é a cremalheira do TX."),
    (79, "Comprimento do braço de direção (pino mestre à barra)", 0.10, "m", "SUS-601", ARB,
     "Placeholder: valor do caso de teste TC-KIN-109 (e = 0,10 m); não é o braço do TX."),
    (80, "Relação de direção estimada pela cremalheira", f"=360/DEGREES(ASIN({K}$B$78/{K}$B$79))", "-", "SUS-601", None,
     "Calculada: SR ≈ 360 / asin(c/braço), Milliken p. 718 (o livro imprime asin(c/braço)/360, que é o inverso da razão "
     "definida na p. 716; EDR-052). Vale com a barra quase perpendicular ao braço. Compare com 06_SUSPENSION B44."),
    (82, "Roll steer dianteiro (ε_f)", 0.0, "°/°", "SUS-415", ARB,
     "Placeholder: sem hardpoints reais. ε > 0 esterça para a direita numa rolagem para a direita (Gillespie p. 220). "
     "Python: kinematics_model.roll_steer_coefficient (na geometria sintética, −0,125 °/°)."),
    (83, "Roll steer traseiro (ε_r)", 0.0, "°/°", "SUS-415", ARB, "Placeholder: sem hardpoints reais (ver B82)."),
    (84, "Gradiente de rolagem ∂φ/∂a_y", f"=IF(AND(ISNUMBER({L}$B$21),{E}$B$11<>0),DEGREES({L}$B$21)/{E}$B$11,\"N/A\")",
     "°/g", "SUS-415", None,
     "Calculada: φ (09_LOADS B21) / a_y em g (08 B11). O φ do modelo é linear em a_y (SUS-103), então a razão é o gradiente."),
    (85, "Contribuição do roll steer ao gradiente de subesterço", f"=IF(ISNUMBER({K}$B$84),({K}$B$82-{K}$B$83)*{K}$B$84,\"N/A\")",
     "°/g", "SUS-415", None, "Calculada: K = (ε_f − ε_r)·∂φ/∂a_y (Gillespie eq. 6-57); K > 0 = subesterço."),
]
ORIGENS_DIR = [("07_KINEMATICS", r, org, fonte) for r, _, _, _, _, org, fonte in DIR_LINHAS]


def _direcao() -> list[dict]:
    a = []
    for r, txt in DIR_CABECALHOS.items():
        a.append(_a("07_KINEMATICS", f"A{r}", txt, _REF52, ("07_KINEMATICS", 30, "A")))
    for r, rot, val, un, sid, org, fonte in DIR_LINHAS:
        st_b = ("07_KINEMATICS", 10, "B") if org is not None else ("05_VEHICLE", 23, "B")
        a += [_a("07_KINEMATICS", f"A{r}", rot, _REF52, ("07_KINEMATICS", 10, "A")),
              _a("07_KINEMATICS", f"B{r}", val, _REF52, st_b),
              _a("07_KINEMATICS", f"C{r}", un, _REF52, ("07_KINEMATICS", 10, "C")),
              _a("07_KINEMATICS", f"D{r}", sid, _REF52, ("07_KINEMATICS", 10, "D")),
              _a("07_KINEMATICS", f"F{r}", fonte, _REF52, ("07_KINEMATICS", 10, "C"))]
        if org is not None:
            st = ("07_KINEMATICS", 10, "C")
            a += [_a("07_KINEMATICS", f"E{r}", org, "EDR-035", st),
                  _a("07_KINEMATICS", f"G{r}", ("COPIA", f"B{r}"), "EDR-035: valor da geração (oculto)", st),
                  _a("07_KINEMATICS", f"H{r}", org, "EDR-035: origem da geração (oculto)", st)]
    # IC lateral (entradas do anti-dive/anti-squat): a definição que gerou o IC tem de ficar dita na Fonte (EDR-053).
    contrato = (" Ao preencher, diga na Fonte de que definição veio o IC: linha de força pelo trabalho virtual (Python, "
                "kinematics_model.force_lines_vw; EDR-053) ou tela do OptimumKinematics (definição a confirmar). Na geometria "
                "sintética as duas dão 38,6 % e cerca de 93 % de anti-dive.")
    fontes = {10: "Placeholder geométrico para testar a fórmula; não é hardpoint real. Fase 2: REFERENCE do OptimumKinematics.",
              11: _KIN, 12: _KIN, 13: _KIN}
    a += [_a("07_KINEMATICS", f"F{r}", txt + contrato, "EDR-053 (conselho de 2026-10-19; usuário seguiu a recomendação)")
          for r, txt in fontes.items()]
    return a


# EDR-039 reaberto (usuário, 2026-10-19): o .tir é PAC2002 do Adams, cujos parâmetros "must always be valid for the TYDEX
# W-axis system (ISO oriented)" (Adams 2024.1, Using the PAC2002 Tire Model, p. 7). Sistema de eixos declarado = ISO_TYDEX.
EIXO_TIR = "ISO_TYDEX"
TIR_PROVISORIO = AQUI.parent / "data" / "Hoosier_R25B_13x7_TTC_provisional.tir"
_REF39 = "EDR-039 reaberto (usuário, 2026-10-19): .tir PAC2002 = TYDEX W (ISO), help do Adams p. 7"


def _eixo_iso() -> list[dict]:
    """10_TIRE declara ISO_TYDEX; referências do Python do caso VERIF recalculadas no mesmo eixo; textos que diziam SAE."""
    import math
    import tire_model as tm
    ds = tm.load_dataset(TIR_PROVISORIO, source_axis_system=EIXO_TIR)
    fx, fy, _ = tm.wheel_forces(ds, "FR", 700.0, math.radians(3), math.radians(1.5), 0.05)   # caso VERIF (build_workbook)
    mu = tm.mu_effective(fx, fy, 700.0)
    return [
        _a("10_TIRE", "B11", EIXO_TIR, _REF39, comentario=(
            "Palavra-chave lida pelas fórmulas. ISO_TYDEX: identidade (o .tir PAC2002 vale no eixo W do TYDEX, orientação ISO: "
            "help do Adams 2024.1, PAC2002, p. 7). SAE_CALSPAN: α_arquivo = −α_ISO; Fy_ISO = −Fy_arquivo; γ_arquivo = −γ_projeto, "
            "para dado bruto em SAE. Qualquer outro texto => dataset INVÁLIDO. Ver tire_model.AXIS_TRANSFORMS e EDR-039.")),
        _a("10_TIRE", "B10", "CONFIRMED — .tir PAC2002 em TYDEX W (ISO), help do Adams p. 7", _REF39, comentario=(
            "V3.0 (EDR-039 reaberto, 2026-10-19): o arquivo declara TIRE_VERSION PAC2002 2018 e compatibilidade com Adams Tire; o "
            "help do Adams diz que os parâmetros do arquivo valem sempre no eixo W do TYDEX (orientação ISO). Confere com a "
            "avaliação do arquivo: Fy < 0 para α > 0 e Fy < 0 para γ > 0 (Pacejka, App. 1, p. 612: só no ISO o camber positivo "
            "desloca Fy para baixo). Até a V2.4 o arquivo era declarado SAE_CALSPAN, com o sinal do camber inferido.")),
        _a("13_VALIDATION", "B37", fx, _REF39 + "; tire_model.wheel_forces(ISO_TYDEX, FR, 700 N, 3°, 1,5°, 0,05)"),
        _a("13_VALIDATION", "B38", fy, _REF39 + "; idem"),
        _a("13_VALIDATION", "B39", mu, _REF39 + "; idem"),
        _a("14_DASHBOARD", "B38", "ISO (TYDEX W, help do Adams PAC2002); EDR-039", _REF39),
    ]
# FB-03, fase 2 (aprovada pelo usuário em 2026-10-19): pior origem ao vivo nos resultados principais do 14_DASHBOARD.
DASH_RESULTADOS = (7, 8, 9, 10, 13, 14, 15, 16, 17, 20, 21, 22, 23, 24, 33, 34, 35, 36, 37)
DASH_CABECALHOS = (5,)             # linha vazia acima da 1ª seção (as linhas de seção são mescladas A:E)


def _formula_pior_origem(celulas: list[tuple[str, str]]) -> str:
    """Maior índice na lista 'Origens' entre as células de origem; origem vazia ou fora da lista conta como '? A confirmar'.

    A lista 'Origens' (00_README) está na ordem inversa da pior origem (raizes.PIOR_PRIMEIRO), então pior = maior índice.
    """
    n = len(ORIGENS_LISTA)
    termos = [f"IFERROR(MATCH('{aba}'!${re.sub(r'[0-9]', '', c)}${re.sub(r'[A-Z]', '', c)},{NOME_LISTA},0),{n})"
              for aba, c in celulas]
    return f"=INDEX({NOME_LISTA},MAX({','.join(termos)}))"


def _feedback2_pior_origem(wb, registro: list[dict]) -> None:
    """Coluna E do 14_DASHBOARD: pior origem das entradas de cada resultado, calculada ao vivo pela coluna Origem.

    As células de origem de cada resultado saem do grafo de fórmulas na geração (python/raizes.py); a fórmula lê os
    rótulos ao vivo, mas a lista de células é a da geração (TC-WB30-014 confere que ela cobre todas as raízes).
    """
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import PatternFill
    from raizes import PIOR_PRIMEIRO, Grafo
    assert tuple(reversed(PIOR_PRIMEIRO)) == ORIGENS_LISTA, "ordem da lista 'Origens' deixou de ser a inversa da pior origem"
    ref = "FB-03 fase 2 / EDR-035 (usuário, 2026-10-19)"
    d = wb["14_DASHBOARD"]
    g = Grafo(wb)
    for r in DASH_CABECALHOS:
        d[f"E{r}"]._style = copy.copy(d["A6"]._style)
        d[f"E{r}"].value = "Coluna E: pior origem das entradas (ao vivo)"
        registro.append({"aba": "14_DASHBOARD", "cel": f"E{r}", "antes": None, "novo": d[f"E{r}"].value, "ref": ref})
    for r in DASH_RESULTADOS:
        celulas = g.celulas_origem(("14_DASHBOARD", f"B{r}"))
        if not celulas:
            raise RuntimeError(f"14_DASHBOARD B{r}: nenhuma raiz com coluna Origem")
        d[f"E{r}"]._style = copy.copy(d[f"D{r}"]._style)
        d[f"E{r}"].value = _formula_pior_origem(celulas)
        registro.append({"aba": "14_DASHBOARD", "cel": f"E{r}", "antes": None, "novo": d[f"E{r}"].value, "ref": ref})
    d.column_dimensions["E"].width = 30
    faixa = f"E{DASH_RESULTADOS[0]}:E{DASH_RESULTADOS[-1]}"
    for simb, cor in COR_ORIGEM.items():
        d.conditional_formatting.add(faixa, FormulaRule(formula=[f'LEFT($E{DASH_RESULTADOS[0]},1)="{simb}"'],
                                                        fill=PatternFill("solid", start_color=cor, end_color=cor)))
    registro.append({"aba": "14_DASHBOARD", "cel": f"{faixa} (formatação)", "antes": None, "novo": "cores por origem",
                     "ref": ref})


LINHA_LKYG = 177   # 10_TIRE: logo abaixo do ε (176); B86:B175 são os coeficientes da V2.4


def _lkyg(wb, layout: dict, registro: list[dict]) -> None:
    """λKyγ (LKYG) nos termos de camber de SHy e SVy (Adams 2024.1, PAC2002, eq. 41-42; EDR-042; usuário, 2026-10-20:
    "Pode corrigir"). Com LKYG = 1 no .tir provisório, nenhum valor muda."""
    import tire_model as tm
    ref = "EDR-042 (LKYG, Adams PAC2002 eq. 41-42; usuário 2026-10-20)"
    t = wb["10_TIRE"]
    prow, rows = layout["prow"], layout["rows"]
    lkyg = tm.load_dataset(TIR_PROVISORIO, source_axis_system=EIXO_TIR).p["LKYG"]
    r = LINHA_LKYG
    for col, val, src in (("A", "LKYG", "A175"), ("B", lkyg, "B175"), ("G", "-", "G175"), ("H", "SUS-501", "H175"),
                          ("I", "Escala da rigidez de camber Kyγ0 (termos de camber de SHy e SVy)", "I175")):
        c = t[f"{col}{r}"]
        c._style = copy.copy(t[src]._style)
        registro.append({"aba": "10_TIRE", "cel": f"{col}{r}", "antes": c.value, "novo": val, "ref": ref})
        c.value = val
    L = f"$B${r}"
    phy3, lmuy = f"$B${prow['PHY3']}", f"$B${prow['LMUY']}"
    for col in "BCDEF":
        gy = f"{col}${rows['gy']}"
        for linha, velho, novo in ((rows["SHy"], f"{phy3}*{gy})", f"{phy3}*{gy}*{L})"),
                                   (rows["SVy"], f"*{gy})*{lmuy})", f"*{gy}*{L})*{lmuy})")):
            c = t[f"{col}{linha}"]
            antes = c.value
            if antes.count(velho) != 1:
                raise RuntimeError(f"10_TIRE {col}{linha}: fórmula inesperada {antes!r}")
            c.value = antes.replace(velho, novo)
            registro.append({"aba": "10_TIRE", "cel": f"{col}{linha}", "antes": antes, "novo": c.value, "ref": ref})
    j = t["J84"]
    antes = j.value
    j.value = antes.replace("(B86:B175)", "(B86:B175 e o LKYG em B177)")
    registro.append({"aba": "10_TIRE", "cel": "J84", "antes": antes, "novo": j.value, "ref": ref})


def _eixo_iso_textos(wb, registro: list[dict]) -> None:
    """Textos e comentários que diziam SAE (EDR-039 reaberto); nenhum valor nem fórmula muda aqui."""
    def troca(aba, cel, velho, novo_txt):
        c = wb[aba][cel]
        assert velho in c.value, f"{aba}!{cel}: texto esperado não encontrado"
        antes = c.value
        c.value = c.value.replace(velho, novo_txt)
        registro.append({"aba": aba, "cel": cel, "antes": antes, "novo": c.value, "ref": _REF39})
    troca("00_README", "A37", "arquivo em SAE (TTC, declarado) transformado para ISO, sinal de camber não confirmado",
          "arquivo em ISO (eixo W do TYDEX, como manda o help do Adams para o PAC2002; EDR-039)")
    troca("04_DECISION_REGISTER", "E17",
          "Sistema de eixos do arquivo = SAE (Calspan/TTC), DECLARADO pelo usuario (dados baixados do TTC), transformado para ISO "
          "(alpha_arq=-alpha, Fy_ISO=-Fy_arq; camber: empuxo para o lado da inclinacao, INFERIDO: IA do TTC nao verificado).",
          "Sistema de eixos do arquivo = ISO (eixo W do TYDEX): o .tir e PAC2002 do Adams, cujos parametros valem sempre no "
          "TYDEX W (help do Adams 2024.1, p. 7; EDR-039 reaberto na V3.0). Ate a V2.4 era declarado SAE (Calspan/TTC), com o "
          "sinal do camber inferido.")
    c = wb["10_TIRE"]["B21"]
    antes = c.comment.text if c.comment else None
    texto = ("Dataset é de pneu Right (TYRESIDE). Rodas L usam espelhamento (H4; equivale ao espelhamento do Adams em relação "
             "a α = 0, help p. 73-74): α_R=−α e Fy_L=−Fy_R; o camber do projeto (+ = topo para fora) vale nos dois lados. "
             "Depois, transformação de eixos do arquivo (ISO_TYDEX = identidade). VERIF = pneu direito, caso fixo de verificação "
             "contra Python.")
    cm = Comment(texto, AUTOR); cm.width, cm.height = 340, 130
    c.comment = cm
    registro.append({"aba": "10_TIRE", "cel": "B21 (comentário)", "antes": antes, "novo": texto, "ref": _REF39})


def _feedback2_formatacao(wb, registro: list[dict]) -> None:
    """Parte da 2ª rodada da revisão metodológica que não é valor de célula: mesclagem do item 7 do guia, versão e changelog."""
    m = wb["00_README"]
    m.merge_cells(f"A{LINHA_NA_TBD}:E{LINHA_NA_TBD}")
    for r in DIR_CABECALHOS:
        wb["07_KINEMATICS"].merge_cells(f"A{r}:E{r}")
    marca = "EDR-035, EDR-050, EDR-051)"
    assert marca in m["A3"].value, "texto da versão em 00_README!A3 mudou"
    m["A3"].value = m["A3"].value.replace(marca, marca[:-1] + "; 2ª rodada da revisão metodológica: hipótese do CG não suspenso como "
                                          "entrada, centro da roda separado, significado de N/A e TBD)")
    c = wb["15_CHANGELOG"]
    for col in "ABCDE":
        c[f"{col}18"]._style = copy.copy(c[f"{col}15"]._style)
    c["A18"].value = "2026-10"
    c["B18"].value = ("V3.0, 2a rodada da revisao metodologica: razao h_ns / raio do pneu como entrada (05_VEHICLE B35/B36, Milliken p. 672); "
                      "05 B16/B17 = raio x razao; 07_KINEMATICS B14 (centro da roda traseira) = raio do pneu, sem passar pelo CG nao "
                      "suspenso; item 7 do guia (N/A e TBD); 14_DASHBOARD coluna E: pior origem das entradas de cada resultado, "
                      "ao vivo (FB-03 fase 2); 10_TIRE: eixo do .tir = ISO_TYDEX (EDR-039 reaberto, help do Adams PAC2002); 07_KINEMATICS "
                      "l.71-85: direcao (Ackermann, relacao pela cremalheira, roll steer; SUS-414, SUS-415, SUS-601; EDR-052).")
    c["C18"].value = ("Decisoes do usuario (2026-10-19). EDR-050 (estendido), EDR-035 (N/A e TBD). Registro celula a celula: "
                      "docs/registro-de-versoes-V3_0.md. Validation: test_workbook_suspensao (TC-WB30-013 estendido). NOT_VALIDATED fisicamente.")
    registro += [
        {"aba": "00_README", "cel": f"A{LINHA_NA_TBD}:E{LINHA_NA_TBD} (mesclagem)", "antes": None, "novo": "A:E", "ref": "EDR-035"},
        {"aba": "00_README", "cel": "A3", "antes": "Versão do workbook: V3.0 ...", "novo": "(acrescenta a 2ª rodada da revisão metodológica)",
         "ref": "EDR-035/050"},
        {"aba": "15_CHANGELOG", "cel": "A18:C18", "antes": None, "novo": "(linha V3.0, 2ª rodada da revisão metodológica)", "ref": "EDR-035/050"},
    ]


def _ky_formula(antiga: str, lfzo: str, lky: str) -> str:
    """Remove o segundo λFz0 ('*LFZO*LKY' -> '*LKY') sem tocar no resto da fórmula."""
    alvo = f"*{lfzo}*{lky})"
    if antiga.count(alvo) != 1:
        raise RuntimeError(f"fórmula de Ky inesperada: {antiga!r}")
    return antiga.replace(alvo, f"*{lky})")


def aplicar(src: Path, out: Path, layout: dict) -> list[dict]:
    wb = openpyxl.load_workbook(src)
    prow, rows = layout["prow"], layout["rows"]
    lfzo, lky = f"$B${prow['LFZO']}", f"$B${prow['LKY']}"
    registro = []
    for alt in alteracoes(rows["Ky"]):
        ws = wb[alt["aba"]]
        c = ws[alt["cel"]]
        antes = c.value
        novo = alt["novo"]
        if alt["aba"] == "10_TIRE" and alt["cel"] in {f"{col}{rows['Ky']}" for col in "BCDEF"}:
            novo = _ky_formula(antes, lfzo, lky)
        if isinstance(novo, tuple) and novo[0] == "COPIA":   # coluna G: valor da geração (EDR-035)
            novo = ws[novo[1]].value
        if alt["estilo"]:
            aba_e, linha_e, col_e = alt["estilo"]
            c._style = copy.copy(wb[aba_e][f"{col_e}{linha_e}"]._style)
        c.value = novo
        if alt["comentario"]:
            cm = Comment(alt["comentario"], AUTOR)
            cm.width, cm.height = 340, 120
            c.comment = cm
        registro.append({**alt, "antes": antes, "novo": novo})

    # Valores de mola, MR, barra e RC herdados do T09/E03: só referência e teste de formulação (usuário, 2026-10-19;
    # EDR-012: T09 nunca é baseline; EDR-022: valor demonstrativo rotulado). Muda só comentário e um aviso; nenhum valor.
    sus = wb["06_SUSPENSION"]
    aviso = ("DEMONSTRATIVO (EDR-012/022): valor herdado do T09/E03, usado só para referência, comparação e teste das "
             "fórmulas (confirmado pelo usuário em 2026-10-19). Não é o TX/E05. ")
    for cel in ("B4", "B5", "B6", "B7", "B10", "B11", "B14", "B15"):
        c = sus[cel]
        antes = c.comment.text if c.comment else None
        texto = aviso + (antes or "")
        cm = Comment(texto, AUTOR)
        cm.width, cm.height = 360, 150
        c.comment = cm
        registro.append({"aba": "06_SUSPENSION", "cel": f"{cel} (comentário)", "antes": antes, "novo": texto, "ref": "EDR-012/022"})
    m = wb["00_README"]
    m["A3"].value = ("Versão do workbook: V3.0 (em desenvolvimento até concluir as correções de feedback; revisão metodológica contra Milliken/Gillespie/Seward/Pacejka: pneu em série, "
                     "termo gravitacional da rolagem, massa suspensa por eixo, anti com sinal, Ky com um λFz0; EDR-045 a EDR-049) "
                     "| Status: DRAFT — cross-check, ainda NÃO REFERENCE-validated")
    tire = wb["10_TIRE"]
    tire[f"B{rows['m_eqv']}"].value = ("Fy0 confere com Seward eq. 5.5-5.13 (PAC'96/2002); combinado (Gxα, Gyκ, SVyκ) com Pacejka 2006 eq. 4.E50-4.E67; "
                                        "Ky com um único λFz0 (EDR-049). Reproduz o MFeval MF5.2. NÃO validado contra medição nem OptimumTire")
    registro.append({"aba": "10_TIRE", "cel": f"B{rows['m_eqv']}", "antes": "(texto V2.4)", "novo": "(texto V3.0)", "ref": "EDR-042/049"})
    registro.append({"aba": "00_README", "cel": "A3", "antes": "Versão do workbook: V2.4 ...", "novo": "Versão do workbook: V3.0 ...", "ref": "EDR-043"})

    c = wb["15_CHANGELOG"]
    for col in "ABCDE":
        c[f"{col}16"]._style = copy.copy(c[f"{col}15"]._style)
    c["A16"].value = "2026-10"
    c["B16"].value = ("V3.0: revisao metodologica contra Milliken (RCVD), Gillespie, Seward e Pacejka. Pneu em serie na frequencia natural e na rigidez de rolagem "
                      "(K_T provisorio do .tir; 06_SUSPENSION l.46-52, B22/B23, B32/B33; 11_DESIGN B8/B9/B20); SUS-103 com termo gravitacional e CG suspenso em a_s "
                      "(09_LOADS B18-B21; 13_VALIDATION B6); massa suspensa por eixo (05_VEHICLE l.25-27; 06 B28/B29); anti-dive/anti-squat com sinal e ponto de reacao "
                      "do carro (freio no cubo, motor no chassi; 06 B39-B41, 07 B11/B13/B14); Ky com um lambda_Fz0 (10_TIRE).")
    c["C16"].value = ("Pedido do usuario (2026-10-19): seguir a literatura, principalmente o Milliken. EDR-045 a EDR-049. Registro celula a celula: "
                      "docs/registro-de-versoes-V3_0.md. Validation: suspensao 54/54 (inclui exemplos da Milliken p. 683, p. 592, Tab. 16.3, Seward 4.1, Gillespie 5-8), "
                      "Excel x Python (test_workbook e test_workbook_suspensao). NOT_VALIDATED fisicamente. Autor: nao registrado.")
    registro.append({"aba": "15_CHANGELOG", "cel": "A16:C16", "antes": None, "novo": "(linha V3.0)", "ref": "EDR-043"})

    _rodada2_formatacao(wb, registro)
    _lkyg(wb, layout, registro)
    _feedback2_formatacao(wb, registro)
    _feedback2_pior_origem(wb, registro)
    _eixo_iso_textos(wb, registro)
    wb.calculation.fullCalcOnLoad = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return registro


AMARELO = "00FFF2CC"
COR_ORIGEM = {"●": "00CCE3F0", "◐": "00FAE2B3", "○": "00F0D7E5", "?": "00E0E0E0"}   # tons claros da paleta Okabe-Ito


def _rodada2_formatacao(wb, registro: list[dict]) -> None:
    """Parte da 1ª rodada de feedback que não é valor de célula: fundo amarelo, comentários migrados, lista, cores e avisos (EDR-035)."""
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    def reg(aba, cel, antes, novo, ref):
        registro.append({"aba": aba, "cel": cel, "antes": antes, "novo": novo, "ref": ref})

    # FB-02: o fundo amarelo "hipótese-chave" sai; a célula fica com o estilo de entrada (fonte azul), como as demais.
    for ws in wb.worksheets:
        for linha in ws.iter_rows():
            for c in linha:
                if c.fill.fill_type and c.fill.fgColor.rgb == AMARELO:
                    c.fill = PatternFill(fill_type=None)
                    reg(ws.title, f"{c.coordinate} (estilo)", "fundo amarelo (hipótese-chave)", "sem fundo", "FB-02 (usuário, 2026-10-19)")
    # FB-03: o texto dos comentários das entradas passou para a coluna Fonte; o comentário sai para não haver duas cópias.
    migrados = [(aba, f"B{r}") for aba, r, _, _ in _todas_origens()] + [("10_TIRE", "B23"), ("10_TIRE", "B24"), ("10_TIRE", "B86")]
    for aba, cel in migrados:
        c = wb[aba][cel]
        if c.comment is not None:
            reg(aba, f"{cel} (comentário)", c.comment.text, "(removido; texto na coluna Fonte)", "EDR-035")
            c.comment = None
    # Lista de origens (nome definido) e colunas.
    rot = f"'00_README'!$A${LINHAS_ORIGENS.start}:$A${LINHAS_ORIGENS.stop - 1}"
    wb.defined_names[NOME_LISTA] = DefinedName(NOME_LISTA, attr_text=rot)
    reg("(pasta)", f"nome definido {NOME_LISTA}", None, rot, "EDR-035")
    quebra = Alignment(wrap_text=True, vertical="top")

    def lista(ws, celulas):
        dv = DataValidation(type="list", formula1=f"={NOME_LISTA}", allow_blank=True, showErrorMessage=True,
                            errorTitle="Origem fora da lista", error="Escolha uma das origens da lista (00_README, 'Como preencher').")
        ws.add_data_validation(dv)
        for cel in celulas:
            dv.add(cel)

    def cores(ws, col, linha0, linha1):
        faixa = f"{col}{linha0}:{col}{linha1}"
        for simb, cor in COR_ORIGEM.items():
            ws.conditional_formatting.add(faixa, FormulaRule(formula=[f'LEFT(${col}{linha0},1)="{simb}"'],
                                                             fill=PatternFill("solid", start_color=cor, end_color=cor)))

    for aba in ABAS_ORIGEM:
        ws = wb[aba]
        linhas = [r for a_, r, org, _ in _todas_origens() if a_ == aba and org is not None]
        todas = [r for a_, r, _, _ in _todas_origens() if a_ == aba]
        r0, r1 = 3, max(ws.max_row, max(todas))
        ws.column_dimensions["E"].width, ws.column_dimensions["F"].width = 30, 80
        ws.column_dimensions["G"].hidden = ws.column_dimensions["H"].hidden = True
        for r in todas:
            ws[f"F{r}"].alignment = quebra
        lista(ws, [f"E{r}" for r in linhas])
        # Avisos primeiro (param a avaliação das cores): valor mudou e origem não; "Dado do TX" sem fonte.
        alerta = PatternFill("solid", start_color="00D55E00", end_color="00D55E00")
        ws.conditional_formatting.add(f"E{r0}:E{r1}", FormulaRule(
            formula=[f"AND(ISNUMBER($G{r0}),$B{r0}<>$G{r0},$E{r0}=$H{r0})"], fill=alerta,
            font=Font(color="00FFFFFF", bold=True), stopIfTrue=True))
        ws.conditional_formatting.add(f"F{r0}:F{r1}", FormulaRule(
            formula=[f'AND(LEFT($E{r0},1)="●",LEN(TRIM($F{r0}))=0)'],
            fill=PatternFill("solid", start_color="00FF0000", end_color="00FF0000"), stopIfTrue=True))
        cores(ws, "E", r0, r1)
        reg(aba, f"E{r0}:F{r1} (formatação)", None, "lista 'Origens' em E; cores por origem; avisos (E: valor mudou e origem não; "
            "F: Dado do TX sem fonte); G/H ocultas", "EDR-035")
    t = wb["10_TIRE"]
    t.column_dimensions["J"].width = 80
    for r, _, _ in TIRE_ORIGENS:
        t[f"J{r}"].alignment = quebra
    lista(t, [f"I{r}" for r, _, _ in TIRE_ORIGENS])
    cores(t, "I", 20, 176)
    reg("10_TIRE", "I20:J176 (formatação)", None, "lista 'Origens' em I; cores por origem", "EDR-035")
    # Versão e changelog.
    m = wb["00_README"]
    assert "EDR-045 a EDR-049)" in m["A3"].value, "texto da versão em 00_README!A3 mudou"
    m["A3"].value = m["A3"].value.replace("EDR-045 a EDR-049)", "EDR-045 a EDR-049; 1ª rodada de feedback (da equipe): origem de cada "
                                                               "entrada, valores derivados como fórmula, guia de preenchimento; "
                                                               "EDR-035, EDR-050, EDR-051)")
    # Texto corrido em A:E mesclado, como o resto da 00_README; a tabela de origens tem o significado em B:D.
    for r in list(range(LINHA_COMO_PREENCHER, LINHAS_ORIGENS.start - 1)) + list(range(LINHAS_ORIGENS.stop, LINHAS_ORIGENS.stop + 4)):
        m.merge_cells(f"A{r}:E{r}")
    for r in range(LINHAS_ORIGENS.start - 1, LINHAS_ORIGENS.stop):
        m.merge_cells(f"B{r}:D{r}")
        m[f"B{r}"].alignment = m[f"E{r}"].alignment = quebra
        m.row_dimensions[r].height = 30
    reg("00_README", "A3", "Versão do workbook: V3.0 ...", "(acrescenta a 1ª rodada de feedback)", "EDR-035/050/051")
    reg("00_README", f"A{LINHA_COMO_PREENCHER}:E{LINHAS_ORIGENS.stop + 3} (mesclagem)", None, "A:E nas linhas de texto; B:D na tabela de origens", "FB-06")
    c = wb["15_CHANGELOG"]
    for col in "ABCDE":
        c[f"{col}17"]._style = copy.copy(c[f"{col}15"]._style)
    c["A17"].value = "2026-10"
    c["B17"].value = ("V3.0, 1a rodada de feedback (da equipe): coluna Origem (lista) e Fonte em cada entrada (05, 06, 07, 08, 11; 10_TIRE I/J), "
                      "com cores e aviso de origem desatualizada; fundo amarelo retirado e legenda corrigida; valores derivados como formula "
                      "(05_VEHICLE B4, B9, B10, B16, B17; 08 B5; entradas novas 05 B30:B34 e 08 B11); guia 'Como preencher' na 00_README.")
    c["C17"].value = ("Feedback da equipe (2026-10-14) e decisoes do usuario (2026-10-19). EDR-025 (A), EDR-035, EDR-050, EDR-051. Registro celula "
                      "a celula: docs/registro-de-versoes-V3_0.md. Validation: test_workbook_suspensao (TC-WB30-010..013). NOT_VALIDATED fisicamente.")
    reg("15_CHANGELOG", "A17:C17", None, "(linha V3.0, 1ª rodada de feedback)", "EDR-035/050/051")


def escrever_registro(registro: list[dict], caminho: Path) -> None:
    def fmt(v):
        if v is None:
            return "(vazia)"
        s = str(v).replace("|", "\\|")
        return f"`{s}`" if s.startswith("=") else s
    linhas = ["| Aba | Célula | Antes | Depois | Fonte / EDR |", "|---|---|---|---|---|"]
    for r in registro:
        linhas.append(f"| {r['aba']} | {r['cel']} | {fmt(r['antes'])} | {fmt(r['novo'])} | {r['ref']} |")
    caminho.parent.mkdir(parents=True, exist_ok=True)
    caminho.write_text("\n".join(linhas) + "\n", encoding="utf-8", newline="\n")


def main() -> None:
    import json
    ap = argparse.ArgumentParser()
    ap.add_argument("src", nargs="?", default=str(AQUI.parent / "excel" / "FSAE_EV_Suspension_Fase1_V2_4.xlsx"))
    ap.add_argument("out", nargs="?", default=str(AQUI.parent / "build" / "FSAE_EV_Suspension_Fase1_V3_0.xlsx"))
    ap.add_argument("--layout", default=str(AQUI.parent / "data" / "layout_V2_4.json"))
    ap.add_argument("--registro", default=None)
    a = ap.parse_args()
    if Path(a.out).resolve() == Path(a.src).resolve():
        raise SystemExit("saída igual à origem: recusado (nunca sobrescrever a planilha entregue)")
    layout = json.loads(Path(a.layout).read_text(encoding="utf-8"))
    reg = aplicar(Path(a.src), Path(a.out), layout)
    if a.registro:
        escrever_registro(reg, Path(a.registro))
    print("salvo", a.out, "células alteradas", len(reg))


if __name__ == "__main__":
    main()
