"""Caso de equilíbrio demonstrativo lido da planilha e a pior origem dos dados (V4.0, incremento 1).

`caso_da_planilha` lê as entradas de uma CÓPIA já recalculada (Excel ou LibreOffice, EDR-044) e monta o
`CasoEquilibrio`. `pior_origem_do_caso` abre a planilha SEM valores (só fórmulas) e acha o pior rótulo de origem
entre as raízes das entradas e o pneu (origem lida da 10_TIRE, I84); `pior_origem_entradas` faz o mesmo sem o
pneu. Nenhuma função salva ou recalcula a planilha; a viva só é aberta para leitura.

ENTRADAS_CASO mapeia cada chave de entrada ao nome definido da V4.0 (data/nomes_V4_0.csv); a leitura resolve o nome na
própria planilha (`wb.defined_names`). Uma cópia da V3.0, sem nomes definidos, cai no endereço da mesma tabela
(`endereco_da_entrada`), então as ferramentas que passam a V3.0 viva continuam funcionando.
Toe, camber estático e ganho de camber ficam em 0: não há dado do carro (EDR-036).
"""
from __future__ import annotations

import math
import warnings
from pathlib import Path

import openpyxl

import suspension_model as sm
import tire_model as tm
from equilibrio import CasoEquilibrio
from raizes import Grafo, origem_da_raiz, PIOR_PRIMEIRO, SEM_ROTULO
from atualizar_planilha import CONF
from nomes_v4_0 import ler_tabela

TABELA_NOMES = Path(__file__).resolve().parent.parent / "data" / "nomes_V4_0.csv"

ENTRADAS_CASO: dict[str, str] = {
    "m": "res_05_massa_total_com_piloto",
    "h": "ent_05_altura_do_cg_total",
    "L": "ent_05_wheelbase_l",
    "tF": "ent_05_front_track_t_f",
    "tR": "ent_05_rear_track_t_r",
    "a": "res_05_cg_eixo_dianteiro_a",
    "b": "res_05_cg_eixo_traseiro_b",
    "muF": "ent_05_massa_nao_suspensa_dianteira",
    "muR": "ent_05_massa_nao_suspensa_traseira",
    "huF": "res_05_altura_cg_nao_suspenso_dianteiro",
    "huR": "res_05_altura_cg_nao_suspenso_traseiro",
    "KsF": "ent_06_spring_rate_dianteiro",
    "KsR": "ent_06_spring_rate_traseiro",
    "MRF": "ent_06_motion_ratio_dianteiro_wheel_spring",
    "MRR": "ent_06_motion_ratio_traseiro_wheel_spring",
    "KarbF": "ent_06_k_arb_dianteiro",
    "KarbR": "ent_06_k_arb_traseiro",
    "zF": "ent_06_roll_center_dianteiro_rh_estatico",
    "zR": "ent_06_roll_center_traseiro_rh_estatico",
    "KTF": "ent_06_rigidez_vertical_do_pneu_dianteiro_k_t",
    "KTR": "ent_06_rigidez_vertical_do_pneu_traseiro_k_t",
    "DF": "ent_08_downforce_total",
    "bal": "ent_08_aero_balance_fracao_dianteira",
}

_ENDERECOS_V30: dict[str, tuple[str, str]] = {}


def _enderecos_da_tabela() -> dict[str, tuple[str, str]]:
    """nome -> (aba, célula) da tabela data/nomes_V4_0.csv (os endereços da V3.0 e da V4.0), lida uma vez."""
    if not _ENDERECOS_V30:
        _ENDERECOS_V30.update({n.nome: (n.aba, n.celula) for n in ler_tabela(TABELA_NOMES)})
    return _ENDERECOS_V30


def endereco_da_entrada(wb, chave: str) -> tuple[str, str]:
    """(aba, célula) da entrada `chave`: resolve o nome definido na planilha `wb`; sem o nome (cópia da V3.0), usa o
    endereço da tabela de nomes. Nome que aponta para algo que não é uma única célula é erro."""
    nome = ENTRADAS_CASO[chave]
    dn = wb.defined_names.get(nome)
    if dn is None:
        return _enderecos_da_tabela()[nome]
    dest = list(dn.destinations)
    cel = dest[0][1].replace("$", "") if len(dest) == 1 else ""
    if not cel or ":" in cel:
        raise ValueError(f"o nome {nome} não aponta para uma única célula ({dn.attr_text})")
    return dest[0][0], cel


def enderecos_das_entradas(wb) -> dict[str, tuple[str, str]]:
    """chave -> (aba, célula) de todas as entradas do caso na planilha `wb` (ver `endereco_da_entrada`)."""
    return {k: endereco_da_entrada(wb, k) for k in ENTRADAS_CASO}


def caso_da_planilha(xlsx_recalculado: Path, tir: Path) -> CasoEquilibrio:
    """Monta o CasoEquilibrio a partir de uma cópia recalculada; ValueError se alguma entrada não tem valor.

    Lê pelos nomes definidos da V4.0 (ou posterior); numa cópia da V3.0, sem nomes, pelos endereços da tabela.
    """
    wb = openpyxl.load_workbook(xlsx_recalculado, data_only=True)
    x: dict[str, float] = {}
    for chave in ENTRADAS_CASO:
        aba, cel = endereco_da_entrada(wb, chave)
        v = wb[aba][cel].value
        if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v):
            raise ValueError("a cópia precisa ser recalculada (Excel ou LibreOffice) antes da leitura: "
                             f"{aba}!{cel} sem valor")
        x[chave] = float(v)
    ds = tm.load_dataset(str(tir), source="caso da planilha", source_axis_system="ISO_TYDEX")
    return CasoEquilibrio(_cfg(x), ds, downforce_N=x["DF"], aero_balance_front=x["bal"])


def _cfg(x: dict) -> sm.VehicleConfig:
    """VehicleConfig das entradas do caso (chaves de ENTRADAS_CASO), com K_φ pelo SUS-206/207 (pneu em série)."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        kw_f, kw_r = sm.wheel_rate(x["KsF"], x["MRF"]), sm.wheel_rate(x["KsR"], x["MRR"])
    kphi_f = sm.roll_stiffness_axle(kw_f, x["tF"], x["KarbF"], x["KTF"])
    kphi_r = sm.roll_stiffness_axle(kw_r, x["tR"], x["KarbR"], x["KTR"])
    return sm.VehicleConfig(
        mass_total=x["m"], cg_height=x["h"], wheelbase=x["L"], track_front=x["tF"], track_rear=x["tR"],
        cg_to_front=x["a"], cg_to_rear=x["b"], unsprung_mass_front=x["muF"], unsprung_mass_rear=x["muR"],
        unsprung_cg_height_front=x["huF"], unsprung_cg_height_rear=x["huR"], roll_center_front=x["zF"],
        roll_center_rear=x["zR"], roll_stiffness_front=kphi_f, roll_stiffness_rear=kphi_r)


# ---------------------------------------------------------------------------------------------------------------------
# V4.1 (incremento 4): o caso que a 16_EQUILIBRIO calcula, lido por nomes
# ---------------------------------------------------------------------------------------------------------------------
def _valor_por_nome(wb, nome: str):
    import aba_equilibrio as ae
    aba, cel = ae._destino(wb, nome)
    return wb[aba][cel].value


def valores_v41(wb) -> dict:
    """chave -> valor (cópia recalculada, data_only) das entradas que a 16_EQUILIBRIO usa: ENTRADAS16, V40 (inclui as 23
    de ENTRADAS_CASO), "coef" (lista dos coeficientes da 10_TIRE B86:B175 e B177) e "coef_nomes" (nome do .tir ->
    valor, para os coeficientes que o .tir também tem)."""
    import aba_equilibrio as ae
    x = {e.chave: _valor_por_nome(wb, ae.NOME_ENT[e.chave]) for e in ae.ENTRADAS16}
    x.update({k: _valor_por_nome(wb, nome) for k, nome in ae.V40.items()})
    t = wb["10_TIRE"]
    x["coef"] = [t[f"B{r}"].value for r in list(range(86, 176)) + [177]]
    x["coef_nomes"] = {str(t[f"A{r}"].value).strip(): t[f"B{r}"].value for r in list(range(86, 176)) + [177]}
    return x


def caso_da_planilha_v41(xlsx_recalculado: Path, tir: Path) -> CasoEquilibrio:
    """O CasoEquilibrio que a 16_EQUILIBRIO da V4.1 calcula, lido por nomes de uma cópia recalculada.

    - Entradas da V4.0 pelos nomes de ENTRADAS_CASO; as novas pelos nomes ent_16_* (aba_equilibrio.NOME_ENT); o camber
      estático da 10_TIRE (B24:E24, °), o ganho de camber em rolagem derivado do SUS-404 (07_KINEMATICS B23/B24, °/mm,
      por kinematics_model.camber_gradient_from_body; autorização 7 do plano) e o % Ackermann da 07 B75 (h2).
    - Pneu: o .tir `tir`, com o sistema de eixos da 10_TIRE B11 e os coeficientes da 10_TIRE (B86:B177) por cima dos do
      arquivo, para que o caso seja o que a planilha avalia (uma cópia de teste pode escalar, por exemplo, LMUY).
    - Chaves (spec 7.4): o raio entra com h2, h3, h8 ou h9 (no Python, o raio liga o cos δ_ref); h2 usa o % Ackermann;
      h7, o CG lateral; h8 troca a downforce fixa por CzA e ρ (downforce_N = 0); h9 usa CdA, Crr e ρ.
    - Entrada inválida para a planilha (chave ligada com entrada vazia, texto numa entrada numérica, h3 desligada com
      h2/h8/h9): ValueError com o mesmo motivo que a 16 mostra (aba_equilibrio.VERIFICACOES, na mesma ordem); idem para
      a instabilidade de rolagem e o Ackermann fora do alcance.
    - Sentido "Direita": devolve equilibrio.espelhar do caso (a curva à direita do carro vista como curva à esquerda do
      carro espelhado), que é o que a planilha mostra (cantos no referencial espelhado; FL = roda interna).
    """
    import math as _m
    import aba_equilibrio as ae
    from equilibrio import espelhar, esterco_rodas
    from kinematics_model import camber_gradient_from_body
    wb = openpyxl.load_workbook(xlsx_recalculado, data_only=True)
    try:
        x = valores_v41(wb)
    finally:
        wb.close()
    motivo = ae.motivo_das_entradas(x)
    if motivo is not None:
        raise ValueError(motivo)
    h = {k: x[k] == 1 for k in ("h2", "h3", "h7", "h8", "h9")}
    base = {k: float(x[k]) for k in ENTRADAS_CASO if not (k == "DF" and h["h8"])}
    base.setdefault("DF", 0.0)
    cfg = _cfg(base)
    try:
        sm.lateral_dfz(cfg, 1.0)
    except ValueError:
        raise ValueError(ae.MSG_DEN) from None
    ds = tm.load_dataset(str(tir), source="caso da planilha V4.1", source_axis_system=x["eixos"])
    for k, v in x["coef_nomes"].items():
        if k in ds.p:
            ds.p[k] = float(v)
    raio = float(x["R"]) if any(h[k] for k in ("h2", "h3", "h8", "h9")) else None
    gF = camber_gradient_from_body(*(2 * [_m.radians(float(x["dgF"])) * 1000.0]), cfg.track_front)
    gR = camber_gradient_from_body(*(2 * [_m.radians(float(x["dgR"])) * 1000.0]), cfg.track_rear)
    caso = CasoEquilibrio(
        cfg, ds, downforce_N=0.0 if h["h8"] else base["DF"], aero_balance_front=base["bal"],
        toe_rad={c: _m.radians(float(x[f"toe{c}"])) for c in ae.CANTOS},
        camber_estatico_rad={c: _m.radians(float(x[f"gs{c}"])) for c in ae.CANTOS},
        ganho_camber={"FL": gF["L"], "FR": gF["R"], "RL": gR["L"], "RR": gR["R"]},
        raio_m=raio, ackermann_pct=float(x["pct"]) if h["h2"] else None,
        cg_lateral_m=float(x["cg"]) if h["h7"] else 0.0,
        cza_m2=float(x["cza"]) if h["h8"] else None,
        rho_kg_m3=float(x["rho"]) if (h["h8"] or h["h9"]) else None,
        cda_m2=float(x["cda"]) if h["h9"] else None, crr=float(x["crr"]) if h["h9"] else None)
    if caso.ackermann_pct not in (None, 0.0):
        try:
            esterco_rodas(caso)
        except ValueError:
            raise ValueError(ae.MSG_ACK) from None
    return espelhar(caso) if x["sentido"] == "Direita" else caso


def pior_origem_do_caso_v41(xlsx_formulas: Path) -> str:
    """Pior origem que a 16_EQUILIBRIO mostra (res_16_pior_origem...), pela mesma regra: raízes preenchidas dos nós da
    V4.0 que o limite usa (raizes.Grafo), a origem do pneu (10_TIRE I84) e as entradas da 16 pela coluna E, cada termo
    condicional só quando a chave o usa (aba_equilibrio.termos_pior_origem). Abre a planilha só para leitura, sem
    data_only (as chaves são valores digitados). Origem vazia ou fora da lista conta como "? A confirmar"."""
    import aba_equilibrio as ae
    wb = openpyxl.load_workbook(xlsx_formulas, data_only=False)
    try:
        termos = ae.termos_pior_origem(wb)
        chaves = {k: _valor_por_nome(wb, ae.NOME_ENT[k]) for k in ("h2", "h3", "h7", "h8", "h9")}
        return ae.pior_origem_de_valores(termos, lambda aba, cel: wb[aba][cel].value, chaves)
    finally:
        wb.close()


def _origens_entradas(wb) -> set[str]:
    """Rótulos de origem das raízes preenchidas de ENTRADAS_CASO (raiz sem rótulo conta como "? A confirmar")."""
    grafo = Grafo(wb)
    achados = set()
    for no in enderecos_das_entradas(wb).values():
        for aba, cel in grafo.raizes_preenchidas(no):
            rotulo = origem_da_raiz(wb, aba, cel)
            achados.add(CONF if rotulo == SEM_ROTULO else rotulo)
    return achados


def _origem_pneu(wb) -> str:
    """Origem do pneu na própria planilha: a célula de origem dos coeficientes do .tir na 10_TIRE
    (raizes.celula_origem("10_TIRE", "B86") → I84). Sem rótulo conta como "? A confirmar"."""
    rotulo = origem_da_raiz(wb, "10_TIRE", "B86")   # lê raizes.celula_origem("10_TIRE", "B86"), a I84
    return CONF if rotulo == SEM_ROTULO else rotulo


def pior_origem_entradas(xlsx_formulas: Path) -> str:
    """Pior rótulo de origem (ordem de raizes.PIOR_PRIMEIRO) entre as raízes das entradas, sem o pneu.

    Abre a planilha só para leitura, sem data_only. Raiz sem rótulo conta como "? A confirmar".
    """
    wb = openpyxl.load_workbook(xlsx_formulas, data_only=False)
    achados = _origens_entradas(wb)
    if not achados:
        raise ValueError("nenhuma raiz preenchida nas entradas do caso: a pior origem não está definida")
    return min(achados, key=PIOR_PRIMEIRO.index)


def pior_origem_do_caso(xlsx_formulas: Path) -> str:
    """Pior rótulo de origem (ordem de raizes.PIOR_PRIMEIRO) entre as raízes das entradas e o pneu.

    A origem do pneu é lida da célula de origem dos coeficientes na 10_TIRE (I84), não de uma constante: se a equipe
    trocar o .tir e o rótulo, o caso acompanha. Abre a planilha só para leitura, sem data_only. Raiz sem rótulo conta
    como "? A confirmar".
    """
    wb = openpyxl.load_workbook(xlsx_formulas, data_only=False)
    achados = _origens_entradas(wb) | {_origem_pneu(wb)}
    return min(achados, key=PIOR_PRIMEIRO.index)
