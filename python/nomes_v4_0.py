"""Tabela de nomes definidos da V4.0 (incremento 3, fundação: nomes definidos e dados por chave).

Cada célula de entrada da V3.0 (`atualizar_planilha.celulas_de_entrada()`) e cada célula com fórmula das abas de
cálculo (05, 06, 07, 08, 09, 10, 11, 13 e 14) recebe um nome `<papel>_<aba2>_<slug>`:
- papel: `ent` (valor de entrada), `org` (coluna Origem, E ou I na 10_TIRE), `fte` (coluna Fonte, F ou J na 10_TIRE),
  `res` (resultado com fórmula);
- aba2: os dois dígitos da aba;
- slug: o rótulo da coluna A da linha, sem acento, em minúsculas, com todo caractere fora de [a-z0-9] trocado por `_`,
  `_` repetidos colapsados, sem `_` nas pontas e com até 40 caracteres. Na 10_TIRE, as células de caso (colunas B a F
  das linhas com casos FL, FR, RL, RR e VERIF) levam também a coluna (`_b`…`_f`).
Em colisão, todos os nomes iguais ganham `_<célula>` (por exemplo `_c86`), o que torna o resultado independente da
ordem. Rótulo vazio dá slug vazio, e o nome vira `<papel>_<aba2>_<célula>` (a mesma regra da colisão).
`sus_id` guarda os IDs SUS que já aparecem em algum texto da linha (nenhum ID é criado); vários vão separados por `;`.

A tabela é gerada do arquivo da tag v3.0 (extraído para uma pasta temporária), aberto só para leitura.
Uso: uv run python -X utf8 python/nomes_v4_0.py --v30 <arquivo> --saida data/nomes_V4_0.csv
"""
from __future__ import annotations

import argparse
import csv
import re
import unicodedata
from dataclasses import dataclass, fields
from pathlib import Path
from typing import Iterable

import openpyxl
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.formula import ArrayFormula

from atualizar_planilha import ABAS_ORIGEM, celulas_de_entrada

ABAS_RES = ("05", "06", "07", "08", "09", "10", "11", "13", "14")   # abas cujas fórmulas viram `res`
PAPEIS = ("ent", "org", "fte", "res")
CABECALHO = ("nome", "aba", "celula", "papel", "sus_id", "rotulo")
TAM_SLUG = 40
COLUNAS_CASO = "BCDEF"   # 10_TIRE: FL, FR, RL, RR e VERIF

# Regras do Excel: começa por letra; só letras ASCII, dígitos e `_`; até 255 caracteres.
NOME_VALIDO = re.compile(r"[A-Za-z][A-Za-z0-9_]{0,254}")
_RE_A1 = re.compile(r"([A-Za-z]{1,3})(\d{1,7})")
_RE_R1C1 = re.compile(r"[Rr]\d*(?:[Cc]\d*)?|[Cc]\d*")
_RE_SUS = re.compile(r"\bSUS-\d{3}(?:-[FR])?\b")
_RE_CEL = re.compile(r"([A-Z]+)(\d+)")


@dataclass(frozen=True)
class Nome:
    """Uma linha da tabela de nomes: o nome definido, a célula a que ele aponta e de onde veio."""
    nome: str
    aba: str
    celula: str
    papel: str
    sus_id: str
    rotulo: str


def parece_referencia(nome: str) -> bool:
    """True se o nome pode ser lido pelo Excel como referência: A1 (até XFD1048576) ou R1C1 (R, C, RC, R1, C1, R1C1…)."""
    m = _RE_A1.fullmatch(nome)
    if m and column_index_from_string(m.group(1).upper()) <= 16384 and 1 <= int(m.group(2)) <= 1048576:
        return True
    return bool(_RE_R1C1.fullmatch(nome))


def slug(rotulo: str) -> str:
    """Rótulo sem acento, em minúsculas, com fora de [a-z0-9] trocado por `_`, sem `_` repetido nem nas pontas, ≤ 40."""
    s = unicodedata.normalize("NFKD", rotulo).encode("ascii", "ignore").decode("ascii").lower()
    s = re.sub(r"[^a-z0-9]+", "_", s).strip("_")
    return s[:TAM_SLUG].rstrip("_")


def _partes(celula: str) -> tuple[str, int]:
    m = _RE_CEL.fullmatch(celula)
    if not m:
        raise ValueError(f"célula inválida: {celula!r}")
    return m.group(1), int(m.group(2))


def _chave(aba: str, celula: str) -> tuple[str, int, int]:
    col, lin = _partes(celula)
    return aba, lin, column_index_from_string(col)


def _eh_formula(v) -> bool:
    return isinstance(v, ArrayFormula) or (isinstance(v, str) and v.startswith("="))


def _papel_entrada(aba: str, celula: str) -> str:
    col, _ = _partes(celula)
    if aba in ABAS_ORIGEM:
        return {"E": "org", "F": "fte"}.get(col, "ent")
    if aba == "10_TIRE":
        return {"I": "org", "J": "fte"}.get(col, "ent")
    return "ent"


def _texto(v) -> str:
    """Texto de uma célula como rótulo: None e fórmula viram vazio; número vira o texto dele."""
    if v is None or _eh_formula(v):
        return ""
    return str(v)


def tabela_nomes(caminho_v30: Path) -> list[Nome]:
    """Tabela de nomes da V3.0 em `caminho_v30` (o arquivo da tag v3.0), ordenada por (aba, linha, coluna)."""
    wb = openpyxl.load_workbook(caminho_v30)
    try:
        celulas: dict[tuple[str, str], str] = {k: _papel_entrada(*k) for k in celulas_de_entrada()}
        for ws in wb.worksheets:
            if ws.title[:2] not in ABAS_RES:
                continue
            for lin in ws.iter_rows():
                for c in lin:
                    if _eh_formula(c.value):
                        k = (ws.title, c.coordinate)
                        if k in celulas:
                            raise ValueError(f"{k[0]}!{k[1]} é entrada e tem fórmula")
                        celulas[k] = "res"
        base = []
        for aba, cel in sorted(celulas, key=lambda k: _chave(*k)):
            if aba not in wb.sheetnames:
                raise ValueError(f"aba {aba!r} não existe na V3.0")
            ws = wb[aba]
            col, r = _partes(cel)
            rotulo = _texto(ws[f"A{r}"].value)
            s = slug(rotulo)
            if s and aba == "10_TIRE" and col in COLUNAS_CASO and any(ws[f"{x}{r}"].value is not None for x in "CDEF"):
                s = f"{s}_{col.lower()}"
            papel = celulas[(aba, cel)]
            textos = [str(ws.cell(r, j).value) for j in range(1, ws.max_column + 1)
                      if isinstance(ws.cell(r, j).value, str) and not _eh_formula(ws.cell(r, j).value)]
            ids = list(dict.fromkeys(m for t in textos for m in _RE_SUS.findall(t)))
            nome = f"{papel}_{aba[:2]}_{s}" if s else ""
            base.append([nome, aba, cel, papel, ";".join(ids), rotulo])
    finally:
        wb.close()
    contagem: dict[str, int] = {}
    for b in base:
        contagem[b[0]] = contagem.get(b[0], 0) + 1
    out = []
    for nome, aba, cel, papel, sus, rot in base:
        if not nome:   # rótulo vazio: só a célula distingue
            nome = f"{papel}_{aba[:2]}_{cel.lower()}"
        elif contagem[nome] > 1:
            nome = f"{nome}_{cel.lower()}"
        out.append(Nome(nome, aba, cel, papel, sus, rot))
    vistos: dict[str, Nome] = {}
    for n in out:
        if n.nome in vistos:
            raise ValueError(f"nome repetido depois da regra de colisão: {n.nome} ({vistos[n.nome].aba}!"
                             f"{vistos[n.nome].celula} e {n.aba}!{n.celula})")
        if not NOME_VALIDO.fullmatch(n.nome) or parece_referencia(n.nome):
            raise ValueError(f"nome inválido para o Excel: {n.nome!r} ({n.aba}!{n.celula})")
        vistos[n.nome] = n
    return out


def verificar_cobertura(tabela: Iterable[Nome], celulas: Iterable[tuple[str, str]]) -> None:
    """Levanta ValueError, listando as células (`aba!célula`), se alguma de `celulas` não tem nome na tabela."""
    com_nome = {(n.aba, n.celula) for n in tabela}
    faltam = sorted((k for k in set(celulas) if k not in com_nome), key=lambda k: _chave(*k))
    if faltam:
        raise ValueError(f"{len(faltam)} célula(s) sem nome na tabela da V4.0: " +
                         ", ".join(f"{a}!{c}" for a, c in faltam))


def escrever_tabela(nomes: Iterable[Nome], caminho_csv: Path) -> None:
    """Grava a tabela em CSV UTF-8 com fim de linha LF e o cabeçalho nome,aba,celula,papel,sus_id,rotulo."""
    with open(caminho_csv, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(CABECALHO)
        for n in nomes:
            w.writerow([getattr(n, c) for c in CABECALHO])


def ler_tabela(caminho_csv: Path) -> list[Nome]:
    """Lê a tabela gravada por `escrever_tabela`; recusa cabeçalho diferente."""
    with open(caminho_csv, encoding="utf-8", newline="") as f:
        linhas = list(csv.reader(f))
    if not linhas or tuple(linhas[0]) != CABECALHO:
        raise ValueError(f"cabeçalho inesperado em {caminho_csv}: {linhas[0] if linhas else '(vazio)'}")
    assert tuple(f.name for f in fields(Nome)) == CABECALHO
    return [Nome(*lin) for lin in linhas[1:]]


def main() -> None:
    p = argparse.ArgumentParser(description="Gera a tabela de nomes definidos da V4.0 a partir da V3.0 (tag v3.0).")
    p.add_argument("--v30", type=Path, required=True, help="arquivo da V3.0 extraído da tag v3.0")
    p.add_argument("--saida", type=Path, required=True, help="CSV de saída (por exemplo data/nomes_V4_0.csv)")
    a = p.parse_args()
    if a.saida.suffix.lower() != ".csv":
        raise SystemExit(f"a saída deve ser um .csv: {a.saida}")
    tabela = tabela_nomes(a.v30)
    escrever_tabela(tabela, a.saida)
    cont = {pp: sum(1 for n in tabela if n.papel == pp) for pp in PAPEIS}
    print(f"{len(tabela)} nomes gravados em {a.saida}: {cont}")


if __name__ == "__main__":
    main()
