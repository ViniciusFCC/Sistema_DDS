"""Planilha viva da equipe no Google Drive: localizar, copiar, identificar a versão e achar a referência (EDR-063).

A viva é o arquivo que a equipe preenche (hoje a cópia da V4.1 no Drive, depois a V4.2…). Ela muda sem commit; por isso a
suíte só a verifica de forma informativa, e as ferramentas (`conferir_viva`, `retrato_viva`, `migrar_viva`) seguem estas regras:
- só leitura: o arquivo do Drive nunca é gravado, movido, renomeado nem salvo. Toda ferramenta trabalha numa CÓPIA local
  (`copia_local`), numa pasta temporária;
- nunca pelo LibreOffice (remove as aspas das abas iniciadas por dígito);
- viva bloqueada (aberta por alguém), placeholder "só online" do Drive (0 byte), arquivo ausente ou Drive desmontado:
  `VivaIlegivel`, que as CLIs transformam no código de saída 3. Nunca PASS.

A versão da viva é a que está na 00_README A3 ("Versão do workbook: V4.1 ..."). A referência dessa versão é o arquivo da
tag Git dela (imutável); antes de a tag existir (V4.2 até a Task 7), é o arquivo de `excel/`.
"""
from __future__ import annotations

import hashlib
import os
import re
import shutil
import subprocess
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import openpyxl

import origens

RAIZ = Path(__file__).resolve().parent.parent
PASTA_DRIVE = Path(r"H:\.shortcut-targets-by-id\1Kytf4APy0OtH5Aq5PqlOwH5UY65YMX5U\T11 - Tupã\2. Projeto\Mecanica\Suspensao"
                   r"\Documentacao\Planilha DDS")
PADRAO_VIVA = "FSAE_EV_Suspension_Fase1_V*.xlsx"
RE_VERSAO = re.compile(r"Versão do workbook: (V\d+\.\d+)\b")
RE_VERSAO_NUM = re.compile(r"V(\d+)\.(\d+)")
RE_EDR = re.compile(r"EDR-\d+")
RE_STATUS_EDR = re.compile(r"^\| (EDR-\d{3}) \|.*\| ([A-Z_]+) \|\s*$", re.M)


class VivaIlegivel(Exception):
    """A viva (ou a referência) não pôde ser lida: ausente, de 0 byte, bloqueada, Drive desmontado, versão desconhecida.
    As CLIs saem com código 3 e nunca contam como aprovado."""


def localizar_viva() -> Path | None:
    """A viva a verificar: a variável `DDS_VIVA` (caminho do arquivo) primeiro; senão o único `FSAE_EV_Suspension_Fase1_V*.xlsx`
    da `PASTA_DRIVE`. None se não há `DDS_VIVA` e a pasta não está montada ou não tem o arquivo. Mais de um arquivo na pasta:
    ValueError (a ferramenta não adivinha qual é a viva). Não abre nenhum arquivo."""
    env = os.environ.get("DDS_VIVA")
    if env:
        return Path(env)
    try:
        achados = sorted(PASTA_DRIVE.glob(PADRAO_VIVA)) if PASTA_DRIVE.is_dir() else []
    except OSError:
        return None
    if len(achados) > 1:
        raise ValueError(f"mais de uma planilha viva em {PASTA_DRIVE}: {[a.name for a in achados]} "
                         "(aponte DDS_VIVA para a certa)")
    return achados[0] if achados else None


def dentro(caminho: Path, pasta: Path) -> bool:
    """True se `caminho` é a própria `pasta` ou está em qualquer nível abaixo dela (só compara caminhos; não abre nada)."""
    c, p = Path(caminho).resolve(), Path(pasta).resolve()
    return c == p or p in c.parents


def recusa_de_saida(alvo: Path, entradas: list[Path]) -> str | None:
    """Guarda comum das CLIs que gravam (EDR-063): mensagem de recusa se `alvo` é um dos arquivos de `entradas` (gravar por cima
    da viva ou de uma planilha de entrada), está em qualquer nível da árvore `excel/` ou dentro de `PASTA_DRIVE` (a viva do Drive é
    só leitura); None se pode gravar. Só compara caminhos: a recusa vem antes de abrir ou escrever qualquer arquivo."""
    alvo = Path(alvo)
    if alvo.resolve() in {Path(e).resolve() for e in entradas}:
        return f"saída por cima de um arquivo de entrada: recusado ({alvo})"
    if dentro(alvo, RAIZ / "excel"):
        return f"saída em excel/ recusada ({alvo}): a cópia para excel/ é um commit próprio"
    if dentro(alvo, PASTA_DRIVE):
        return f"saída na pasta do Drive recusada ({alvo}): a viva do Drive é só leitura"
    return None


def sha256(arquivo: Path) -> str:
    h = hashlib.sha256()
    with open(arquivo, "rb") as f:
        for bloco in iter(lambda: f.read(1 << 20), b""):
            h.update(bloco)
    return h.hexdigest()


def copia_local(viva: Path, pasta: Path) -> Path:
    """Copia a viva para `pasta` (leitura do original, gravação só na cópia) e devolve a cópia. 0 byte, arquivo ausente,
    bloqueio, erro de leitura, cópia de tamanho diferente do original ou arquivo que não é um pacote xlsx: VivaIlegivel."""
    viva, pasta = Path(viva), Path(pasta)
    try:
        tamanho = viva.stat().st_size
    except OSError as e:
        raise VivaIlegivel(f"viva ausente ou inacessível ({viva}): {e}") from e
    if tamanho == 0:
        raise VivaIlegivel(f"viva de 0 byte (placeholder do Drive 'só online' ou arquivo vazio): {viva}")
    destino = pasta / viva.name
    try:
        pasta.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(viva, destino)   # abre o original só para leitura
    except OSError as e:
        raise VivaIlegivel(f"viva bloqueada ou ilegível ({viva}): {e}") from e
    if destino.stat().st_size != tamanho:
        raise VivaIlegivel(f"cópia da viva incompleta ({destino.stat().st_size} de {tamanho} bytes): {viva}")
    if not zipfile.is_zipfile(destino):
        raise VivaIlegivel(f"a viva não é um pacote xlsx legível: {viva}")
    return destino


def versao(xlsx: Path) -> str:
    """Versão da planilha, lida da 00_README A3 ("V4.1"). VivaIlegivel se não houver."""
    try:
        wb = openpyxl.load_workbook(xlsx, read_only=True)
        try:
            a3 = wb["00_README"]["A3"].value
        finally:
            wb.close()
    except Exception as e:   # arquivo corrompido, aba ausente…
        raise VivaIlegivel(f"não consegui ler a versão de {xlsx}: {e}") from e
    m = RE_VERSAO.match(a3 or "")
    if not m:
        raise VivaIlegivel(f"00_README!A3 sem \"Versão do workbook: Vx.y\": {str(a3)[:80]!r}")
    return m.group(1)


@dataclass(frozen=True)
class Referencia:
    """O que a conferência de uma versão precisa: tag Git, arquivo em `excel/`, células de entrada, rótulos de Origem,
    número de fórmulas e se há EDR citado nas fichas de meta (V4.2 em diante)."""
    versao: str
    tag: str
    caminho_excel: str
    entradas: Callable[[], set]
    rotulos: tuple[str, ...]
    formulas: int
    confere_edr: bool = False


def _entradas_v41() -> set:
    from aba_equilibrio import celulas_de_entrada_v41
    return celulas_de_entrada_v41()


def _entradas_v42() -> set:
    from aba_metas import celulas_de_entrada_v42
    return celulas_de_entrada_v42()


# O número de fórmulas repete tools/comparar_planilhas.FORMULAS_ESPERADAS_V41 e _V42 (o TC-WB42-004 confere os dois).
REFERENCIAS: dict[str, Referencia] = {
    "V4.1": Referencia("V4.1", "v4.1", "excel/FSAE_EV_Suspension_Fase1_V4_1.xlsx", _entradas_v41,
                       (origens.TX_V41, origens.T09_V41, *origens.ORIGENS_V42[2:]), 117569),
    "V4.2": Referencia("V4.2", "v4.2", "excel/FSAE_EV_Suspension_Fase1_V4_2.xlsx", _entradas_v42,
                       origens.ORIGENS_V42, 117627, confere_edr=True),
}
# Troca de rótulos de Origem entre versões (a decisão 1 do EDR-062); a migração e a base dos conflitos usam o mesmo mapa.
MAPAS_ROTULO: dict[tuple[str, str], dict[str, str]] = {("V4.1", "V4.2"): origens.MAPA_DADO_V41_V42}


def referencia_da(ver: str) -> Referencia:
    """Referência cadastrada da versão `ver` ("V4.1"); VivaIlegivel se a versão não tem."""
    if ver not in REFERENCIAS:
        raise VivaIlegivel(f"versão {ver} sem referência cadastrada (conhecidas: {sorted(REFERENCIAS)})")
    return REFERENCIAS[ver]


def _tag_existe(tag: str) -> bool:
    return subprocess.run(["git", "rev-parse", "-q", "--verify", f"refs/tags/{tag}"], cwd=RAIZ,
                          capture_output=True).returncode == 0


def arquivo_de_referencia(ver: str, pasta: Path) -> Path:
    """Arquivo de referência da versão `ver`: o blob da tag (extraído em `pasta`, binário); sem a tag, o de `excel/`.
    VivaIlegivel se nenhum dos dois existe."""
    ref = referencia_da(ver)
    if _tag_existe(ref.tag):
        destino = Path(pasta) / f"referencia_{ref.tag}.xlsx"
        r = subprocess.run(["git", "show", f"{ref.tag}:{ref.caminho_excel}"], cwd=RAIZ, capture_output=True)
        if r.returncode != 0:
            raise VivaIlegivel(f"a tag {ref.tag} não tem {ref.caminho_excel}")
        destino.write_bytes(r.stdout)
        return destino
    local = RAIZ / ref.caminho_excel
    if not local.exists():
        raise VivaIlegivel(f"sem referência para {ver}: a tag {ref.tag} não existe e {ref.caminho_excel} também não")
    return local


def edrs_citados(xlsx: Path) -> list[tuple[str, str]]:
    """EDRs citados nas fichas de meta: (célula, EDR) de 03_TARGETS N5:N11 (campo "EDR do compromisso": o texto todo, que
    pode estar no formato errado, exceto vazio e "TBD") e F5:F11 (Decision: cada `EDR-ddd…` do texto). Vazio se a aba não
    tem essas colunas (V4.1)."""
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    try:
        if "03_TARGETS" not in wb.sheetnames:
            return []
        ws = wb["03_TARGETS"]
        out: list[tuple[str, str]] = []
        for r in range(5, 12):
            n = ws[f"N{r}"].value
            if n is not None and str(n).strip() not in ("", "TBD"):
                out.append((f"03_TARGETS!N{r}", str(n).strip()))
            f = ws[f"F{r}"].value
            if isinstance(f, str):
                out += [(f"03_TARGETS!F{r}", e) for e in RE_EDR.findall(f)]
        return out
    finally:
        wb.close()


def edrs_aprovados(edr_md: Path) -> dict[str, str]:
    """ID -> status da tabela de índice do registro de decisões (`docs/Engineering_Decision_Register.md`)."""
    return dict(RE_STATUS_EDR.findall(Path(edr_md).read_text(encoding="utf-8")))
