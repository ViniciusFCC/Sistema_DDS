"""Abas do equilíbrio no limite da V4.1 (incremento 4): 16_EQUILIBRIO, 17_EQ_PAR e 18_EQ_PAR_FINA.

A análise de par do Milliken (§7.1, p. 283-286) em fórmulas do Excel, com as regras do EDR-056 (entrada do pneu
limitada à faixa do .tir, com aviso; Fz <= 0 = descolamento; NEUTRO só dentro da resolução numérica medida por
refinamento de grade). O oráculo é `equilibrio.analise_par`; cada família de fórmula (catálogo F1-F12, F16 e a pior
origem da F17, plano docs/superpowers/plans/2026-10-21-v4-1-inc4-equilibrio-na-planilha.md) reproduz a função citada:

- 16_EQUILIBRIO: entradas novas (toe, raio, CG lateral, aero, tração, sentido e chaves h2, h3, h7, h8, h9), a
  verificação das entradas (chave ligada com entrada vazia = N/A com o motivo "<entrada> a confirmar"), o caso efetivo
  (curva à esquerda equivalente; com "Direita", o espelho de `equilibrio.espelhar`), a bisseção do Ackermann e o
  resultado no limite (saturações, resolução, veredito, C_N, limitações, tempo de volta, pior origem);
- 17_EQ_PAR: grade pedida, a_y = k·0,05 g (k = 0…40) e 25 valores de α;
- 18_EQ_PAR_FINA: grade refinada, a_y = k·0,025 g (k = 0…80) e 49 valores de α (`analise_par`: passo/2, 2·n − 1).

Decisões (plano, "Decisões do plano"; registrar no EDR-060):
- um sentido de curva por vez (decisão 2); diferencial aberto declarado em texto (decisão 3);
- pico de Fx (h9) pela forma fechada |Dx| + SVx, válida quando o κ do pico cai em [max(0, KPUMIN), KPUMAX] (decisão 4).
  O κ do pico não é calculado: como φ(u) = (1 − Ex)·u + Ex·atan(u) é crescente (Ex <= 1), κ_pico está na faixa se e
  só se φ(u(κ_min)) <= tan(π/(2·Cx)) <= φ(u(KPUMAX)), com u(κ) = Bx·(κ + SHx). Fora disso: N/A com o motivo;
- o Python liga o cos δ_ref junto com o raio (EDR-057, item 5: h3 = cos δ_ref com δ_ref = L/R), sem chave própria.
  Para a planilha nunca mostrar um resultado que o oráculo não reproduz, h2, h8 ou h9 ligada com a h3 desligada dá
  N/A com o motivo (em `VERIFICACOES`);
- nomes: toda entrada nova tem ent/org/fte_16_<slug>; toda fórmula das linhas rotuladas da 16 tem res_16_<slug>; nas
  abas de grade, os resultados do resumo têm res_17_/res_18_<slug> (as células por linha da grade e a bisseção do
  Ackermann são cálculo auxiliar, sem nome). Regras de nome do EDR-058 (`nomes_v4_0`).

SI nas células de cálculo; ° e g só nas entradas de toe e nas células de exibição. Convenções do CLAUDE.md (ISO 8855;
camber + = topo para fora; toe + = toe-out, EDR-002). Nenhuma função salva a planilha viva nem a de referência:
`copia_v41` trabalha numa pasta temporária a partir do arquivo da tag v4.0.
"""
from __future__ import annotations

import copy
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter as COL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from atualizar_planilha import ABAS_ORIGEM, COR_ORIGEM, NOME_LISTA, ORIGENS_LISTA, CONF, celulas_de_entrada
from caso_planilha import ENTRADAS_CASO
from nomes_v4_0 import NOME_VALIDO, Nome, ler_tabela, parece_referencia, slug
from raizes import Grafo, linha

RAIZ = Path(__file__).resolve().parent.parent
TAG_V40 = "v4.0"
CAMINHO_V40 = "excel/FSAE_EV_Suspension_Fase1_V4_0.xlsx"
TABELA_V40 = RAIZ / "data" / "nomes_V4_0.csv"

ABA16, ABA17, ABA18 = "16_EQUILIBRIO", "17_EQ_PAR", "18_EQ_PAR_FINA"
CANTOS = ("FL", "FR", "RL", "RR")
AY_MAX_G, PASSO_G, N_ALFA = 2.0, 0.05, 25        # grade padrão de equilibrio.analise_par
# aba -> (passo [g] como texto da fórmula, número de linhas, n_alfa, grade pedida?)
GRADES = {ABA17: ("0.05", 41, N_ALFA, True), ABA18: ("0.025", 81, 2 * N_ALFA - 1, False)}
EPS = "0.000001"          # tire_model.EPS (o mesmo valor da 10_TIRE B176; constante da fórmula, não raiz)
MENOS_INF = "-1E+300"     # −inf da regra do Python (ay_valido None) nas comparações do veredito
PRIMEIRA_LINHA = 24       # primeira linha de dados nas abas de grade
N_BISSECAO = 200          # iterações da bisseção do Ackermann (as mesmas de equilibrio._esterco_ackermann)

# ---------------------------------------------------------------------------------------------------------------------
# Nomes da V4.0 usados nas fórmulas (data/nomes_V4_0.csv)
# ---------------------------------------------------------------------------------------------------------------------
V40 = {**ENTRADAS_CASO,
       "ms": "res_05_massa_suspensa_total", "hs": "res_05_altura_do_cg_suspenso_h_s",
       "msF": "res_05_massa_suspensa_eixo_dianteiro_mb_l_m_uf", "msR": "res_05_massa_suspensa_eixo_traseiro_ma_l_m_ur",
       "KphiF": "res_06_roll_stiffness_dianteira_k_phif_k_wt2_2", "KphiR": "res_06_roll_stiffness_traseira_k_phir_k_wt2_2_k",
       "zs": "res_09_altura_do_roll_axis_no_cg_suspenso_z_s_e",
       "stFL": "res_09_fz_estatico_fl", "stFR": "res_09_fz_estatico_fr",
       "stRL": "res_09_fz_estatico_rl", "stRR": "res_09_fz_estatico_rr",
       "dgF": "ent_07_camber_gain_dianteiro_dcamber_dz", "dgR": "ent_07_camber_gain_traseiro_dcamber_dz",
       "pct": "res_07_ackermann",
       "gsFL": "ent_10_camber_gama_projeto_topo_para_fora_b", "gsFR": "ent_10_camber_gama_projeto_topo_para_fora_c",
       "gsRL": "ent_10_camber_gama_projeto_topo_para_fora_d", "gsRR": "ent_10_camber_gama_projeto_topo_para_fora_e",
       "s_a": "res_10_s_alfa_alfa_arquivo_s_alfa_alfa_iso", "s_f": "res_10_s_fy_fy_iso_s_fy_fy_arquivo",
       "s_g": "res_10_s_gama_gama_arquivo_s_gama_gama_projeto", "eixos": "ent_10_sistema_de_eixos_do_arquivo_declarado_sa"}
# Resultados da V4.0 que as fórmulas usam além das 23 entradas do caso (o "base válida" confere que são números).
V40_DERIVADOS = ("ms", "hs", "msF", "msR", "KphiF", "KphiR", "zs", "stFL", "stFR", "stRL", "stRR")
# Coeficientes do .tir na 10_TIRE (B86:B175 e B177); as fórmulas usam os nomes ent_10_<coeficiente>.
COEF_TIR = ("FNOMIN", "FZMIN", "FZMAX", "ALPMIN", "ALPMAX", "CAMMIN", "CAMMAX", "KPUMIN", "KPUMAX", "LFZO", "LCX",
            "LMUX", "LEX", "LKX", "LHX", "LVX", "LGAX", "LCY", "LMUY", "LEY", "LKY", "LHY", "LVY", "LGAY", "LKYG",
            "PCX1", "PDX1", "PDX2", "PDX3", "PEX1", "PEX2", "PEX3", "PEX4", "PKX1", "PKX2", "PKX3", "PHX1", "PHX2",
            "PVX1", "PVX2", "PCY1", "PDY1", "PDY2", "PDY3", "PEY1", "PEY2", "PEY3", "PEY4", "PKY1", "PKY2", "PKY3",
            "PHY1", "PHY2", "PHY3", "PVY1", "PVY2", "PVY3", "PVY4")


def nome_coef(k: str) -> str:
    """Nome definido da V4.0 do coeficiente `k` do .tir (10_TIRE)."""
    return f"ent_10_{k.lower()}"


# ---------------------------------------------------------------------------------------------------------------------
# Entradas novas (16_EQUILIBRIO, colunas B = valor, C = unidade, D = referência, E = Origem, F = Fonte)
# ---------------------------------------------------------------------------------------------------------------------
_FONTE_CHAVE = ("Chave de hipótese (spec 7.4): 1 = ligada, 0 = desligada. Escolha de modelo, não dado do carro: "
                "não entra na pior origem. ")


@dataclass(frozen=True)
class Entrada:
    chave: str
    rotulo: str
    padrao: object
    unidade: str
    ref: str
    fonte: str
    validacao: str            # lista_sentido | chave | numero | positivo | nao_negativo


ENTRADAS16 = (
    Entrada("sentido", "Sentido da curva", "Esquerda", "-", "h7 (spec 7.4); decisão 2 do plano",
            "Escolha da análise: Esquerda (padrão) ou Direita. Com Direita, as fórmulas analisam a curva à esquerda do "
            "carro espelhado (equilibrio.espelhar): FL da análise = FR real (roda interna). Troque o sentido para ver o "
            "outro.", "lista_sentido"),
    *(Entrada(f"toe{c}", f"Toe {c}", 0, "°", "EDR-002; EDR-036",
              "Sem dado (EDR-036): 0 até medir. Toe + = toe-out (EDR-002), em graus.", "numero") for c in CANTOS),
    Entrada("R", "Raio do skidpad R", None, "m", "spec 7.1 e 7.2",
            "A confirmar (regulamento). Vazio = sem dado. Exigido por h2, h3, h8 e h9 e pelo tempo de volta.", "positivo"),
    Entrada("cg", "CG lateral (y_CG)", 0, "m", "h7 (spec 7.4)",
            "Sem dado (EDR-036): 0 até medir. + = para a esquerda (ISO 8855). Só entra com a h7 ligada.", "numero"),
    Entrada("cza", "CzA", None, "m²", "h8 (spec 7.4)", "A confirmar. Só entra com a h8 ligada.", "numero"),
    Entrada("rho", "ρ (densidade do ar)", None, "kg/m³", "h8, h9 (spec 7.4)",
            "A confirmar. Entra com a h8 ou a h9 ligada.", "positivo"),
    Entrada("cda", "CdA", None, "m²", "h9 (spec 7.4); EDR-057", "A confirmar. Só entra com a h9 ligada.", "nao_negativo"),
    Entrada("crr", "Crr (resistência ao rolamento)", None, "-", "h9 (spec 7.4); EDR-057",
            "A confirmar. Só entra com a h9 ligada.", "nao_negativo"),
    Entrada("h2", "Chave h2: Ackermann (SUS-414)", 0, "0/1", "spec 7.4, h2",
            _FONTE_CHAVE + "Usa o % Ackermann da 07_KINEMATICS (B75); exige R e a h3.", "chave"),
    Entrada("h3", "Chave h3: cos δ_ref na frente", 0, "0/1", "spec 7.1, item 3; 7.4, h3",
            _FONTE_CHAVE + "δ_ref = L/R; exige R.", "chave"),
    Entrada("h7", "Chave h7: CG lateral", 0, "0/1", "spec 7.4, h7", _FONTE_CHAVE + "Usa o CG lateral acima.", "chave"),
    Entrada("h8", "Chave h8: aero com V = √(a_y·R)", 0, "0/1", "spec 7.4, h8; EDR-013",
            _FONTE_CHAVE + "Exige R, CzA, ρ e a h3; substitui a downforce fixa da 08_VEHICLE_STATE (EDR-013).", "chave"),
    Entrada("h9", "Chave h9: tração na traseira", 0, "0/1", "spec 7.4, h9; EDR-057",
            _FONTE_CHAVE + "Exige R, CdA, Crr, ρ e a h3; elipse de atrito normalizada (EDR-057), diferencial aberto (h10).",
            "chave"),
)
LINHA_ENTRADA0 = 4
LINHA_ENTRADA = {e.chave: LINHA_ENTRADA0 + i for i, e in enumerate(ENTRADAS16)}


def _nome_ent(e: Entrada, papel: str = "ent") -> str:
    return f"{papel}_16_{slug(e.rotulo)}"


ENTRADAS_V41: list[tuple[str, str, str, object, str, str]] = [
    (_nome_ent(e), ABA16, f"B{LINHA_ENTRADA[e.chave]}", e.padrao, e.unidade, e.rotulo) for e in ENTRADAS16]
NOME_ENT = {e.chave: _nome_ent(e) for e in ENTRADAS16}


def celulas_de_entrada_v41() -> set[tuple[str, str]]:
    """Células de entrada da V4.1: as da V4.0 (`atualizar_planilha.celulas_de_entrada()`) e, para cada entrada nova da
    16_EQUILIBRIO, valor (B), Origem (E) e Fonte (F)."""
    return celulas_de_entrada() | {(ABA16, f"{c}{LINHA_ENTRADA[e.chave]}") for e in ENTRADAS16 for c in "BEF"}


# ---------------------------------------------------------------------------------------------------------------------
# Verificação das entradas (a mesma lista na planilha e em caso_planilha.caso_da_planilha_v41)
# ---------------------------------------------------------------------------------------------------------------------
@dataclass(frozen=True)
class Verificacao:
    """Uma verificação da cadeia de "Situação das entradas" (a primeira que falha é o motivo do N/A).

    tipo: lista_sentido | chave | num | positivo | nao_negativo | eixos | coef | derivados | h3 | den | ack.
    nome: chave de V40 ou de ENTRADAS16 (None nas verificações sem um nome só). texto: base da mensagem; "num" dá
    "<texto> a confirmar" (vazio) ou "<texto> não é número". cond: quando a verificação vale (CONDICOES; None = sempre).
    """
    tipo: str
    nome: str | None
    texto: str
    cond: str | None = None


MSG_H3 = ("h3 (cos δ_ref) desligada com h2, h8 ou h9 ligada: no Python o raio R liga o cos δ_ref (EDR-057, item 5); "
          "ligue a h3")
MSG_DEN = "instabilidade de rolagem: m_s·g·h2 >= K_φF + K_φR (SUS-103, Milliken p. 682)"
MSG_ACK = "L/R fora do alcance da geometria de Ackermann (SUS-414); revise o raio"
MSG_EIXOS = "sistema de eixos do pneu (10_TIRE B11) inválido"
MSG_COEF = "coeficiente do pneu (10_TIRE B86:B177) não é número"
MSG_DERIVADOS = "resultado da V4.0 sem valor numérico (05_VEHICLE, 06_SUSPENSION ou 09_LOADS)"

# Condição -> (fórmula do Excel com {chaves}, função do Python sobre o dicionário de valores por chave)
def _um(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and v == 1


CONDICOES = {
    "raio": ("OR({h2}=1,{h3}=1,{h8}=1,{h9}=1)", lambda x: any(_um(x[h]) for h in ("h2", "h3", "h8", "h9"))),
    "h2": ("{h2}=1", lambda x: _um(x["h2"])),
    "h7": ("{h7}=1", lambda x: _um(x["h7"])),
    "h8": ("{h8}=1", lambda x: _um(x["h8"])),
    "h9": ("{h9}=1", lambda x: _um(x["h9"])),
    "h8h9": ("OR({h8}=1,{h9}=1)", lambda x: _um(x["h8"]) or _um(x["h9"])),
    "nao_h8": ("{h8}<>1", lambda x: not _um(x["h8"])),
    "h3": ("AND(OR({h2}=1,{h8}=1,{h9}=1),{h3}<>1)", lambda x: (_um(x["h2"]) or _um(x["h8"]) or _um(x["h9"]))
           and not _um(x["h3"])),
}


def _rotulos_v40() -> dict[str, str]:
    """nome -> "rótulo (aba célula)" das entradas da V4.0, para as mensagens."""
    return {n.nome: f"{n.rotulo} ({n.aba} {n.celula})" for n in ler_tabela(TABELA_V40)}


def _verificacoes() -> tuple[Verificacao, ...]:
    rot = _rotulos_v40()
    v = [Verificacao("lista_sentido", "sentido", "sentido da curva inválido (use Esquerda ou Direita)")]
    v += [Verificacao("chave", h, f"chave {h} deve ser 0 ou 1") for h in ("h2", "h3", "h7", "h8", "h9")]
    v += [Verificacao("num", f"toe{c}", f"toe {c}") for c in CANTOS]
    v += [Verificacao("eixos", "eixos", MSG_EIXOS), Verificacao("coef", None, MSG_COEF)]
    v += [Verificacao("num", k, rot[V40[k]]) for k in ("dgF", "dgR", "gsFL", "gsFR", "gsRL", "gsRR")]
    v += [Verificacao("num", k, rot[V40[k]], "nao_h8" if k == "DF" else None) for k in ENTRADAS_CASO]
    v += [Verificacao("derivados", None, MSG_DERIVADOS), Verificacao("h3", None, MSG_H3, "h3")]
    v += [Verificacao("num", "R", "raio do skidpad R", "raio"), Verificacao("positivo", "R", "raio do skidpad R", "raio")]
    v += [Verificacao("num", "pct", rot[V40["pct"]], "h2"), Verificacao("num", "cg", "CG lateral", "h7")]
    v += [Verificacao("num", "cza", "CzA", "h8"), Verificacao("num", "rho", "ρ", "h8h9"),
          Verificacao("positivo", "rho", "ρ", "h8h9")]
    v += [Verificacao("num", "cda", "CdA", "h9"), Verificacao("nao_negativo", "cda", "CdA", "h9"),
          Verificacao("num", "crr", "Crr", "h9"), Verificacao("nao_negativo", "crr", "Crr", "h9")]
    v += [Verificacao("den", None, MSG_DEN), Verificacao("ack", None, MSG_ACK)]
    return tuple(v)


VERIFICACOES = _verificacoes()


def _eh_numero(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and v == v and abs(v) != float("inf")


def motivo_das_entradas(x: dict) -> str | None:
    """Primeira verificação que falha (a mesma ordem e o mesmo texto da planilha), sem as duas que dependem de conta
    (instabilidade de rolagem e alcance do Ackermann, que caso_planilha confere pelo Python). `x`: chave -> valor lido
    da cópia recalculada, para as chaves de ENTRADAS16 e de V40, mais "coef" (lista dos valores dos coeficientes) e
    "eixos" (texto da 10_TIRE B11). None = entradas válidas."""
    from tire_model import AXIS_TRANSFORMS
    for vf in VERIFICACOES:
        if vf.cond is not None and not CONDICOES[vf.cond][1](x):
            continue
        val = x.get(vf.nome) if vf.nome else None
        if vf.tipo == "lista_sentido" and val not in ("Esquerda", "Direita"):
            return vf.texto
        if vf.tipo == "chave" and not (_eh_numero(val) and val in (0, 1)):
            return vf.texto
        if vf.tipo == "num":
            if val is None:
                return f"{vf.texto} a confirmar"
            if not _eh_numero(val):
                return f"{vf.texto} não é número"
        if vf.tipo == "positivo" and _eh_numero(val) and not val > 0:
            return f"{vf.texto} deve ser > 0"
        if vf.tipo == "nao_negativo" and _eh_numero(val) and not val >= 0:
            return f"{vf.texto} deve ser >= 0"
        if vf.tipo == "eixos" and val not in AXIS_TRANSFORMS:
            return vf.texto
        if vf.tipo == "coef" and not all(_eh_numero(c) for c in x["coef"]):
            return vf.texto
        if vf.tipo == "h3":
            return vf.texto        # a condição já é a falha
    return None


# ---------------------------------------------------------------------------------------------------------------------
# Textos de número (vírgula decimal, como equilibrio._num)
# ---------------------------------------------------------------------------------------------------------------------
def _f(x: str, casas: int) -> str:
    """format(x, f'.{casas}f') com vírgula: FIXED sem separador de milhar; o SUBSTITUTE cobre os dois locais."""
    return f'SUBSTITUTE(FIXED({x},{casas},TRUE),".",",")'


def _g(x: str) -> str:
    """format(x, 'g') com vírgula, para a_y/G = k·passo (no máximo 3 casas): ROUND a 6 casas e texto do número."""
    return f'SUBSTITUTE(ROUND({x},6)&"",".",",")'


# ---------------------------------------------------------------------------------------------------------------------
# Layout das abas de grade (17, 18)
# ---------------------------------------------------------------------------------------------------------------------
TERMOS = ("dfz", "gy", "Ky", "SHy", "muy", "Dy", "By", "SVy", "Eyn", "Ey0", "Eyp")
TRACAO = ("Dx", "SVx", "SHx", "Ex", "Bx", "fxmax", "fat")
LIMITES = ("FZMIN", "FZMAX", "CAMMIN", "CAMMAX")
RODAS_EIXO = {"F": ("FL", "FR"), "R": ("RL", "RR")}


def _colunas(n_alfa: int, pedida: bool) -> dict:
    nomes = ["k", "ay_g", "ay", "phi", "DF", "dFF", "dFR"]
    nomes += [f"{g}_{c}" for g in ("Fz", "gam", "Fzu", "gf") for c in CANTOS]
    nomes += [(t, c) for c in CANTOS for t in TERMOS]
    nomes += ["fxr"] + [(t, c) for c in ("RL", "RR") for t in TRACAO]
    for e in "FR":
        nomes += [f"pre_{e}"] + [("fy", e, i) for i in range(n_alfa)]
        nomes += [f"{n}_{e}" for n in ("cap", "idx", "mot", "req", "folga", "pref", "cand")]
    if pedida:
        nomes += [("x", c, lim) for c in CANTOS for lim in LIMITES]
    return {n: COL(j + 1) for j, n in enumerate(nomes)}


def _cabecalho(n) -> str:
    return n if isinstance(n, str) else "_".join(str(x) for x in n)


RESUMO_GRADE = (   # chave, rótulo, unidade (linhas 4..13 das abas de grade)
    ("S_F", "Saturação do eixo dianteiro S_f", "m/s²"),
    ("S_R", "Saturação do eixo traseiro S_r", "m/s²"),
    ("POS_F", "Linha k da saturação dianteira", "-"),
    ("POS_R", "Linha k da saturação traseira", "-"),
    ("INV_F", "Primeira linha com capacidade dianteira inválida (posição)", "-"),
    ("INV_R", "Primeira linha com capacidade traseira inválida (posição)", "-"),
    ("V_F", "Último a_y válido do eixo dianteiro", "m/s²"),
    ("V_R", "Último a_y válido do eixo traseiro", "m/s²"),
    ("MOT_F", "Motivo de o eixo dianteiro não saturar", "-"),
    ("MOT_R", "Motivo de o eixo traseiro não saturar", "-"),
)
LINHA_RESUMO0 = 4
LINHA_ALFA = 16        # i (16), α_ref (17), α no referencial do arquivo por roda (18-21)


def _nome_res(aba: str, rotulo: str) -> str:
    return f"res_{aba[:2]}_{slug(rotulo)}"


def _nomes_resumo(aba: str) -> dict[str, tuple[str, str]]:
    """chave -> (nome, célula) do resumo de uma aba de grade."""
    return {k: (_nome_res(aba, rot), f"B{LINHA_RESUMO0 + i}") for i, (k, rot, _) in enumerate(RESUMO_GRADE)}


# ---------------------------------------------------------------------------------------------------------------------
# Layout da 16_EQUILIBRIO
# ---------------------------------------------------------------------------------------------------------------------
@dataclass
class _Linha16:
    tipo: str                 # secao | texto | res
    chave: str = ""
    rotulo: str = ""
    formula: str = ""         # com {chaves}
    unidade: str = ""
    ref: str = ""
    familia: str = ""
    pior: bool = False        # coluna E com a pior origem viva (resultado do limite)
    obs: str = ""


def _faixa(aba: str, col: str, n_linhas: int) -> str:
    return f"'{aba}'!${col}${PRIMEIRA_LINHA}:${col}${PRIMEIRA_LINHA + n_linhas - 1}"


def _cond_verif(vf: Verificacao) -> str:
    """Fórmula de uma verificação: "-" quando passa, a mensagem quando falha (sem o "=")."""
    n = f"{{{vf.nome}}}" if vf.nome else ""
    if vf.tipo == "lista_sentido":
        corpo = f'IF(OR(EXACT({n},"Esquerda"),EXACT({n},"Direita")),"-","{vf.texto}")'   # EXACT: com caixa, como o Python
    elif vf.tipo == "chave":
        corpo = f'IF(ISNUMBER({n}),IF(OR({n}=0,{n}=1),"-","{vf.texto}"),"{vf.texto}")'
    elif vf.tipo == "num":
        corpo = f'IF(ISBLANK({n}),"{vf.texto} a confirmar",IF(ISNUMBER({n}),"-","{vf.texto} não é número"))'
    elif vf.tipo == "positivo":
        corpo = f'IF(ISNUMBER({n}),IF({n}>0,"-","{vf.texto} deve ser > 0"),"-")'
    elif vf.tipo == "nao_negativo":
        corpo = f'IF(ISNUMBER({n}),IF({n}>=0,"-","{vf.texto} deve ser >= 0"),"-")'
    elif vf.tipo == "eixos":
        corpo = f'IF(AND(ISNUMBER({{s_a}}),ISNUMBER({{s_f}}),ISNUMBER({{s_g}}),OR(EXACT({{eixos}},"SAE_CALSPAN"),EXACT({{eixos}},"ISO_TYDEX"))),"-","{vf.texto}")'
    elif vf.tipo == "coef":
        corpo = (f"IF(AND(COUNT('10_TIRE'!$B$86:$B$175)=90,ISNUMBER('10_TIRE'!$B$177)),\"-\",\"{vf.texto}\")")
    elif vf.tipo == "derivados":
        corpo = f'IF({{base_ok}}=1,"-","{vf.texto}")'
    elif vf.tipo == "h3":
        corpo = f'"{vf.texto}"'
    elif vf.tipo == "den":
        corpo = f'IF(ISNUMBER({{den}}),IF({{den}}<=0,"{vf.texto}","-"),"-")'
    elif vf.tipo == "ack":
        corpo = f'IF({{ack_alcance}}=0,"{vf.texto}","-")'
    else:
        raise ValueError(vf.tipo)
    if vf.cond is not None:
        corpo = f'IF({CONDICOES[vf.cond][0]},{corpo},"-")'
    return corpo


def _layout16() -> list[_Linha16]:
    L: list[_Linha16] = []
    sec = lambda t: L.append(_Linha16("secao", rotulo=t))
    txt = lambda t: L.append(_Linha16("texto", rotulo=t))

    def res(chave, rotulo, formula, unidade="-", familia="", ref="", pior=False, obs=""):
        L.append(_Linha16("res", chave, rotulo, formula, unidade, ref, familia, pior, obs))

    ok = "{ok}=1"
    # Verificação das entradas (F1)
    sec("Verificação das entradas (chave ligada com entrada vazia = N/A com o motivo; nunca 0 em silêncio)")
    for i, vf in enumerate(VERIFICACOES):
        res(f"vf{i}", f"Verificação: {vf.texto}", "=" + _cond_verif(vf), familia="F1", ref="Review Focus 1 e 5")
    # A 1ª falha acumulada fica na coluna G de cada verificação (sem um IF de 56 níveis; o Excel aceita 64).
    res("status", "Situação das entradas", f'=IF({{G_vf{len(VERIFICACOES) - 1}}}="-","OK",{{G_vf{len(VERIFICACOES) - 1}}})',
        familia="F1", obs="\"OK\" ou o motivo do N/A (a 1ª verificação que falha; coluna G = 1ª falha até a linha).")
    # Caso efetivo (F1)
    sec("Caso efetivo: curva à esquerda equivalente (com \"Direita\", o espelho de equilibrio.espelhar)")
    res("ok", "Caso válido (situação = OK)", '=IF({status}="OK",1,0)', "0/1", "F1")
    res("dir", "Curva à direita (análise no espelho)", '=IF(EXACT({sentido},"Direita"),1,0)', "0/1", "F1")
    res("raio", "Raio ativo (h2, h3, h8 ou h9 ligada)", "=IF(" + CONDICOES["raio"][0] + ",1,0)", "0/1", "F1")
    isn = ",".join(f"ISNUMBER({{{k}}})" for k in V40_DERIVADOS)
    res("base_ok", "Resultados da V4.0 usados com valor numérico", f"=IF(AND({isn}),1,0)", "0/1", "F1")
    res("dref", "δ_ref = L/R", f'=IF(AND({ok},{{raio}}=1),{{L}}/{{R}},"-")', "rad", "F3", "spec 7.1, item 3")
    res("cosf", "cos δ_ref (frente; 1 sem raio)", f'=IF({ok},IF({{raio}}=1,COS({{L}}/{{R}}),1),"N/A")', "-", "F3",
        "equilibrio._cos_dianteira")
    res("tanf", "tan δ_ref (arrasto do pneu dianteiro, h9)",
        f'=IF({ok},IF(AND({{raio}}=1,{{h9}}=1),TAN({{L}}/{{R}}),0),"N/A")', "-", "F9", "equilibrio.fx_exigido_traseira")
    res("Cy", "Cy = PCY1·LCY", f'=IF({ok},{{P_PCY1}}*{{P_LCY}},"N/A")', "-", "F7", "tire_model.pure_fy0")
    res("fz0", "F'z0 = LFZO·FNOMIN", f'=IF({ok},{{P_LFZO}}*{{P_FNOMIN}},"N/A")', "N", "F7", "tire_model.pure_fy0")
    # Ganho de camber (F2)
    sec("Ganho de camber em rolagem ∂γ/∂φ (F2; Gillespie eq. 6-50; SUS-404 da 07_KINEMATICS, autorização 7 do plano)")
    for e, t in (("F", "tF"), ("R", "tR")):
        nome = "dianteiro" if e == "F" else "traseiro"
        res(f"dg{e}r", f"dγ/dz {nome} (07_KINEMATICS, °/mm → rad/m)", f'=IF({ok},RADIANS({{dg{e}}})*1000,"N/A")',
            "rad/m", "F2", "kinematics_model.camber_gradient_from_body")
        l, r = RODAS_EIXO[e]
        res(f"g{l}", f"∂γ/∂φ {l} (real) = −(t/2)·dγ/dz − 1", f'=IF({ok},-0.5*{{{t}}}*{{dg{e}r}}-1,"N/A")', "rad/rad",
            "F2", "kinematics_model.camber_gradient_from_body")
        res(f"g{r}", f"∂γ/∂φ {r} (real) = (t/2)·dγ/dz + 1", f'=IF({ok},0.5*{{{t}}}*{{dg{e}r}}+1,"N/A")', "rad/rad",
            "F2", "kinematics_model.camber_gradient_from_body")
    # Entradas efetivas (F1)
    sec("Entradas efetivas da análise (F1): com \"Direita\", troca L↔R e inverte o CG lateral e o ganho de camber")
    troca = {"FL": "FR", "FR": "FL", "RL": "RR", "RR": "RL"}
    for c in CANTOS:
        res(f"toe_{c}", f"Toe {c} da análise", f'=IF({ok},RADIANS(IF({{dir}}=1,{{toe{troca[c]}}},{{toe{c}}})),"N/A")',
            "rad", "F1", "equilibrio.espelhar")
    for c in CANTOS:
        res(f"gs_{c}", f"Camber estático {c} da análise (10_TIRE B24:E24)",
            f'=IF({ok},RADIANS(IF({{dir}}=1,{{gs{troca[c]}}},{{gs{c}}})),"N/A")', "rad", "F1", "equilibrio.espelhar")
    for c in CANTOS:
        res(f"ge_{c}", f"∂γ/∂φ {c} da análise", f'=IF({ok},IF({{dir}}=1,-{{g{troca[c]}}},{{g{c}}}),"N/A")', "rad/rad",
            "F1", "equilibrio.espelhar")
    res("cge", "CG lateral da análise", f'=IF({ok},IF({{h7}}=1,IF({{dir}}=1,-{{cg}},{{cg}}),0),"N/A")', "m", "F1",
        "equilibrio.espelhar")
    # Deriva por roda (F3) e grade de α (F8)
    sec("Deriva por roda Δα (F3): toe (EDR-002) + Ackermann (h2, SUS-414, δ_ref = L/R) e grade de α (F8)")
    res("pcte", "% Ackermann usado (h2)", f'=IF({ok},IF({{h2}}=1,{{pct}},"-"),"N/A")', "%", "F3", "SUS-414")
    interna = "{do}+{pct}/100*(ATAN(1/(1/TAN({do})-{tF}/{L}))-{do})"
    res("do", "δ da roda externa (FR da análise)",
        f'=IF(AND({ok},{{h2}}=1),IF({{pct}}=0,{{L}}/{{R}},0.5*({{bis_lo}}+{{bis_hi}})),"-")', "rad", "F3",
        "equilibrio._esterco_ackermann")
    res("di", "δ da roda interna (FL da análise)", f'=IF(AND({ok},{{h2}}=1),IF({{pct}}=0,{{L}}/{{R}},{interna}),"-")',
        "rad", "F3", "equilibrio._esterco_ackermann")
    for c, s in (("FL", "-"), ("FR", ""), ("RL", "-"), ("RR", "")):
        if c in ("FL", "FR"):
            d = "{di}" if c == "FL" else "{do}"
            f_ = f'=IF({ok},IF({{h2}}=1,{s}{{toe_{c}}}+(-({d}-{{dref}})),{s}{{toe_{c}}}),"N/A")'
        else:
            f_ = f'=IF({ok},{s}{{toe_{c}}},"N/A")'
        res(f"da_{c}", f"Δα {c} (toe + Ackermann)", f_, "rad", "F3", "equilibrio.deriva_total")
    absd = ",".join(f"ABS({{da_{c}}})" for c in CANTOS)
    res("amp", "Amplitude da grade de α A = min(|ALPMIN|, ALPMAX) − max|Δα|",
        f'=IF({ok},MIN(ABS({{P_ALPMIN}}),{{P_ALPMAX}})-MAX({absd}),"N/A")', "rad", "F8", "equilibrio.grade_alfa")
    # Cargas (F4)
    sec("Cargas por roda (F4): SUS-101 + SUS-103 (três massas, termo gravitacional) + aero + CG lateral (h7)")
    res("h2r", "h2 = h_s − z_s (braço de rolagem)", '=IF({base_ok}=1,{hs}-{zs},"N/A")', "m", "F4",
        "suspension_model.lateral_dfz")
    res("den", "K_φF + K_φR − m_s·g·h2", '=IF({base_ok}=1,{KphiF}+{KphiR}-{ms}*9.81*{h2r},"N/A")', "N·m/rad", "F4",
        "suspension_model.lateral_dfz")
    res("cF", "m_s,F·z_F + m_uF·h_uF", '=IF({base_ok}=1,{msF}*{zF}+{muF}*{huF},"N/A")', "kg·m", "F4",
        "suspension_model.lateral_dfz")
    res("cR", "m_s,R·z_R + m_uR·h_uR", '=IF({base_ok}=1,{msR}*{zR}+{muR}*{huR},"N/A")', "kg·m", "F4",
        "suspension_model.lateral_dfz")
    res("dFyF", "ΔFz do CG lateral por roda, frente (W_F·y/t_F)", f'=IF({ok},({{stFL}}+{{stFR}})*{{cge}}/{{tF}},"N/A")',
        "N", "F4", "suspension_model.static_fz_corner_cg_lateral")
    res("dFyR", "ΔFz do CG lateral por roda, traseira (W_R·y/t_R)", f'=IF({ok},({{stRL}}+{{stRR}})*{{cge}}/{{tR}},"N/A")',
        "N", "F4", "suspension_model.static_fz_corner_cg_lateral")
    res("aero", "Aerodinâmica usada", '=IF({h8}=1,"0,5·ρ·CzA·a_y·R por linha (h8; substitui a downforce fixa da '
        '08_VEHICLE_STATE, EDR-013)","downforce fixa da 08_VEHICLE_STATE (EDR-013)")', "-", "F4", "equilibrio.downforce")
    txt("Diferencial: aberto (h10), o único tipo que a planilha implementa (spec 7.4); Fx igual nas duas rodas traseiras.")
    # Resultado no limite (F11, F12, F16)
    sec("Resultado no limite (grade pedida 0,05 g e 25 α; resolução pela grade refinada 0,025 g e 49 α; EDR-056)")
    for e in "FR":
        nome = "dianteiro" if e == "F" else "traseiro"
        res(f"S_{e}", f"Saturação do eixo {nome} S_{e.lower()} (grade pedida)", f"={{a17_S_{e}}}", "m/s²", "F16",
            "equilibrio.analise_par", True)
    for e in "FR":
        res(f"S2_{e}", f"S_{e.lower()} na grade refinada", f"={{a18_S_{e}}}", "m/s²", "F16", "equilibrio._resolucao", True)
    for e in "FR":
        res(f"V_{e}", f"Último a_y válido V_{e.lower()} (grade pedida)", f"={{a17_V_{e}}}", "m/s²", "F16",
            "equilibrio._saturacao", True)
    for e in "FR":
        res(f"M_{e}", f"Motivo de o eixo {'dianteiro' if e == 'F' else 'traseiro'} não saturar", f"={{a17_MOT_{e}}}",
            "-", "F16", "equilibrio._saturacao", True)
    for e in "FR":
        res(f"dif_{e}", f"|S_{e.lower()} − S_{e.lower()} refinada| (ou o passo, se só uma grade satura)",
            f'=IF(AND(ISNUMBER({{S_{e}}}),ISNUMBER({{S2_{e}}})),ABS({{S_{e}}}-{{S2_{e}}}),'
            f'IF(OR(ISNUMBER({{S_{e}}}),ISNUMBER({{S2_{e}}})),0.05*9.81,"-"))', "m/s²", "F12", "equilibrio._resolucao", True)
    res("resol", "Resolução numérica das saturações (EDR-056, item 2)",
        '=IF(AND({dif_F}="-",{dif_R}="-"),"N/A",MAX(IF(ISNUMBER({dif_F}),{dif_F},0),IF(ISNUMBER({dif_R}),{dif_R},0),'
        '1E-9*9.81))', "m/s²", "F12", "equilibrio._resolucao", True)
    res("push1", "Só a frente satura e a traseira segue válida até S_f + resolução",
        '=IF(ISNUMBER({S_F}),IF({V_R}>={S_F}+{resol},1,0),0)', "0/1", "F12", "equilibrio.analise_par", True)
    res("loose1", "Só a traseira satura e a frente segue válida até S_r + resolução",
        '=IF(ISNUMBER({S_R}),IF({V_F}>={S_R}+{resol},1,0),0)', "0/1", "F12", "equilibrio.analise_par", True)
    res("ver", "Veredito no limite (PUSH, LOOSE, NEUTRO ou N/A)",
        '=IF({ok}<>1,"N/A",IF(AND(ISNUMBER({S_F}),ISNUMBER({S_R})),IF(ABS({S_R}-{S_F})<={resol},"NEUTRO",'
        'IF({S_R}-{S_F}>0,"PUSH","LOOSE")),IF({push1}=1,"PUSH",IF({loose1}=1,"LOOSE","N/A"))))', "-", "F12",
        "equilibrio.analise_par (EDR-056, itens 2 e 3; Milliken p. 320)", True)
    res("mot", "Motivo do N/A",
        '=IF({ver}<>"N/A","-",IF({ok}<>1,{status},IF(ISNUMBER({S_F}),{M_R},IF(ISNUMBER({S_R}),{M_F},'
        'IF({V_F}<={V_R},{M_F},{M_R})))))', "-", "F12", "equilibrio.analise_par", True)
    res("aylim", "a_y limite (a menor saturação)",
        '=IF({ok}<>1,"N/A",IF(AND(ISNUMBER({S_F}),ISNUMBER({S_R})),MIN({S_F},{S_R}),IF(ISNUMBER({S_F}),{S_F},'
        'IF(ISNUMBER({S_R}),{S_R},"N/A"))))', "m/s²", "F16", "equilibrio.analise_par", True)
    n17 = GRADES[ABA17][1]
    ay17 = _faixa(ABA17, "{c17_ay}", n17)
    res("kcn", "Linha k do eixo saturado (intervalo do C_N)",
        '=IF({ver}="PUSH",{a17_POS_F},IF({ver}="LOOSE",{a17_POS_R},"-"))', "-", "F12", "equilibrio.analise_par", True)
    fR, fF = _faixa(ABA17, "{c17_folga_R}", n17), _faixa(ABA17, "{c17_folga_F}", n17)
    res("ma", "Folga do outro eixo na linha k",
        f'=IF({{ver}}="PUSH",INDEX({fR},{{kcn}}),IF({{ver}}="LOOSE",INDEX({fF},{{kcn}}),"-"))', "N", "F12",
        "equilibrio.analise_par", True)
    res("mb", "Folga do outro eixo na linha k + 1",
        f'=IF({{ver}}="PUSH",INDEX({fR},{{kcn}}+1),IF({{ver}}="LOOSE",INDEX({fF},{{kcn}}+1),"-"))', "N", "F12",
        "equilibrio.analise_par", True)
    res("la", "a_y da linha k", f'=IF(ISNUMBER({{kcn}}),INDEX({ay17},{{kcn}}),"-")', "m/s²", "F12",
        "equilibrio.analise_par", True)
    res("lb", "a_y da linha k + 1", f'=IF(ISNUMBER({{kcn}}),INDEX({ay17},{{kcn}}+1),"-")', "m/s²", "F12",
        "equilibrio.analise_par", True)
    res("fint", "Folga do outro eixo em a_y limite (interpolada)",
        '=IF(ISNUMBER({kcn}),{ma}+({mb}-{ma})*({aylim}-{la})/({lb}-{la}),"-")', "N", "F12", "equilibrio.analise_par", True)
    res("cn", "C_N residual no limite (Milliken §8.4, p. 311-312)",
        '=IF({ver}="PUSH",{fint}*{b}/({m}*9.81*{L}),IF({ver}="LOOSE",-{fint}*{a}/({m}*9.81*{L}),"N/A"))', "-", "F12",
        "equilibrio.analise_par (spec 7.2)", True)
    for k, rot in (("S_F", "S_f"), ("S_R", "S_r"), ("aylim", "a_y limite"), ("resol", "Resolução")):
        res(f"{k}_g", f"{rot} em g", f'=IF(ISNUMBER({{{k}}}),{{{k}}}/9.81,"N/A")', "g", "F16", "exibição", True)
    # Limitações (EDR-056, item 1): uma por roda e limite, o pior caso das linhas usadas no veredito
    for c in CANTOS:
        for lim in LIMITES:
            x = _faixa(ABA17, f"{{c17_x_{c}_{lim}}}", n17)
            p = f"MATCH(MAX({x}),{x},0)"
            if lim.startswith("FZ"):
                fz = _faixa(ABA17, f"{{c17_Fz_{c}}}", n17)
                t = (f'"Fz_{c} = "&{_f(f"INDEX({fz},{p})", 1)}&" N → "&{_f("{P_" + lim + "}", 1)}&" N ({lim})"')
            else:
                gm = _faixa(ABA17, f"{{c17_gam_{c}}}", n17)
                t = (f'"γ_{c} = "&{_f(f"DEGREES(INDEX({gm},{p}))", 2)}&"° → "'
                     f'&{_f("DEGREES({P_" + lim + "}/{s_g})", 2)}&"° ({lim})"')
            res(f"lim_{c}_{lim}", f"Limitação {c} {lim}", f'=IF({{ok}}<>1,"-",IF(MAX({x})<0,"-",{t}))', "-", "F16",
                "equilibrio._limitacoes", True)
    junta = "&".join(f'IF({{lim_{c}_{lim}}}="-","",{{lim_{c}_{lim}}}&"; ")' for c in CANTOS for lim in LIMITES)
    res("lims", "Limitações da entrada do pneu usadas no veredito (EDR-056, item 1)",
        f'=IF(LEN({junta})=0,"nenhuma",LEFT({junta},LEN({junta})-2))', "-", "F16", "equilibrio._limitacoes", True)
    res("tvolta", "Tempo de volta no skidpad t = 2π·√(R/a_y limite)",
        '=IF(ISNUMBER({aylim}),IF(ISBLANK({R}),"N/A: raio do skidpad R a confirmar",IF(ISNUMBER({R}),IF({R}>0,'
        '2*PI()*SQRT({R}/{aylim}),"N/A: raio do skidpad R deve ser > 0"),"N/A: raio do skidpad R não é número")),"N/A")',
        "s", "F16", "spec 7.2", True)
    res("pior", "Pior origem das entradas (ao vivo)", "={PIOR}", "-", "F17", "raizes.Grafo; EDR-035, fase 2", True)
    return L


_LAYOUT16 = _layout16()
LINHA_LAYOUT0 = LINHA_ENTRADA0 + len(ENTRADAS16) + 1     # primeira linha depois das entradas


def _linhas16() -> dict[str, int]:
    """chave -> linha de cada resultado da 16_EQUILIBRIO."""
    out = {}
    for i, ln in enumerate(_LAYOUT16):
        if ln.tipo == "res":
            out[ln.chave] = LINHA_LAYOUT0 + i
    return out


LINHAS16 = _linhas16()
COL_BIS = "H"   # bisseção do Ackermann: colunas H (j), I (lo), J (hi), K (mid), L (teste)


def _nomes_v41() -> list[Nome]:
    out = []
    for e in ENTRADAS16:
        r = LINHA_ENTRADA[e.chave]
        for papel, col in (("ent", "B"), ("org", "E"), ("fte", "F")):
            out.append(Nome(_nome_ent(e, papel), ABA16, f"{col}{r}", papel, "", e.rotulo))
    base = [(ln.chave, _nome_res(ABA16, ln.rotulo), f"B{LINHAS16[ln.chave]}", ln.rotulo)
            for ln in _LAYOUT16 if ln.tipo == "res"]
    for aba in (ABA17, ABA18):
        for k, rot, _ in RESUMO_GRADE:
            nome, cel = _nomes_resumo(aba)[k]
            base.append((f"{aba}:{k}", nome, cel, rot))
    cont: dict[str, int] = {}
    for _, nome, _, _ in base:
        cont[nome] = cont.get(nome, 0) + 1
    for chave, nome, cel, rot in base:
        aba = chave.split(":")[0] if ":" in chave else ABA16
        if cont[nome] > 1:
            nome = f"{nome}_{cel.lower()}"
        out.append(Nome(nome, aba, cel, "res", "", rot))
    vistos = set()
    for n in out:
        if n.nome in vistos or not NOME_VALIDO.fullmatch(n.nome) or parece_referencia(n.nome):
            raise ValueError(f"nome inválido ou repetido na V4.1: {n.nome}")
        vistos.add(n.nome)
    return out


NOMES_V41: list[Nome] = _nomes_v41()
_NOME_POR_CELULA = {(n.aba, n.celula): n.nome for n in NOMES_V41}
NOME16 = {ln.chave: _NOME_POR_CELULA[(ABA16, f"B{LINHAS16[ln.chave]}")] for ln in _LAYOUT16 if ln.tipo == "res"}
NOME_RESUMO = {(aba, k): _NOME_POR_CELULA[(aba, cel)] for aba in (ABA17, ABA18)
               for k, (_, cel) in _nomes_resumo(aba).items()}


# ---------------------------------------------------------------------------------------------------------------------
# Pior origem viva (F17, só a parte usada pelo limite): raízes de raizes.Grafo mais as entradas da 16
# ---------------------------------------------------------------------------------------------------------------------
# Nós da V4.0 cujas raízes entram na pior origem, com a condição (None = sempre). O 08 B8 (downforce fixa) só conta
# com a h8 desligada; o % Ackermann só com a h2 ligada.
def _nos_pior_origem() -> list[tuple[str, str | None]]:
    nos = [(V40[k], "nao_h8" if k == "DF" else None) for k in ENTRADAS_CASO]
    nos += [(V40[k], None) for k in ("dgF", "dgR", "gsFL", "gsFR", "gsRL", "gsRR")]
    nos += [(V40["pct"], "h2")]
    return nos


ENT16_PIOR = (("toeFL", None), ("toeFR", None), ("toeRL", None), ("toeRR", None), ("R", "raio"), ("cg", "h7"),
              ("cza", "h8"), ("rho", "h8h9"), ("cda", "h9"), ("crr", "h9"))


def _origem_da_raiz(aba: str, cel: str) -> tuple[str, str]:
    """Célula de Origem de uma raiz: coluna E nas abas de entrada; na 10_TIRE, I23-I25 para os casos de carga e I84
    para o resto (coeficientes e metadados do pneu: a origem do pneu, como caso_planilha._origem_pneu)."""
    r = linha(cel)
    if aba in ABAS_ORIGEM:
        return aba, f"E{r}"
    if aba == "10_TIRE":
        return aba, f"I{r}" if r in (23, 24, 25) else "I84"
    raise ValueError(f"raiz {aba}!{cel} sem coluna de Origem conhecida")


def _destino(wb, nome: str) -> tuple[str, str]:
    dn = wb.defined_names.get(nome)
    if dn is None:
        raise KeyError(f"nome definido {nome} não existe na planilha")
    dest = list(dn.destinations)
    if len(dest) != 1 or ":" in dest[0][1]:
        raise ValueError(f"o nome {nome} não aponta para uma única célula ({dn.attr_text})")
    return dest[0][0], dest[0][1].replace("$", "")


def termos_pior_origem(wb) -> list[tuple[str, str, frozenset | None]]:
    """(aba, célula de Origem, condições) da pior origem do limite; condições = None (sempre) ou o conjunto de chaves
    de CONDICOES (vale se alguma valer). Raízes vazias não entram (raizes.Grafo.raizes_preenchidas), como no 14_DASHBOARD;
    a origem do pneu (10_TIRE I84) entra sempre; as entradas da 16 entram pela coluna E, com a condição de uso."""
    g = Grafo(wb)
    termos: dict[tuple[str, str], frozenset | None] = {("10_TIRE", "I84"): None}

    def add(cel, cond):
        if cel in termos and termos[cel] is None:
            return
        if cond is None:
            termos[cel] = None
        else:
            termos[cel] = (termos.get(cel) or frozenset()) | {cond}

    for nome, cond in _nos_pior_origem():
        for raiz in g.raizes_preenchidas(_destino(wb, nome)):
            add(_origem_da_raiz(*raiz), cond)
    for k, cond in ENT16_PIOR:
        add((ABA16, f"E{LINHA_ENTRADA[k]}"), cond)
    return [(aba, cel, termos[(aba, cel)]) for aba, cel in sorted(termos, key=lambda t: (t[0], linha(t[1]), t[1]))]


def _formula_pior(termos) -> str:
    """INDEX(Origens, MAX(...)) como em atualizar_planilha._formula_pior_origem, com IF(cond, termo, 1) nas condicionais
    (1 = "● Dado do TX", neutro no MAX). Sem o "=" e com {chaves} das condições."""
    n = len(ORIGENS_LISTA)
    partes = []
    for aba, cel, conds in termos:
        col, r = cel[0], cel[1:]
        t = f"IFERROR(MATCH('{aba}'!${col}${r},{NOME_LISTA},0),{n})"
        if conds:
            cs = [CONDICOES[c][0] for c in sorted(conds)]
            t = f"IF({cs[0] if len(cs) == 1 else 'OR(' + ','.join(cs) + ')'},{t},1)"
        partes.append(t)
    return f"INDEX({NOME_LISTA},MAX({','.join(partes)}))"


def pior_origem_de_valores(termos, origem_de, chaves: dict) -> str:
    """Pior origem pela mesma regra da fórmula: origem_de(aba, cel) -> texto da célula; chaves: chave -> valor."""
    pior = 1
    for aba, cel, conds in termos:
        if conds and not any(CONDICOES[c][1](chaves) for c in conds):
            continue
        v = origem_de(aba, cel)
        pior = max(pior, ORIGENS_LISTA.index(v) + 1 if v in ORIGENS_LISTA else len(ORIGENS_LISTA))
    return ORIGENS_LISTA[pior - 1]


# ---------------------------------------------------------------------------------------------------------------------
# Montagem
# ---------------------------------------------------------------------------------------------------------------------
class _Subst(dict):
    def __missing__(self, k):
        raise KeyError(f"chave de fórmula desconhecida: {{{k}}}")


@dataclass
class _Ctx:
    wb: object
    registro: list = field(default_factory=list)
    mapa: dict = field(default_factory=dict)

    def put(self, ws, cel: str, familia: str, formula: str):
        texto = formula.format_map(self.mapa) if "{" in formula else formula
        ws[cel] = texto
        self.registro.append({"aba": ws.title, "cel": cel, "familia": familia, "formula": texto})


def _mapa_base(cols17: dict, termos) -> dict:
    m = _Subst()
    partes = [dict(V40), {f"P_{k}": nome_coef(k) for k in COEF_TIR}, {e.chave: NOME_ENT[e.chave] for e in ENTRADAS16},
              dict(NOME16), {f"a{aba[:2]}_{k}": nome for (aba, k), nome in NOME_RESUMO.items()},
              {f"c17_{_cabecalho(n)}": letra for n, letra in cols17.items()},
              {f"G_{k}": f"'{ABA16}'!$G${r}" for k, r in LINHAS16.items() if k.startswith("vf")}]
    for p in partes:
        repetidas = set(m) & set(p)
        if repetidas:
            raise ValueError(f"chaves de fórmula repetidas: {sorted(repetidas)}")
        m.update(p)
    # estado (lo, hi) depois das N_BISSECAO iterações: linha j = N_BISSECAO do bloco (a 1ª linha é j = 0)
    m["bis_lo"] = f"'{ABA16}'!$I${LINHA_ENTRADA0 + N_BISSECAO}"
    m["bis_hi"] = f"'{ABA16}'!$J${LINHA_ENTRADA0 + N_BISSECAO}"
    m["ack_alcance"] = f"'{ABA16}'!$L${LINHA_ENTRADA0 - 1}"
    pior = _formula_pior(termos)
    m["PIOR"] = pior.format_map(m)
    return m


def _estilos(wb) -> dict:
    v = wb["05_VEHICLE"]
    return {"titulo": v["A1"]._style, "secao": v["A3"]._style, "rotulo": v["A4"]._style, "ent": v["B14"]._style,
            "un": v["C14"]._style, "id": v["D14"]._style, "org": v["E14"]._style, "fte": v["F14"]._style,
            "calc": v["B23"]._style, "cab": v["E2"]._style}


def _estilo(c, est, chave):
    c._style = copy.copy(est[chave])


def _aba16(ctx: _Ctx, est: dict) -> None:
    ws = ctx.wb.create_sheet(ABA16)
    ws["A1"] = "16_EQUILIBRIO — equilíbrio no limite: análise de par (Milliken §7.1, p. 283-286) com as regras do EDR-056"
    _estilo(ws["A1"], est, "titulo")
    ws["A2"] = ("Resultados DEMONSTRATIVOS enquanto a pior origem não for \"● Dado do TX\" (D16). Oráculo: "
                "python/equilibrio.py (analise_par). Grades nas abas 17_EQ_PAR e 18_EQ_PAR_FINA.")
    for cel, t in (("E2", "Origem (entradas) / pior origem (resultados)"), ("F2", "Fonte / observação")):
        ws[cel] = t
        _estilo(ws[cel], est, "cab")
    ws["A3"] = "Entradas do equilíbrio no limite (V4.1): toe, raio, CG lateral, aero, tração e chaves de hipótese"
    _estilo(ws["A3"], est, "secao")
    for e in ENTRADAS16:
        r = LINHA_ENTRADA[e.chave]
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = e.rotulo, e.padrao, e.unidade, e.ref
        ws[f"E{r}"], ws[f"F{r}"] = CONF, e.fonte
        for col, k in (("A", "rotulo"), ("B", "ent"), ("C", "un"), ("D", "id"), ("E", "org"), ("F", "fte")):
            _estilo(ws[f"{col}{r}"], est, k)
    for i, ln in enumerate(_LAYOUT16):
        r = LINHA_LAYOUT0 + i
        if ln.tipo == "secao":
            ws[f"A{r}"] = ln.rotulo
            _estilo(ws[f"A{r}"], est, "secao")
            continue
        if ln.tipo == "texto":
            ws[f"A{r}"] = ln.rotulo
            continue
        ws[f"A{r}"], ws[f"C{r}"], ws[f"D{r}"] = ln.rotulo, ln.unidade, f"{ln.familia}; {ln.ref}".strip("; ")
        _estilo(ws[f"A{r}"], est, "rotulo")
        ctx.put(ws, f"B{r}", ln.familia, ln.formula)
        _estilo(ws[f"B{r}"], est, "calc")
        if ln.chave.startswith("vf"):   # 1ª falha acumulada das verificações (coluna G)
            i = int(ln.chave[2:])
            ant = f'IF(G{r - 1}<>"-",G{r - 1},B{r})' if i > 0 else f"=B{r}"
            ctx.put(ws, f"G{r}", "F1", ant if i == 0 else "=" + ant)
        if ln.pior:
            ctx.put(ws, f"E{r}", "F17", "={PIOR}")
        if ln.obs:
            ws[f"F{r}"] = ln.obs
    _bissecao(ctx, ws)
    _validacoes16(ws)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 62, 22
    ws.column_dimensions["D"].width, ws.column_dimensions["E"].width, ws.column_dimensions["F"].width = 30, 30, 80
    for r in LINHA_ENTRADA.values():
        ws[f"F{r}"].alignment = Alignment(wrap_text=True, vertical="top")


def _bissecao(ctx: _Ctx, ws) -> None:
    """Bisseção de equilibrio._esterco_ackermann em colunas fixas (H:L), N_BISSECAO iterações a partir de
    lo = 1e-12 e hi = atan(L/t)·(1 − 1e-9); a raiz de (δo + δi(δo))/2 = δ_ref é única (o lado esquerdo cresce com δo)."""
    r0 = LINHA_ENTRADA0
    pre = "AND({ok_pre},{h2}=1,{pct}<>0)"
    ok_pre = ("AND({raio}=1,ISNUMBER({R}),{R}>0,ISNUMBER({pct}),ISNUMBER({L}),ISNUMBER({tF}),{base_ok}=1)")
    pre = pre.replace("{ok_pre}", ok_pre)
    interna = lambda d: f"{d}+{{pct}}/100*(ATAN(1/(1/TAN({d})-{{tF}}/{{L}}))-{d})"
    ws[f"H{r0 - 2}"] = "Bisseção do Ackermann (SUS-414; equilibrio._esterco_ackermann): j, lo, hi, meio, (meio + δi)/2 < δ_ref"
    ws[f"H{r0 - 1}"] = "alcance"
    hi0 = "ATAN({L}/{tF})*(1-1E-9)"
    ctx.put(ws, f"L{r0 - 1}", "F3", f'=IF({pre},IF(({hi0}+{interna(hi0)})/2<{{L}}/{{R}},0,1),1)')
    for j in range(N_BISSECAO + 1):
        r = r0 + j
        ws[f"H{r}"] = j
        if j == 0:
            ctx.put(ws, f"I{r}", "F3", f'=IF({pre},1E-12,"-")')
            ctx.put(ws, f"J{r}", "F3", f'=IF({pre},{hi0},"-")')
        else:
            ctx.put(ws, f"I{r}", "F3", f'=IF({pre},IF(L{r - 1},K{r - 1},I{r - 1}),"-")')
            ctx.put(ws, f"J{r}", "F3", f'=IF({pre},IF(L{r - 1},J{r - 1},K{r - 1}),"-")')
        if j < N_BISSECAO:
            ctx.put(ws, f"K{r}", "F3", f'=IF({pre},0.5*(I{r}+J{r}),"-")')
            ctx.put(ws, f"L{r}", "F3", f'=IF({pre},(K{r}+{interna(f"K{r}")})/2<{{L}}/{{R}},FALSE)')


def _validacoes16(ws) -> None:
    def dv(**kw):
        d = DataValidation(allow_blank=True, showErrorMessage=True, errorTitle="Entrada inválida", **kw)
        ws.add_data_validation(d)
        return d

    por_tipo = {
        "lista_sentido": dv(type="list", formula1='"Esquerda,Direita"', error="Escolha Esquerda ou Direita."),
        "chave": dv(type="list", formula1='"0,1"', error="Chave de hipótese: 0 (desligada) ou 1 (ligada)."),
        "numero": dv(type="decimal", operator="between", formula1="-1E+300", formula2="1E+300",
                     error="Digite um número (sem unidade)."),
        "positivo": dv(type="decimal", operator="greaterThan", formula1="0", error="Digite um número > 0 (sem unidade)."),
        "nao_negativo": dv(type="decimal", operator="greaterThanOrEqual", formula1="0",
                           error="Digite um número >= 0 (sem unidade)."),
    }
    org = dv(type="list", formula1=f"={NOME_LISTA}", error="Escolha uma das origens da lista (00_README, 'Como preencher').")
    for e in ENTRADAS16:
        r = LINHA_ENTRADA[e.chave]
        por_tipo[e.validacao].add(f"B{r}")
        org.add(f"E{r}")
    r0, r1 = LINHA_ENTRADA0, LINHA_ENTRADA0 + len(ENTRADAS16) - 1
    for simb, cor in COR_ORIGEM.items():
        ws.conditional_formatting.add(f"E{r0}:E{r1}", FormulaRule(formula=[f'LEFT($E{r0},1)="{simb}"'],
                                                                   fill=PatternFill("solid", start_color=cor, end_color=cor)))


def _aba_grade(ctx: _Ctx, est: dict, aba: str) -> None:
    passo, n_lin, n_alfa, pedida = GRADES[aba]
    ws = ctx.wb.create_sheet(aba)
    cl = _colunas(n_alfa, pedida)
    ws["A1"] = (f"{aba} — grade {'pedida' if pedida else 'refinada'} da análise de par: a_y = k·{passo.replace('.', ',')} g "
                f"(k = 0…{n_lin - 1}), {n_alfa} valores de α (equilibrio.analise_par)")
    _estilo(ws["A1"], est, "titulo")
    ws["A2"] = ("Cálculo auxiliar: uma linha por a_y. Fz e γ por roda (F4, F5), entrada limitada à faixa do .tir (F6, "
                "EDR-056), termos do PAC2002 (F7), Fy do eixo por α (F8), tração (F9), capacidade (F10), folga e saturação "
                "(F11). Com \"Direita\", os cantos são os do carro espelhado (FL = roda interna).")
    ws["A3"] = "Resumo da grade"
    _estilo(ws["A3"], est, "secao")
    rs = _nomes_resumo(aba)
    a = lambda n: f"{cl[n]}{PRIMEIRA_LINHA}:{cl[n]}{PRIMEIRA_LINHA + n_lin - 1}"
    a1 = lambda n: f"{cl[n]}{PRIMEIRA_LINHA + 1}:{cl[n]}{PRIMEIRA_LINHA + n_lin - 1}"   # k = 1…n
    ult = n_lin
    resumo = {}
    for e in "FR":
        cand, pref, ay, mot = a1(f"cand_{e}"), a(f"pref_{e}"), a("ay"), a(f"mot_{e}")
        S, POS, INV = (f"{{a{aba[:2]}_{x}_{e}}}" for x in ("S", "POS", "INV"))
        resumo[f"S_{e}"] = f'=IF(COUNT({cand})=0,"N/A",MIN({cand}))'
        resumo[f"POS_{e}"] = f'=IF(ISNUMBER({S}),MATCH({S},{cand},0),"-")'
        resumo[f"INV_{e}"] = f'=IF(INDEX({pref},{ult}),"-",MATCH(FALSE,{pref},0))'
        resumo[f"V_{e}"] = (f'=IF(ISNUMBER({S}),INDEX({ay},{POS}+1),IF(ISNUMBER({INV}),IF({INV}=1,{MENOS_INF},'
                            f'INDEX({ay},{INV}-1)),INDEX({ay},{ult})))')
        resumo[f"MOT_{e}"] = (f'=IF(ISNUMBER({S}),"-",IF(ISNUMBER({INV}),INDEX({mot},{INV})&" em a_y = "&'
                              f'{_g(f"INDEX({ay},{INV})/9.81")}&" g","não satura até "&{_g(f"INDEX({ay},{ult})/9.81")}&" g"))')
    for i, (k, rot, un) in enumerate(RESUMO_GRADE):
        r = LINHA_RESUMO0 + i
        ws[f"A{r}"], ws[f"C{r}"], ws[f"D{r}"] = rot, un, "F11; equilibrio._saturacao"
        _estilo(ws[f"A{r}"], est, "rotulo")
        ctx.put(ws, rs[k][1], "F11", resumo[k])
        _estilo(ws[rs[k][1]], est, "calc")
        ctx.put(ws, f"E{r}", "F17", "={PIOR}")
    # Grade de α (F8): α_ref_i = −A·i/(n − 1) e α no referencial do arquivo por roda, s_α·espelho·(α_ref + Δα)
    ra = LINHA_ALFA
    ws[f"A{ra - 1}"] = "Grade de α (F8; equilibrio.grade_alfa) e α no referencial do arquivo por roda (tire_model.wheel_forces)"
    _estilo(ws[f"A{ra - 1}"], est, "secao")
    ws[f"A{ra}"], ws[f"A{ra + 1}"] = "i", "α_ref [rad]"
    alfa_arq = {}
    for i in range(n_alfa):
        c = COL(2 + i)
        ws[f"{c}{ra}"] = i
        ctx.put(ws, f"{c}{ra + 1}", "F8", f'=IF(ISNUMBER({{amp}}),-{{amp}}*{i}/{n_alfa - 1},"N/A")')
        for j, roda in enumerate(CANTOS):
            rr = ra + 2 + j
            ws[f"A{rr}"] = f"α arquivo {roda} [rad]"
            x = f"{c}${ra + 1}+{{da_{roda}}}"
            corpo = f"{{s_a}}*(-({x}))" if roda.endswith("L") else f"{{s_a}}*({x})"
            ctx.put(ws, f"{c}{rr}", "F8", f'=IF(ISNUMBER({c}${ra + 1}),{corpo},"N/A")')
            alfa_arq[(roda, i)] = f"${c}${rr}"
    # Tabela por linha
    rc = PRIMEIRA_LINHA - 1
    for n, letra in cl.items():
        ws[f"{letra}{rc}"] = _cabecalho(n)
    for k in range(n_lin):
        _linha_grade(ctx, ws, cl, k, PRIMEIRA_LINHA + k, passo, n_alfa, pedida, alfa_arq)
    ws.freeze_panes = f"D{PRIMEIRA_LINHA}"
    ws.column_dimensions["A"].width = 58


def _mf(T: dict, x: str) -> str:
    """Fórmula mágica do PAC2002 (tire_model._mf + SVy) com o Ey do sinal de x = α_arquivo + SHy."""
    Ey = f"CHOOSE(SIGN({x})+2,{T['Eyn']},{T['Ey0']},{T['Eyp']})"
    return (f"{T['Dy']}*SIN({{Cy}}*ATAN({T['By']}*{x}-{Ey}*({T['By']}*{x}-ATAN({T['By']}*{x}))))+{T['SVy']}")


def _linha_grade(ctx, ws, cl, k, r, passo, n_alfa, pedida, alfa_arq):
    C = lambda n: f"{cl[n]}{r}"
    put = lambda n, fam, f: ctx.put(ws, C(n), fam, f)
    ws[C("k")] = k
    put("ay", "F4", f"={k}*{passo}*9.81")
    put("ay_g", "F4", f"={C('ay')}/9.81")
    ay, phi = C("ay"), C("phi")
    put("phi", "F4", f'=IF({{ok}}=1,{{ms}}*{ay}*{{h2r}}/{{den}},"N/A")')
    put("DF", "F4", f'=IF({{ok}}=1,IF({{h8}}=1,0.5*{{rho}}*{{cza}}*(ABS({ay})*{{R}}),{{DF}}),"N/A")')
    put("dFF", "F4", f'=IF(ISNUMBER({phi}),({{KphiF}}*{phi}+{ay}*{{cF}})/{{tF}},"N/A")')
    put("dFR", "F4", f'=IF(ISNUMBER({phi}),({{KphiR}}*{phi}+{ay}*{{cR}})/{{tR}},"N/A")')
    DF = C("DF")
    aero = {"F": f"{DF}*{{bal}}/2", "R": f"{DF}*(1-{{bal}})/2"}
    for c in CANTOS:
        e = c[0]
        dlat, dy = C(f"dF{e}"), f"{{dFy{e}}}"
        s1, s2 = ("-", "+") if c.endswith("L") else ("+", "-")
        put(f"Fz_{c}", "F4", f'=IF(ISNUMBER({phi}),{{st{c}}}{s1}{dlat}+{aero[e]}{s2}{dy},"N/A")')
        put(f"gam_{c}", "F5", f'=IF(ISNUMBER({phi}),{{gs_{c}}}+{{ge_{c}}}*{phi},"N/A")')
        fz, gm = C(f"Fz_{c}"), C(f"gam_{c}")
        put(f"Fzu_{c}", "F6", f'=IF(ISNUMBER({fz}),MIN(MAX({fz},{{P_FZMIN}}),{{P_FZMAX}}),"N/A")')
        put(f"gf_{c}", "F6", f'=IF(ISNUMBER({gm}),MIN(MAX({{s_g}}*{gm},{{P_CAMMIN}}),{{P_CAMMAX}}),"N/A")')
        # Termos do PAC2002 por roda (F7), com Fz e γ limitados (o pneu não extrapola, EDR-056)
        FZ, GF = C(f"Fzu_{c}"), C(f"gf_{c}")
        T = {t: C((t, c)) for t in TERMOS}
        P = lambda nome: f"{{P_{nome}}}"
        expr = {
            "dfz": f"({FZ}-{{fz0}})/{{fz0}}",
            "gy": f"{GF}*{P('LGAY')}",
            "Ky": (f"{P('PKY1')}*{{fz0}}*SIN(2*ATAN({FZ}/({P('PKY2')}*{{fz0}})))*(1-{P('PKY3')}*ABS({T['gy']}))"
                   f"*{P('LKY')}"),
            "SHy": f"({P('PHY1')}+{P('PHY2')}*{T['dfz']})*{P('LHY')}+{P('PHY3')}*{T['gy']}*{P('LKYG')}",
            "muy": f"({P('PDY1')}+{P('PDY2')}*{T['dfz']})*(1-{P('PDY3')}*{T['gy']}*{T['gy']})*{P('LMUY')}",
            "Dy": f"{T['muy']}*{FZ}",
            "By": f"{T['Ky']}/({{Cy}}*{T['Dy']}+{EPS})",
            "SVy": (f"{FZ}*(({P('PVY1')}+{P('PVY2')}*{T['dfz']})*{P('LVY')}"
                    f"+({P('PVY3')}+{P('PVY4')}*{T['dfz']})*{T['gy']}*{P('LKYG')})*{P('LMUY')}"),
        }
        for t, s in (("Eyn", -1), ("Ey0", 0), ("Eyp", 1)):
            expr[t] = (f"MIN(({P('PEY1')}+{P('PEY2')}*{T['dfz']})*(1-({P('PEY3')}+{P('PEY4')}*{T['gy']})*({s}))"
                       f"*{P('LEY')},1)")
        for t in TERMOS:
            put((t, c), "F7", f'=IF(ISNUMBER({FZ}),{expr[t]},"N/A")')
    # Tração na traseira (F9): Fx exigido, pico de Fx pela forma fechada e fator da elipse de atrito
    req_f = C("req_F")
    put("fxr", "F9", f'=IF(AND(ISNUMBER({phi}),{{h9}}=1),0.5*{{rho}}*{{cda}}*ABS({ay})*{{R}}+{{crr}}*{{m}}*9.81'
                     f'+ABS({req_f})*{{tanf}},0)')
    for c in ("RL", "RR"):
        FZ, GF, dfz = C(f"Fzu_{c}"), C(f"gf_{c}"), C(("dfz", c))
        X = {t: C((t, c)) for t in TRACAO}
        on = f"AND({{h9}}=1,ISNUMBER({FZ}))"
        P = lambda nome: f"{{P_{nome}}}"
        gx = f"({GF}*{P('LGAX')})"
        put(("Dx", c), "F9", f'=IF({on},({P("PDX1")}+{P("PDX2")}*{dfz})*(1-{P("PDX3")}*{gx}*{gx})*{P("LMUX")}*{FZ},"-")')
        put(("SVx", c), "F9", f'=IF({on},{FZ}*({P("PVX1")}+{P("PVX2")}*{dfz})*{P("LVX")}*{P("LMUX")},"-")')
        put(("SHx", c), "F9", f'=IF({on},({P("PHX1")}+{P("PHX2")}*{dfz})*{P("LHX")},"-")')
        put(("Ex", c), "F9", (f'=IF({on},MIN(({P("PEX1")}+{P("PEX2")}*{dfz}+{P("PEX3")}*{dfz}^2)'
                              f'*(1-{P("PEX4")}*SIGN({P("KPUMAX")}+{X["SHx"]}))*{P("LEX")},1),"-")'))
        put(("Bx", c), "F9", (f'=IF({on},({FZ}*({P("PKX1")}+{P("PKX2")}*{dfz})*EXP({P("PKX3")}*{dfz})*{P("LKX")})'
                              f'/({P("PCX1")}*{P("LCX")}*{X["Dx"]}+{EPS}),"-")'))
        cx = f"{P('PCX1')}*{P('LCX')}"
        alvo = f"TAN(PI()/(2*{cx}))"
        phi_u = lambda u: f"((1-{X['Ex']})*({u})+{X['Ex']}*ATAN({u}))"
        u_min = f"{X['Bx']}*(MAX(0,{P('KPUMIN')})+{X['SHx']})"
        u_max = f"{X['Bx']}*({P('KPUMAX')}+{X['SHx']})"
        put(("fxmax", c), "F9", (f'=IF({X["Dx"]}="-","-",IF({cx}>1,IF({X["Bx"]}>0,IF(AND({phi_u(u_min)}<={alvo},'
                                 f'{alvo}<={phi_u(u_max)}),ABS({X["Dx"]})+{X["SVx"]},"fora"),"fora"),"fora"))'))
        fxm = X["fxmax"]
        put(("fat", c), "F9", (f'=IF({{h9}}<>1,1,IF({fxm}="-","-",IF({fxm}="fora","fora",IF({C("fxr")}/2>{fxm},'
                               f'"excede",SQRT(1-({C("fxr")}/2/{fxm})^2)))))'))
    # Capacidade e saturação de cada eixo (F6, F8, F10, F11)
    for e, (cL, cR) in RODAS_EIXO.items():
        desc = lambda c: (f'"Fz_{c} = "&{_f(C(f"Fz_{c}"), 1)}&" N ≤ 0: descolamento, fora do domínio do modelo de '
                          f'transferência de carga (SUS-104)"')
        tracao = ""
        if e == "R":
            fl, fr = C(("fat", "RL")), C(("fat", "RR"))
            tracao = (f'IF({{h9}}=1,IF({fl}="fora","pico de Fx fora da forma fechada",IF({fl}="excede",'
                      f'"tração excede o atrito longitudinal em RL",IF({fr}="fora","pico de Fx fora da forma fechada",'
                      f'IF({fr}="excede","tração excede o atrito longitudinal em RR","-")))),"-")')
        resto = f'IF({{amp}}<=0,"toe consome a faixa de α",{tracao or '"-"'})'
        put(f"pre_{e}", "F6", (f'=IF(NOT(ISNUMBER({C("Fz_" + cL)})),"N/A",IF({C("Fz_" + cL)}<=0,{desc(cL)},'
                               f'IF({C("Fz_" + cR)}<=0,{desc(cR)},{resto})))'))
        TL = {t: C((t, cL)) for t in TERMOS}
        TR = {t: C((t, cR)) for t in TERMOS}
        for i in range(n_alfa):
            xL, xR = f"({alfa_arq[(cL, i)]}+{TL['SHy']})", f"({alfa_arq[(cR, i)]}+{TR['SHy']})"
            fyL, fyR = f"-({{s_f}}*({_mf(TL, xL)}))", f"{{s_f}}*({_mf(TR, xR)})"
            if e == "F":
                soma = f"({fyL}+{fyR})*{{cosf}}"
            else:
                soma = f"{C(('fat', cL))}*({fyL})+{C(('fat', cR))}*({fyR})"
            put(("fy", e, i), "F8", f'=IF({C("pre_" + e)}="-",{soma},"N/A")')
        faixa = f"{C(('fy', e, 0))}:{C(('fy', e, n_alfa - 1))}"
        PRE, CAP, IDX, MOT = C(f"pre_{e}"), C(f"cap_{e}"), C(f"idx_{e}"), C(f"mot_{e}")
        REQ, FOL, PREF, CAND = C(f"req_{e}"), C(f"folga_{e}"), C(f"pref_{e}"), C(f"cand_{e}")
        put(f"cap_{e}", "F10", f'=IF({PRE}="-",MAX({faixa}),"N/A")')
        put(f"idx_{e}", "F10", f'=IF({PRE}="-",MATCH({CAP},{faixa},0),"N/A")')
        put(f"mot_{e}", "F10", f'=IF({PRE}<>"-",{PRE},IF({IDX}={n_alfa},"pico na borda da grade de α","-"))')
        braco = "{b}" if e == "F" else "{a}"
        put(f"req_{e}", "F11", f'=IF(ISNUMBER({phi}),{{m}}*{ay}*{braco}/{{L}},"N/A")')
        put(f"folga_{e}", "F11", f'=IF({MOT}="-",{CAP}-{REQ},"N/A")')
        put(f"pref_{e}", "F11", f'={MOT}="-"' if k == 0 else f'=AND({cl["pref_" + e]}{r - 1},{MOT}="-")')
        if k == 0:
            ws[CAND] = "-"
        else:
            m0, m1, a0 = f"{cl['folga_' + e]}{r - 1}", FOL, f"{cl['ay']}{r - 1}"
            put(f"cand_{e}", "F11", f'=IF({PREF},IF(AND({m0}>0,{m1}<=0),{a0}+({ay}-{a0})*{m0}/({m0}-{m1}),"-"),"-")')
    if pedida:   # excesso de cada limitação nas linhas usadas no veredito (equilibrio._limitacoes_detalhe)
        incl = "TRUE" if k == 0 else f'OR(NOT(ISNUMBER({{aylim}})),{cl["ay"]}{r - 1}<{{aylim}})'
        for c in CANTOS:
            mot = C(f"mot_{c[0]}")
            fz, gm = C(f"Fz_{c}"), C(f"gam_{c}")
            casos = {"FZMIN": (f"{fz}<{{P_FZMIN}}", f"ABS({fz}-{{P_FZMIN}})"),
                     "FZMAX": (f"{fz}>{{P_FZMAX}}", f"ABS({fz}-{{P_FZMAX}})"),
                     "CAMMIN": (f"{{s_g}}*{gm}<{{P_CAMMIN}}", f"ABS({gm}-{{P_CAMMIN}}/{{s_g}})"),
                     "CAMMAX": (f"{{s_g}}*{gm}>{{P_CAMMAX}}", f"ABS({gm}-{{P_CAMMAX}}/{{s_g}})")}
            for lim, (cond, exc) in casos.items():
                put(("x", c, lim), "F6", f'=IF(AND({incl},{mot}="-"),IF({cond},{exc},-1),-1)')


def montar(wb) -> list[dict]:
    """Escreve as abas 16_EQUILIBRIO, 17_EQ_PAR e 18_EQ_PAR_FINA no workbook da V4.0 (aberto pelo openpyxl, com os
    nomes definidos da V4.0) e devolve o registro [{aba, cel, familia, formula}] de toda célula com fórmula escrita.
    Não acrescenta os nomes novos (`acrescentar_nomes`) nem salva."""
    for aba in (ABA16, ABA17, ABA18):
        if aba in wb.sheetnames:
            raise ValueError(f"a aba {aba} já existe")
    if wb.sheetnames[-1] != "15_CHANGELOG":
        raise ValueError(f"a última aba deveria ser 15_CHANGELOG; é {wb.sheetnames[-1]}")
    termos = termos_pior_origem(wb)
    ctx = _Ctx(wb)
    ctx.mapa = _mapa_base(_colunas(N_ALFA, True), termos)
    est = _estilos(wb)
    _aba16(ctx, est)
    _aba_grade(ctx, est, ABA17)
    _aba_grade(ctx, est, ABA18)
    return ctx.registro


def acrescentar_nomes(wb) -> None:
    """Nomes definidos da V4.1 (NOMES_V41) no escopo do workbook; recusa nome já existente."""
    existentes = set(wb.defined_names)
    for n in NOMES_V41:
        if n.nome in existentes:
            raise ValueError(f"nome já definido: {n.nome}")
        col = "".join(ch for ch in n.celula if ch.isalpha())
        wb.defined_names[n.nome] = DefinedName(n.nome, attr_text=f"'{n.aba}'!${col}${n.celula[len(col):]}")


_V40_BYTES: bytes | None = None


def extrair_v40(pasta: Path) -> Path:
    """Arquivo da tag v4.0 numa pasta (binário, `git show`); nunca a V4.0 da árvore nem a viva."""
    global _V40_BYTES
    if _V40_BYTES is None:
        _V40_BYTES = subprocess.run(["git", "show", f"{TAG_V40}:{CAMINHO_V40}"], cwd=RAIZ, capture_output=True,
                                    check=True).stdout
    destino = Path(pasta) / "V4_0_tag.xlsx"
    destino.write_bytes(_V40_BYTES)
    return destino


def copia_v41(tmp: Path, alteracoes: dict | None = None, nome: str = "V4_1_copia.xlsx") -> Path:
    """Helper de teste: V4.0 da tag v4.0 + `montar` + nomes da V4.1 + `alteracoes` (nome definido -> valor; None
    esvazia a célula), gravada em `tmp/nome`. Os valores de teste ficam só nessa cópia. Recusa caminho em excel/."""
    tmp = Path(tmp)
    saida = (tmp / nome).resolve()
    if (RAIZ / "excel").resolve() in saida.parents:
        raise ValueError(f"cópia de teste em excel/ recusada: {saida}")
    wb = openpyxl.load_workbook(extrair_v40(tmp))
    try:
        montar(wb)
        acrescentar_nomes(wb)
        for nome_def, valor in (alteracoes or {}).items():
            aba, cel = _destino(wb, nome_def)
            wb[aba][cel].value = valor
        wb.calculation.fullCalcOnLoad = True
        wb.save(saida)
    finally:
        wb.close()
    return saida
