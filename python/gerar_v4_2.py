"""Gera a planilha V4.2 (incremento 5: metas e coerência; projeto T11) a partir da V4.1 da tag v4.1.

Uso:
    uv run python -X utf8 python/gerar_v4_2.py [--v41 ARQ] [--saida build/FSAE_EV_Suspension_Fase1_V4_2.xlsx]
                                               [--registro docs/registro-de-versoes-V4_2-celulas.md]
                                               [--nomes-csv data/nomes_V4_2.csv]

Sem --v41, o arquivo da tag v4.1 é extraído (git show, binário) para uma pasta temporária: nunca se usa a V4.1 da
árvore nem a viva do Drive. A tabela dos nomes que a V4.2 acrescenta vem de data/nomes_V4_2.csv (a tabela completa é a
de data/nomes_V4_0.csv, mais data/nomes_V4_1.csv, mais essa). Nesta etapa do incremento ela só tem o cabeçalho: as
metas (03_TARGETS, 23_LISTAS, 14_DASHBOARD) entram na Task 3. Com --nomes-csv, o CSV é regravado só com o cabeçalho.

O que esta etapa muda em relação à V4.1 (EDR-062; plano docs/superpowers/plans/2026-10-23-v4-2-inc5-metas-e-coerencia.md,
decisão 1 e autorização 5 (a), (b), (d) e (e)):
- dado: toda célula de Origem (colunas E e H das abas de `ABAS_ORIGEM`, I da 10_TIRE, E da 16_EQUILIBRIO) com o valor
  exatamente igual a `● Dado do TX` ou `◐ Dado do T09` passa a `◐ Dado de outro carro (TX, T09)`
  (`origens.MAPA_DADO_V41_V42`). O T11 ainda não tem dado: ninguém vira `● Dado do T11` por mapa;
- lista Origens (00_README A43:A48): os dois primeiros rótulos viram os do T11 e de outro carro; B:D e E dessas duas
  linhas passam a `origens.SIGNIFICADO_V42`. A lista continua em `'00_README'!$A$43:$A$48`, com 6 rótulos nas mesmas
  posições (as fórmulas de pior origem usam o índice);
- regra: os textos que explicam a Origem e o literal `"● Dado do TX"` do `EXACT` das fórmulas 16_EQUILIBRIO B188 e B189
  (prefixo DEMONSTRATIVO do D16) passam a citar o T11 (`origens.MAPA_REGRA_V41_V42`). Nada mais muda nessas fórmulas;
- a emenda da h7 na 21_COBERTURA (texto da spec §7.4, linha 7);
- a versão na 00_README A3 e a linha 22 do 15_CHANGELOG.
O 04_DECISION_REGISTER e o 15_CHANGELOG ficam como estão (registro histórico). As 117569 fórmulas da V4.1 e os 1145
nomes ficam iguais, salvo as duas fórmulas acima. A saída nunca fica em excel/ nem por cima da origem: a recusa vem
antes de abrir qualquer arquivo.
"""
from __future__ import annotations

import argparse
import copy
import subprocess
import tempfile
from pathlib import Path

import openpyxl

import origens
from atualizar_planilha import ABAS_ORIGEM, escrever_registro
from nomes_v4_0 import escrever_tabela, ler_tabela

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
PASTA_EXCEL = RAIZ / "excel"
SAIDA_PADRAO = RAIZ / "build" / "FSAE_EV_Suspension_Fase1_V4_2.xlsx"
CSV_NOMES = RAIZ / "data" / "nomes_V4_2.csv"
TAG_V41 = "v4.1"
CAMINHO_V41 = "excel/FSAE_EV_Suspension_Fase1_V4_1.xlsx"
PLANO = "docs/superpowers/plans/2026-10-23-v4-2-inc5-metas-e-coerencia.md"
REGISTRO = "docs/registro-de-versoes-V4_2-celulas.md"

REF_DADO = ("EDR-062, decisão 1 (substituição de dado: TX e T09 viram \"outro carro\"; o T11 ainda não tem dado); "
            "plano, autorização 5 (a); teste: TC-WB42-001")
REF_LISTA = ("EDR-062, decisão 1 (lista Origens: 6 rótulos nas mesmas posições; mudam os dois primeiros e o seu "
             "significado); plano, autorização 5 (a); teste: TC-WB42-001")
REF_REGRA = ("EDR-062, decisão 1 (substituição de regra: o texto e o literal do D16 passam a falar do T11); "
             "plano, autorização 5 (a) e (b); teste: TC-WB42-001")
REF_H7 = "EDR-060 item 8 (emenda da h7 deferida do incremento 4, spec §7.4 linha 7); plano, autorização 5 (d); teste: TC-WB42-001"
REF_VERSAO = "EDR-062; plano, autorização 5 (e) (padrão da V4.0 e da V4.1); teste: TC-WB42-001"

# colunas de Origem por aba: E e a oculta H (guarda a Origem anterior, para o aviso de troca) nas abas de entrada;
# I na 10_TIRE; E na 16_EQUILIBRIO
COLUNAS_ORIGEM: dict[str, tuple[str, ...]] = {**{a: ("E", "H") for a in ABAS_ORIGEM}, "10_TIRE": ("I",),
                                              "16_EQUILIBRIO": ("E",)}
ABA_LISTA, LINHA_LISTA = "00_README", 43      # A43:A48 = lista Origens

# textos explicativos (valores) que citam o rótulo do TX: substituição de regra. (aba, célula)
TEXTOS_REGRA = (
    ("00_README", "A49"), ("14_DASHBOARD", "A40"), ("16_EQUILIBRIO", "A2"), ("20_EQ_TROCA", "A2"), ("22_CADEIA", "A2"),
)
# as duas fórmulas do D16: só o literal muda
FORMULAS_D16 = (("16_EQUILIBRIO", "B188"), ("16_EQUILIBRIO", "B189"))
# abas que guardam o rótulo antigo de propósito (registro histórico)
ABAS_HISTORICAS = ("04_DECISION_REGISTER", "15_CHANGELOG")
ANTIGOS = (origens.TX_V41, origens.T09_V41, "Dado do TX", "Dado do T09")

# 21_COBERTURA D12 (coluna "Planilha" da hipótese 7): spec §7.4 linha 7. A F12 é a fórmula "Chave agora" e não muda.
H7 = ("21_COBERTURA", "D12")
H7_ANTES = "Sim (entrada, 0 até medir)"
H7_NOVO = ("Sim (entrada, 0 até medir; um sentido por vez, pela entrada \"Sentido da curva\": o outro se vê trocando "
           "o sentido; EDR-060, 2026-10-22)")

TEXTO_A3 = ("Versão do workbook: V4.2 (incremento 5: metas e coerência; projeto T11 (Origens e D16)) | Status: DRAFT — "
            "cross-check, ainda NÃO REFERENCE-validated")
LINHA_CHANGELOG = 22   # a V4.1 usou a 21


def _changelog() -> list[tuple[str, str]]:
    lin = LINHA_CHANGELOG
    return [
        (f"A{lin}", "2026-10"),
        (f"B{lin}", "V4.2, metas e coerencia (incremento 5), etapa do projeto T11: lista Origens com o T11 e \"outro carro\" "
                    "(TX, T09), prefixo DEMONSTRATIVO do D16 pelo T11, emenda da h7 na 21_COBERTURA. Gerada da tag v4.1 por "
                    "python/gerar_v4_2.py."),
        (f"C{lin}", f"EDR-062 (projeto T11: Origens e prefixo D16), EDR-060 (h7). Plano: {PLANO}. Registro celula a celula: "
                    f"{REGISTRO}. NOT_VALIDATED fisicamente."),
    ]


def _regra(texto: str) -> str:
    """Substituição de regra: o rótulo do TX vira o do T11 (com símbolo e, nos textos, sem ele)."""
    for antigo, novo in origens.MAPA_REGRA_V41_V42.items():
        texto = texto.replace(antigo, novo)
    return texto.replace("Dado do TX", "Dado do T11")


def alteracoes_v4_2(wb) -> list[dict]:
    """Células existentes que a V4.2 muda em relação à V4.1 (autorização 5 (a), (b), (d) e (e)), cada uma com aba, cel,
    antes (lido de `wb`, a V4.1 aberta), novo e ref. Não grava nada."""
    out: list[dict] = []
    # (a) dado: as células de Origem com o rótulo antigo (a 00_README A43:A44 é a própria lista, tratada abaixo)
    for aba, colunas in COLUNAS_ORIGEM.items():
        ws = wb[aba]
        for col in colunas:
            for r in range(1, ws.max_row + 1):
                v = ws[f"{col}{r}"].value
                if v in origens.MAPA_DADO_V41_V42:
                    out.append({"aba": aba, "cel": f"{col}{r}", "antes": v, "novo": origens.MAPA_DADO_V41_V42[v],
                                "ref": REF_DADO})
    # (a) a lista Origens e o significado dos dois rótulos novos
    ws = wb[ABA_LISTA]
    esperada = origens.ORIGENS_V42
    for i, rotulo in enumerate(esperada):
        r = LINHA_LISTA + i
        atual = ws[f"A{r}"].value
        if i < 2:
            antigo = (origens.TX_V41, origens.T09_V41)[i]
            if atual != antigo:
                raise ValueError(f"{ABA_LISTA}!A{r}: a origem não é a V4.1 esperada ({atual!r} != {antigo!r})")
            sig, cat = origens.SIGNIFICADO_V42[rotulo]
            out.append({"aba": ABA_LISTA, "cel": f"A{r}", "antes": atual, "novo": rotulo, "ref": REF_LISTA})
            out.append({"aba": ABA_LISTA, "cel": f"B{r}", "antes": ws[f"B{r}"].value, "novo": sig, "ref": REF_LISTA})
            out.append({"aba": ABA_LISTA, "cel": f"E{r}", "antes": ws[f"E{r}"].value, "novo": cat, "ref": REF_LISTA})
        elif atual != rotulo:
            raise ValueError(f"{ABA_LISTA}!A{r}: rótulo inesperado na V4.1 ({atual!r} != {rotulo!r})")
    # (a) regra: textos explicativos
    for aba, cel in TEXTOS_REGRA:
        antes = wb[aba][cel].value
        if not isinstance(antes, str) or "Dado do TX" not in antes:
            raise ValueError(f"{aba}!{cel}: o texto da V4.1 não cita \"Dado do TX\": {antes!r}")
        out.append({"aba": aba, "cel": cel, "antes": antes, "novo": _regra(antes), "ref": REF_REGRA})
    # (b) regra: o literal do EXACT nas duas fórmulas do D16, e nada mais
    for aba, cel in FORMULAS_D16:
        antes = wb[aba][cel].value
        literal = f'"{origens.TX_V41}"'
        if not (isinstance(antes, str) and antes.startswith("=") and antes.count(literal) == 1):
            raise ValueError(f"{aba}!{cel}: fórmula sem o literal {literal} exatamente uma vez: {antes!r}")
        out.append({"aba": aba, "cel": cel, "antes": antes,
                    "novo": antes.replace(literal, f'"{origens.T11}"'), "ref": REF_REGRA})
    # (d) emenda da h7
    antes = wb[H7[0]][H7[1]].value
    if antes != H7_ANTES:
        raise ValueError(f"{H7[0]}!{H7[1]}: texto inesperado na V4.1: {antes!r}")
    out.append({"aba": H7[0], "cel": H7[1], "antes": antes, "novo": H7_NOVO, "ref": REF_H7})
    # (e) versão e changelog
    out.append({"aba": "00_README", "cel": "A3", "antes": wb["00_README"]["A3"].value, "novo": TEXTO_A3,
                "ref": REF_VERSAO})
    for cel, novo in _changelog():
        antes = wb["15_CHANGELOG"][cel].value
        if antes is not None:
            raise ValueError(f"15_CHANGELOG!{cel} já preenchida na origem: {antes!r}")
        out.append({"aba": "15_CHANGELOG", "cel": cel, "antes": None, "novo": novo, "ref": REF_VERSAO})
    return out


def _texto(v) -> str | None:
    """Texto de uma célula, de valor ou de fórmula (inclusive de fórmula matricial)."""
    return v if isinstance(v, str) else getattr(v, "text", None)


def rotulos_antigos_restantes(wb) -> list[str]:
    """Onde ainda há `Dado do TX` ou `Dado do T09` (valor, fórmula, validação ou formatação condicional) fora do
    04_DECISION_REGISTER e do 15_CHANGELOG, descontado o texto de `origens.OUTRO_CARRO`. Lista vazia = limpo."""
    achados = []
    for ws in wb.worksheets:
        if ws.title in ABAS_HISTORICAS:
            continue
        for lin in ws.iter_rows():
            for c in lin:
                t = _texto(c.value)
                if isinstance(t, str) and any(a in t for a in ANTIGOS):
                    achados.append(f"{ws.title}!{c.coordinate}")
        for dv in ws.data_validations.dataValidation:
            if any(a in (s or "") for s in (dv.formula1, dv.formula2) for a in ANTIGOS):
                achados.append(f"{ws.title}: validação {dv.sqref}")
        for rng, regras in ws.conditional_formatting._cf_rules.items():
            if any(a in f for r in regras for f in (r.formula or []) for a in ANTIGOS):
                achados.append(f"{ws.title}: formatação condicional {rng.sqref}")
    return achados


def extrair_v41(pasta: Path) -> Path:
    """Arquivo da tag v4.1 numa pasta (binário, `git show`); nunca a V4.1 da árvore nem a viva."""
    destino = Path(pasta) / "V4_1_tag.xlsx"
    destino.write_bytes(subprocess.run(["git", "show", f"{TAG_V41}:{CAMINHO_V41}"], cwd=RAIZ, capture_output=True,
                                       check=True).stdout)
    return destino


def _recusar(caminho_v41: Path, saida: Path) -> None:
    """Recusa saída dentro de excel/ ou igual à origem. Só compara caminhos: não abre nenhum arquivo."""
    s = saida.resolve()
    if s == Path(caminho_v41).resolve():
        raise SystemExit(f"saída igual à origem: recusado ({saida})")
    if s.parent == PASTA_EXCEL.resolve() or PASTA_EXCEL.resolve() in s.parents:
        raise SystemExit(f"saída em excel/ recusada ({saida}): gere em build/ ou numa pasta temporária; a cópia "
                         "para excel/ é um passo à parte, num commit próprio")


def aplicar(wb, alteracoes: list[dict]) -> list[dict]:
    """Grava `alteracoes` em `wb` (cada "antes" tem de bater com a célula) e a formatação da linha nova do changelog.
    Devolve o registro (cópia das alterações)."""
    changelog = wb["15_CHANGELOG"]
    for col in "ABCDE":   # mesma formatação da linha anterior, como na V4.1
        changelog[f"{col}{LINHA_CHANGELOG}"]._style = copy.copy(changelog[f"{col}{LINHA_CHANGELOG - 1}"]._style)
    registro = []
    for alt in alteracoes:
        c = wb[alt["aba"]][alt["cel"]]
        if c.value != alt["antes"]:
            raise ValueError(f"{alt['aba']}!{alt['cel']}: a origem não é a V4.1 esperada ({c.value!r} != {alt['antes']!r})")
        c.value = alt["novo"]
        registro.append(dict(alt))
    wb.calculation.fullCalcOnLoad = True
    return registro


INTRO_REGISTRO = (
    "# Registro de células da V4.2 (gerado por python/gerar_v4_2.py)\n\n"
    "Gerado da tag v4.1; não edite à mão (regere com `uv run python -X utf8 python/gerar_v4_2.py --registro "
    f"{REGISTRO}`). Cada célula existente que a V4.2 muda em relação à V4.1 (autorização 5 do plano {PLANO}), com o "
    "texto de antes e o de depois: a Origem das células de dado do TX e do T09 (`◐ Dado de outro carro (TX, T09)`), a lista "
    "Origens e o significado dos dois rótulos novos, os textos que explicam a Origem, o literal `\"● Dado do T11\"` do `EXACT` "
    "do D16 (16_EQUILIBRIO B188 e B189), a emenda da h7 (21_COBERTURA D12), a versão (00_README A3) e a linha 22 do "
    "15_CHANGELOG. As 117569 fórmulas da V4.1 e os 1145 nomes ficam iguais.\n\n")


def escrever_registro_v42(registro: list[dict], caminho: Path) -> None:
    """Registro em Markdown: a introdução (INTRO_REGISTRO) e a tabela de `atualizar_planilha.escrever_registro`, LF."""
    caminho = Path(caminho)
    escrever_registro(registro, caminho)
    caminho.write_text(INTRO_REGISTRO + caminho.read_text(encoding="utf-8"), encoding="utf-8", newline="\n")


def gerar(v41: Path | None, saida: Path, registro: Path | None = None) -> None:
    """Abre a V4.1 (`v41`, ou a extraída da tag v4.1), aplica `alteracoes_v4_2`, confere que não sobrou rótulo antigo e
    grava `saida`. Com `registro`, grava também o registro célula a célula (.md). A recusa de `saida` vem antes de abrir
    qualquer arquivo."""
    saida = Path(saida)
    with tempfile.TemporaryDirectory(prefix="dds_v42_") as tmp:
        if v41 is None:
            _recusar(Path(tmp) / "V4_1_tag.xlsx", saida)
            origem = extrair_v41(Path(tmp))
        else:
            _recusar(Path(v41), saida)
            origem = Path(v41)
        if ler_tabela(CSV_NOMES):
            raise NotImplementedError("data/nomes_V4_2.csv com nomes: as metas entram na Task 3")
        wb = openpyxl.load_workbook(origem)
        try:
            reg = aplicar(wb, alteracoes_v4_2(wb))
            sobra = rotulos_antigos_restantes(wb)
            if sobra:
                raise ValueError(f"rótulo antigo (Dado do TX / Dado do T09) ainda presente: {sobra[:5]}")
            saida.parent.mkdir(parents=True, exist_ok=True)
            wb.save(saida)
        finally:
            wb.close()
    if registro is not None:
        escrever_registro_v42(reg, Path(registro))


def main() -> None:
    p = argparse.ArgumentParser(description="Gera a V4.2 (metas e coerência; projeto T11) a partir da V4.1 da tag v4.1.")
    p.add_argument("--v41", type=Path, default=None, help="arquivo da V4.1 (padrão: extraído da tag v4.1)")
    p.add_argument("--saida", type=Path, default=SAIDA_PADRAO)
    p.add_argument("--registro", type=Path, default=None, help="registro célula a célula (.md)")
    p.add_argument("--nomes-csv", type=Path, default=None,
                   help="regrava a tabela de nomes da V4.2 (CSV, só o cabeçalho nesta etapa) antes de gerar")
    a = p.parse_args()
    if a.nomes_csv is not None:
        escrever_tabela([], a.nomes_csv)
    gerar(a.v41, a.saida, a.registro)
    print(f"salvo {a.saida}")


if __name__ == "__main__":
    main()
